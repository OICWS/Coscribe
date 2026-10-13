"""Spike: ppt-master's element emitter (copied unmodified) + our own traversal and packaging.

Usage: walker.py <svg_dir> <out.pptx>
Own code here: the tree walk, group wrapping, background promotion, slide packaging with python-pptx.
"""

from __future__ import annotations

import copy
import re
import sys
from pathlib import Path
from xml.etree import ElementTree as ET

sys.path.insert(0, str(Path(__file__).parent / "vendor"))

from lxml import etree  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Emu  # noqa: E402

from svg_to_pptx.drawingml import elements as E  # noqa: E402
from svg_to_pptx.drawingml.context import ConvertContext, ShapeResult  # noqa: E402
from svg_to_pptx.drawingml.styles import build_effect_xml, get_element_opacity  # noqa: E402
from svg_to_pptx.drawingml.utils import (  # noqa: E402
    EMU_PER_PX,
    SVG_NS,
    _extract_inheritable_styles,
    project_definition_index,
    resolve_url_id,
)

CONVERTERS = {
    "rect": E.convert_rect, "circle": E.convert_circle, "ellipse": E.convert_ellipse,
    "line": E.convert_line, "path": E.convert_path, "polygon": E.convert_polygon,
    "polyline": E.convert_polyline, "text": E.convert_text, "image": E.convert_image,
}
NON_VISUAL = {"defs", "title", "desc", "metadata", "linearGradient", "radialGradient", "filter", "clipPath",
              "marker", "style", "symbol"}
NSDECL = ('xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
          'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
          'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"')


def tag(e: ET.Element) -> str:
    return e.tag.replace(f"{{{SVG_NS}}}", "")


def translate(transform: str) -> tuple[float, float]:
    m = re.fullmatch(r"\s*translate\(\s*([-\d.]+)[\s,]*([-\d.]+)?\s*\)\s*", transform or "")
    if not transform:
        return 0.0, 0.0
    if not m:
        raise ValueError(f"spike handles translate() only, got {transform!r}")
    return float(m.group(1)), float(m.group(2) or 0)


def convert(elem: ET.Element, ctx: ConvertContext) -> ShapeResult | None:
    t = tag(elem)
    if t in NON_VISUAL:
        return None
    if t == "g":
        return convert_group(elem, ctx)
    fn = CONVERTERS.get(t)
    if fn is None:
        raise ValueError(f"unsupported <{t}>")
    return fn(elem, ctx)


def convert_group(elem: ET.Element, ctx: ConvertContext) -> ShapeResult | None:
    dx, dy = translate(elem.get("transform", ""))
    opacity = get_element_opacity(elem) or 1.0
    filter_id = resolve_url_id(elem.get("filter", ""))
    child_ctx = ctx.child(dx=dx, dy=dy, style_overrides=_extract_inheritable_styles(elem),
                          opacity_multiplier=opacity, filter_id=filter_id or ctx.filter_id)
    child_ctx.depth = ctx.depth + 1
    results = [r for r in (convert(c, child_ctx) for c in elem) if r is not None]
    if not results:
        return None
    if len(results) == 1 and not elem.get("id"):
        return results[0]
    boxes = [r.bounds_emu for r in results if r.bounds_emu]
    x0, y0 = min(b[0] for b in boxes), min(b[1] for b in boxes)
    x1, y1 = max(b[2] for b in boxes), max(b[3] for b in boxes)
    w, h = max(x1 - x0, 1), max(y1 - y0, 1)
    gid = ctx.next_id()
    effect = build_effect_xml(ctx.defs[filter_id], child_ctx.opacity_multiplier) if filter_id in ctx.defs else ""
    name = elem.get("id") or f"Group {gid}"
    xml = (f'<p:grpSp><p:nvGrpSpPr><p:cNvPr id="{gid}" name="{name}"/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>'
           f'<p:grpSpPr><a:xfrm><a:off x="{x0}" y="{y0}"/><a:ext cx="{w}" cy="{h}"/>'
           f'<a:chOff x="{x0}" y="{y0}"/><a:chExt cx="{w}" cy="{h}"/></a:xfrm>{effect}</p:grpSpPr>'
           + "\n".join(r.xml for r in results) + "</p:grpSp>")
    return ShapeResult(xml=xml, bounds_emu=(x0, y0, x0 + w, y0 + h))


def background(root: ET.Element) -> str | None:
    first = next((c for c in root if tag(c) not in NON_VISUAL), None)
    if first is None or tag(first) != "rect":
        return None
    if (first.get("x", "0"), first.get("y", "0")) != ("0", "0") or first.get("width") != "1280":
        return None
    fill = first.get("fill", "")
    if not re.fullmatch(r"#[0-9A-Fa-f]{6}", fill):
        return None
    root.remove(first)
    return f'<p:bg><p:bgPr><a:solidFill><a:srgbClr val="{fill[1:].upper()}"/></a:solidFill><a:effectLst/></p:bgPr></p:bg>'


def slide_shapes(svg_path: Path) -> tuple[str | None, list[str]]:
    root = ET.parse(svg_path).getroot()
    defs, _ = project_definition_index(root)
    ctx = ConvertContext(defs=defs, svg_dir=svg_path.parent, resource_root=svg_path.parent,
                         inherited_styles=_extract_inheritable_styles(root))
    bg = background(root)
    out = []
    for child in root:
        r = convert(child, ctx)
        if r is not None:
            out.append(r.xml)
    return bg, out


def main(svg_dir: Path, out: Path) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(1280 * EMU_PER_PX), Emu(720 * EMU_PER_PX)
    blank = prs.slide_layouts[6]
    for svg in sorted(svg_dir.glob("*.svg")):
        slide = prs.slides.add_slide(blank)
        bg, shapes = slide_shapes(svg)
        tree = slide.shapes._spTree
        for xml in shapes:
            frag = etree.fromstring(f"<p:wrap {NSDECL}>{xml}</p:wrap>")
            for el in frag:
                tree.append(el)
        if bg:
            csld = slide._element.cSld
            csld.insert(0, etree.fromstring(bg.replace("<p:bg>", f"<p:bg {NSDECL}>", 1)))
        print(svg.name, len(shapes), "top-level shapes")
    prs.save(out)


if __name__ == "__main__":
    main(Path(sys.argv[1]), Path(sys.argv[2]))
