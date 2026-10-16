"""Automatic checks on one produced workbook, in plain Python: no model is called here.

Each check returns a plain dict so a run can be written as JSON and two runs compared.
`check_workbook` runs the ones a task's `expect.json` asks for. Values are read from a copy
that LibreOffice recalculated (the path `recalc_xlsx` takes), formulas and styles from the file
the agent made, so a workbook of pasted numbers and one of formulas look different.
"""

from __future__ import annotations

import datetime as dt
import re
import shutil
import tempfile
import unicodedata
import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from .render import render_pages
from .tasks import Task, XlsxExpect
from .text_checks import check_facts, find_inventions, plain_number

_DEFAULT_TOLERANCE = 0.005
_DEFAULT_COLUMN_WIDTH = 8.43  # characters, what Excel and LibreOffice give a column nobody sized
_LIST_CAP = 20
_CHART_KINDS = {
    "line": ("lineChart", "line3DChart"),
    "bar": ("barChart", "bar3DChart"),
    "pie": ("pieChart", "pie3DChart", "doughnutChart"),
    "scatter": ("scatterChart",),
    "area": ("areaChart", "area3DChart"),
}
_C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def split_address(address: str) -> tuple[str, str]:
    """`Sheet!B2:B9` or `'My sheet'!B2` into the sheet name and the cell or range."""
    sheet, _, cells = address.rpartition("!")
    if len(sheet) >= 2 and sheet[0] == sheet[-1] == "'":
        sheet = sheet[1:-1].replace("''", "'")
    return sheet, cells.replace("$", "")


def cells_in(workbook: Any, address: str) -> list[Any]:
    """The cells an address names; a sheet that does not exist raises `KeyError`."""
    sheet, cells = split_address(address)
    selection = workbook[sheet][cells]
    if hasattr(selection, "row"):
        return [selection]
    if selection and isinstance(selection[0], tuple):
        return [cell for row in selection for cell in row]
    return list(selection)


def _plain(value: Any) -> Any:
    """What a person reads from a cell: a date as `YYYY-MM-DD`."""
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time(0) else value.isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    return value


def _is_number(value: Any) -> bool:
    return isinstance(value, int | float) and not isinstance(value, bool)


def value_matches(got: Any, spec: dict[str, Any]) -> bool:
    expected = spec["value"]
    if spec.get("type") == "date":
        return isinstance(got, dt.date) and _plain(got) == expected
    if expected is None or expected == "":
        return got is None or got == ""
    if _is_number(expected):
        tolerance = float(spec.get("tol", _DEFAULT_TOLERANCE))
        return _is_number(got) and abs(got - expected) <= tolerance
    return isinstance(got, str) and got == expected


def check_opens(path: Path) -> dict[str, Any]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(str(path))
    except Exception as exc:  # noqa: BLE001 -- any failure to open is the finding
        return {"ok": False, "error": f"{type(exc).__name__}: {exc}", "sheets": []}
    names = list(workbook.sheetnames)
    workbook.close()
    return {"ok": True, "sheets": names}


def recalculate(
    path: Path, scratch: Path, name: str = "recalc.xlsx"
) -> tuple[Path, dict[str, Any]]:
    """A recalculated copy, through the engine the tools use; the file is not touched."""
    from coscribe.tools.office.xlsx.calc import recalculate as engine

    copy = scratch / name
    shutil.copyfile(path, copy)
    result = engine(copy)
    report = {
        "ok": result["status"] == "success",
        "status": result["status"],
        "skipped_reason": result.get("skipped_reason"),
        "formulas": result.get("total_formulas"),
        "errors": result.get("total_errors"),
        "error_locations": result.get("error_locations"),
    }
    return copy, report


def workbook_texts(values: Any) -> list[str]:
    """The text of each sheet as a person reads it: the sheet's name, then one line per row."""
    texts = []
    for sheet in values.worksheets:
        lines = [sheet.title]
        for row in sheet.iter_rows(values_only=True):
            shown = [
                plain_number(v) if _is_number(v) else str(_plain(v))
                for v in row
                if v is not None and v != ""
            ]
            if shown:
                lines.append(" ".join(shown))
        texts.append("\n".join(lines))
    return texts


def check_sheets(names: list[str], wanted: list[str]) -> dict[str, Any]:
    present = {n.casefold() for n in names}
    missing = [w for w in wanted if w.casefold() not in present]
    return {"ok": not missing, "missing": missing}


def check_cells(values: Any, specs: list[dict[str, Any]]) -> dict[str, Any]:
    wrong = []
    for spec in specs:
        try:
            (cell,) = cells_in(values, spec["at"])
        except KeyError:
            wrong.append({"at": spec["at"], "expected": spec["value"], "got": "no such sheet"})
            continue
        if not value_matches(cell.value, spec):
            wrong.append({"at": spec["at"], "expected": spec["value"], "got": _plain(cell.value)})
    return {
        "ok": not wrong,
        "checked": len(specs),
        "matched": len(specs) - len(wrong),
        "wrong": wrong[:_LIST_CAP],
    }


def _is_formula(cell: Any) -> bool:
    return isinstance(cell.value, str) and cell.value.startswith("=")


def _check_kind_of_cell(workbook: Any, ranges: list[str], want_formula: bool) -> dict[str, Any]:
    wrong, total = [], 0
    for address in ranges:
        try:
            cells = cells_in(workbook, address)
        except KeyError:
            wrong.append(f"{address} (no such sheet)")
            continue
        for cell in cells:
            total += 1
            if _is_formula(cell) != want_formula:
                wrong.append(f"{cell.parent.title}!{cell.coordinate}")
    return {"ok": not wrong, "checked": total, "wrong": wrong[:_LIST_CAP]}


def check_formulas(workbook: Any, ranges: list[str]) -> dict[str, Any]:
    """Every cell of the ranges holds a formula: a total typed in as a number is the finding."""
    return _check_kind_of_cell(workbook, ranges, want_formula=True)


def check_constants(workbook: Any, ranges: list[str]) -> dict[str, Any]:
    """Inputs stay values: a formula there means the model computed what the user supplies."""
    return _check_kind_of_cell(workbook, ranges, want_formula=False)


def check_number_formats(workbook: Any, specs: list[dict[str, Any]]) -> dict[str, Any]:
    wrong, total = [], 0
    for spec in specs:
        try:
            cells = cells_in(workbook, spec["at"])
        except KeyError:
            wrong.append({"at": spec["at"], "format": "no such sheet"})
            continue
        for cell in cells:
            total += 1
            shown = cell.number_format
            # `yyyy\-mm\-dd` and `yyyy-mm-dd` show the same: a backslash only escapes a literal.
            plain = shown.replace("\\", "")
            if not any(want in plain for want in spec["any_of"]):
                wrong.append({"at": f"{cell.parent.title}!{cell.coordinate}", "format": shown})
    return {"ok": not wrong, "checked": total, "wrong": wrong[:_LIST_CAP]}


def check_bold(workbook: Any, ranges: list[str]) -> dict[str, Any]:
    wrong, total = [], 0
    for address in ranges:
        try:
            cells = cells_in(workbook, address)
        except KeyError:
            wrong.append(f"{address} (no such sheet)")
            continue
        for cell in cells:
            total += 1
            if not cell.font.b:
                wrong.append(f"{cell.parent.title}!{cell.coordinate}")
    return {"ok": not wrong, "checked": total, "wrong": wrong[:_LIST_CAP]}


def check_used_range(workbook: Any, limits: dict[str, str]) -> dict[str, Any]:
    from openpyxl.utils.cell import range_boundaries

    outside = []
    for sheet_name, cells in limits.items():
        if sheet_name not in workbook.sheetnames:
            outside.append(f"{sheet_name} (no such sheet)")
            continue
        first_col, first_row, last_col, last_row = range_boundaries(cells)
        for row in workbook[sheet_name].iter_rows():
            for cell in row:
                if cell.value is None or cell.value == "":
                    continue
                inside = first_col <= cell.column <= last_col and first_row <= cell.row <= last_row
                if not inside:
                    outside.append(f"{sheet_name}!{cell.coordinate}")
    return {"ok": not outside, "count": len(outside), "outside": outside[:_LIST_CAP]}


def _color(color: Any) -> Any:
    if color is None:
        return None
    if color.type == "rgb":
        return str(color.rgb)[-6:]  # LibreOffice writes the alpha byte as FF, openpyxl as 00
    return (color.type, color.theme if color.type == "theme" else color.indexed, color.tint)


def _side(side: Any) -> Any:
    return None if side is None else side.style


def cell_style(cell: Any) -> dict[str, Any]:
    """What a person sees of a cell besides its value. The spellings LibreOffice changes when it
    saves a workbook (an escaped `\\-` or a quoted `"$"` in a number format, `general` for no
    alignment, the alpha byte of a color) are folded, so the comparison finds a lost style and not
    a rewrite."""
    border = cell.border
    return {
        "number_format": cell.number_format.replace("\\", "").replace('"', ""),
        "bold": bool(cell.font.b),
        "italic": bool(cell.font.i),
        "size": cell.font.sz,
        "font": cell.font.name,
        "font_color": _color(cell.font.color),
        "fill": cell.fill.patternType,
        "fill_color": _color(cell.fill.fgColor) if cell.fill.patternType else None,
        "border": [
            _side(border.left),
            _side(border.right),
            _side(border.top),
            _side(border.bottom),
        ],
        "align": cell.alignment.horizontal or "general",
        "wrap": bool(cell.alignment.wrap_text),
    }


def check_untouched_styles(original: Any, produced: Any, ranges: list[str]) -> dict[str, Any]:
    """The styles of cells the brief does not ask to change, against the workbook the user gave."""
    changed, total = [], 0
    for address in ranges:
        sheet, _ = split_address(address)
        if sheet not in produced.sheetnames:
            changed.append({"at": address, "difference": "the sheet is gone"})
            continue
        for before in cells_in(original, address):
            total += 1
            was, now = cell_style(before), cell_style(produced[sheet][before.coordinate])
            diff = {k: [was[k], now[k]] for k in was if was[k] != now[k]}
            if diff:
                changed.append({"at": f"{sheet}!{before.coordinate}", "difference": diff})
    return {"ok": not changed, "checked": total, "changed": changed[:_LIST_CAP]}


def _count_rules(workbook: Any) -> int:
    return sum(
        len(conditional.rules)
        for sheet in workbook.worksheets
        for conditional in sheet.conditional_formatting
    )


def _count_validations(workbook: Any) -> int:
    return sum(len(sheet.data_validations.dataValidation) for sheet in workbook.worksheets)


def check_features_kept(original: Any, produced: Any, minimum_rules: int) -> dict[str, Any]:
    """The conditional formats and validations of the original are still there."""
    before = {"rules": _count_rules(original), "validations": _count_validations(original)}
    after = {"rules": _count_rules(produced), "validations": _count_validations(produced)}
    return {
        "ok": after["rules"] >= max(before["rules"], minimum_rules)
        and after["validations"] >= before["validations"],
        "before": before,
        "after": after,
    }


def check_original_keeps(original: Path, wanted: list[str]) -> dict[str, Any]:
    """Whether the user's file still holds each feature, by a text in its sheets' XML."""
    with zipfile.ZipFile(original) as archive:
        sheets = "".join(
            archive.read(name).decode("utf-8", "replace")
            for name in archive.namelist()
            if name.startswith("xl/worksheets/sheet")
        )
    lost = [text for text in wanted if text not in sheets]
    return {"ok": not lost, "lost": lost}


def _points_in(reference: str) -> int:
    """The cells a series reference names, however many areas: `(Data!$A$1:$A$3,Data!$B$1:$B$3)`."""
    from openpyxl.utils.cell import range_boundaries

    total = 0
    for part in re.findall(r"\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?", reference):
        first_col, first_row, last_col, last_row = range_boundaries(part.replace("$", ""))
        total += (last_col - first_col + 1) * (last_row - first_row + 1)
    return total


def _text(element: Any) -> str:
    return "" if element is None else "".join(t.text or "" for t in element.iter(f"{_A}t"))


def _chart_facts(chart_xml: bytes) -> dict[str, Any]:
    root = ET.fromstring(chart_xml)  # noqa: S314 -- the agent's own file
    plot = root.find(f"{_C}chart/{_C}plotArea")
    kind, points = None, 0
    if plot is not None:
        for child in plot:
            name = child.tag.removeprefix(_C)
            if name.endswith("Chart"):
                kind = kind or name
                for series in child.iterfind(f"{_C}ser"):
                    for data in (f"{_C}val", f"{_C}yVal"):  # scatter charts name theirs yVal
                        for reference in series.iterfind(f"{data}//{_C}f"):
                            points = max(points, _points_in(reference.text or ""))
    title = root.find(f"{_C}chart/{_C}title")
    axes = [
        _text(axis.find(f"{_C}title"))
        for axis in (plot if plot is not None else [])
        if axis.tag.endswith("Ax") and axis.find(f"{_C}title") is not None
    ]
    return {"kind": kind, "title": _text(title), "axis_titles": axes, "points": points}


def read_charts(path: Path) -> list[dict[str, Any]]:
    with zipfile.ZipFile(path) as archive:
        names = sorted(n for n in archive.namelist() if re.fullmatch(r"xl/charts/chart\d+\.xml", n))
        return [_chart_facts(archive.read(n)) for n in names]


def _chart_fits(chart: dict[str, Any], spec: dict[str, Any]) -> bool:
    kinds = _CHART_KINDS.get(str(spec.get("kind")), ())
    titled = not spec.get("title") or check_facts(chart["title"], [[spec["title"]]], [])["ok"]
    return (
        (not kinds or chart["kind"] in kinds)
        and titled
        and chart["points"] >= int(spec.get("points_min", 0))
        and all(
            check_facts(" | ".join(chart["axis_titles"]), [[t]], [])["ok"]
            for t in spec.get("axis_titles", [])
        )
    )


_SCALES = (1, 1_000, 10_000)


def check_figures(values: Any, specs: list[dict[str, Any]]) -> dict[str, Any]:
    """Each number appears in some cell, in ones, thousands or 万, within a relative tolerance:
    where it sits and in which unit is the model's choice in an open task."""
    shown = [
        cell.value
        for sheet in values.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if _is_number(cell.value)
    ]
    missing = []
    for spec in specs:
        expected, rel = float(spec["value"]), float(spec.get("rel", 0.001))
        wanted = [expected / scale for scale in _SCALES]
        if not any(abs(got - w) <= abs(w) * rel for got in shown for w in wanted):
            missing.append(spec["value"])
    return {"ok": not missing, "found": len(specs) - len(missing), "missing": missing}


def check_input_probe(path: Path, probe: dict[str, Any], scratch: Path) -> dict[str, Any]:
    """Change one assumption, found by its value, on a copy and recalculate: results wired to
    the assumption move to the expected figures, pasted ones do not."""
    from openpyxl import load_workbook

    workbook = load_workbook(str(path))
    found = [
        cell
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if _is_number(cell.value) and abs(cell.value - probe["find"]) <= 1e-9
    ]
    if not found:
        return {"ok": False, "error": f"no constant cell holds {probe['find']}"}
    for cell in found:
        cell.value = probe["set"]
    changed = scratch / "input-probe.xlsx"
    workbook.save(str(changed))
    copy, recalc = recalculate(changed, scratch, "input-probe-recalc.xlsx")
    values = load_workbook(str(copy), data_only=True)
    report = check_figures(values, probe["figures"])
    values.close()
    return {**report, "cells_set": len(found), "recalc_ok": recalc["ok"],
            "ok": report["ok"] and recalc["ok"]}


def check_formula_count(workbook: Any, minimum: int) -> dict[str, Any]:
    count = sum(
        1
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    )
    return {"ok": count >= minimum, "count": count, "minimum": minimum}


def check_charts(charts: list[dict[str, Any]], specs: list[dict[str, Any]]) -> dict[str, Any]:
    missing = [s for s in specs if not any(_chart_fits(c, s) for c in charts)]
    return {"ok": not missing, "found": charts, "missing": missing}


def _offered_values(formula: str, sheet_name: str, values: Any) -> list[str] | None:
    """The choices of a list validation, typed in (`"a,b"`) or read from the cells it points at."""
    formula = formula.strip()
    if formula.startswith('"') and formula.endswith('"'):
        return [v.strip() for v in formula[1:-1].split(",")]
    try:
        reference = formula if "!" in formula else f"{sheet_name}!{formula}"
        return [str(c.value) for c in cells_in(values, reference) if c.value is not None]
    except (KeyError, ValueError):
        return None


def check_validations(workbook: Any, values: Any, specs: list[dict[str, Any]]) -> dict[str, Any]:
    """A dropdown covers every cell of a range and offers exactly the values asked for."""
    missing = []
    for spec in specs:
        sheet_name, _ = split_address(spec["at"])
        if sheet_name not in workbook.sheetnames:
            missing.append({**spec, "reason": "no such sheet"})
            continue
        wanted = {c.coordinate for c in cells_in(workbook, spec["at"])}
        covered: set[str] = set()
        for validation in workbook[sheet_name].data_validations.dataValidation:
            offered = _offered_values(validation.formula1 or "", sheet_name, values)
            if (
                validation.type == "list"
                and offered
                and sorted(offered) == sorted(str(v) for v in spec["values"])
            ):
                for part in str(validation.sqref).split():
                    covered |= {c.coordinate for c in cells_in(workbook, f"{sheet_name}!{part}")}
        left = sorted(wanted - covered)
        if left:
            missing.append({**spec, "reason": f"no matching list on {left[:5]}"})
    return {"ok": not missing, "missing": missing}


def check_probe(path: Path, probe: dict[str, Any], scratch: Path) -> dict[str, Any]:
    """Change inputs on a copy and recalculate it: a model whose totals are formulas moves, one
    that pasted numbers does not."""
    from openpyxl import load_workbook

    workbook = load_workbook(str(path))
    for address, value in probe["set"].items():
        (cell,) = cells_in(workbook, address)
        cell.value = value
    changed = scratch / "probe.xlsx"
    workbook.save(str(changed))
    copy, recalc = recalculate(changed, scratch, "probe-recalc.xlsx")
    values = load_workbook(str(copy), data_only=True)
    report = check_cells(values, probe["expect"])
    values.close()
    return {**report, "recalc_ok": recalc["ok"], "ok": report["ok"] and recalc["ok"]}


def _text_width(text: str) -> float:
    """Columns are measured in characters of the default font; a CJK character takes two."""
    return float(sum(2 if unicodedata.east_asian_width(c) in "WF" else 1 for c in text))


def _column_widths(sheet: Any) -> dict[int, float]:
    from openpyxl.utils import column_index_from_string

    widths: dict[int, float] = {}
    for letter, dimension in sheet.column_dimensions.items():
        if dimension.width:
            first = column_index_from_string(letter)
            for index in range(first, (dimension.max or first) + 1):
                widths[index] = float(dimension.width)
    return widths


def clipped_texts(workbook: Any, values: Any) -> list[str]:
    """Text that is wider than its column while the cell to its right is taken, so a reader sees
    `Order Dat` for `Order Date`. Wrapped and merged cells are left out; text that is centred or
    right-aligned can spill to the left and is still counted, so this errs towards reporting."""
    clipped = []
    for sheet in workbook.worksheets:
        widths = _column_widths(sheet)
        shown = values[sheet.title]
        for row in shown.iter_rows():
            for cell in row:
                text = cell.value
                if not isinstance(text, str) or not text.strip():
                    continue
                if workbook[sheet.title][cell.coordinate].alignment.wrap_text:
                    continue
                if any(cell.coordinate in merged for merged in sheet.merged_cells.ranges):
                    continue
                width = widths.get(cell.column, _DEFAULT_COLUMN_WIDTH)
                neighbour = shown.cell(row=cell.row, column=cell.column + 1).value
                if _text_width(text) > width + 0.5 and neighbour not in (None, ""):
                    clipped.append(f"{sheet.title}!{cell.coordinate} {text[:30]!r}")
    return clipped


def hash_cells(pdf: Path) -> int:
    """Cells printed as `###` because the number does not fit its column; LibreOffice draws them
    the way Excel does."""
    import pdfplumber

    with pdfplumber.open(str(pdf)) as document:
        text = "\n".join(page.extract_text() or "" for page in document.pages)
    return len(re.findall(r"(?<![^\s])#{3,}(?![^\s])", text))


def check_legibility(workbook: Any, values: Any, pdf: Path | None) -> dict[str, Any]:
    """Every number and label can be read as printed: nothing is `###` and no text is cut off."""
    clipped = clipped_texts(workbook, values)
    hashes = None if pdf is None else hash_cells(pdf)
    return {
        "ok": not clipped and not hashes,
        "hash_cells": hashes,
        "clipped": clipped[:_LIST_CAP],
        "clipped_count": len(clipped),
    }


def _guarded(report: dict[str, Any], key: str, check: Any) -> None:
    """One check that cannot run says why instead of ending the report."""
    try:
        report[key] = check()
    except Exception as exc:  # noqa: BLE001 -- the failure is the finding
        report[key] = {"ok": False, "error": f"{type(exc).__name__}: {exc}"}


def _run_checks(
    task: Task,
    expect: XlsxExpect,
    path: Path,
    formulas: Any,
    values: Any,
    scratch: Path,
    out_dir: Path,
    report: dict[str, Any],
    workspace: Path,
) -> None:
    from openpyxl import load_workbook

    texts = workbook_texts(values)
    _guarded(report, "facts", lambda: check_facts("\n".join(texts), expect.facts, expect.absent))
    _guarded(
        report,
        "inventions",
        lambda: find_inventions(texts, task.source_texts(), expect.allowed_numbers, "sheets"),
    )
    if expect.sheets:
        _guarded(report, "sheets", lambda: check_sheets(list(formulas.sheetnames), expect.sheets))
    if expect.cells:
        _guarded(report, "cells", lambda: check_cells(values, expect.cells))
    if expect.formulas:
        _guarded(report, "formulas", lambda: check_formulas(formulas, expect.formulas))
    if expect.constants:
        _guarded(report, "constants", lambda: check_constants(formulas, expect.constants))
    if expect.formulas_min:
        _guarded(
            report, "formulas_min", lambda: check_formula_count(formulas, expect.formulas_min)
        )
    if expect.figures:
        _guarded(report, "figures", lambda: check_figures(values, expect.figures))
    if expect.number_formats:
        _guarded(
            report, "number_formats", lambda: check_number_formats(formulas, expect.number_formats)
        )
    if expect.bold:
        _guarded(report, "bold", lambda: check_bold(formulas, expect.bold))
    if expect.used_range:
        _guarded(report, "used_range", lambda: check_used_range(formulas, expect.used_range))
    if expect.original:
        given = next(p for p in task.inputs if p.name == expect.original)
        original = load_workbook(str(given))
        if expect.untouched:
            _guarded(
                report,
                "untouched",
                lambda: check_untouched_styles(original, formulas, expect.untouched),
            )
        _guarded(
            report,
            "features",
            lambda: check_features_kept(original, formulas, expect.conditional_formats_min),
        )
        if expect.original_keeps:
            in_workspace = workspace / expect.original
            keeps = expect.original_keeps
            _guarded(report, "original_keeps", lambda: check_original_keeps(in_workspace, keeps))
    if expect.charts:
        _guarded(report, "charts", lambda: check_charts(read_charts(path), expect.charts))
    if expect.validations:
        _guarded(
            report, "validations", lambda: check_validations(formulas, values, expect.validations)
        )
    if expect.input_probe:
        input_probe = expect.input_probe
        _guarded(
            report, "input_probe", lambda: check_input_probe(path, input_probe, scratch)
        )
    if expect.probe:
        probe = expect.probe
        _guarded(report, "probe", lambda: check_probe(path, probe, scratch))
    _guarded(report, "render", lambda: render_pages(path, out_dir))
    pdf = out_dir / report["render"]["pdf"] if report["render"].get("pdf") else None
    _guarded(report, "legibility", lambda: check_legibility(formulas, values, pdf))


def check_workbook(task: Task, path: Path, workspace: Path, out_dir: Path) -> dict[str, Any]:
    """The checks a task's expectations ask for; one that cannot run says why and the rest go on."""
    from openpyxl import load_workbook

    expect = task.expect
    assert isinstance(expect, XlsxExpect)
    report: dict[str, Any] = {"output": str(path.relative_to(workspace))}
    report["opens"] = check_opens(path)
    if not report["opens"]["ok"]:
        return report
    formulas = load_workbook(str(path))
    with tempfile.TemporaryDirectory(prefix="coscribe_eval_") as scratch_name:
        scratch = Path(scratch_name)
        try:
            copy, report["recalc"] = recalculate(path, scratch)
        except Exception as exc:  # noqa: BLE001 -- the values are then read as the agent saved them
            copy, report["recalc"] = path, {"ok": False, "error": f"{type(exc).__name__}: {exc}"}
        values = load_workbook(str(copy), data_only=True)
        try:
            _run_checks(task, expect, path, formulas, values, scratch, out_dir, report, workspace)
        finally:
            values.close()
            formulas.close()
    return report
