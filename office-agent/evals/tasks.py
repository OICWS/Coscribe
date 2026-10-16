"""An evaluation task: what a user types, the files they give, and what the result must hold."""

from __future__ import annotations

import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

KINDS = ("pptx", "xlsx", "docx")
PPTX_TASKS = Path(__file__).resolve().parent / "pptx" / "tasks"
TASK_ROOTS = {kind: Path(__file__).resolve().parent / kind / "tasks" for kind in KINDS}

# Plain-text inputs; the other kinds `input_text` knows are read with their own library.
_TEXT_SUFFIXES = {".md", ".txt", ".csv", ".json"}


@dataclass(frozen=True)
class Expect:
    """`facts`: each entry must appear in the deck; a list inside it holds alternative spellings."""

    facts: list[list[str]]
    pages_min: int
    pages_max: int
    absent: list[str] = field(default_factory=list)
    allowed_numbers: list[str] = field(default_factory=list)
    pictures_min: int = 0
    template: str | None = None  # an input deck whose layouts the result must keep
    pictures_of: str | None = None  # an input deck whose every picture the result must show
    notes: float = 0.0  # the share of slides that must carry speaker notes


@dataclass(frozen=True)
class XlsxExpect:
    """What a workbook must hold. A cell address is `Sheet!B2`, a range `Sheet!B2:B13`.

    `cells`: `{"at", "value", "tol"?}`, read after LibreOffice recalculated a copy; a number is
    compared within `tol` (default 0.005), a date as `YYYY-MM-DD`, text exactly.
    `formulas` / `constants`: ranges whose cells must (not) be formulas.
    `number_formats`: `{"at", "any_of"}`, each cell's format must contain one of the strings.
    `bold`: ranges whose cells must be bold.
    `used_range`: per sheet, the range nothing may be written outside of.
    `original` and `untouched`: for an edit, the input workbook and the ranges whose styles must
    be as they were. `original_keeps`: text that must still be in the original's sheets in the
    workspace afterwards (a feature an edit could drop: kept in place, or the file left alone).
    `probe`: `{"set": {address: value}, "expect": [cells]}`, inputs changed on a copy that is
    then recalculated.
    `charts`: `{"kind", "title", "points_min", "axis_titles"}`; `validations`: `{"at", "values"}`.
    For an open task, where the model chooses the layout: `formulas_min`, at least that many
    formula cells; `figures`, `{"value", "rel"?}` numbers some cell must show, in ones, thousands
    or 万; `input_probe`, `{"find", "set", "figures"}`, the constant cell holding `find` set to
    `set` on a copy, which is recalculated and must then show `figures`.
    """

    facts: list[list[str]]
    absent: list[str] = field(default_factory=list)
    allowed_numbers: list[str] = field(default_factory=list)
    sheets: list[str] = field(default_factory=list)
    cells: list[dict[str, object]] = field(default_factory=list)
    formulas: list[str] = field(default_factory=list)
    constants: list[str] = field(default_factory=list)
    number_formats: list[dict[str, object]] = field(default_factory=list)
    bold: list[str] = field(default_factory=list)
    used_range: dict[str, str] = field(default_factory=dict)
    original: str | None = None
    untouched: list[str] = field(default_factory=list)
    original_keeps: list[str] = field(default_factory=list)
    conditional_formats_min: int = 0
    formulas_min: int = 0
    figures: list[dict[str, object]] = field(default_factory=list)
    input_probe: dict[str, object] | None = None
    probe: dict[str, object] | None = None
    charts: list[dict[str, object]] = field(default_factory=list)
    validations: list[dict[str, object]] = field(default_factory=list)


@dataclass(frozen=True)
class DocxExpect:
    """What a document must hold.

    `facts`, `absent`, `allowed_numbers`: as for a deck, over the text a reader sees once
    tracked changes are accepted. `pages`: `[min, max]`, optional.
    `headings`: `{"text", "level"?}`, found in this order among the document's headings (a
    paragraph in a Title or Heading style; Title is level 0), other headings allowed between.
    `toc`: a real table-of-contents field is there.
    `tables`: `{"rows_min", "cols_min", "header"?}`, each satisfied by a different table.
    `images_min`; `image_fit`: no picture wider than the text area.
    `captions`: `{"style", "starts_with", "below_image"?}`.
    `tracked`: `{"min_ins", "min_del", "rejected_equals_original"}`: the edit is a redline, and
    rejecting all of it gives the input back.
    `original` with `styles` / `header_has` / `footer_has` / `style_uses`: the user's template or
    document: its styles are still defined, its header and footer text are still there, and
    paragraphs use the named styles at least that often. `keep_styles`: every style the original
    used is still used.
    """

    facts: list[list[str]]
    absent: list[str] = field(default_factory=list)
    allowed_numbers: list[str] = field(default_factory=list)
    pages: list[int] | None = None
    headings: list[dict[str, object]] = field(default_factory=list)
    toc: bool = False
    tables: list[dict[str, object]] = field(default_factory=list)
    images_min: int = 0
    image_fit: bool = False
    captions: list[dict[str, object]] = field(default_factory=list)
    tracked: dict[str, object] | None = None
    original: str | None = None
    styles: list[str] = field(default_factory=list)
    header_has: list[str] = field(default_factory=list)
    footer_has: list[str] = field(default_factory=list)
    style_uses: dict[str, int] = field(default_factory=dict)
    keep_styles: bool = False


@dataclass(frozen=True)
class Task:
    name: str
    directory: Path
    brief: str
    inputs: list[Path]
    expect: Expect | XlsxExpect | DocxExpect
    kind: str = "pptx"

    def source_texts(self) -> list[str]:
        """Everything a number in the deck may legitimately come from."""
        return [self.brief, *(t for t in map(input_text, self.inputs) if t)]


def input_text(path: Path) -> str:
    """The text of an input file a person would read; an image or an unknown kind has none."""
    suffix = path.suffix.lower()
    if suffix in _TEXT_SUFFIXES:
        return path.read_text(encoding="utf-8")
    if suffix == ".pptx":
        from .pptx_checks import read_deck

        return read_deck(path).all_text()
    if suffix == ".docx":
        from docx import Document

        document = Document(str(path))
        cells = (c.text for t in document.tables for row in t.rows for c in row.cells)
        return "\n".join([*(p.text for p in document.paragraphs), *cells])
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(path, data_only=True, read_only=True)
        return "\n".join(
            " ".join("" if v is None else str(v) for v in row)
            for sheet in workbook.worksheets
            for row in sheet.iter_rows(values_only=True)
        )
    if suffix == ".pdf":
        import pdfplumber

        with pdfplumber.open(str(path)) as pdf:
            return "\n".join(page.extract_text() or "" for page in pdf.pages)
    return ""


def _parse_expect(raw: dict[str, object], where: Path) -> Expect:
    pages = raw.get("pages")
    if not (isinstance(pages, list) and len(pages) == 2 and 0 < pages[0] <= pages[1]):
        raise ValueError(f'{where}: "pages" must be [min, max] with 0 < min <= max')
    facts = _facts(raw, where, required=True)
    absent = raw.get("absent", [])
    allowed = raw.get("allowed_numbers", [])
    if not isinstance(absent, list) or not isinstance(allowed, list):
        raise ValueError(f'{where}: "absent" and "allowed_numbers" must be lists')
    return Expect(
        facts=facts,
        pages_min=int(pages[0]),
        pages_max=int(pages[1]),
        absent=[str(a) for a in absent],
        allowed_numbers=[str(a) for a in allowed],
        pictures_min=int(raw.get("pictures_min", 0)),  # type: ignore[call-overload]
        template=str(raw["template"]) if raw.get("template") else None,
        pictures_of=str(raw["pictures_of"]) if raw.get("pictures_of") else None,
        notes=float(raw.get("notes", 0)),  # type: ignore[arg-type]
    )


def _facts(raw: dict[str, object], where: Path, required: bool) -> list[list[str]]:
    facts_raw = raw.get("facts", [])
    if not isinstance(facts_raw, list) or (required and not facts_raw):
        raise ValueError(f'{where}: "facts" must be a non-empty list')
    facts = [[str(f)] if isinstance(f, str) else [str(a) for a in f] for f in facts_raw]
    if any(not alternatives or not all(alternatives) for alternatives in facts):
        raise ValueError(f"{where}: a fact must be a non-empty string or a list of them")
    return facts


def _list_of(raw: dict[str, object], key: str, kind: type, where: Path) -> list[Any]:
    value = raw.get(key, [])
    if not isinstance(value, list) or not all(isinstance(v, kind) for v in value):
        raise ValueError(f'{where}: "{key}" must be a list of {kind.__name__}')
    return value


def _address_is_valid(address: object) -> bool:
    if not isinstance(address, str):
        return False
    sheet, _, cells = address.rpartition("!")
    return bool(sheet and cells)


def _parse_xlsx_expect(raw: dict[str, object], where: Path) -> XlsxExpect:
    cells = _list_of(raw, "cells", dict, where)
    formulas = _list_of(raw, "formulas", str, where)
    constants = _list_of(raw, "constants", str, where)
    number_formats = _list_of(raw, "number_formats", dict, where)
    charts = _list_of(raw, "charts", dict, where)
    validations = _list_of(raw, "validations", dict, where)
    for spec in (*cells, *number_formats, *validations):
        if not _address_is_valid(spec.get("at")):
            raise ValueError(f'{where}: "at" must look like Sheet!B2 or Sheet!B2:B9: {spec}')
    for spec in cells:
        if "value" not in spec:
            raise ValueError(f'{where}: a cell check needs a "value": {spec}')
        if ":" in spec["at"]:
            raise ValueError(f"{where}: a cell check names one cell, not a range: {spec}")
    for spec in number_formats:
        if not spec.get("any_of") or not isinstance(spec["any_of"], list):
            raise ValueError(f'{where}: a number format check needs "any_of": {spec}')
    bold = _list_of(raw, "bold", str, where)
    for address in (*formulas, *constants, *bold):
        if not _address_is_valid(address):
            raise ValueError(f"{where}: {address!r} must look like Sheet!B2 or Sheet!B2:B9")
    used_range = raw.get("used_range", {})
    if not isinstance(used_range, dict) or not all(
        isinstance(k, str) and isinstance(v, str) for k, v in used_range.items()
    ):
        raise ValueError(f'{where}: "used_range" must map a sheet name to a range')
    probe = raw.get("probe")
    if probe is not None and not (
        isinstance(probe, dict)
        and isinstance(probe.get("set"), dict)
        and isinstance(probe.get("expect"), list)
    ):
        raise ValueError(f'{where}: "probe" needs "set" (address to value) and "expect" (cells)')
    untouched = _list_of(raw, "untouched", str, where)
    original = raw.get("original")
    if (untouched or raw.get("conditional_formats_min")) and not original:
        raise ValueError(f'{where}: "untouched" and "conditional_formats_min" need "original"')
    facts = _facts(raw, where, required=False)
    figures = _list_of(raw, "figures", dict, where)
    input_probe = raw.get("input_probe")
    if input_probe is not None and not (
        isinstance(input_probe, dict)
        and {"find", "set", "figures"} <= input_probe.keys()
    ):
        raise ValueError(f'{where}: "input_probe" needs "find", "set" and "figures"')
    if not (facts or cells or formulas or probe or figures):
        raise ValueError(f'{where}: expects nothing: give "facts", "cells", "formulas" or "probe"')
    return XlsxExpect(
        facts=facts,
        absent=_list_of(raw, "absent", str, where),
        allowed_numbers=_list_of(raw, "allowed_numbers", str, where),
        sheets=_list_of(raw, "sheets", str, where),
        cells=cells,
        formulas=formulas,
        constants=constants,
        number_formats=number_formats,
        bold=bold,
        used_range=used_range,
        original=str(original) if original else None,
        untouched=untouched,
        original_keeps=_list_of(raw, "original_keeps", str, where),
        conditional_formats_min=int(raw.get("conditional_formats_min", 0)),  # type: ignore[call-overload]
        formulas_min=int(raw.get("formulas_min", 0)),  # type: ignore[call-overload]
        figures=figures,
        input_probe=input_probe,
        probe=probe,
        charts=charts,
        validations=validations,
    )


def _parse_docx_expect(raw: dict[str, object], where: Path) -> DocxExpect:
    facts = _facts(raw, where, required=False)
    headings = _list_of(raw, "headings", dict, where)
    tables = _list_of(raw, "tables", dict, where)
    captions = _list_of(raw, "captions", dict, where)
    for spec in headings:
        if not spec.get("text"):
            raise ValueError(f'{where}: a heading check needs "text": {spec}')
    for spec in tables:
        if not (spec.get("rows_min") and spec.get("cols_min")):
            raise ValueError(f'{where}: a table check needs "rows_min" and "cols_min": {spec}')
    for spec in captions:
        if not (spec.get("style") and spec.get("starts_with")):
            raise ValueError(f'{where}: a caption check needs "style" and "starts_with": {spec}')
    pages = raw.get("pages")
    if pages is not None and not (
        isinstance(pages, list) and len(pages) == 2 and 0 < pages[0] <= pages[1]
    ):
        raise ValueError(f'{where}: "pages" must be [min, max] with 0 < min <= max')
    tracked = raw.get("tracked")
    if tracked is not None and not isinstance(tracked, dict):
        raise ValueError(f'{where}: "tracked" must be an object')
    style_uses = raw.get("style_uses", {})
    if not isinstance(style_uses, dict) or not all(isinstance(v, int) for v in style_uses.values()):
        raise ValueError(f'{where}: "style_uses" must map a style name to a count')
    original = raw.get("original")
    template_keys = ("styles", "header_has", "footer_has", "style_uses", "keep_styles")
    if any(raw.get(key) for key in template_keys) and not original:
        raise ValueError(f'{where}: {", ".join(template_keys)} need "original"')
    if isinstance(tracked, dict) and tracked.get("rejected_equals_original") and not original:
        raise ValueError(f'{where}: "rejected_equals_original" needs "original"')
    asks = (
        facts,
        headings,
        tables,
        captions,
        tracked,
        pages,
        raw.get("images_min"),
        raw.get("image_fit"),
        raw.get("toc"),
        *(raw.get(key) for key in template_keys),
    )
    if not any(asks):
        raise ValueError(f"{where}: expects nothing: give facts, headings, tables or similar")
    return DocxExpect(
        facts=facts,
        absent=_list_of(raw, "absent", str, where),
        allowed_numbers=_list_of(raw, "allowed_numbers", str, where),
        pages=[int(pages[0]), int(pages[1])] if pages else None,
        headings=headings,
        toc=bool(raw.get("toc", False)),
        tables=tables,
        images_min=int(raw.get("images_min", 0)),  # type: ignore[call-overload]
        image_fit=bool(raw.get("image_fit", False)),
        captions=captions,
        tracked=tracked,
        original=str(original) if original else None,
        styles=_list_of(raw, "styles", str, where),
        header_has=_list_of(raw, "header_has", str, where),
        footer_has=_list_of(raw, "footer_has", str, where),
        style_uses=style_uses,
        keep_styles=bool(raw.get("keep_styles", False)),
    )


def load_task(directory: Path, kind: str = "pptx") -> Task:
    brief_path = directory / "brief.md"
    expect_path = directory / "expect.json"
    if not brief_path.is_file() or not expect_path.is_file():
        raise ValueError(f"{directory}: needs brief.md and expect.json")
    brief = brief_path.read_text(encoding="utf-8").strip()
    if not brief:
        raise ValueError(f"{brief_path} is empty")
    inputs_dir = directory / "inputs"
    inputs = sorted(p for p in inputs_dir.rglob("*") if p.is_file()) if inputs_dir.is_dir() else []
    raw = json.loads(expect_path.read_text(encoding="utf-8"))
    parse = {"xlsx": _parse_xlsx_expect, "docx": _parse_docx_expect}.get(kind, _parse_expect)
    expect = parse(raw, expect_path)
    if getattr(expect, "original", None) and expect.original not in {p.name for p in inputs}:  # type: ignore[union-attr]
        raise ValueError(f"{expect_path}: original {expect.original!r} is not among the inputs")  # type: ignore[union-attr]
    return Task(
        name=directory.name,
        directory=directory,
        brief=brief,
        inputs=inputs,
        expect=expect,
        kind=kind,
    )


def discover(selector: str = "all", root: Path | None = None, kind: str = "pptx") -> list[Task]:
    """`all`, or comma-separated task names or number prefixes (`03`, `03-proposal`)."""
    if kind not in TASK_ROOTS:
        raise ValueError(f"no such kind: {kind}; use one of {', '.join(KINDS)}")
    root = root or TASK_ROOTS[kind]
    directories = sorted(p for p in root.iterdir() if p.is_dir()) if root.is_dir() else []
    if selector != "all":
        wanted = [s.strip() for s in selector.split(",") if s.strip()]
        chosen = [
            d for d in directories if any(d.name == w or d.name.startswith(f"{w}-") for w in wanted)
        ]
        unknown = [
            w
            for w in wanted
            if not any(d.name == w or d.name.startswith(f"{w}-") for d in directories)
        ]
        if unknown:
            raise ValueError(f"no such task: {', '.join(unknown)}")
        directories = chosen
    return [load_task(d, kind) for d in directories]
