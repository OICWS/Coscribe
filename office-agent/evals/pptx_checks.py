"""Automatic checks on one produced deck, in plain Python: no model is called here.

Each check returns a plain dict so a run can be written as JSON and two runs compared.
`check_deck` runs them all; the pure ones (facts, inventions, page count, the editability
probe) are unit-tested, the render and the layout checks need LibreOffice.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from .render import render_pages as render_deck
from .tasks import Task
from .text_checks import check_facts, find_inventions, normalise, plain_number  # noqa: F401


@dataclass
class SlideText:
    title: str
    texts: list[str] = field(default_factory=list)


@dataclass
class Deck:
    slides: list[SlideText]

    def all_text(self) -> str:
        return "\n".join(t for s in self.slides for t in [s.title, *s.texts])


def _leaf_shapes(shapes: Any) -> Any:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _leaf_shapes(shape.shapes)
        else:
            yield shape


def _chart_texts(chart: Any) -> list[str]:
    texts: list[str] = []
    if chart.has_title:
        texts.append(chart.chart_title.text_frame.text)
    for plot in chart.plots:
        texts.extend(str(c) for c in plot.categories)
        for series in plot.series:
            texts.append(str(series.name))
            texts.extend(plain_number(v) for v in series.values if v is not None)
    return texts


def read_deck(path: Path) -> Deck:
    """Every piece of text a person would see on each slide, groups, tables and charts included."""
    from pptx import Presentation

    slides: list[SlideText] = []
    for slide in Presentation(str(path)).slides:
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None else ""
        texts: list[str] = []
        for shape in _leaf_shapes(slide.shapes):
            if title_shape is not None and shape._element is title_shape._element:
                continue
            if shape.has_text_frame:
                texts.extend(p.text for p in shape.text_frame.paragraphs if p.text.strip())
            if getattr(shape, "has_table", False) and shape.has_table:
                texts.extend(cell.text for row in shape.table.rows for cell in row.cells)
            if getattr(shape, "has_chart", False) and shape.has_chart:
                texts.extend(_chart_texts(shape.chart))
        if slide.has_notes_slide:
            texts.append(slide.notes_slide.notes_text_frame.text)
        slides.append(SlideText(title=title, texts=texts))
    return Deck(slides)


def check_opens(path: Path) -> dict[str, Any]:
    try:
        deck = read_deck(path)
    except Exception as exc:  # noqa: BLE001 -- any failure to open is the finding
        return {"ok": False, "error": f"{type(exc).__name__}: {exc}", "slides": 0}
    return {"ok": True, "slides": len(deck.slides)}


def check_inventions(deck: Deck, source_texts: list[str], allowed: list[str]) -> dict[str, Any]:
    units = ["\n".join([slide.title, *slide.texts]) for slide in deck.slides]
    return find_inventions(units, source_texts, allowed, "slides")


def check_page_count(count: int, low: int, high: int) -> dict[str, Any]:
    return {"ok": low <= count <= high, "count": count, "expected": [low, high]}


def count_pictures(path: Path) -> int:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    return sum(
        1
        for slide in Presentation(str(path)).slides
        for shape in _leaf_shapes(slide.shapes)
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    )


def picture_thumbnails(path: Path) -> list[Any]:
    """Every picture of a deck, small and in colour, to tell the same image re-encoded or
    cropped from a different one."""
    import io

    from PIL import Image
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    return [
        Image.open(io.BytesIO(shape.image.blob)).convert("RGB").resize((32, 32))
        for slide in Presentation(str(path)).slides
        for shape in _leaf_shapes(slide.shapes)
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    ]


# Mean channel difference (0-255) under which two thumbnails are one picture; different
# pictures measured 32-73 apart, the same one re-saved 0.
_SAME_PICTURE = 8.0


def check_pictures_kept(original: list[Any], result: list[Any]) -> dict[str, Any]:
    from PIL import ImageChops, ImageStat

    def distance(a: Any, b: Any) -> float:
        return float(sum(ImageStat.Stat(ImageChops.difference(a, b)).mean) / 3)

    missing = [
        n
        for n, picture in enumerate(original, start=1)
        if not any(distance(picture, other) < _SAME_PICTURE for other in result)
    ]
    return {"ok": not missing, "missing": missing, "of": len(original)}


def check_pictures(count: int, minimum: int) -> dict[str, Any]:
    return {"ok": count >= minimum, "count": count, "minimum": minimum}


def check_notes(path: Path, share: float) -> dict[str, Any]:
    """Slides whose speaker notes hold text, against the share the brief asks for."""
    from pptx import Presentation

    slides = list(Presentation(str(path)).slides)
    noted = sum(
        1 for s in slides if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip()
    )
    return {"ok": noted >= share * len(slides), "with_notes": noted, "slides": len(slides)}


def layout_names(path: Path) -> set[str]:
    from pptx import Presentation

    return {layout.name for layout in Presentation(str(path)).slide_layouts}


def check_template_kept(deck_layouts: set[str], template_layouts: set[str]) -> dict[str, Any]:
    """A deck made on the user's template carries the template's layouts; one rebuilt from
    scratch does not."""
    missing = sorted(template_layouts - deck_layouts)
    return {"ok": not missing, "missing_layouts": missing}


def check_editability(path: Path, workspace: Path) -> dict[str, Any]:
    """What coscribe's own tools can do with the deck: every slide must have a title that
    `read_pptx` shows and a shape list that `list_pptx_shapes` returns."""
    from coscribe.tools.presentations import build_presentation_tools

    tools = {t.__name__: t for t in build_presentation_tools(workspace)}
    relative = str(path.relative_to(workspace))
    result: dict[str, Any] = {"untitled_slides": [], "shape_list_errors": []}
    try:
        overview = tools["read_pptx"](path=relative)
    except Exception as exc:  # noqa: BLE001
        return {**result, "ok": False, "error": f"read_pptx: {type(exc).__name__}: {exc}"}
    pages = overview.split("\n\n---\n\n")
    for number, page in enumerate(pages, start=1):
        heading = page.split("\n", 1)[0]
        if not heading.startswith("## ") or not heading[3:].strip():
            result["untitled_slides"].append(number)
    for number in range(1, len(pages) + 1):
        try:
            tools["list_pptx_shapes"](path=relative, slide=number)
        except Exception as exc:  # noqa: BLE001 -- a crash on one slide is what is measured
            result["shape_list_errors"].append(
                {"slide": number, "error": f"{type(exc).__name__}: {exc}"}
            )
    result["ok"] = not result["untitled_slides"] and not result["shape_list_errors"]
    return result


def check_layout(path: Path, pdf: Path | None) -> dict[str, Any]:
    """coscribe's own QA functions, as its tools use them after a write."""
    from pptx import Presentation

    from coscribe.tools.office.pptx import qa

    prs = Presentation(str(path))
    overflow = qa.check_overflow(pdf) if pdf is not None and pdf.is_file() else None
    overlaps = qa.check_text_overlaps(prs)
    low_contrast = qa.check_low_contrast(prs)
    plain = qa.check_missing_visual_elements(prs)
    return {
        "overflow_words": None if overflow is None else len(overflow),
        "overflow": overflow,
        "text_overlaps": len(overlaps),
        "low_contrast": len(low_contrast),
        "slides_without_visual": plain,
    }


def check_deck(task: Task, path: Path, workspace: Path, out_dir: Path) -> dict[str, Any]:
    """All checks for one deck; a check that cannot run says why instead of failing the rest."""
    report: dict[str, Any] = {"deck": str(path.relative_to(workspace))}
    report["opens"] = check_opens(path)
    if not report["opens"]["ok"]:
        return report
    deck = read_deck(path)
    expect = task.expect
    report["facts"] = check_facts(deck.all_text(), expect.facts, expect.absent)
    report["inventions"] = check_inventions(deck, task.source_texts(), expect.allowed_numbers)
    report["pages"] = check_page_count(len(deck.slides), expect.pages_min, expect.pages_max)
    if expect.pictures_min:
        report["pictures"] = check_pictures(count_pictures(path), expect.pictures_min)
    if expect.pictures_of:
        original = next(p for p in task.inputs if p.name == expect.pictures_of)
        report["pictures_kept"] = check_pictures_kept(
            picture_thumbnails(original), picture_thumbnails(path)
        )
    if expect.template:
        template = next(p for p in task.inputs if p.name == expect.template)
        report["template"] = check_template_kept(layout_names(path), layout_names(template))
    if expect.notes:
        report["notes"] = check_notes(path, expect.notes)
    report["editability"] = check_editability(path, workspace)
    report["render"] = render_deck(path, out_dir)
    pdf = out_dir / report["render"]["pdf"] if report["render"].get("pdf") else None
    try:
        report["layout"] = check_layout(path, pdf)
    except Exception as exc:  # noqa: BLE001
        report["layout"] = {"error": f"{type(exc).__name__}: {exc}"}
    return report
