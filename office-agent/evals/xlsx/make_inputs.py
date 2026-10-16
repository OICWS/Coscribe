"""Builds the input files of the Excel tasks (the messy CSVs, the workbooks to edit), so they are
reproducible rather than opaque. The outputs are committed.

    python -m evals.xlsx.make_inputs
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import random
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

TASKS = Path(__file__).resolve().parent / "tasks"

ORDERS_CSV = """\
Order Date,Customer ,Category,Amount
2026/01/05,Acme Ltd ,hardware,"$1,200.50"
5 Jan 2026,Beta GmbH,Software ,980
2026-01-09,Acme Ltd,SOFTWARE,$75

2026-01-12,Cobalt Inc,Hardware,"$2,310.00"
2026-01-03,Delta Co,services,450.25
2026-01-12,Cobalt Inc,Hardware,"$2,310.00"
2026-01-15,Beta GmbH,services ,"1,100"
2026-01-18,Delta Co,Software,$620.75
2026-01-20,Acme Ltd,hardware,"$845.00"
2026-01-22,Echo LLC,Services,300
"""

VISITS_CSV = """\
Month,Visitors
Jan,1200
Feb,1350
Mar,1280
Apr,1500
May,1720
Jun,1650
Jul,1580
Aug,1800
Sep,2010
Oct,1950
Nov,2200
Dec,2480
"""

SALES_CSV = """\
区域,销售额(元)
华东,12850000
华北,9420000
华南,8760000
西南,4310000
东北,2180000
"""

_HEADER_FILL = PatternFill("solid", fgColor="1F3864")
_THIN = Side(style="thin", color="808080")


def _header(sheet, columns: int) -> None:  # type: ignore[no-untyped-def]
    for column in range(1, columns + 1):
        cell = sheet.cell(row=1, column=column)
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
        cell.border = Border(bottom=_THIN)


def transactions(path: Path) -> None:
    rows = [
        ((1, 6), "Food", 120.50),
        ((1, 9), "Food", 98.20),
        ((1, 15), "Rent", 1500),
        ((1, 21), "Transport", 45),
        ((1, 27), "Transport", 60),
        ((2, 3), "Food", 110),
        ((2, 10), "Food", 87.30),
        ((2, 15), "Rent", 1500),
        ((2, 20), "Transport", 52.50),
        ((3, 4), "Food", 134.80),
        ((3, 15), "Rent", 1500),
        ((3, 18), "Transport", 38),
        ((3, 24), "Transport", 71.25),
        ((3, 29), "Food", 76),
    ]
    book = Workbook()
    sheet = book.active
    sheet.title = "Data"
    sheet.append(["Date", "Category", "Amount"])
    for (month, day), category, amount in rows:
        sheet.append([dt.datetime(2026, month, day), category, amount])
    _header(sheet, 3)
    for row in range(2, len(rows) + 2):
        sheet.cell(row=row, column=1).number_format = "yyyy-mm-dd"
        sheet.cell(row=row, column=3).number_format = "#,##0.00"
    sheet.column_dimensions["A"].width = 14
    sheet.column_dimensions["B"].width = 14
    sheet.column_dimensions["C"].width = 12
    book.save(path)


def shop(path: Path) -> None:
    book = Workbook()
    orders = book.active
    orders.title = "Orders"
    orders.append(["Order ID", "SKU", "Qty"])
    for order in [
        (1001, "B-200", 3),
        (1002, "A-100", 10),
        (1003, "C-300", 2),
        (1004, "A-100", 5),
        (1005, "D-400", 1),
        (1006, "B-200", 7),
        (1007, "C-300", 4),
        (1008, "E-500", 6),
    ]:
        orders.append(list(order))
    prices = book.create_sheet("Prices")
    prices.append(["SKU", "Product", "Unit price"])
    for price in [
        ("E-500", "Backpack", 39.0),
        ("C-300", "Stapler", 7.95),
        ("A-100", "Pen", 1.2),
        ("D-400", "Desk lamp", 24.9),
        ("B-200", "Notebook", 3.5),
        ("F-600", "Whiteboard", 55.0),
    ]:
        prices.append(list(price))
    for sheet, columns in ((orders, 3), (prices, 3)):
        _header(sheet, columns)
        for column in "ABC":
            sheet.column_dimensions[column].width = 14
    for row in range(2, 7):
        prices.cell(row=row, column=3).number_format = "0.00"
    for row in range(2, 10):
        orders.cell(row=row, column=3).number_format = "0"
    book.save(path)


def budget(path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "Budget"
    sheet.append(["Item", "Budget", "Actual"])
    for line in [
        ("Marketing", 12000, 13450),
        ("Travel", 5000, 4200),
        ("Software", 8000, 8000),
        ("Training", 3000, 3650),
        ("Office", 2500, 2275),
        ("Events", 6000, 7100),
    ]:
        sheet.append(list(line))
    _header(sheet, 3)
    for row in range(2, 8):
        for column in (2, 3):
            cell = sheet.cell(row=row, column=column)
            cell.number_format = '"$"#,##0'
            cell.border = Border(bottom=_THIN)
        sheet.cell(row=row, column=1).font = Font(name="Calibri", size=11, italic=True)
    sheet.conditional_formatting.add(
        "C2:C7",
        CellIsRule(
            operator="greaterThan", formula=["6000"], fill=PatternFill("solid", bgColor="F4B6B6")
        ),
    )
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 12
    sheet.column_dimensions["C"].width = 12
    sheet.freeze_panes = "A2"
    book.save(path)


def tickets(path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "Tickets"
    sheet.append(["ID", "Title", "Status", "Owner"])
    for number, (title, owner) in enumerate(
        [
            ("Login page times out", "Mia"),
            ("Export to PDF drops tables", "Ben"),
            ("Dark mode contrast", "Mia"),
            ("Slow search on large files", "Noor"),
            ("Typo in onboarding mail", "Ben"),
            ("Calendar sync fails", "Noor"),
            ("Add CSV import", "Mia"),
            ("Crash on empty sheet", "Ben"),
            ("Update the help page", "Noor"),
            ("Rename Team to Workspace", "Mia"),
        ],
        start=101,
    ):
        sheet.append([f"T-{number}", title, None, owner])
    _header(sheet, 4)
    sheet.column_dimensions["A"].width = 8
    sheet.column_dimensions["B"].width = 30
    sheet.column_dimensions["C"].width = 14
    sheet.column_dimensions["D"].width = 10
    book.save(path)


_PRICES = {"智能门锁": 1299, "扫地机器人": 2499, "空气净化器": 1899}


def sales_ledger() -> str:
    """480 order lines over 2026 with a seeded generator, so the totals in expect.json hold."""
    rng = random.Random(9)
    out = io.StringIO()
    writer = csv.writer(out, lineterminator="\n")
    writer.writerow(["日期", "区域", "产品线", "销量", "单价", "销售额"])
    for month in range(1, 13):
        for _ in range(40):
            day = rng.randint(1, 28)
            region = rng.choice(["华东", "华东", "华南", "华北", "西南"])
            product = rng.choice(list(_PRICES))
            units = rng.randint(5, 60) + (month * 2 if region == "华南" else 0)
            price = _PRICES[product]
            writer.writerow(
                [f"2026-{month:02d}-{day:02d}", region, product, units, price, units * price]
            )
    return out.getvalue()


_PLAN = [
    ("Jan", 42000, 40500),
    ("Feb", 45000, 46800),
    ("Mar", 47000, 47000),
    ("Apr", 50000, 48100),
    ("May", 52000, 53900),
    ("Jun", 55000, 56250),
]
# A sparkline lives in an extension openpyxl does not read: a workbook it loads and saves again
# loses it, which is what task 11 is about.
_SPARKLINE = (
    '<extLst><ext uri="{05C60535-1F16-4fd2-B633-F4F36F0B64E0}" '
    'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
    '<x14:sparklineGroups xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main">'
    "<x14:sparklineGroup><x14:sparklines><x14:sparkline><xm:f>Plan!C2:C7</xm:f>"
    "<xm:sqref>H2</xm:sqref></x14:sparkline></x14:sparklines></x14:sparklineGroup>"
    "</x14:sparklineGroups></ext></extLst></worksheet>"
)


def plan(path: Path) -> None:
    """A user's half-year plan: styled, a rule, frozen panes and a sparkline in H2."""
    book = Workbook()
    sheet = book.active
    sheet.title = "Plan"
    sheet.append(["Month", "Target", "Actual"])
    for line in _PLAN:
        sheet.append(list(line))
    sheet["H1"] = "Trend"
    _header(sheet, 3)
    for row in range(2, 8):
        for column in (2, 3):
            cell = sheet.cell(row=row, column=column)
            cell.number_format = '"$"#,##0'
            cell.border = Border(bottom=_THIN)
    sheet.conditional_formatting.add(
        "C2:C7",
        CellIsRule(operator="lessThan", formula=["B2"], fill=PatternFill("solid", bgColor="F4B6B6")),
    )
    for column, width in (("A", 12), ("B", 12), ("C", 12), ("H", 16)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    book.save(path)
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    sheet_xml = parts["xl/worksheets/sheet1.xml"].decode()
    parts["xl/worksheets/sheet1.xml"] = sheet_xml.replace("</worksheet>", _SPARKLINE).encode()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


_REGIONS = ["North", "South", "East", "West"]
_ITEMS = {"Desk": 320, "Chair": 145, "Lamp": 38, "Shelf": 210, "Monitor arm": 89}


def orders_2025() -> str:
    """5,000 orders over 2025 with a seeded generator, so the totals in expect.json hold."""
    rng = random.Random(25)
    out = io.StringIO()
    writer = csv.writer(out, lineterminator="\n")
    writer.writerow(["Order", "Date", "Region", "Product", "Qty", "Unit price"])
    start = dt.date(2025, 1, 1)
    for number in range(1, 5001):
        day = start + dt.timedelta(days=rng.randint(0, 364))
        product = rng.choice(list(_ITEMS))
        writer.writerow(
            [f"SO-{number:05d}", day.isoformat(), rng.choice(_REGIONS), product,
             rng.randint(1, 12), _ITEMS[product]]
        )
    return out.getvalue()


def main() -> None:
    def put(task: str, name: str) -> Path:
        folder = TASKS / task / "inputs"
        folder.mkdir(parents=True, exist_ok=True)
        return folder / name

    put("01-clean-csv", "orders_export.csv").write_text(ORDERS_CSV, encoding="utf-8")
    transactions(put("02-monthly-summary", "transactions.xlsx"))
    shop(put("03-lookup-two-sheets", "shop.xlsx"))
    put("04-line-chart", "monthly_visits.csv").write_text(VISITS_CSV, encoding="utf-8")
    budget(put("05-edit-budget", "budget_q1.xlsx"))
    put("07-chinese-formats", "各区域销售.csv").write_text(SALES_CSV, encoding="utf-8")
    tickets(put("08-validation-list", "tickets.xlsx"))
    put("09-sales-report", "销售明细.csv").write_text(sales_ledger(), encoding="utf-8")
    plan(put("11-rich-workbook", "plan_2026.xlsx"))
    put("12-big-ledger", "orders_2025.csv").write_text(orders_2025(), encoding="utf-8")


if __name__ == "__main__":
    main()
