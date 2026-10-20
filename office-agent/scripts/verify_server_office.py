#!/usr/bin/env python3
"""Drives the server's Office tools through its WebSocket with a stub model: a deck build with a
native chart and table, a deck read, an Excel write and read, a Word write and read, a web page
read.

    python scripts/verify_server_office.py                         dev server
    python scripts/verify_server_office.py --cmd <sidecar exe>     the PyInstaller sidecar

A library the packaged build left out (a lazy import PyInstaller's analysis cannot see) shows up
here as a failed tool call, which an import-time check cannot show. Needs the dev environment
for the client side (websockets, python-docx) and `scripts/stub_llm.py` for the model.
"""

from __future__ import annotations

import argparse
import json
import os
import socket
import subprocess
import sys
import tempfile
import threading
import time
import urllib.error
import urllib.request
import zipfile
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any

from websockets.sync.client import connect

ROOT = Path(__file__).resolve().parents[1]
TOKEN = "verify-server-office"

_SVG = (
    'xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1280 720" font-family="Microsoft YaHei, Arial"'
)
_CHART = """<g id="revenue" data-pptx-replace-with="chart"><metadata type="application/json">
{"x": 64, "y": 160, "width": 560, "height": 480, "type": "column",
 "categories": ["East", "South", "North"],
 "series": [{"name": "Q3 revenue", "values": [5410, 3980, 2120]}]}
</metadata></g>"""
_TABLE = """<g id="regions" data-pptx-replace-with="table"><metadata type="application/json">
{"x": 680, "y": 180, "width": 536, "height": 200, "schema": "ppt-master.semantic-table.v2",
 "columns": ["Region", "Growth"], "rows": [["East", {"text": "+12.2%", "align": "r"}]]}
</metadata></g>"""
_PAGE = f'<svg {_SVG}><text id="title" x="64" y="100" font-size="40">East grows fastest</text>'


def _free_port() -> int:
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        return int(s.getsockname()[1])


def _wait_http(url: str, proc: subprocess.Popen[bytes], timeout: float = 120) -> None:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        if proc.poll() is not None:
            raise RuntimeError(f"{url}: the process exited with {proc.returncode}")
        try:
            urllib.request.urlopen(url, timeout=1).read()
            return
        except urllib.error.HTTPError:
            return
        except OSError:
            time.sleep(0.1)
    raise RuntimeError(f"{url} was not ready in {timeout} s")


class _Turns:
    """One WebSocket conversation: a message in, the tool results out, until the model's "OK"."""

    def __init__(self, port: int) -> None:
        self.ws = connect(f"ws://127.0.0.1:{port}/ws/verify-office?token={TOKEN}", max_size=None)

    def call(self, tool: str, arguments: dict[str, Any], via_use_tool: bool = False) -> Any:
        if via_use_tool:
            tool, arguments = "use_tool", {"name": tool, "arguments": arguments}
        self.ws.send(
            json.dumps(
                {"type": "user_message", "text": f"[[stub-tool {tool} {json.dumps(arguments)}]]"}
            )
        )
        results: list[dict[str, Any]] = []
        deadline = time.monotonic() + 180
        while time.monotonic() < deadline:
            event = json.loads(self.ws.recv(timeout=180))
            if event["type"] == "approval_required":
                self.ws.send(
                    json.dumps({"type": "approval_response", "id": event["id"], "approved": True})
                )
            elif event["type"] == "tool_result":
                results.append(event)
            elif event["type"] == "error":
                raise RuntimeError(event["message"])
            elif event["type"] == "agent_message" and results:
                break
        else:
            raise RuntimeError(f"{tool}: the turn did not finish")
        (result,) = [r for r in results if r["tool_name"] != "use_tool"] or results[-1:]
        if result["is_error"]:
            raise RuntimeError(f"{tool} failed: {str(result['result'])[:500]}")
        return result["result"]


def _serve_page() -> tuple[ThreadingHTTPServer, int]:
    class Handler(BaseHTTPRequestHandler):
        body = b"<html><head><title>T</title></head><body><main><h1>Hello</h1></main></body></html>"

        def do_HEAD(self) -> None:
            self.send_response(200)
            self.send_header("Content-Type", "text/html")
            self.send_header("Content-Length", str(len(self.body)))
            self.end_headers()

        def do_GET(self) -> None:
            self.do_HEAD()
            self.wfile.write(self.body)

        def log_message(self, format: str, *args: object) -> None:
            pass

    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    return server, server.server_address[1]


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--cmd", nargs="+", default=[sys.executable, "-m", "coscribe.web.app"])
    ap.add_argument("--verbose", action="store_true", help="show the server's own log")
    args = ap.parse_args()

    from docx import Document

    with tempfile.TemporaryDirectory() as tmp_name:
        tmp = Path(tmp_name)
        workspace = tmp / "workspace"
        pages = workspace / "deck-pages"
        pages.mkdir(parents=True)
        (pages / "01-chart.svg").write_text(f"{_PAGE}{_CHART}{_TABLE}</svg>", encoding="utf-8")
        document = Document()
        document.add_heading("Report", 1)
        document.add_paragraph("A paragraph of text.")
        document.save(str(workspace / "in.docx"))

        providers = tmp / "providers.json"
        llm_port, port = _free_port(), _free_port()
        providers.write_text(
            json.dumps(
                {
                    "providers": {
                        "stub": {"base_url": f"http://127.0.0.1:{llm_port}/v1", "api_key": "x"}
                    }
                }
            )
        )
        env = {
            # web_search's page reader sends every request through a configured proxy, NO_PROXY
            # or not, so the local test page is only reachable without one.
            **{k: v for k, v in os.environ.items() if "proxy" not in k.lower()},
            "HOME": str(tmp),
            "APPDATA": str(tmp),
            "USERPROFILE": str(tmp),
            "COSCRIBE_DEFAULT_MODEL": "stub:fake",
            "COSCRIBE_PROVIDERS_CONFIG_PATH": str(providers),
            "COSCRIBE_WORKSPACE_ROOT": str(workspace),
            "COSCRIBE_LAUNCH_TOKEN": TOKEN,
        }
        quiet = {} if args.verbose else {"stdout": subprocess.DEVNULL, "stderr": subprocess.DEVNULL}
        llm = subprocess.Popen(
            [sys.executable, str(ROOT / "scripts" / "stub_llm.py"), "--port", str(llm_port)]
        )
        web, web_port = _serve_page()
        server = subprocess.Popen([*args.cmd, "--port", str(port)], cwd=tmp, env=env, **quiet)
        try:
            _wait_http(f"http://127.0.0.1:{port}/", server)
            turns = _Turns(port)
            checks: list[tuple[str, bool]] = []

            turns.call(
                "build_pptx_from_pages",
                {"pages_dir": "deck-pages", "path": "deck.pptx"},
                via_use_tool=True,
            )
            with zipfile.ZipFile(workspace / "deck.pptx") as package:
                names = package.namelist()
            checks.append(
                ("deck built, chart with its workbook", any("embeddings/" in n for n in names))
            )
            checks.append(("deck has a chart part", any("charts/chart" in n for n in names)))
            read = turns.call("read_pptx", {"path": "deck.pptx"})
            checks.append(
                ("deck read", "East grows fastest" in json.dumps(read, ensure_ascii=False))
            )

            turns.call(
                "write_xlsx_cells",
                {
                    "path": "out.xlsx",
                    "blocks": [{"sheet": "S", "at": "A1", "rows": [["a", 1], ["b", 2]]}],
                },
                via_use_tool=True,
            )
            read = turns.call("read_xlsx", {"path": "out.xlsx"})
            checks.append(
                ("workbook written and read", "b" in json.dumps(read, ensure_ascii=False))
            )

            read = turns.call("read_docx", {"path": "in.docx"})
            checks.append(
                ("Word read", "A paragraph of text." in json.dumps(read, ensure_ascii=False))
            )
            turns.call(
                "write_docx", {"path": "out.docx", "content": "# Title\n\n- one"}, via_use_tool=True
            )
            checks.append(("Word written", (workspace / "out.docx").exists()))

            page = turns.call("read_web_page", {"url": f"http://127.0.0.1:{web_port}/"})
            checks.append(("web page read", "Hello" in json.dumps(page, ensure_ascii=False)))
        finally:
            server.terminate()
            llm.terminate()
            web.shutdown()
            try:
                server.wait(10)
            except subprocess.TimeoutExpired:
                server.kill()

    for name, ok in checks:
        print(f"{'ok  ' if ok else 'FAIL'} {name}")
    if not all(ok for _, ok in checks):
        sys.exit(1)


if __name__ == "__main__":
    main()
