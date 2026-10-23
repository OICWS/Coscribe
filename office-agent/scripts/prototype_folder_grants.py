"""Prototype for issue #560, part 2: how the folders reach the container.

Run on Windows only, after prototype_script_sandbox.py proved the container itself:

    python scripts/prototype_folder_grants.py

Questions it answers, each printed as a table row or a timing:

1. Capabilities instead of the container's own SID. Each folder gets an ACE for a custom
   capability SID (DeriveCapabilitySidsFromName) once, when the user adds it; a run is then given
   the capabilities of its conversation's folders. A folder of another conversation, granted to a
   capability this run does not hold, must stay closed.
2. A read-write folder without the delete right: which everyday writes still work (overwrite,
   replace-by-rename, the Office libraries' saves) and which fail. This decides whether "delete
   goes to the recycle bin" can be enforced by the OS.
3. What writing the ACEs costs on a large folder, and taking them away again.
4. Folders where ACLs do not work: an SMB share and an exFAT volume.

It records results; the contract rows (marked) fail the run when they differ.
"""

from __future__ import annotations

import json
import os
import shutil
import subprocess
import sys
import tempfile
import time
from pathlib import Path

assert sys.platform == "win32", "the AppContainer prototype runs on Windows only"

sys.path.insert(0, str(Path(__file__).parent))
import prototype_script_sandbox as box  # noqa: E402

PROFILE = "coscribe.script-sandbox.prototype.caps"

PROBE = r"""
import json, os, shutil, sys

cfg = json.loads(sys.argv[1])
nd, dl, other, ro = cfg["nodelete"], cfg["delete"], cfg["other"], cfg["ro"]


def run(name, fn):
    try:
        detail = fn()
        print(json.dumps({"check": name, "ok": True, "detail": str(detail)[:400]}), flush=True)
    except BaseException as exc:
        print(json.dumps({"check": name, "ok": False, "detail": f"{type(exc).__name__}: {exc}"[:400]}), flush=True)


def write(path, text="x", mode="w"):
    with open(path, mode, encoding="utf-8") as fh:
        fh.write(text)
    return path


def p(folder, name):
    return os.path.join(folder, name)


run("cap: write own folder", lambda: write(p(nd, "new.txt")))
run("cap: read other conversation's folder", lambda: open(p(other, "theirs.txt")).read())
run("cap: write other conversation's folder", lambda: write(p(other, "mine.txt")))
run("cap: read ro folder", lambda: open(p(ro, "input.txt")).read())
run("cap: write ro folder", lambda: write(p(ro, "new.txt")))

run("no-delete: overwrite existing", lambda: write(p(nd, "existing.txt"), "new"))
run("no-delete: append", lambda: write(p(nd, "existing.txt"), "more", "a"))
run("no-delete: mkdir", lambda: os.mkdir(p(nd, "sub")))
run("no-delete: remove file", lambda: os.remove(p(nd, "existing.txt")))
run("no-delete: rmdir empty", lambda: os.rmdir(p(nd, "sub")))
run("no-delete: rename in folder", lambda: os.rename(p(nd, "new.txt"), p(nd, "renamed.txt")))
run("no-delete: replace existing (temp then rename)", lambda: (write(p(nd, "tmp.txt")), os.replace(p(nd, "tmp.txt"), p(nd, "existing.txt"))))


def openpyxl_save():
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active["A1"] = 1
    wb.save(p(nd, "book.xlsx"))
    wb.active["A1"] = 2
    wb.save(p(nd, "book.xlsx"))
    return openpyxl.load_workbook(p(nd, "book.xlsx")).active["A1"].value


def pandas_save():
    import pandas as pd
    pd.DataFrame({"a": [1]}).to_excel(p(nd, "frame.xlsx"))
    pd.DataFrame({"a": [2]}).to_excel(p(nd, "frame.xlsx"))
    pd.DataFrame({"a": [1]}).to_csv(p(nd, "frame.csv"))
    return pd.read_excel(p(nd, "frame.xlsx"))["a"][0]


def docx_save():
    import docx
    doc = docx.Document()
    doc.add_paragraph("one")
    doc.save(p(nd, "doc.docx"))
    doc.add_paragraph("two")
    doc.save(p(nd, "doc.docx"))
    return len(docx.Document(p(nd, "doc.docx")).paragraphs)


def pptx_save():
    import pptx
    deck = pptx.Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(p(nd, "deck.pptx"))
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(p(nd, "deck.pptx"))
    return len(pptx.Presentation(p(nd, "deck.pptx")).slides)


def pypdf_save():
    import pypdf
    for pages in (1, 2):
        writer = pypdf.PdfWriter()
        for _ in range(pages):
            writer.add_blank_page(100, 100)
        with open(p(nd, "out.pdf"), "wb") as fh:
            writer.write(fh)
    writer = pypdf.PdfWriter()
    writer.add_blank_page(100, 100)
    writer.write(p(nd, "out2.pdf"))
    writer.write(p(nd, "out2.pdf"))
    return len(pypdf.PdfReader(p(nd, "out.pdf")).pages)


def pillow_save():
    from PIL import Image
    for color in ("red", "blue"):
        Image.new("RGB", (8, 8), color).save(p(nd, "img.png"))
        Image.new("RGB", (8, 8), color).save(p(nd, "img.jpg"))
    return Image.open(p(nd, "img.png")).getpixel((0, 0))


def matplotlib_save():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    for n in (1, 2):
        plt.plot(range(n + 1))
        plt.savefig(p(nd, "chart.png"))
        plt.savefig(p(nd, "chart.pdf"))
        plt.close()
    return os.path.getsize(p(nd, "chart.png"))


run("no-delete: openpyxl save twice", openpyxl_save)
run("no-delete: pandas to_excel twice + csv", pandas_save)
run("no-delete: python-docx save twice", docx_save)
run("no-delete: python-pptx save twice", pptx_save)
run("no-delete: pypdf write twice (handle and path)", pypdf_save)
run("no-delete: Pillow png and jpg twice", pillow_save)
run("no-delete: matplotlib savefig png and pdf twice", matplotlib_save)

run("delete-allowed: remove file", lambda: os.remove(p(dl, "existing.txt")))
run("delete-allowed: rmtree", lambda: (os.makedirs(p(dl, "a/b")), write(p(dl, "a/b/c.txt")), shutil.rmtree(p(dl, "a"))))

for name, target in cfg.get("extra", {}).items():
    run(f"{name}: write", lambda t=target: write(os.path.join(t, "from-container.txt")))
    run(f"{name}: list", lambda t=target: os.listdir(t))
"""

LIBRARIES = ("openpyxl", "pandas", "python-docx", "python-pptx", "pypdf", "Pillow", "matplotlib")

# Contract rows: True must succeed, False must fail. Rows not here are recorded only.
EXPECTED = {
    "cap: write own folder": True,
    "cap: read other conversation's folder": False,
    "cap: write other conversation's folder": False,
    "cap: read ro folder": True,
    "cap: write ro folder": False,
    "no-delete: overwrite existing": True,
    "no-delete: remove file": False,
    "delete-allowed: remove file": True,
    "delete-allowed: rmtree": True,
}


def icacls(*args: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run(["icacls", *args], capture_output=True, text=True)


def powershell(command: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        ["powershell.exe", "-NoProfile", "-NonInteractive", "-Command", command],
        capture_output=True,
        text=True,
    )


def large_folder_timing(base: Path, sid: str, folders: int) -> None:
    tree = base / "large"
    t0 = time.monotonic()
    for d in range(folders):
        sub = tree / f"d{d:04}"
        sub.mkdir(parents=True)
        for f in range(100):
            (sub / f"f{f:03}.txt").write_bytes(b"x")
    took = time.monotonic() - t0
    print(f"\n{folders * 100:,} files in {folders} folders created in {took:.1f}s")
    t0 = time.monotonic()
    out = icacls(str(tree), "/grant", f"*{sid}:(OI)(CI)(M)", "/Q")
    print(f"grant on the tree: {time.monotonic() - t0:.1f}s, exit {out.returncode}")
    t0 = time.monotonic()
    out = icacls(str(tree), "/remove:g", f"*{sid}", "/Q")
    print(f"remove from the tree: {time.monotonic() - t0:.1f}s, exit {out.returncode}")
    shutil.rmtree(tree, ignore_errors=True)


def smb_share(base: Path, sid: str) -> tuple[str, Path] | None:
    folder = base / "shared"
    folder.mkdir()
    out = powershell(
        f"New-SmbShare -Name coscribeproto -Path '{folder}' -FullAccess Everyone | Out-Null"
    )
    if out.returncode:
        print(f"SMB share not created: {out.stderr.strip()[:300]}")
        return None
    grant = icacls(str(folder), "/grant", f"*{sid}:(OI)(CI)(M)")
    print(f"SMB share folder granted to the capability: exit {grant.returncode}")
    return r"\\localhost\coscribeproto", folder


def exfat_volume(base: Path) -> str | None:
    vhd = base / "exfat.vhdx"
    script = base / "diskpart.txt"
    script.write_text(
        f'create vdisk file="{vhd}" maximum=64 type=expandable\n'
        "attach vdisk\n"
        "create partition primary\n"
        'format fs=exfat quick label="CSPROTO"\n'
        "assign letter=X\n",
        encoding="ascii",
    )
    try:
        out = subprocess.run(["diskpart", "/s", str(script)], capture_output=True, text=True)
    except OSError as exc:
        # diskpart needs an administrator; a standard user's run skips the volume.
        print(f"exFAT volume not created: {exc}")
        return None
    if out.returncode or not Path("X:\\").exists():
        print(f"exFAT volume not created: {out.stdout.strip()[-400:]}")
        return None
    acl = icacls("X:\\", "/grant", "*S-1-15-2-1:(OI)(CI)(RX)")
    print(f"icacls on exFAT: exit {acl.returncode} {(acl.stdout + acl.stderr).strip()[:200]}")
    return "X:\\"


def detach_exfat(base: Path) -> None:
    script = base / "detach.txt"
    script.write_text(
        f'select vdisk file="{base / "exfat.vhdx"}"\ndetach vdisk\n', encoding="ascii"
    )
    subprocess.run(["diskpart", "/s", str(script)], capture_output=True)


def main() -> int:
    base = Path(tempfile.mkdtemp(prefix="coscribe-grants-")).resolve()
    venv = base / "script-env"
    interpreter, prefix = box.base_interpreter(base)
    subprocess.run([str(interpreter), "-m", "venv", str(venv)], check=True)
    py = venv / "Scripts" / "python.exe"
    subprocess.run([str(py), "-m", "pip", "install", "-q", *LIBRARIES], check=True)

    folders = {name: base / name for name in ("nodelete", "delete", "other", "ro", "scratch")}
    for folder in folders.values():
        folder.mkdir()
    (folders["nodelete"] / "existing.txt").write_text("old", encoding="utf-8")
    (folders["delete"] / "existing.txt").write_text("old", encoding="utf-8")
    (folders["other"] / "theirs.txt").write_text("theirs", encoding="utf-8")
    (folders["ro"] / "input.txt").write_text("reference", encoding="utf-8")
    probe = folders["scratch"] / "probe.py"
    probe.write_text(PROBE, encoding="utf-8")

    container, container_text = box.container_sid(PROFILE)
    caps = {name: box.capability_sid(f"coscribe-proto-{name}") for name in ("runtime", *folders)}
    print(f"container {container_text}")
    for name, (_, text) in caps.items():
        print(f"capability {name}: {text}")

    # (RX,W) is read plus write-data, append, extended attributes and attributes: no DELETE on
    # the file and no DELETE_CHILD on the folder, so the only way to remove a file is a broker.
    rights = {"nodelete": "RX,W", "delete": "M", "other": "M", "ro": "RX", "scratch": "M"}
    grants: list[tuple[Path, str]] = [
        (venv, caps["runtime"][1]),
        (prefix, caps["runtime"][1]),
    ]
    for path, sid in grants:
        icacls(str(path), "/grant", f"*{sid}:(OI)(CI)(RX)", "/Q")
    for name, folder in folders.items():
        icacls(str(folder), "/grant", f"*{caps[name][1]}:(OI)(CI)({rights[name]})", "/Q")
        grants.append((folder, caps[name][1]))

    extra: dict[str, str] = {}
    share = smb_share(base, caps["nodelete"][1])
    if share:
        extra["SMB share (granted)"] = share[0]
        grants.append((share[1], caps["nodelete"][1]))
    exfat = exfat_volume(base)
    if exfat:
        extra["exFAT volume (no ACLs)"] = exfat

    held = [caps[name][0] for name in ("runtime", "nodelete", "delete", "ro", "scratch")]
    cfg = {name: str(folder) for name, folder in folders.items()} | {"extra": extra}
    env = {"PYTHONIOENCODING": "utf-8"}
    failed = 0
    try:
        log = base / "probe.log"
        _, handle = box.run_in_container(
            container, [str(py), str(probe), json.dumps(cfg)], folders["nodelete"], env, log, held
        )
        print(f"\nprobe exit code {box.wait(handle, 600)}\n")
        rows = {}
        for line in log.read_text(encoding="utf-8", errors="replace").splitlines():
            try:
                row = json.loads(line)
                rows[row["check"]] = row
            except (ValueError, KeyError):
                print(f"  | {line}")
        for check, row in rows.items():
            want = EXPECTED.get(check)
            got = row["ok"]
            mark = "   " if want is None else ("ok " if got == want else "BAD")
            failed += want is not None and got != want
            print(f"{mark} {check:48} {'allow' if got else 'deny':6} {row['detail']}")
        missing = set(EXPECTED) - set(rows)
        failed += len(missing)
        for check in sorted(missing):
            print(f"BAD {check:48} (no result)")
        if exfat:
            print(f"\nfiles on the exFAT volume: {os.listdir(exfat)}")
        # Desktop and OneDrive folders can hold hundreds of thousands of files.
        for count in (200, 2000):
            large_folder_timing(base, caps["other"][1], count)
    finally:
        if share:
            powershell("Remove-SmbShare -Name coscribeproto -Force")
        if exfat:
            detach_exfat(base)
        for path, sid in grants:
            icacls(str(path), "/remove:g", f"*{sid}", "/Q")
        box.userenv.DeleteAppContainerProfile(PROFILE)
        shutil.rmtree(base, ignore_errors=True)

    print(f"\n{failed} contract result(s) differ")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
