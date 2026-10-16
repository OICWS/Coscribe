"""Writing a workbook from block and object operations: execute, guard, recalculate, check.

openpyxl saves the workbook, marked to recalculate when Excel opens it; IronCalc then writes each
formula's value into the file, so every viewer shows numbers and the checks read them (decisions
0011 and 0012). No other program rewrites the user's file."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from .blocks import parse_blocks, parse_ops
from .calc import recalculate
from .column_layout import explicit_width_starts, widen_after_recalc, widen_for_values
from .engine_objects import apply_ops
from .engine_openpyxl import conflicts, write_block
from .findings import Written, check
from .guard import copy_path, inventory, load_watching, losses

_RESULTS_LISTED = 40


def _open(path: Path) -> tuple[Any, list[str], Written]:
    from openpyxl import Workbook

    written = Written()
    if not path.exists():
        workbook = Workbook()
        for sheet in list(workbook.worksheets):
            workbook.remove(sheet)
        return workbook, [], written
    workbook, dropped = load_watching(path)
    written.sheets_with_content = {
        sheet.title
        for sheet in workbook.worksheets
        if any(cell.value is not None for row in sheet.iter_rows() for cell in row)
    }
    return workbook, dropped, written


def _save(workbook: Any, path: Path) -> Path:
    workbook.calculation.fullCalcOnLoad = True
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".~{path.stem}.coscribe{path.suffix}")
    workbook.save(str(temporary))
    return temporary


def _place(temporary: Path, path: Path, existed: bool, lost: list[str]) -> Path:
    target = copy_path(path) if existed and lost else path
    temporary.replace(target)
    return target


def _formula_cells(workbook: Any, written: Written) -> dict[str, list[str]]:
    return {
        name: [c for c, _ in cells if str(workbook[name][c].value).startswith("=")]
        for name, cells in written.cells.items()
        if name in workbook.sheetnames
    }


def _widen(target: Path, formulas: dict[str, list[str]]) -> None:
    """Columns wide enough for the computed numbers, which exist only once the values are in the
    file; only its `<cols>` are patched, so the values stay."""
    for name, cells in formulas.items():
        widen_after_recalc(target, name, cells)


def _results(workbook: Any, values: Any, written: Written) -> dict[str, Any]:
    """The computed value of each formula written, so a wrong one shows in this call."""
    listed: dict[str, Any] = {}
    total = 0
    for name, cells in written.cells.items():
        for coordinate, _ in cells:
            if name in workbook.sheetnames and str(workbook[name][coordinate].value).startswith(
                "="
            ):
                total += 1
                if len(listed) < _RESULTS_LISTED:
                    listed[f"{name}!{coordinate}"] = values[name][coordinate].value
    return {"formula_results": listed, "formulas_not_listed": total - len(listed)}


def _finish(
    path: Path, workbook: Any, dropped: list[str], written: Written, deleted: bool
) -> dict[str, Any]:
    from openpyxl import load_workbook

    existed = path.exists()
    before = inventory(path) if existed else None
    temporary = _save(workbook, path)
    try:
        lost = [] if before is None else losses(before, inventory(temporary), dropped, deleted)
        target = _place(temporary, path, existed, lost)
    finally:
        temporary.unlink(missing_ok=True)
    recalc = recalculate(target)
    if recalc.get("status") != "skipped":
        _widen(target, _formula_cells(workbook, written))
    saved = load_workbook(str(target))
    values = load_workbook(str(target), data_only=True)
    report = {
        "path": str(target),
        "saved_as_copy": target != path,
        "lost_if_saved_over_original": lost,
        "recalc": recalc,
        **_results(saved, values, written),
        "findings": check(saved, values, target, written),
    }
    return report


def _fit(workbook: Any, path: Path, written: Written) -> None:
    """Widen only the columns just written to, so a user's own layout stays as it was."""
    for name, cells in written.cells.items():
        sheet = workbook[name]
        starts = explicit_width_starts(path, name) if path.exists() else set()
        widen_for_values(sheet, None, [sheet[c] for c, _ in cells], starts)


def write_cells(
    path: Path, blocks: Any, styles: Any = None, overwrite: bool = False
) -> dict[str, Any]:
    """`write_xlsx_cells`: blocks of cells with their styles, then recalculated and checked."""
    parsed = parse_blocks(blocks, styles)
    workbook, dropped, written = _open(path)
    clashing = [] if overwrite else conflicts(workbook, parsed)
    if clashing:
        raise ValueError(
            f"{len(clashing)} cells already hold something else, e.g. {clashing[:5]}; "
            "read them, then write again with overwrite=true if they should be replaced"
        )
    for block in parsed:
        for cell, looked in write_block(workbook, block):
            written.add(block.sheet, cell.coordinate, looked)
    _fit(workbook, path, written)
    return _finish(path, workbook, dropped, written, deleted=False)


def modify_workbook(path: Path, ops: Any) -> dict[str, Any]:
    """`modify_xlsx`: sheets, rows and columns, charts, rules, print setup and names."""
    parsed = parse_ops(ops)
    if not path.exists() and not any(op["op"] == "add_sheet" for op in parsed):
        raise ValueError(f"{path.name} does not exist; write cells to it first")
    workbook, dropped, written = _open(path)
    apply_ops(workbook, parsed)
    for op in parsed:
        if op.get("sheet") in workbook.sheetnames:
            written.cells.setdefault(op["sheet"], [])
    deleted = any(op["op"] == "delete_sheet" for op in parsed)
    return _finish(path, workbook, dropped, written, deleted)
