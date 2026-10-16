"""The block operations executed with openpyxl: the engine that needs no Excel."""

from __future__ import annotations

from copy import copy
from typing import TYPE_CHECKING, Any

from .blocks import Block, Cell

if TYPE_CHECKING:
    from openpyxl import Workbook
    from openpyxl.worksheet.worksheet import Worksheet


def sheet_for(workbook: Workbook, name: str) -> Worksheet:
    if name in workbook.sheetnames:
        found: Worksheet = workbook[name]
        return found
    created: Worksheet = workbook.create_sheet(name)
    return created


def conflicts(workbook: Workbook, blocks: list[Block]) -> list[str]:
    """Cells a write would change that already hold something else; nothing is written over
    them unless the caller says so."""
    found = []
    for block in blocks:
        if block.sheet not in workbook.sheetnames:
            continue
        sheet = workbook[block.sheet]
        for i, row in enumerate(block.cells):
            for j, cell in enumerate(row):
                there = sheet.cell(block.row + i, block.col + j)
                if there.value not in (None, "") and there.value != cell.value:
                    found.append(f"{block.sheet}!{there.coordinate}={there.value!r}")
    return found


def _font(old: Any, style: dict[str, Any]) -> Any:
    from openpyxl.styles import Font

    underline = old.underline
    if "underline" in style:
        underline = "single" if style["underline"] else None
    return Font(
        name=style.get("font", old.name),
        size=style.get("size", old.size),
        bold=style.get("bold", old.bold),
        italic=style.get("italic", old.italic),
        underline=underline,
        color=style["color"] if "color" in style else old.color,
        vertAlign=old.vertAlign,
        strike=old.strike,
    )


def _alignment(old: Any, style: dict[str, Any]) -> Any:
    from openpyxl.styles import Alignment

    return Alignment(
        horizontal=style.get("align", old.horizontal),
        vertical=style.get("valign", old.vertical),
        wrap_text=style.get("wrap", old.wrap_text),
        indent=style.get("indent", old.indent),
        text_rotation=old.text_rotation,
        shrink_to_fit=old.shrink_to_fit,
    )


def _border(old: Any, sides: dict[str, str]) -> Any:
    from openpyxl.styles import Border, Side

    def side(name: str) -> Any:
        if name not in sides:
            return getattr(old, name)
        kind: Any = None if sides[name] == "none" else sides[name]
        return Side(style=kind, color="808080")

    return Border(left=side("left"), right=side("right"), top=side("top"), bottom=side("bottom"))


def apply_style(target: Any, style: dict[str, Any]) -> None:
    from openpyxl.styles import PatternFill

    if style.keys() & {"font", "size", "bold", "italic", "underline", "color"}:
        target.font = _font(target.font, style)
    if style.keys() & {"align", "valign", "wrap", "indent"}:
        target.alignment = _alignment(target.alignment, style)
    if "fill" in style:
        target.fill = PatternFill("solid", fgColor=style["fill"])
    if "border" in style:
        target.border = _border(target.border, style["border"])
    if "num" in style:
        target.number_format = style["num"]


def _copy_look(source: Any, target: Any) -> None:
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.number_format = source.number_format


def _put(sheet: Worksheet, row: int, col: int, cell: Cell, block: Block) -> Any:
    from openpyxl.comments import Comment

    target = sheet.cell(row, col)
    target.value = cell.value
    if cell.text:
        target.data_type = "s"
    if block.like_row is not None:
        _copy_look(sheet.cell(block.like_row, col), target)
    elif block.like is not None:
        _copy_look(sheet.cell(*block.like), target)
    apply_style(target, cell.style)
    if cell.note:
        target.comment = Comment(cell.note, "coscribe")
    return target


def _fill_down(sheet: Worksheet, block: Block) -> list[Any]:
    """Repeat the block's last row down to `fill_down_to`, formulas moved with their row."""
    from openpyxl.formula.translate import Translator

    last = block.row + len(block.cells) - 1
    written = []
    for col in range(block.col, block.col + len(block.cells[-1])):
        source = sheet.cell(last, col)
        for row in range(last + 1, (block.fill_down_to or last) + 1):
            target = sheet.cell(row, col)
            value = source.value
            if source.data_type == "f" and isinstance(value, str):
                value = Translator(value, origin=source.coordinate).translate_formula(
                    target.coordinate
                )
            target.value = value
            if source.data_type == "s":
                target.data_type = "s"
            _copy_look(source, target)
            written.append(target)
    return written


def write_block(workbook: Workbook, block: Block) -> list[tuple[Any, bool]]:
    """Writes one block; returns each written cell with whether it was given a look (a style
    of its own, `like_row` or `like`), which the edit-style check needs."""
    sheet = sheet_for(workbook, block.sheet)
    written = []
    for i, row in enumerate(block.cells):
        for j, cell in enumerate(row):
            target = _put(sheet, block.row + i, block.col + j, cell, block)
            looked = bool(cell.style) or block.like_row is not None or block.like is not None
            written.append((target, looked))
    written += [(target, True) for target in _fill_down(sheet, block)]
    return written
