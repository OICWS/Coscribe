"""The column widths `write_xlsx` gives the sheet it writes, and which ones it may change later."""

from __future__ import annotations

import json
import zipfile
from copy import copy
from dataclasses import dataclass, field, replace
from pathlib import Path
from typing import TYPE_CHECKING, Any
from xml.etree import ElementTree as ET

from .column_widths import DEFAULT_WIDTH, MAX_WIDTH, needed_width

if TYPE_CHECKING:
    from openpyxl.worksheet.worksheet import Worksheet

# A width write_xlsx chose and one a person or template chose look the same in the file (and
# LibreOffice drops `bestFit` when it saves), so the columns write_xlsx sized are listed in a
# custom document property, which LibreOffice and Excel keep.
FITTED_PROPERTY = "coscribe.fitted_columns"
_MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_REL_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
_RELS = "{http://schemas.openxmlformats.org/package/2006/relationships}"


@dataclass
class Span:
    """One `<col>` element: a run of columns with the same width and look."""

    min: int
    max: int
    width: float | None = None
    look: dict[str, Any] = field(default_factory=dict)
    style: Any = None


@dataclass
class Fit:
    owned: set[int] = field(default_factory=set)
    widths: dict[int, float] = field(default_factory=dict)


def _sheet_part(archive: zipfile.ZipFile, sheet_name: str) -> str | None:
    workbook = ET.fromstring(archive.read("xl/workbook.xml"))
    rel_id = next(
        (s.get(_REL_ID) for s in workbook.iter(f"{_MAIN}sheet") if s.get("name") == sheet_name),
        None,
    )
    rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    target = next(
        (r.get("Target", "") for r in rels.iter(f"{_RELS}Relationship") if r.get("Id") == rel_id),
        None,
    )
    if rel_id is None or not target:
        return None
    return target.lstrip("/") if target.startswith("/") else f"xl/{target}"


def explicit_width_starts(path: Path, sheet_name: str) -> set[int]:
    """The first column of each `<col>` that carries a width; openpyxl gives one without a width
    (only hidden, styled or outlined) a width of 13, so the file itself is read."""
    with zipfile.ZipFile(path) as archive:
        part = _sheet_part(archive, sheet_name)
        if part is None:
            return set()
        starts = set()
        with archive.open(part) as stream:
            for _, element in ET.iterparse(stream):
                if element.tag.endswith("}col") and element.get("width") is not None:
                    starts.add(int(element.get("min", "0")))
                elif element.tag.endswith("}sheetData"):
                    break
                element.clear()
    return starts


def _fitted_map(workbook: Any) -> dict[str, list[list[int]]]:
    try:
        value = json.loads(workbook.custom_doc_props[FITTED_PROPERTY].value)
    except (KeyError, ValueError, TypeError):
        return {}
    return value if isinstance(value, dict) else {}


def fitted_columns(workbook: Any, sheet_name: str) -> set[int]:
    return {
        column
        for first, last in _fitted_map(workbook).get(sheet_name, [])
        for column in range(first, last + 1)
    }


def record_fitted(workbook: Any, sheet_name: str, columns: set[int]) -> None:
    from openpyxl.packaging.custom import StringProperty

    runs: list[list[int]] = []
    for column in sorted(columns):
        if runs and runs[-1][1] == column - 1:
            runs[-1][1] = column
        else:
            runs.append([column, column])
    fitted = _fitted_map(workbook)
    fitted[sheet_name] = runs
    if FITTED_PROPERTY in workbook.custom_doc_props.names:
        del workbook.custom_doc_props[FITTED_PROPERTY]
    workbook.custom_doc_props.append(
        StringProperty(name=FITTED_PROPERTY, value=json.dumps(fitted, ensure_ascii=False))
    )


def forget_fitted(workbook: Any, sheet_name: str, columns: range) -> None:
    fitted = fitted_columns(workbook, sheet_name)
    if fitted & set(columns):
        record_fitted(workbook, sheet_name, fitted - set(columns))


def carried_spans(old: Worksheet, width_starts: set[int], fitted: set[int]) -> list[Span]:
    """The old sheet's columns, as they go onto the sheet that replaces it: everything kept, but
    a width write_xlsx chose itself is dropped so the new content sizes the column again."""
    from openpyxl.utils import column_index_from_string

    spans = []
    dimensions: dict[str, Any] = dict(old.column_dimensions)
    for letter, dimension in dimensions.items():
        first = dimension.min or column_index_from_string(letter)
        last = dimension.max or first
        mine = all(column in fitted for column in range(first, last + 1))
        explicit = first in width_starts and not mine
        spans.append(
            Span(
                first,
                last,
                float(dimension.width or 0) if explicit else None,
                {
                    "hidden": bool(dimension.hidden),
                    "outlineLevel": dimension.outlineLevel or 0,
                    "collapsed": bool(dimension.collapsed),
                },
                copy(dimension._style) if dimension.has_style else None,
            )
        )
    return spans


def _set_width(spans: list[Span], column: int, width: float) -> None:
    for i, span in enumerate(spans):
        if span.min <= column <= span.max:
            pieces = [replace(span, max=column - 1)] if span.min < column else []
            pieces.append(replace(span, min=column, max=column, width=width))
            if column < span.max:
                pieces.append(replace(span, min=column + 1))
            spans[i : i + 1] = pieces
            return
    spans.append(Span(column, column, width))


def _kept_width(spans: list[Span], column: int) -> float | None:
    return next((s.width for s in spans if s.min <= column <= s.max and s.width is not None), None)


def _empty(value: Any) -> bool:
    return value is None or value == ""


def _needed(sheet: Worksheet, spans: list[Span]) -> dict[int, float]:
    needed: dict[int, float] = {}
    last = sheet.max_column
    spilling = []
    for row in sheet.iter_rows():
        for column, cell in enumerate(row, start=1):
            # A formula has no value before LibreOffice recalculates the saved file; its header
            # and the constants beside it size the column.
            if cell.data_type == "f":
                continue
            width = needed_width(cell.value, cell.number_format)
            if isinstance(cell.value, str) and column < last and _empty(row[column].value):
                spilling.append((row, column, width))
                continue
            needed[column] = max(needed.get(column, 0.0), width)

    def shown(column: int) -> float:
        kept = _kept_width(spans, column)
        return (
            kept
            if kept is not None
            else min(max(needed.get(column, 0.0), DEFAULT_WIDTH), MAX_WIDTH)
        )

    # Text runs on over the empty cells to its right, up to the sheet's last column (where a
    # printed page cuts it); only text longer than that run widens its own column.
    for row, column, width in spilling:
        room, right = shown(column), column + 1
        while right <= last and _empty(row[right - 1].value):
            room += shown(right)
            right += 1
        if width > room:
            needed[column] = max(needed.get(column, 0.0), width)
    return needed


def _apply(sheet: Worksheet, spans: list[Span]) -> None:
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.dimensions import ColumnDimension

    for span in sorted(spans, key=lambda s: s.min):
        if span.width is None and not (any(span.look.values()) or span.style):
            continue
        letter = get_column_letter(span.min)
        dimension: Any = ColumnDimension(
            sheet, index=letter, width=span.width or 0, min=span.min, max=span.max, **span.look
        )
        if span.style is not None:
            dimension._style = copy(span.style)
        sheet.column_dimensions[letter] = dimension


def fit_columns(sheet: Worksheet, spans: list[Span]) -> Fit:
    """Widen every column of `sheet` whose longest shown value does not fit; a column with a
    kept width in `spans` stays as it is."""
    fit = Fit()
    needed = _needed(sheet, spans)
    for column in range(1, sheet.max_column + 1):
        if _kept_width(spans, column) is not None:
            continue
        fit.owned.add(column)
        if needed.get(column, 0.0) > DEFAULT_WIDTH:
            fit.widths[column] = round(min(needed[column], MAX_WIDTH), 2)
            _set_width(spans, column, fit.widths[column])
    _apply(sheet, spans)
    return fit
