"""Column widths that fit what a sheet shows, measured the way Excel draws a cell."""

from __future__ import annotations

import unicodedata
from typing import TYPE_CHECKING, Any

if TYPE_CHECKING:
    from openpyxl.worksheet.worksheet import Worksheet

# Excel's and LibreOffice's width of a column nobody sized, in characters of the default font.
DEFAULT_WIDTH = 8.43
# The cell's own margins and a bold header eat about two characters of the width.
_PADDING = 2.0
# One long note must not push the rest of the sheet off the page; past this the text is cut,
# which a reader can still widen, rather than every other column being scrolled away.
MAX_WIDTH = 50.0
# Excel shows a General number in at most this many characters, rounding decimals or switching
# to scientific notation rather than printing `###`.
_GENERAL_MAX_CHARS = 11
_DIGIT_CODES = "0#?"


def text_width(text: str) -> float:
    """Width in characters of the default font; a CJK or full-width character takes two."""
    widest = 0
    for line in text.split("\n"):
        widest = max(widest, sum(2 if unicodedata.east_asian_width(c) in "WF" else 1 for c in line))
    return float(widest)


def _sections(number_format: str) -> list[str]:
    sections: list[str] = []
    current: list[str] = []
    quoted = False
    for char in number_format:
        if char == '"':
            quoted = not quoted
        if char == ";" and not quoted:
            sections.append("".join(current))
            current = []
        else:
            current.append(char)
    sections.append("".join(current))
    return sections


def _split_codes(section: str) -> tuple[str, str, str]:
    """The digit codes (`#,##0.00%`) and the text printed before and after them (`$`, `元`)."""
    codes: list[str] = []
    prefix: list[str] = []
    suffix: list[str] = []
    i = 0
    while i < len(section):
        char = section[i]
        literal = suffix if codes else prefix
        if char == '"':
            end = section.find('"', i + 1)
            end = len(section) if end < 0 else end
            literal.append(section[i + 1 : end])
            i = end + 1
        elif char == "[":
            end = section.find("]", i)
            end = len(section) if end < 0 else end
            inner = section[i + 1 : end]
            if inner.startswith("$"):
                literal.append(inner[1:].split("-")[0])
            i = end + 1
        elif char in "\\_*":
            # `\x` prints x, `_x` leaves a gap as wide as x, `*x` only fills spare room.
            literal.append("" if char == "*" else section[i + 1 : i + 2])
            i += 2
        else:
            (codes if char in _DIGIT_CODES + ".,%" else literal).append(char)
            i += 1
    return "".join(codes), "".join(prefix), "".join(suffix)


def _general(value: float) -> str:
    text = str(value) if isinstance(value, int) else f"{value:.15g}"
    return text[:_GENERAL_MAX_CHARS]


def _formatted(value: float, number_format: str) -> str:
    sections = _sections(number_format)
    section = sections[0]
    if value < 0 and len(sections) > 1:
        section = sections[1]
    elif value == 0 and len(sections) > 2:
        section = sections[2]
    if section.strip().lower() in ("general", ""):
        return _general(value)
    codes, prefix, suffix = _split_codes(section)
    sign = "-" if value < 0 and len(sections) == 1 else ""
    if not any(code in codes for code in _DIGIT_CODES):
        return sign + prefix + suffix
    percent = codes.count("%")
    shown = abs(value) * 100**percent
    integer_codes, _, decimal_codes = codes.replace("%", "").partition(".")
    decimals = sum(decimal_codes.count(code) for code in _DIGIT_CODES)
    grouping = "," if "," in integer_codes.rstrip(",") else ""
    return f"{sign}{prefix}{shown:{grouping}.{decimals}f}{'%' * percent}{suffix}"


def displayed_text(value: Any, number_format: str = "General") -> str:
    """The text a cell shows: a number with its number format applied."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int | float):
        return _formatted(value, number_format or "General")
    return str(value)


def explicit_column_widths(sheet: Worksheet) -> dict[str, float]:
    """The width set on each column of `sheet`, one entry per column letter."""
    from openpyxl.utils import column_index_from_string, get_column_letter

    widths: dict[str, float] = {}
    for letter, dimension in sheet.column_dimensions.items():
        if not dimension.width:
            continue
        first = dimension.min or column_index_from_string(letter)
        for index in range(first, (dimension.max or first) + 1):
            widths[get_column_letter(index)] = float(dimension.width)
    return widths


def fit_column_widths(sheet: Worksheet, kept: dict[str, float]) -> None:
    """Widen every column of `sheet` whose longest shown value does not fit the default width;
    a column in `kept` gets that width instead, whatever its content."""
    from openpyxl.utils import get_column_letter

    needed: dict[str, float] = {}
    for row in sheet.iter_rows():
        for column, (cell, right) in enumerate(zip(row, [*row[1:], None], strict=True), start=1):
            # A formula's value only exists after LibreOffice recalculates the saved file; its
            # header or neighbouring constants size the column instead.
            if cell.data_type == "f":
                continue
            # Text runs on over an empty cell to its right, so a title above a table needs no
            # width of its own. The last column is still sized: a printed page cuts it off.
            if isinstance(cell.value, str) and right is not None and right.value in (None, ""):
                continue
            letter = get_column_letter(column)
            width = text_width(displayed_text(cell.value, cell.number_format))
            needed[letter] = max(needed.get(letter, 0.0), width + _PADDING if width else 0.0)
    for letter, width in needed.items():
        if letter not in kept and width > DEFAULT_WIDTH:
            sheet.column_dimensions[letter].width = round(min(width, MAX_WIDTH), 2)
    for letter, width in kept.items():
        sheet.column_dimensions[letter].width = width
