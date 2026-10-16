from __future__ import annotations

import zipfile
from collections.abc import Sequence
from pathlib import Path
from typing import Any, Optional

from ..._file_locks import locked_by_path
from ..._workspace import WorkspaceScope
from .build import modify_workbook, write_cells
from .calc import recalculate
from .csv_blocks import expand_blocks
from .helpers import _render_sheet_table
from .overview import sheet_layout


def _formulas_or_none(path: Path) -> Any:
    """The workbook with its formulas, for the layout; a file openpyxl reads only for its values
    still reads, without one."""
    from openpyxl import load_workbook

    try:
        return load_workbook(str(path))
    except (KeyError, ValueError, TypeError, AttributeError, OSError, zipfile.BadZipFile):
        return None


def _with_layout(table: str, sheet: Any) -> str:
    layout = sheet_layout(sheet) if hasattr(sheet, "iter_cols") else ""
    return f"{table}\n\n{layout}" if layout else table


class SpreadsheetToolkit:
    def __init__(
        self,
        root: str | Path,
        *,
        state_dir: str | Path | None = None,
        extra_readable: Sequence[str | Path] = (),
        extra_writable: Sequence[str | Path] = (),
    ) -> None:
        # `state_dir` keeps the constructor every office toolkit shares, which the approval
        # preview builds each of the same way (conversation/edit_preview.py).
        del state_dir
        self._scope = WorkspaceScope(
            root, extra_readable=extra_readable, extra_writable=extra_writable
        )

    def _check_readable(self, path: str) -> Path:
        file_path = self._scope.resolve(path)
        if not file_path.exists():
            raise ValueError(f"File does not exist: {path}")
        if not file_path.is_file():
            raise ValueError(f"Path is not a file: {path}")
        return file_path

    @locked_by_path
    def read_xlsx(self, path: str, sheet: Optional[str] = None) -> str:  # noqa: UP045
        from openpyxl import load_workbook

        file_path = self._check_readable(path)
        workbook = load_workbook(str(file_path), data_only=True)
        formulas = _formulas_or_none(file_path)
        if sheet is not None:
            if sheet not in workbook.sheetnames:
                raise ValueError(
                    f"Sheet {sheet!r} not found. Available sheets: {workbook.sheetnames}"
                )
            return _with_layout(_render_sheet_table(workbook[sheet]), formulas and formulas[sheet])
        sections = []
        for name in workbook.sheetnames:
            table = _with_layout(_render_sheet_table(workbook[name]), formulas and formulas[name])
            sections.append(f"## {name}\n\n{table}")
        return "\n\n".join(sections)

    @locked_by_path
    def recalc_xlsx(self, path: str) -> dict[str, object]:
        file_path = self._check_readable(path)
        result = recalculate(file_path)
        return {
            "path": self._scope.relative(file_path),
            "status": result["status"],
            "skipped_reason": result.get("skipped_reason"),
            "total_formulas": result.get("total_formulas"),
            "total_errors": result.get("total_errors"),
            "error_locations": result.get("error_locations"),
            "unchecked": result.get("unchecked"),
        }

    @locked_by_path
    def write_xlsx_cells(
        self,
        path: str,
        blocks: list[dict[str, Any]],
        styles: dict[str, Any] | None = None,
        overwrite: bool = False,
    ) -> dict[str, Any]:
        file_path = self._scope.resolve(path, write=True)
        report = write_cells(file_path, expand_blocks(self._scope, blocks), styles, overwrite)
        return {**report, "path": self._scope.relative(Path(report["path"]))}

    @locked_by_path
    def modify_xlsx(self, path: str, ops: list[dict[str, Any]]) -> dict[str, Any]:
        report = modify_workbook(self._scope.resolve(path, write=True), ops)
        return {**report, "path": self._scope.relative(Path(report["path"]))}
