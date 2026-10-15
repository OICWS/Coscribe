"""`modify_xlsx`'s operations executed with openpyxl: sheets, rows and columns, charts,
conditional formats, validations, print setup and names."""

from __future__ import annotations

from collections.abc import Callable
from typing import TYPE_CHECKING, Any

if TYPE_CHECKING:
    from openpyxl import Workbook
    from openpyxl.worksheet.worksheet import Worksheet


_PALETTE = ["1F4E79", "ED7D31", "70AD47", "FFC000", "5B9BD5", "A5A5A5", "7030A0", "C00000"]


def _sheet(workbook: Workbook, name: str) -> Worksheet:
    if name not in workbook.sheetnames:
        raise ValueError(f"No sheet {name!r}; the workbook has {workbook.sheetnames}")
    sheet: Worksheet = workbook[name]
    return sheet


def _bounds(ref: str) -> tuple[int, int, int, int]:
    from openpyxl.utils.cell import range_boundaries

    min_col, min_row, max_col, max_row = range_boundaries(ref.replace("$", ""))
    if min_col is None or min_row is None or max_col is None or max_row is None:
        raise ValueError(f"{ref!r} is not a range like 'A1:C13'")
    return min_col, min_row, max_col, max_row


def _split_ref(ref: str, default_sheet: str) -> tuple[str, str]:
    """`'Sales data'!A1:C13` or `A1:C13` -> (sheet, range)."""
    if "!" not in ref:
        return default_sheet, ref
    sheet, cells = ref.rsplit("!", 1)
    return sheet.strip("'"), cells


def _add_sheet(workbook: Workbook, op: dict[str, Any]) -> None:
    if op["name"] in workbook.sheetnames:
        raise ValueError(f"A sheet {op['name']!r} already exists")
    sheet = workbook.create_sheet(op["name"], op.get("index"))
    if op.get("tab_color"):
        sheet.sheet_properties.tabColor = op["tab_color"]


def _rename_sheet(workbook: Workbook, op: dict[str, Any]) -> None:
    _sheet(workbook, op["sheet"]).title = op["name"]


def _delete_sheet(workbook: Workbook, op: dict[str, Any]) -> None:
    workbook.remove(_sheet(workbook, op["sheet"]))


def _tab_color(workbook: Workbook, op: dict[str, Any]) -> None:
    _sheet(workbook, op["sheet"]).sheet_properties.tabColor = op["color"]


def _cols(workbook: Workbook, op: dict[str, Any]) -> None:
    from openpyxl.utils import column_index_from_string, get_column_letter

    sheet = _sheet(workbook, op["sheet"])
    first, _, last = str(op["cols"]).partition(":")
    start, end = column_index_from_string(first), column_index_from_string(last or first)
    for index in range(start, end + 1):
        dimension = sheet.column_dimensions[get_column_letter(index)]
        if "width" in op:
            dimension.width = float(op["width"])
        if "hidden" in op:
            dimension.hidden = bool(op["hidden"])


def _rows(workbook: Workbook, op: dict[str, Any]) -> None:
    sheet = _sheet(workbook, op["sheet"])
    first, _, last = str(op["rows"]).partition(":")
    for index in range(int(first), int(last or first) + 1):
        dimension = sheet.row_dimensions[index]
        if "height" in op:
            dimension.height = float(op["height"])
        if "hidden" in op:
            dimension.hidden = bool(op["hidden"])


def _freeze(workbook: Workbook, op: dict[str, Any]) -> None:
    _sheet(workbook, op["sheet"]).freeze_panes = op.get("at") or None


def _merge(workbook: Workbook, op: dict[str, Any]) -> None:
    sheet = _sheet(workbook, op["sheet"])
    if op.get("undo"):
        sheet.unmerge_cells(op["range"])
    else:
        sheet.merge_cells(op["range"])


def _chart_object(kind: str, stacked: bool, labels: bool) -> Any:
    from openpyxl.chart import AreaChart, BarChart, LineChart, PieChart, ScatterChart
    from openpyxl.chart.label import DataLabelList

    # Flags left unset read as shown in LibreOffice: series, category and value pile up.
    shown = (
        DataLabelList(
            showVal=True,
            showSerName=False,
            showCatName=False,
            showLegendKey=False,
            showPercent=False,
        )
        if labels
        else None
    )
    if kind in ("column", "bar"):
        return BarChart(
            barDir="col" if kind == "column" else "bar",
            grouping="stacked" if stacked else "clustered",
            overlap=100 if stacked else None,
            dLbls=shown,
        )
    kinds = {"line": LineChart, "pie": PieChart, "area": AreaChart, "scatter": ScatterChart}
    return kinds[kind](dLbls=shown)


def _chart_series(chart: Any, source: Worksheet, ref: str, kind: str) -> None:
    from openpyxl.chart import Reference, Series

    min_col, min_row, max_col, max_row = _bounds(ref)
    if max_col <= min_col or max_row <= min_row:
        raise ValueError(f"A chart's source {ref!r} needs a header row, categories and a series")
    categories = Reference(source, min_col=min_col, min_row=min_row + 1, max_row=max_row)
    if kind == "scatter":
        for col in range(min_col + 1, max_col + 1):
            values = Reference(source, min_col=col, min_row=min_row, max_row=max_row)
            chart.series.append(Series(values, categories, title_from_data=True))
        return
    data = Reference(source, min_col=min_col + 1, min_row=min_row, max_col=max_col, max_row=max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)


def _chart(workbook: Workbook, op: dict[str, Any]) -> None:
    sheet = _sheet(workbook, op["sheet"])
    source_name, ref = _split_ref(op["source"], op["sheet"])
    chart = _chart_object(op["kind"], bool(op.get("stacked")), bool(op.get("labels")))
    _chart_series(chart, _sheet(workbook, source_name), ref, op["kind"])
    chart.title = op.get("title") or None
    if op["kind"] != "pie":
        chart.x_axis.title = op.get("x_title") or None
        chart.y_axis.title = op.get("y_title") or None
        chart.x_axis.delete = chart.y_axis.delete = False
    # A series left to the theme's colour is drawn invisibly once LibreOffice has resaved the
    # workbook (a line chart rendered with no line), so every series gets a colour.
    if op["kind"] != "pie":
        colors = [*op.get("colors", []), *_PALETTE]
        for series, color in zip(chart.series, colors, strict=False):
            series.graphicalProperties.solidFill = color
            series.graphicalProperties.line.solidFill = color
    width, height = op.get("size", [16, 8])
    chart.width, chart.height = float(width), float(height)
    sheet.add_chart(chart, op["at"])


def _scale(colors: list[str]) -> Any:
    from openpyxl.formatting.rule import ColorScale, FormatObject
    from openpyxl.styles.colors import Color

    kinds = [FormatObject(type="min")]
    if len(colors) == 3:
        kinds.append(FormatObject(type="percentile", val=50))
    kinds.append(FormatObject(type="max"))
    return ColorScale(cfvo=kinds, color=[Color(rgb=c) for c in colors])


def _conditional_rule(op: dict[str, Any]) -> Any:
    from openpyxl.formatting.rule import DataBar, FormatObject, Rule
    from openpyxl.styles import Font, PatternFill
    from openpyxl.styles.colors import Color
    from openpyxl.styles.differential import DifferentialStyle

    fill = PatternFill("solid", bgColor=op["fill"]) if op.get("fill") else None
    font = Font(color=op["color"]) if op.get("color") else None
    style = DifferentialStyle(fill=fill, font=font)
    rule = op["rule"]
    if rule == "color_scale":
        colors = (op.get("colors") or ["F8696B", "FFEB84", "63BE7B"])[:3]
        return Rule(type="colorScale", colorScale=_scale(colors))
    if rule == "data_bar":
        color = Color(rgb=(op.get("colors") or ["638EC6"])[0])
        bar = DataBar(cfvo=[FormatObject(type="min"), FormatObject(type="max")], color=color)
        return Rule(type="dataBar", dataBar=bar)
    if rule == "cell":
        values = [str(v) for v in op.get("values", [])]
        operator = op.get("operator", "greaterThan")
        return Rule(type="cellIs", operator=operator, formula=values, dxf=style)
    if rule == "formula":
        return Rule(type="expression", formula=[str(op["formula"]).lstrip("=")], dxf=style)
    return Rule(type="top10", rank=int(op.get("rank", 10)), dxf=style)


def _conditional(workbook: Workbook, op: dict[str, Any]) -> None:
    _sheet(workbook, op["sheet"]).conditional_formatting.add(op["range"], _conditional_rule(op))


def _validation(workbook: Workbook, op: dict[str, Any]) -> None:
    from openpyxl.worksheet.datavalidation import DataValidation

    if "list" in op:
        formula = '"' + ",".join(str(v) for v in op["list"]) + '"'
    else:
        formula = str(op["source"]) if str(op["source"]).startswith("=") else "=" + op["source"]
    message = str(op["message"]) if op.get("message") else None
    validation = DataValidation(
        type="list",
        formula1=formula,
        allow_blank=True,
        error=message,
        showErrorMessage=message is not None,
    )
    validation.add(op["range"])
    _sheet(workbook, op["sheet"]).add_data_validation(validation)


def _print(workbook: Workbook, op: dict[str, Any]) -> None:
    sheet = _sheet(workbook, op["sheet"])
    if "landscape" in op:
        sheet.page_setup.orientation = "landscape" if op["landscape"] else "portrait"
    if op.get("fit_width"):
        from openpyxl.worksheet.properties import PageSetupProperties

        sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = int(op["fit_width"]), 0
    if op.get("repeat_rows"):
        sheet.print_title_rows = str(op["repeat_rows"])
    if op.get("area"):
        sheet.print_area = str(op["area"])


def _name(workbook: Workbook, op: dict[str, Any]) -> None:
    from openpyxl.workbook.defined_name import DefinedName

    ref = str(op["ref"]).lstrip("=")
    workbook.defined_names[str(op["name"])] = DefinedName(str(op["name"]), attr_text=ref)


APPLY: dict[str, Callable[[Workbook, dict[str, Any]], None]] = {
    "add_sheet": _add_sheet,
    "rename_sheet": _rename_sheet,
    "delete_sheet": _delete_sheet,
    "tab_color": _tab_color,
    "cols": _cols,
    "rows": _rows,
    "freeze": _freeze,
    "merge": _merge,
    "chart": _chart,
    "conditional": _conditional,
    "validation": _validation,
    "print": _print,
    "name": _name,
}


def apply_ops(workbook: Workbook, ops: list[dict[str, Any]]) -> None:
    from openpyxl.utils.exceptions import CellCoordinatesException

    for index, op in enumerate(ops):
        try:
            APPLY[op["op"]](workbook, op)
        except (ValueError, KeyError, TypeError, CellCoordinatesException) as exc:
            raise ValueError(f"ops[{index}] ({op['op']}): {exc}") from None
