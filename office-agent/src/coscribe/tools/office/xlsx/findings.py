"""The checker run after every block write, on the recalculated copy: what a reader would see
wrong, by cell, so the model can fix it in the next call. Only findings the model can act on
are reported."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from .column_layout import column_width, explicit_width_starts, sheet_default_width, sheet_spans
from .column_widths import needed_width

_LIMIT = 25
# A column of width w shows about 7w + 5 pixels; a default row is 20 pixels.
_PX_PER_CM = 96 / 2.54
_ROW_PX = 20


@dataclass
class Written:
    """What a write touched, for the checks that look only at it."""

    cells: dict[str, list[tuple[str, bool]]] = field(default_factory=dict)
    sheets_with_content: set[str] = field(default_factory=set)

    def add(self, sheet: str, coordinate: str, looked: bool) -> None:
        self.cells.setdefault(sheet, []).append((coordinate, looked))


def _is_formula(value: Any) -> bool:
    return isinstance(value, str) and value.startswith("=")


def _is_number(value: Any) -> bool:
    return isinstance(value, int | float) and not isinstance(value, bool)


def _too_narrow(sheet: Any, shown: Any, path: Path) -> list[str]:
    """Numbers shown as `###` and text cut off by a filled cell to its right."""
    from openpyxl.utils import get_column_letter

    spans = sheet_spans(sheet, explicit_width_starts(path, sheet.title))
    default = sheet_default_width(sheet)
    merged = {(r, c) for rng in sheet.merged_cells.ranges for r, c in rng.cells}
    found = []
    for row in sheet.iter_rows():
        for cell in row:
            value = shown.cell(cell.row, cell.column).value
            if value in (None, "") or (cell.row, cell.column) in merged:
                continue
            width = column_width(spans, cell.column, default)
            need = needed_width(value, cell.number_format)
            if need <= width or cell.alignment.wrap_text:
                continue
            where = f"{sheet.title}!{cell.coordinate}"
            if not isinstance(value, str):
                found.append(f"{where} prints as ### (needs width {need:.0f}, has {width:.0f})")
            elif sheet.cell(cell.row, cell.column + 1).value not in (None, ""):
                letter = get_column_letter(cell.column)
                found.append(f"{where} text is cut off; widen column {letter} or wrap it")
    return found


def _pattern_breaks(sheet: Any) -> list[str]:
    """A typed number between formulas in its row or column: a total pasted where the rest
    are computed, which goes stale when an input changes."""
    found = []
    for row in sheet.iter_rows():
        for cell in row:
            if not _is_number(cell.value):
                continue
            r, c = cell.row, cell.column
            up, down = sheet.cell(r - 1, c).value if r > 1 else None, sheet.cell(r + 1, c).value
            left = sheet.cell(r, c - 1).value if c > 1 else None
            right = sheet.cell(r, c + 1).value
            if (_is_formula(up) and _is_formula(down)) or (
                _is_formula(left) and _is_formula(right)
            ):
                found.append(f"{sheet.title}!{cell.coordinate} is a typed number between formulas")
    return found


def _chart_cells(sheet: Any, chart: Any, spans: list[Any], default: float) -> tuple[int, ...]:
    """The rows and columns a chart covers: (first row, last row, first col, last col)."""
    anchor = chart.anchor
    start = getattr(anchor, "_from", None)
    if start is None:
        return ()
    if getattr(anchor, "to", None) is not None:
        return start.row + 1, anchor.to.row + 1, start.col + 1, anchor.to.col + 1
    width_px, height_px = chart.width * _PX_PER_CM, chart.height * _PX_PER_CM
    col, used = start.col + 1, 0.0
    while used < width_px:
        used += 7 * column_width(spans, col, default) + 5
        col += 1
    rows = int(height_px // _ROW_PX) + 1
    return start.row + 1, start.row + rows, start.col + 1, col - 1


def _charts_over_cells(sheet: Any, path: Path) -> list[str]:
    spans = sheet_spans(sheet, explicit_width_starts(path, sheet.title))
    default = sheet_default_width(sheet)
    found = []
    for chart in getattr(sheet, "_charts", []):
        cells = _chart_cells(sheet, chart, spans, default)
        if not cells:
            continue
        top, bottom, first, last = cells
        covered = [
            cell.coordinate
            for row in sheet.iter_rows(min_row=top, max_row=bottom, min_col=first, max_col=last)
            for cell in row
            if cell.value not in (None, "")
        ]
        if covered:
            title = f" covers {len(covered)} filled cells ({', '.join(covered[:3])}...)"
            found.append(f"{sheet.title}: a chart{title}; anchor it beside the table")
    return found


def _look(cell: Any) -> tuple[Any, ...]:
    return (cell.number_format, bool(cell.font.bold), cell.fill.fgColor.rgb if cell.fill else None)


def _header_row(sheet: Any, column: int, headers: dict[int, int]) -> int:
    """The first filled row of a column: its header, which data below rightly looks unlike."""
    if column not in headers:
        headers[column] = next(
            (
                r
                for r in range(1, sheet.max_row + 1)
                if sheet.cell(r, column).value not in ("", None)
            ),
            0,
        )
    return headers[column]


def _unstyled_edits(workbook: Any, written: Written) -> list[str]:
    """A new cell under a filled one that did not take its look, in a sheet that had content
    before: a row added to a user's table that looks unlike the rows above it."""
    found = []
    for name, cells in written.cells.items():
        if name not in written.sheets_with_content or name not in workbook.sheetnames:
            continue
        sheet = workbook[name]
        headers: dict[int, int] = {}
        for coordinate, looked in cells:
            cell = sheet[coordinate]
            if looked or cell.value in (None, "") or cell.row == 1:
                continue
            above = sheet.cell(cell.row - 1, cell.column)
            if above.value in (None, "") or above.row == _header_row(sheet, cell.column, headers):
                continue
            if _look(above) != _look(cell):
                found.append(
                    f"{name}!{coordinate} does not look like {above.coordinate} above it; "
                    "give the block `like_row` or `like`"
                )
    return found


def check(workbook: Any, shown: Any, path: Path, written: Written) -> list[str]:
    """`workbook` holds the formulas, `shown` the recalculated values of the same file."""
    found = []
    for name in written.cells:
        if name not in workbook.sheetnames:
            continue
        sheet = workbook[name]
        found += _too_narrow(sheet, shown[name], path)
        found += _pattern_breaks(sheet)
        found += _charts_over_cells(sheet, path)
    found += _unstyled_edits(workbook, written)
    return found[:_LIMIT]
