"""A slide of a deck built on the user's template, as the Slides panel draws it: the template's
background, master and layout shapes and what the page took over from a template slide lie under
the page, so the panel shows what PowerPoint shows."""

from __future__ import annotations

import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches

from coscribe.tools.office.pptx import slide_svg
from coscribe.tools.office.pptx.engine import build_deck

from .test_deck_engine_template import _page, _template

_TABLE = (
    '<g id="risks" data-pptx-replace-with="table"><metadata type="application/json">'
    '{"x": 64, "y": 300, "width": 560, "height": 120, "schema": "ppt-master.semantic-table.v2",'
    ' "columns": ["风险", "负责人"], "rows": [["范围蔓延", "张明"]]}'
    "</metadata></g>"
)


def _build(tmp_path: Path, *pages: str) -> Path:
    folder = tmp_path / "pages"
    folder.mkdir()
    for i, page in enumerate(pages, start=1):
        (folder / f"{i:02d}.svg").write_text(page, encoding="utf-8")
    report = build_deck(folder, tmp_path / "deck.pptx", template=tmp_path / "t.pptx")
    assert report.ok, [e for p in report.pages for e in p.errors]
    return tmp_path / "deck.pptx"


def test_the_template_slides_decoration_lies_under_the_page(tmp_path: Path) -> None:
    _template(tmp_path / "t.pptx")
    deck = _build(tmp_path, _page('data-template-slide="1"'))

    svg = slide_svg(deck, 0, tmp_path)

    panel = svg.find('fill="#1E2761"')
    logo = svg.find("data:image/png;base64,")
    title = svg.find("第三季度汇报")
    assert 0 < panel < title
    assert 0 < logo < title


def test_the_masters_shapes_and_background_are_drawn(tmp_path: Path) -> None:
    _template(tmp_path / "t.pptx")
    prs = Presentation(str(tmp_path / "t.pptx"))
    master = prs.slide_masters[0]
    # python-pptx adds shapes to slides only: draw the stripe on a slide and move it.
    stripe = prs.slides[0].shapes.add_shape(1, Emu(12000000), 0, Emu(192000), prs.slide_height)
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = RGBColor(0x00, 0x8E, 0xCF)
    master.shapes._spTree.append(stripe.element)
    master.background.fill.solid()
    master.background.fill.fore_color.rgb = RGBColor(0xF4, 0xF4, 0xF4)
    prs.save(str(tmp_path / "t.pptx"))
    deck = _build(tmp_path, _page('data-layout="Blank"'))

    svg = slide_svg(deck, 0, tmp_path)

    assert re.search(r'<rect x="0" y="0" width="1280\.00" height="720\.00" fill="#F4F4F4"', svg)
    assert 'fill="#008ECF"' in svg


def test_a_table_taken_over_from_a_template_slide_leaves_the_pages_own_table_its_drawing(
    tmp_path: Path,
) -> None:
    _template(tmp_path / "t.pptx")
    prs = Presentation(str(tmp_path / "t.pptx"))
    history = prs.slides[1].shapes.add_table(2, 2, Inches(1), Inches(5), Inches(4), Inches(1))
    history.table.cell(0, 0).text = "版本"
    prs.save(str(tmp_path / "t.pptx"))
    deck = _build(tmp_path, _page('data-template-slide="2"', _TABLE))

    svg = slide_svg(deck, 0, tmp_path)

    page_table = svg[svg.find('id="risks"') :]
    assert "范围蔓延" in page_table
    assert "版本" not in page_table


def test_a_deck_without_a_template_gets_only_the_white_slide_under_it(tmp_path: Path) -> None:
    folder = tmp_path / "pages"
    folder.mkdir()
    (folder / "01.svg").write_text(_page(), encoding="utf-8")
    assert build_deck(folder, tmp_path / "deck.pptx").ok

    svg = slide_svg(tmp_path / "deck.pptx", 0, tmp_path)

    beneath = svg.split('<text id="title"', 1)[0]
    assert beneath.endswith(
        '<g><rect x="0" y="0" width="1280.00" height="720.00" fill="#FFFFFF" /></g>'
    )


def test_custom_paths_gradients_and_rotation_are_drawn(tmp_path: Path) -> None:
    _template(tmp_path / "t.pptx")
    prs = Presentation(str(tmp_path / "t.pptx"))
    cover = prs.slides[0]
    width = prs.slide_width
    builder = cover.shapes.build_freeform(Emu(width // 2), 0, scale=1.0)
    builder.add_line_segments([(width, 0), (width, Emu(3000000))])
    wedge = builder.convert_to_shape()
    wedge.fill.gradient()
    wedge.fill.gradient_stops[0].color.rgb = RGBColor(0x0B, 0x4F, 0x4C)
    wedge.fill.gradient_stops[1].color.rgb = RGBColor(0x1A, 0x9C, 0x98)
    oval = cover.shapes.add_shape(9, Inches(9), Inches(4), Inches(2), Inches(2))
    oval.rotation = 20
    prs.save(str(tmp_path / "t.pptx"))
    deck = _build(tmp_path, _page('data-template-slide="1"'))

    svg = slide_svg(deck, 0, tmp_path)

    assert '<path d="M640.00,0.00 L1280.00,0.00 L1280.00,314.96 Z"' in svg
    assert 'stop-color="#0B4F4C"' in svg and 'fill="url(#tl-grad-1)"' in svg
    assert re.search(r'transform="rotate\(20\.00 [\d.]+ [\d.]+\)"><ellipse', svg)
