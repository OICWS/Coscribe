"""`edit_pptx_text(find=..., replace=...)`: a phrase replaced in any shape (group members, table
cells, placeholders) on one slide or all, each character that stays keeping its run's format."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches

from coscribe.tools.presentations import build_presentation_tools


@pytest.fixture(autouse=True)
def _no_libreoffice(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr("shutil.which", lambda name: None)


def _edit(root: Path, **kwargs: Any) -> dict[str, Any]:
    tools = {tool.__name__: tool for tool in build_presentation_tools(root)}  # type: ignore[attr-defined]
    return tools["edit_pptx_text"](path="deck.pptx", **kwargs)  # type: ignore[no-any-return]


def _deck(root: Path) -> None:
    """Slide 1: a title, a group holding "售价 399 元" ("399" bold and orange, " 元" plain).
    Slide 2: a table whose cell says "399 元", and a text box saying "首发 399 元"."""
    prs = Presentation()
    first = prs.slides.add_slide(prs.slide_layouts[5])
    first.shapes.title.text = "云朵 Pro 399 元"
    group = first.shapes.add_group_shape()
    paragraph = group.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1)).text_frame
    label, price, unit = (paragraph.paragraphs[0].add_run() for _ in range(3))
    label.text, price.text, unit.text = "售价 ", "399", " 元"
    price.font.bold = True
    price.font.color.rgb = RGBColor(0xC2, 0x41, 0x0C)
    second = prs.slides.add_slide(prs.slide_layouts[6])
    table = second.shapes.add_table(1, 2, Inches(1), Inches(1), Inches(4), Inches(1)).table
    table.cell(0, 1).text = "399 元"
    second.shapes.add_textbox(Inches(1), Inches(3), Inches(4), Inches(1)).text = "首发 399 元"
    prs.save(str(root / "deck.pptx"))


def _runs(root: Path) -> list[Any]:
    group = list(Presentation(str(root / "deck.pptx")).slides[0].shapes)[1]
    return list(group.shapes[0].text_frame.paragraphs[0].runs)


def test_a_phrase_in_a_group_member_is_replaced_keeping_its_format(tmp_path: Path) -> None:
    _deck(tmp_path)

    result = _edit(tmp_path, slide=1, find="399", replace="349")

    label, price, unit = _runs(tmp_path)
    assert (label.text, price.text, unit.text) == ("售价 ", "349", " 元")
    assert price.font.bold and str(price.font.color.rgb) == "C2410C"
    assert unit.font.bold is None
    assert {(r["shape"], r["count"]) for r in result["replaced"]} == {("0", 1), ("1/0", 1)}


def test_a_phrase_across_runs_takes_the_format_it_starts_in(tmp_path: Path) -> None:
    _deck(tmp_path)

    _edit(tmp_path, slide=1, find="售价 399", replace="现价 349")

    label, price, unit = _runs(tmp_path)
    assert (label.text, price.text, unit.text) == ("现价 349", "", " 元")
    assert label.font.bold is None


def test_every_slide_with_slide_zero_table_cells_included(tmp_path: Path) -> None:
    _deck(tmp_path)

    result = _edit(tmp_path, slide=0, find="399 元", replace="349 元")

    second = Presentation(str(tmp_path / "deck.pptx")).slides[1]
    table, box = list(second.shapes)
    assert table.table.cell(0, 1).text == "349 元"
    assert box.text_frame.text == "首发 349 元"
    assert sorted(r["slide"] for r in result["replaced"]) == [1, 1, 2, 2]
    # "399" and " 元" are separate runs in the group, so they still match together.
    assert "".join(run.text for run in _runs(tmp_path)) == "售价 349 元"


def test_remove_deletes_the_phrase(tmp_path: Path) -> None:
    _deck(tmp_path)

    _edit(tmp_path, slide=2, find="首发 ", remove=True)

    assert list(Presentation(str(tmp_path / "deck.pptx")).slides[1].shapes)[1].text == "399 元"


def test_a_forgotten_replacement_deletes_nothing(tmp_path: Path) -> None:
    _deck(tmp_path)

    with pytest.raises(ValueError, match="remove=true"):
        _edit(tmp_path, slide=0, find="399")

    assert "".join(run.text for run in _runs(tmp_path)) == "售价 399 元"


def test_a_line_break_in_the_replacement_is_refused(tmp_path: Path) -> None:
    _deck(tmp_path)

    with pytest.raises(ValueError, match="content="):
        _edit(tmp_path, slide=1, find="399", replace="349\n限时")


def _broken_line(root: Path) -> None:
    """One paragraph: "售价", a line break, "399 元"."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.add_run().text = "售价"
    paragraph.add_line_break()
    paragraph.add_run().text = "399 元"
    prs.save(str(root / "deck.pptx"))


def test_a_phrase_never_matches_across_a_line_break(tmp_path: Path) -> None:
    _broken_line(tmp_path)

    with pytest.raises(ValueError, match="line break"):
        _edit(tmp_path, slide=1, find="售价399", replace="现价349")

    _edit(tmp_path, slide=1, find="399", replace="349")

    shape = list(Presentation(str(tmp_path / "deck.pptx")).slides[0].shapes)[0]
    assert shape.text_frame.text == "售价\v349 元"


def test_a_phrase_that_is_not_there_says_how_to_look(tmp_path: Path) -> None:
    _deck(tmp_path)

    with pytest.raises(ValueError, match="'499' is not on any slide. read_pptx"):
        _edit(tmp_path, slide=0, find="499", replace="1")


@pytest.mark.parametrize(
    "kwargs",
    [
        {"slide": 1, "find": "399", "title": "新标题"},
        {"slide": 1},
        {"slide": 0, "title": "新标题"},
    ],
)
def test_find_is_one_way_and_slide_zero_needs_it(tmp_path: Path, kwargs: dict[str, Any]) -> None:
    _deck(tmp_path)

    with pytest.raises(ValueError):
        _edit(tmp_path, **kwargs)
