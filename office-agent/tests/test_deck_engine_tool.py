"""`build_pptx_from_pages`, the deck engine's tool, and the instructions that send new decks to it.

The tool's description is the page contract the model writes to, so a page using everything it
allows must compile without an error; the rest pins the tool's path rules and its refusal.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
from PIL import Image
from pptx import Presentation

from coscribe.tools import NEW_DECK_INSTRUCTIONS
from coscribe.tools.office.pptx.engine.tool import build_engine_tools
from coscribe.tools.skills import load_builtin_skills

_SVG = (
    'xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1280 720" font-family="Microsoft YaHei, Arial"'
)

# One of everything the tool's description allows.
EVERYTHING = f"""<svg {_SVG}>
  <defs>
    <linearGradient id="band" x1="0" y1="0" x2="1" y2="0">
      <stop offset="0%" stop-color="#1A73E8"/><stop offset="100%" stop-color="#0D47A1"/>
    </linearGradient>
    <radialGradient id="glow" cx="0.5" cy="0.5" r="0.5">
      <stop offset="0%" stop-color="#FFFFFF" stop-opacity="0.4"/>
      <stop offset="100%" stop-color="#FFFFFF" stop-opacity="0"/>
    </radialGradient>
    <filter id="soft" x="-15%" y="-20%" width="130%" height="150%">
      <feDropShadow dx="0" dy="6" stdDeviation="8" flood-color="#000000" flood-opacity="0.10"/>
    </filter>
    <clipPath id="round"><circle cx="1100" cy="560" r="80"/></clipPath>
    <marker id="arrow" markerWidth="10" markerHeight="10" refX="9" refY="5" orient="auto">
      <path d="M0,0 L10,5 L0,10 Z" fill="#5B6B7B"/>
    </marker>
  </defs>
  <rect x="0" y="0" width="1280" height="720" fill="#F7F9FC"/>
  <text id="title" x="64" y="90" font-size="40" font-weight="bold"
    fill="#1E2A36">一页全部元素</text>
  <rect x="64" y="120" width="1152" height="8" fill="url(#band)"/>
  <g id="card" transform="translate(64 160)">
    <rect x="0" y="0" width="360" height="200" rx="12" fill="#FFFFFF" filter="url(#soft)"/>
    <text x="24" y="60" font-size="20" fill="#5B6B7B">华东 <tspan font-weight="bold"
      fill="#1A73E8">+12.2%</tspan><tspan x="24" dy="30">第二行</tspan></text>
  </g>
  <circle cx="600" cy="260" r="40" fill="#1A73E8" fill-opacity="0.12"/>
  <ellipse cx="760" cy="260" rx="60" ry="30" fill="url(#glow)" stroke="#1A73E8"/>
  <line x1="480" y1="400" x2="900" y2="400" stroke="#5B6B7B" stroke-width="2"
    marker-end="url(#arrow)"/>
  <polyline points="480,500 560,460 640,480 720,430" fill="none" stroke="#0D47A1"
    stroke-opacity="0.8" stroke-width="3"/>
  <polygon points="980,200 1040,300 920,300" fill="#F2A541" opacity="0.9"/>
  <path d="M 64 600 C 200 540 300 660 440 600 Q 520 560 600 600 Z" fill="#E8F0FE"/>
  <image href="../photo.png" x="1020" y="480" width="160" height="160" clip-path="url(#round)"/>
</svg>"""


def _tool(workspace: Path, **kwargs: Any) -> Any:
    (tool,) = build_engine_tools(workspace, **kwargs)
    return tool


def _write(folder: Path, name: str, svg: str) -> None:
    folder.mkdir(parents=True, exist_ok=True)
    (folder / name).write_text(svg, encoding="utf-8")


def _title_page(title: str) -> str:
    title_text = f'<text id="title" x="64" y="90" font-size="40" fill="#1E2A36">{title}</text>'
    return f"<svg {_SVG}>{title_text}</svg>"


def test_a_page_using_everything_the_contract_allows_compiles_cleanly(tmp_path: Path) -> None:
    Image.new("RGB", (160, 160), "#C00000").save(tmp_path / "photo.png")
    _write(tmp_path / "deck-pages", "01-all.svg", EVERYTHING)

    result = _tool(tmp_path)("deck-pages", "deck.pptx")

    assert result["ok"], result
    assert "errors" not in result["pages"][0]
    assert result["path"] == "deck.pptx"
    slide = Presentation(str(tmp_path / "deck.pptx")).slides[0]
    assert slide.shapes.title.text == "一页全部元素"
    assert "card" in {shape.name for shape in slide.shapes}


def test_sources_feed_the_number_check(tmp_path: Path) -> None:
    _write(
        tmp_path / "pages",
        "1.svg",
        f'<svg {_SVG}><text id="title" x="64" y="90" font-size="40" fill="#1E2A36">'
        "营收 5,410 与 9,999</text></svg>",
    )
    (tmp_path / "notes.md").write_text("华东 5410 万元", encoding="utf-8")

    result = _tool(tmp_path)("pages", "deck.pptx", sources=["notes.md"])

    (warning,) = result["pages"][0]["warnings"]
    assert "9999" in warning and "5410" not in warning


@pytest.mark.parametrize(
    ("pages_dir", "path", "sources", "message"),
    [
        ("missing", "deck.pptx", None, "Not a folder"),
        ("pages", "deck.ppt", None, "must end in .pptx"),
        ("pages", "deck.pptx", ["nowhere.docx"], "does not exist"),
    ],
)
def test_bad_arguments_are_refused_before_anything_is_written(
    tmp_path: Path, pages_dir: str, path: str, sources: list[str] | None, message: str
) -> None:
    _write(tmp_path / "pages", "1.svg", _title_page("一"))

    with pytest.raises(ValueError, match=message):
        _tool(tmp_path)(pages_dir, path, sources=sources)
    assert not (tmp_path / "deck.pptx").exists()


def test_pages_or_a_deck_outside_the_workspace_are_refused(tmp_path: Path) -> None:
    workspace = tmp_path / "work"
    _write(tmp_path / "elsewhere", "1.svg", _title_page("一"))
    _write(workspace / "pages", "1.svg", _title_page("一"))
    tool = _tool(workspace)

    with pytest.raises(PermissionError):
        tool(str(tmp_path / "elsewhere"), "deck.pptx")
    with pytest.raises(PermissionError):
        tool("pages", str(tmp_path / "deck.pptx"))


def test_pages_in_a_trusted_folder_take_images_from_that_folder(tmp_path: Path) -> None:
    trusted = tmp_path / "trusted"
    trusted.mkdir()
    Image.new("RGB", (40, 20), "#C00000").save(trusted / "photo.png")
    _write(
        trusted / "pages",
        "1.svg",
        f'<svg {_SVG}><image href="../photo.png" x="10" y="200" width="400" height="200"/></svg>',
    )

    result = _tool(tmp_path / "work", extra_writable=[trusted])(
        str(trusted / "pages"), str(trusted / "deck.pptx")
    )

    assert result["ok"], result


def test_a_deck_changed_after_its_build_is_refused_with_the_way_out(tmp_path: Path) -> None:
    _write(tmp_path / "pages", "1.svg", _title_page("一"))
    tool = _tool(tmp_path)
    tool("pages", "deck.pptx")
    prs = Presentation(str(tmp_path / "deck.pptx"))
    prs.slides[0].shapes.title.text = "用户改过"
    prs.save(str(tmp_path / "deck.pptx"))

    with pytest.raises(ValueError, match="edit_pptx_"):
        tool("pages", "deck.pptx")
    assert Presentation(str(tmp_path / "deck.pptx")).slides[0].shapes.title.text == "用户改过"

    assert tool("pages", "deck.pptx", replace_changes=True)["ok"]


def test_the_instructions_send_new_decks_to_the_skill_and_the_tool_that_exist() -> None:
    (skill,) = [s for s in load_builtin_skills() if "build_pptx_from_pages" in s.description]

    assert f'load_skill("{skill.name}")' in NEW_DECK_INSTRUCTIONS
    assert build_engine_tools(Path("."))[0].__name__ in NEW_DECK_INSTRUCTIONS


def test_a_scaled_group_with_text_is_reported_as_the_contract_says(tmp_path: Path) -> None:
    _write(
        tmp_path / "pages",
        "1.svg",
        f'<svg {_SVG}><g id="card" transform="translate(64 160) scale(2)">'
        '<text x="0" y="40" font-size="20" fill="#1E2A36">华东</text></g></svg>',
    )

    result = _tool(tmp_path)("pages", "deck.pptx")

    assert not result["ok"]
    assert any("card" in e for e in result["pages"][0]["errors"])
