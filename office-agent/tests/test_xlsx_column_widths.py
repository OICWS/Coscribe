"""`write_xlsx_cells` widens the columns it writes to what they show, keeps a width that was
already set, and `column_widths` measures a cell the way Excel draws it."""

from __future__ import annotations

import datetime as dt
import functools
import re
import shutil
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from coscribe.tools.office.xlsx import build_spreadsheet_tools
from coscribe.tools.office.xlsx.column_layout import _patch_cols, widen_for_values
from coscribe.tools.office.xlsx.column_widths import (
    DEFAULT_WIDTH,
    MAX_WIDTH,
    displayed_text,
    text_width,
)
from coscribe.tools.office.xlsx.helpers import _coerce_cell
from evals.xlsx_checks import clipped_texts

FILE = "book.xlsx"


def _tools(root: Path) -> dict[str, object]:
    return {tool.__name__: tool for tool in build_spreadsheet_tools(root)}  # type: ignore[attr-defined]


def _write(root: Path, content: str, sheet_name: str = "Orders") -> None:
    """A pipe table written over the sheet from A1."""
    _tools(root)["write_xlsx_cells"](
        path=FILE, blocks=[{"sheet": sheet_name, "rows": _rows(content)}], overwrite=True
    )


def _rows(content: str) -> list[list[object]]:
    return [
        [_coerce_cell(c.strip()) if c.strip() else None for c in line.strip().strip("|").split("|")]
        for line in content.strip().splitlines()
        if not set(line.strip()) <= set("|- :")
    ]


def _edit(root: Path, *, path: str, sheet_name: str, start_cell: str, content: str) -> None:
    _tools(root)["write_xlsx_cells"](
        path=path,
        blocks=[{"sheet": sheet_name, "at": start_cell, "rows": _rows(content)}],
        overwrite=True,
    )


def _set_width(root: Path, *, path: str, sheet_name: str, columns: str, width: float) -> None:
    _tools(root)["modify_xlsx"](
        path=path, ops=[{"op": "cols", "sheet": sheet_name, "cols": columns, "width": width}]
    )


def _merge(root: Path, *, path: str, sheet_name: str, cell_range: str) -> None:
    _tools(root)["modify_xlsx"](
        path=path, ops=[{"op": "merge", "sheet": sheet_name, "range": cell_range}]
    )


def _width(root: Path, letter: str, sheet: str = "Orders") -> float:
    return load_workbook(root / FILE)[sheet].column_dimensions[letter].width


def _cols(root: Path) -> list[dict[str, str]]:
    """The `<col>` elements of the first sheet as the file has them."""
    with zipfile.ZipFile(root / FILE) as archive:
        xml = archive.read("xl/worksheets/sheet1.xml").decode()
    return [dict(re.findall(r'(\w+)="([^"]*)"', col)) for col in re.findall(r"<col [^>]*>", xml)]


def _book_with_cols(root: Path, cols: str) -> None:
    """A sheet Orders whose `<cols>` is exactly `cols`, as a person's file can have it and
    openpyxl cannot write it."""
    book = Workbook()
    assert book.active is not None
    book.active.title = "Orders"
    book.save(root / FILE)
    with zipfile.ZipFile(root / FILE) as archive:
        entries = [(info, archive.read(info.filename)) for info in archive.infolist()]
    with zipfile.ZipFile(root / FILE, "w") as archive:
        for info, data in entries:
            if info.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b"<sheetData", f"<cols>{cols}</cols><sheetData".encode())
            archive.writestr(info, data)


def test_a_label_wider_than_the_default_column_is_no_longer_cut_off(tmp_path: Path) -> None:
    _write(tmp_path, "| SKU | Unit price | Open tickets |\n| --- | --- | --- |\n| A1 | 3.5 | 12 |")

    book = load_workbook(tmp_path / FILE)
    assert clipped_texts(book, book) == []
    assert _width(tmp_path, "C") >= text_width("Open tickets")


@pytest.mark.parametrize(
    ("value", "fmt", "shown"),
    [
        (12850000, '#,##0"元"', "12,850,000元"),
        (0.3425, "0.0%", "34.2%"),
        (-1234.5, "$#,##0;($#,##0)", "($1,234)"),
        (1285.0, '#,##0.00"万"', "1,285.00万"),
        (123456789012, "General", "12345678901"),
        (12345678, "0.00E+00", "1.23E+07"),
        (1234567, '#,##0,"K"', "1,235K"),
        (-5, '0" a;b";(0)', "(5)"),
        (5, '0" a;b";(0)', "5 a;b"),
        (-5, "0\\;;(0)", "(5)"),
    ],
)
def test_a_number_is_measured_with_its_number_format(value: float, fmt: str, shown: str) -> None:
    assert displayed_text(value, fmt) == shown


def test_a_format_that_is_not_rendered_is_measured_by_its_own_length() -> None:
    assert text_width(displayed_text(0.5, "# ??/??")) >= len("# ??/??")
    assert text_width(displayed_text(46000, "yyyy-mm-dd")) >= len("2025-12-10")


def test_a_cjk_character_counts_twice_and_one_long_note_is_capped(tmp_path: Path) -> None:
    _write(tmp_path, f"| 销售额（万元） | abcdefghijklmn | {'long note ' * 20} |\n| 1 | 2 | 3 |")

    assert text_width("销售额（万元）") == 14
    assert _width(tmp_path, "A") == _width(tmp_path, "B") >= 14
    assert _width(tmp_path, "C") == MAX_WIDTH


def test_a_label_widens_its_column_only_when_it_cannot_run_on(tmp_path: Path) -> None:
    title = "Quarterly sales report for all regions"
    _write(tmp_path, f"| {title} | | | | |\n| Region | Q1 | Q2 | Q3 | Q4 |")
    assert _cols(tmp_path) == []

    label = "Sales for all regions together"
    _write(tmp_path, f"| {label} | | x |\n| Region | Q1 | Q2 |")
    assert _width(tmp_path, "A") >= text_width(label)


def test_a_label_alone_in_its_row_widens_when_nothing_to_its_right_holds_it(
    tmp_path: Path,
) -> None:
    name = "Northwind Traders International"
    _write(tmp_path, f"| Customer |\n| {name} |")
    assert _width(tmp_path, "A") == text_width(name) + 2

    _write(tmp_path, f"| Customer | Notes |\n| {name} | |", sheet_name="Notes")
    assert _width(tmp_path, "A", sheet="Notes") == text_width(name) + 2


def test_writing_cells_widens_text_longer_than_the_empty_cells_beside_it(
    tmp_path: Path,
) -> None:
    note = "Open orders: see the list below for details"
    _write(tmp_path, "| a | b |\n| 1 | 2 |")

    edit = functools.partial(_edit, tmp_path)
    edit(path=FILE, sheet_name="Orders", start_cell="A3", content=f"| {note} |")

    assert _width(tmp_path, "A") == text_width(note) + 2


def test_writing_a_sheet_again_widens_a_column_but_never_narrows_it(tmp_path: Path) -> None:
    _write(tmp_path, "| Item | Price |\n| Lamp | 1 |")
    assert _cols(tmp_path) == []

    _write(tmp_path, "| Item | Price |\n| Lamp | 123456789.5 |")
    wide = _width(tmp_path, "B")
    assert wide >= len("123456789.5")

    _write(tmp_path, "| Item | Price |\n| Lamp | 1 |")
    assert _width(tmp_path, "B") == wide


def test_a_width_set_earlier_is_kept_unless_the_content_needs_more(tmp_path: Path) -> None:
    content = "| Product | Unit price |\n| Desk lamp | 24.5 |"
    _write(tmp_path, content)
    _set_width(tmp_path, path=FILE, sheet_name="Orders", columns="A:B", width=30)
    _write(tmp_path, content)
    assert _width(tmp_path, "A") == _width(tmp_path, "B") == 30

    _set_width(tmp_path, path=FILE, sheet_name="Orders", columns="B", width=6)
    _write(tmp_path, content)
    assert 6 < _width(tmp_path, "B") and _width(tmp_path, "B") >= text_width("Unit price")


def test_a_template_range_keeps_its_width_where_the_content_fits(tmp_path: Path) -> None:
    _book_with_cols(
        tmp_path, '<col min="1" max="1" width="40"/><col min="2" max="16384" width="7"/>'
    )

    _write(tmp_path, "| Product | Unit price | x |\n| Desk lamp | 24.5 | 3 |")

    assert [(c["min"], c["max"], c["width"]) for c in _cols(tmp_path)] == [
        ("1", "1", "40"),
        ("2", "2", "12"),
        ("3", "16384", "7"),
    ]


def test_widths_after_recalculation_widen_the_saved_file_and_never_narrow_it(
    tmp_path: Path,
) -> None:
    _book_with_cols(tmp_path, '<col min="3" max="4" width="9" customWidth="1"/>')

    _patch_cols(tmp_path / FILE, "Orders", {2: 20.0, 3: 5.0, 4: 30.0}, DEFAULT_WIDTH)

    assert [(c["min"], c["max"], c["width"]) for c in _cols(tmp_path)] == [
        ("2", "2", "20"),
        ("3", "3", "9"),
        ("4", "4", "30"),
    ]
    load_workbook(tmp_path / FILE)


@pytest.mark.real_libreoffice
@pytest.mark.skipif(shutil.which("soffice") is None, reason="needs LibreOffice")
def test_a_column_of_formulas_is_sized_from_the_computed_values(tmp_path: Path) -> None:
    _write(tmp_path, "| Item | Sum |\n| a | =123456789012/7 |")

    assert load_workbook(tmp_path / FILE, data_only=True)["Orders"]["B2"].value
    (sized,) = [c for c in _cols(tmp_path) if c["min"] == "2"]
    assert float(sized["width"]) >= len("17636684145")

    _edit(tmp_path, path=FILE, sheet_name="Orders", start_cell="C1", content="| Total |\n| =B2*1 |")
    (sized,) = [c for c in _cols(tmp_path) if c["min"] == "3"]
    assert float(sized["width"]) >= len("17636684145")


def test_writing_cells_widens_the_columns_it_wrote(tmp_path: Path) -> None:
    _book_with_cols(
        tmp_path,
        '<col min="1" max="3" width="14" customWidth="1"/><col min="7" max="7" hidden="1"/>',
    )

    _edit(
        tmp_path,
        path=FILE,
        sheet_name="Orders",
        start_cell="D1",
        content="| Product | Unit price | Line total | Notes |\n| Desk lamp | 24.9 | =E2*2 | x |",
    )

    cols = [(c["min"], c["max"], c.get("width"), c.get("hidden")) for c in _cols(tmp_path)]
    assert cols == [
        ("1", "3", "14", None),
        ("4", "4", "11", None),
        ("5", "5", "12", None),
        ("6", "6", "12", None),
        ("7", "7", None, "1"),
    ]


def test_writing_cells_leaves_merged_text_and_text_with_room_beside_it(tmp_path: Path) -> None:
    _write(tmp_path, "| a | b | c |\n| 1 | 2 | 3 |")
    _merge(tmp_path, path=FILE, sheet_name="Orders", cell_range="A3:C3")

    _edit(
        tmp_path,
        path=FILE,
        sheet_name="Orders",
        start_cell="A3",
        content="| 123456789012 |\n| Open orders: see below |",
    )

    assert _cols(tmp_path) == []


@pytest.mark.parametrize(
    ("value", "fmt", "shown"),
    [
        (1285, 'General"元"', "1285元"),
        (-3, 'General" pcs";(General)', "(3)"),
        (1234567, '#,##0.0,"K"', "1,234.6K"),
        (1234567890, '#,##0.0,,"M"', "1,234.6M"),
        (-5, "0;;0", ""),
        (5, ";;;", ""),
        (1, "0000000000", "0000000001"),
        (12345678, "000-00-0000", "012-34-5678"),
        (5, "000,000", "000,005"),
        (12345.67, "#,##0 %", "1,234,567 %"),
    ],
)
def test_general_inside_a_format_scaling_and_empty_sections(
    value: float, fmt: str, shown: str
) -> None:
    assert displayed_text(value, fmt) == shown


def test_a_date_is_measured_by_its_number_format() -> None:
    day = dt.datetime(2026, 9, 30, 14, 5)
    assert displayed_text(day, "yyyy-mm-dd") == "yyyy-mm-dd"
    assert text_width(displayed_text(day, "mmmm d, yyyy")) == len("September d, yyyy")
    assert text_width(displayed_text(day, "General")) == len("2026-09-30 14:05:00")
    assert text_width(displayed_text(dt.date(2026, 9, 30), "General")) == len("2026-09-30")


def test_a_column_set_narrower_than_the_default_is_widened_to_its_values(tmp_path: Path) -> None:
    _write(tmp_path, "| Item | Qty |\n| Lamp | 1 |")
    _set_width(tmp_path, path=FILE, sheet_name="Orders", columns="B", width=3)

    _edit(tmp_path, path=FILE, sheet_name="Orders", start_cell="B2", content="| 12345 |")

    assert _width(tmp_path, "B") == len("12345") + 2


def test_a_sheet_wide_default_width_counts_as_the_width_a_column_has(tmp_path: Path) -> None:
    book = Workbook()
    assert book.active is not None
    book.active.title = "Orders"
    book.active.sheet_format.defaultColWidth = 20
    book.save(tmp_path / FILE)
    edit = functools.partial(_edit, tmp_path)

    edit(path=FILE, sheet_name="Orders", start_cell="B2", content="| 1234567.25 |")
    edit(path=FILE, sheet_name="Orders", start_cell="A3", content=f"| {'a note ' * 4} |")

    assert _cols(tmp_path) == []


def test_measuring_a_cell_does_not_create_its_neighbours() -> None:
    sheet = Workbook().active
    assert sheet is not None
    sheet["A1"] = "a label"
    sheet["C1"] = 1

    widen_for_values(sheet, None, [sheet["A1"]], set())

    assert sorted(sheet._cells) == [(1, 1), (1, 3)]
