"""Native charts and tables on a page: a group marked `data-pptx-replace-with` holding its data
as JSON becomes a PowerPoint chart (data in an embedded workbook) or table, and the checks see
its numbers."""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from pptx import Presentation

from coscribe.tools.office.pptx.engine import check_deck, compile_pages

_SVG = (
    'xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1280 720" font-family="Microsoft YaHei, Arial"'
)
_CHART = """<g id="revenue" data-pptx-replace-with="chart"><metadata type="application/json">
{"x": 64, "y": 160, "width": 560, "height": 480, "type": "column",
 "categories": ["华东", "华南", "华北"],
 "series": [{"name": "Q3 营收（万元）", "values": [5410, 3980, 2120]}],
 "data_labels": {"number_format": "#,##0"}, "style": {"colors": ["#C2410C"]}}
</metadata></g>"""
_TABLE = """<g id="regions" data-pptx-replace-with="table"><metadata type="application/json">
{"x": 680, "y": 180, "width": 536, "height": 200, "schema": "ppt-master.semantic-table.v2",
 "columns": ["区域", "同比"], "rows": [["华东", {"text": "+12.2%", "align": "r"}]]}
</metadata></g>"""


def _page(*body: str) -> str:
    title = '<text id="title" x="64" y="100" font-size="40">华东增长最快</text>'
    return f"<svg {_SVG}>{title}{''.join(body)}</svg>"


def _shapes(path: Path) -> dict[str, object]:
    return {shape.name: shape for shape in Presentation(str(path)).slides[0].shapes}


def test_a_chart_marker_becomes_a_chart_with_its_data_in_a_workbook(tmp_path: Path) -> None:
    out = tmp_path / "deck.pptx"

    report = compile_pages([("p1", _page(_CHART))], out)

    assert report.ok, report.errors()
    chart = _shapes(out)["revenue"].chart  # type: ignore[attr-defined]
    assert list(chart.plots[0].categories) == ["华东", "华南", "华北"]
    assert list(chart.series[0].values) == [5410, 3980, 2120]
    with zipfile.ZipFile(out) as package:
        (workbook,) = [n for n in package.namelist() if n.startswith("ppt/embeddings/")]
        sheet = load_workbook(io.BytesIO(package.read(workbook))).active
    assert [cell.value for cell in sheet["B"]] == ["Q3 营收（万元）", 5410, 3980, 2120]


def test_charts_on_two_pages_each_keep_their_own_parts(tmp_path: Path) -> None:
    out = tmp_path / "deck.pptx"

    compile_pages([("p1", _page(_CHART)), ("p2", _page(_CHART.replace("5410", "6000")))], out)

    first, second = (
        next(s for s in slide.shapes if s.has_chart).chart
        for slide in Presentation(str(out)).slides
    )
    assert first.series[0].values[0] == 5410 and second.series[0].values[0] == 6000


def test_a_table_marker_becomes_a_table(tmp_path: Path) -> None:
    out = tmp_path / "deck.pptx"

    compile_pages([("p1", _page(_TABLE))], out)

    table = _shapes(out)["regions"].table  # type: ignore[attr-defined]
    assert [[cell.text for cell in row.cells] for row in table.rows] == [
        ["区域", "同比"],
        ["华东", "+12.2%"],
    ]


def test_a_broken_marker_is_an_error_naming_it_and_the_rest_is_built(tmp_path: Path) -> None:
    out = tmp_path / "deck.pptx"

    report = compile_pages([("p1", _page(_CHART.replace('"column"', '"map"'), _TABLE))], out)

    assert any(error.startswith('p1: <g id="revenue">') for error in report.errors())
    assert set(_shapes(out)) == {"title", "regions"}


def test_the_number_check_reads_the_chart_and_the_table(tmp_path: Path) -> None:
    report = compile_pages([("p1", _page(_CHART, _TABLE))], tmp_path / "deck.pptx")

    check_deck(report, sources=["华东 5410，华南 3980，同比 +12.2%"])

    assert not report.pages[0].errors, report.pages[0].errors
    (warning,) = report.pages[0].warnings
    assert "numbers no source file has: 2120" in warning
