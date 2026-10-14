import zipfile
from pathlib import Path
from typing import Any

import pytest
from lxml import etree
from openpyxl import Workbook

from coscribe.tools._xlsx_lookup_formulas import (
    _NO_MATCH,
    _NOT_GIVEN,
    _SML_NS,
    _UNRESOLVABLE,
    FormulaError,
    _as_int_arg,
    _excel_eq,
    _format_cell_value,
    _is_one_argument_list,
    _nearest_index,
    _parse_literal,
    _patch_cell,
    _resolve_arg,
    _resolve_reference,
    _Sentinel,
    _sheet_archive_paths,
    _split_top_level_args,
    _wildcard_match,
    evaluate_and_patch_lookup_formulas,
    evaluate_formula,
    evaluate_xlookup,
    evaluate_xmatch,
    parse_single_call,
    xmatch_index,
)


def test_parse_single_call_recognizes_a_plain_call() -> None:
    assert parse_single_call('=XLOOKUP("Bob",A2:A3,B2:B3)') == (
        "XLOOKUP",
        ['"Bob"', "A2:A3", "B2:B3"],
    )


def test_parse_single_call_strips_xlfn_prefix_case_insensitively() -> None:
    assert parse_single_call('=_xlfn.xlookup("Bob",A2:A3,B2:B3)') == (
        "XLOOKUP",
        ['"Bob"', "A2:A3", "B2:B3"],
    )


def test_parse_single_call_rejects_a_non_top_level_call() -> None:
    assert parse_single_call('=1+XLOOKUP("Bob",A2:A3,B2:B3)') is None
    assert parse_single_call('=XLOOKUP("Bob",A2:A3,B2:B3)+1') is None


def test_parse_single_call_rejects_an_unsupported_function() -> None:
    assert parse_single_call("=SUM(A1:A3)") is None


def test_parse_single_call_splits_args_respecting_nested_parens_and_quotes() -> None:
    # A quoted string containing a comma, and a nested function call with
    # its own comma, must not fracture the top-level argument split.
    name, args = parse_single_call('=XLOOKUP("a,b",A:A,IF(B1>0,C:C,D:D))')  # type: ignore[misc]
    assert name == "XLOOKUP"
    assert args == ['"a,b"', "A:A", "IF(B1>0,C:C,D:D)"]


def test_xmatch_index_exact_match() -> None:
    assert xmatch_index("Bob", ["Alice", "Bob", "Carol"], match_mode=0, search_mode=1) == 1


def test_xmatch_index_is_case_insensitive_for_text() -> None:
    assert xmatch_index("bob", ["Alice", "Bob", "Carol"], match_mode=0, search_mode=1) == 1


def test_xmatch_index_number_and_its_string_form_do_not_match() -> None:
    from coscribe.tools._xlsx_lookup_formulas import _NO_MATCH

    assert xmatch_index(42, ["42"], match_mode=0, search_mode=1) is _NO_MATCH


def test_xmatch_index_approximate_next_smaller() -> None:
    # match_mode=-1: exact match or the largest value less than lookup_value.
    assert xmatch_index(5, [1, 3, 7, 9], match_mode=-1, search_mode=1) == 1  # value 3


def test_xmatch_index_approximate_next_larger() -> None:
    assert xmatch_index(5, [1, 3, 7, 9], match_mode=1, search_mode=1) == 2  # value 7


def test_xmatch_index_wildcard_match() -> None:
    assert xmatch_index("A*", ["Bob", "Apple", "Ant"], match_mode=2, search_mode=1) == 1


def test_xmatch_index_reverse_search_returns_the_last_match() -> None:
    assert xmatch_index("x", ["x", "y", "x"], match_mode=0, search_mode=-1) == 2


def _values_workbook(rows: list[list[object]]):
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    return wb


def test_evaluate_formula_xlookup_scalar() -> None:
    wb = _values_workbook([["Alice", 9], ["Bob", 7], ["Carol", 5]])
    result = evaluate_formula('=XLOOKUP("Bob",A1:A3,B1:B3)', "Sheet", wb)
    assert result == 7


def test_evaluate_formula_xlookup_with_fallback() -> None:
    wb = _values_workbook([["Alice", 9], ["Bob", 7]])
    result = evaluate_formula('=XLOOKUP("Nobody",A1:A2,B1:B2,"n/a")', "Sheet", wb)
    assert result == "n/a"


def test_evaluate_formula_xlookup_no_fallback_is_formula_error() -> None:
    wb = _values_workbook([["Alice", 9], ["Bob", 7]])
    result = evaluate_formula('=XLOOKUP("Nobody",A1:A2,B1:B2)', "Sheet", wb)
    assert isinstance(result, FormulaError)
    assert result.code == "#N/A"


def test_evaluate_formula_xlookup_mismatched_range_lengths_is_value_error() -> None:
    wb = _values_workbook([["Alice", 9], ["Bob", 7], ["Carol", 5]])
    result = evaluate_formula('=XLOOKUP("Bob",A1:A3,B1:B2)', "Sheet", wb)
    assert isinstance(result, FormulaError)
    assert result.code == "#VALUE!"


def test_evaluate_formula_xmatch() -> None:
    wb = _values_workbook([["Alice"], ["Bob"], ["Carol"]])
    result = evaluate_formula('=XMATCH("Carol",A1:A3)', "Sheet", wb)
    assert result == 3


def test_evaluate_formula_declines_a_non_reference_argument() -> None:
    # SUM(A1:A2) as an argument isn't a plain literal or reference this
    # module resolves -- must decline (not guess), same contract as a
    # too-complex whole formula.
    wb = _values_workbook([["Alice", 9], ["Bob", 7]])
    result = evaluate_formula("=XLOOKUP(SUM(A1:A1),A1:A2,B1:B2)", "Sheet", wb)
    assert result is _UNRESOLVABLE


def test_evaluate_and_patch_lookup_formulas_end_to_end(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(["Alice", 9])
    ws.append(["Bob", 7])
    ws["D1"] = '=_xlfn.XLOOKUP("Bob",A1:A2,B1:B2)'
    path = tmp_path / "t.xlsx"
    wb.save(str(path))

    report = evaluate_and_patch_lookup_formulas(path)

    assert report["formulas_evaluated"] == 1
    assert report["results"] == {"Sheet!D1": "7"}

    from openpyxl import load_workbook

    reopened = load_workbook(str(path), data_only=True)
    assert reopened.active["D1"].value == 7


def test_evaluate_and_patch_lookup_formulas_resolves_the_correct_sheet_in_a_multi_sheet_workbook(
    tmp_path: Path,
) -> None:
    # Real bug caught building this: openpyxl writes relationship Targets
    # as absolute-from-package-root ("/xl/worksheets/sheet1.xml"), not
    # relative -- a naive `target if target.startswith("xl/") else
    # f"xl/{target}"` check misses the leading "/" and produces a bogus
    # "xl//xl/worksheets/sheet1.xml" path that matches no real archive
    # entry, silently patching nothing. A single-sheet workbook can't
    # catch this (there's nothing to mismatch against); a second sheet
    # can and does.
    wb = Workbook()
    data_sheet = wb.active
    data_sheet.title = "Data"
    data_sheet.append(["Alice", 9])
    data_sheet.append(["Bob", 7])
    data_sheet["D1"] = '=_xlfn.XLOOKUP("Bob",A1:A2,B1:B2)'
    other_sheet = wb.create_sheet("Other")
    other_sheet["A1"] = "untouched"
    path = tmp_path / "t.xlsx"
    wb.save(str(path))

    report = evaluate_and_patch_lookup_formulas(path)

    assert report == {"formulas_evaluated": 1, "results": {"Data!D1": "7"}}

    from openpyxl import load_workbook

    reopened = load_workbook(str(path), data_only=True)
    assert reopened["Data"]["D1"].value == 7
    assert reopened["Other"]["A1"].value == "untouched"


def test_evaluate_and_patch_lookup_formulas_is_a_no_op_without_matches(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.append([1, 2, "=A1+B1"])
    path = tmp_path / "t.xlsx"
    original_bytes = None
    wb.save(str(path))
    original_bytes = path.read_bytes()

    report = evaluate_and_patch_lookup_formulas(path)

    assert report == {"formulas_evaluated": 0, "results": {}}
    assert path.read_bytes() == original_bytes


def test_evaluate_formula_xlookup_over_a_one_cell_range() -> None:
    wb = _values_workbook([["a", 1]])
    assert evaluate_formula('=XLOOKUP("a",A1:A1,B1:B1)', "Sheet", wb) == 1
    assert evaluate_formula('=XMATCH("a",A1:A1)', "Sheet", wb) == 1


def test_evaluate_formula_xlookup_of_an_empty_result_cell_is_zero() -> None:
    wb = _values_workbook([["a", 1], ["b", None]])
    assert evaluate_formula('=XLOOKUP("b",A1:A2,B1:B2)', "Sheet", wb) == 0


def test_xmatch_index_approximate_match_on_text() -> None:
    assert xmatch_index("b", ["a", "c"], match_mode=-1, search_mode=1) == 0
    assert xmatch_index("b", ["a", "c"], match_mode=1, search_mode=1) == 1
    assert (
        xmatch_index("B", ["a", "b"], match_mode=-1, search_mode=1) == 1
    )  # exact, case-insensitive
    assert (
        xmatch_index("b", ["a", 5], match_mode=1, search_mode=1) is _NO_MATCH
    )  # numbers never take part


def test_two_calls_joined_by_an_operator_are_left_alone() -> None:
    wb = _values_workbook([["k"]])
    text = '=XLOOKUP(A1,B:B,C:C,"na")&XLOOKUP("k",B:B,C:C)'
    assert parse_single_call(text) is None
    assert evaluate_formula(text, "Sheet", wb) is _UNRESOLVABLE


def test_parse_single_call_keeps_parentheses_inside_a_quoted_sheet_name() -> None:
    assert parse_single_call("=XLOOKUP(\"a\",'Sales)'!A:A,'Sales)'!B:B)") == (
        "XLOOKUP",
        ['"a"', "'Sales)'!A:A", "'Sales)'!B:B"],
    )


# -- the argument list: where it splits and where it stops being one call --


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("", []),
        ("   ", []),
        ("a", ["a"]),
        ("a,b,c", ["a", "b", "c"]),
        ("a,", ["a", ""]),
        (",b", ["", "b"]),
        ('"a,b",c', ['"a,b"', "c"]),
        ('"say ""hi, there""",c', ['"say ""hi, there"""', "c"]),
        ('"",c', ['""', "c"]),
        ('"a"', ['"a"']),
        ("'My,Sheet'!A:A,B:B", ["'My,Sheet'!A:A", "B:B"]),
        ("'it''s'!A1,B1", ["'it''s'!A1", "B1"]),
        ("IF(a,b),c", ["IF(a,b)", "c"]),
        ("F(G(a,b),c),d", ["F(G(a,b),c)", "d"]),
        ('IF(a,")"),c', ['IF(a,")")', "c"]),
        ("(a,b", ["(a,b"]),
        ("a),b", ["a),b"]),
        ('"a"b', ['"a"b']),
    ],
)
def test_split_top_level_args_breaks_only_at_top_level_commas(
    text: str, expected: list[str]
) -> None:
    assert _split_top_level_args(text) == expected


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("", True),
        ("a,b", True),
        ("F(a),G(b)", True),
        ("a)&F(b", False),
        ("a)", False),
        ("(a", False),
        ("((a)", False),
        ('"a)",b', True),
        ('"(",b', True),
        ("'Sales)'!A:A", True),
        ("'(Sales'!A:A", True),
        ('"it\'s)",b', True),
        ("'say \"hi)'!A:A", True),
        ('")" )', False),
        ('"a" ) (', False),
    ],
)
def test_one_argument_list_means_the_call_closes_at_its_last_paren(
    text: str, expected: bool
) -> None:
    assert _is_one_argument_list(text) is expected


@pytest.mark.parametrize(
    "text",
    [
        "XLOOKUP(a,b,c)",
        "=XLOOKUP(a,b,c",
        "=XLOOKUP a,b,c)",
        "=SUM(A1:A3)",
        "=1+XLOOKUP(a,b,c)",
        "=_xlfn._xlws.SORT(A1:A3)",
        "=XLOOKUP(a)&XLOOKUP(b)",
        "=9XLOOKUP(a,b)",
        "=",
        "",
    ],
)
def test_parse_single_call_declines_what_is_not_one_supported_call(text: str) -> None:
    assert parse_single_call(text) is None


def test_parse_single_call_accepts_every_spelling_the_file_may_hold() -> None:
    expected = ("XMATCH", ["a", "B:B"])

    assert parse_single_call("=XMATCH(a,B:B)") == expected
    assert parse_single_call("=xmatch(a,B:B)") == expected
    assert parse_single_call("=_xlfn.XMATCH(a,B:B)") == expected
    assert parse_single_call("=_XLFN._XLWS.xmatch(a,B:B)") == expected
    assert parse_single_call("=_xlfn._xlws.XMATCH(a,B:B)") == expected
    assert parse_single_call("=  XMATCH( a , B:B )  ") == expected
    assert parse_single_call("=XMATCH(a,\nB:B)") == expected
    assert parse_single_call("=XLOOKUP()") == ("XLOOKUP", [])
    assert parse_single_call('=XLOOKUP("a)",B:B)') == ("XLOOKUP", ['"a)"', "B:B"])


# -- literals and references --


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ('"abc"', "abc"),
        ('""', ""),
        ('"a""b"', 'a"b'),
        ('"a""b""c"', 'a"b"c'),
        ("TRUE", True),
        ("true", True),
        ("False", False),
        ("42", 42),
        ("-7", -7),
        ("1.5", 1.5),
        ("1e3", 1000.0),
        ("3.0", 3.0),
    ],
)
def test_parse_literal_reads_strings_booleans_and_numbers(text: str, expected: object) -> None:
    parsed = _parse_literal(text)

    assert parsed == expected
    assert type(parsed) is type(expected)


@pytest.mark.parametrize("text", ["", '"', '"abc', 'abc"', "abc", "A1", "1+1", "'a'", "1.2.3"])
def test_parse_literal_declines_anything_else(text: str) -> None:
    assert _parse_literal(text) is _UNRESOLVABLE


def _grid() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Main"
    for row in ([1, 2, 3], [4, 5, 6], [7, 8, 9]):
        ws.append(row)
    other = wb.create_sheet("My Sheet")
    for row in (["x", "y"], ["z", "w"]):
        other.append(row)
    third = wb.create_sheet("Third")
    third["A1"] = "t"
    return wb


@pytest.mark.parametrize(
    ("ref", "current", "expected"),
    [
        ("B2", "Main", [5]),
        ("b2", "Main", [5]),
        ("$B$2", "Main", [5]),
        ("C3", "Main", [9]),
        ("D4", "Main", [None]),
        ("A1:B2", "Main", [1, 2, 4, 5]),
        ("$A$1:$B$2", "Main", [1, 2, 4, 5]),
        ("B2:C3", "Main", [5, 6, 8, 9]),
        ("A2:C2", "Main", [4, 5, 6]),
        ("B1:B3", "Main", [2, 5, 8]),
        ("A:A", "Main", [1, 4, 7]),
        ("B:C", "Main", [2, 3, 5, 6, 8, 9]),
        ("$B:$B", "Main", [2, 5, 8]),
        ("2:2", "Main", [4, 5, 6]),
        ("1:2", "Main", [1, 2, 3, 4, 5, 6]),
        ("$2:$3", "Main", [4, 5, 6, 7, 8, 9]),
        ("'My Sheet'!A1", "Main", ["x"]),
        ("'My Sheet'!A1:B2", "Main", ["x", "y", "z", "w"]),
        ("'My Sheet'!B:B", "Main", ["y", "w"]),
        ("Third!A1", "Main", ["t"]),
        ("A1", "Third", ["t"]),
        ("A1", "My Sheet", ["x"]),
        ("Main!C1", "Third", [3]),
        ("AA1", "Main", [None]),
        ("B12", "Main", [None]),
        ("A10:B12", "Main", [None] * 6),
    ],
)
def test_resolve_reference_reads_cells_in_reading_order(
    ref: str, current: str, expected: list[object]
) -> None:
    assert _resolve_reference(ref, current, _grid()) == expected


@pytest.mark.parametrize("ref", ["Nowhere!A1", "'No Sheet'!A1", "A1+B1", "SUM(A1)", "", "1", "A"])
def test_resolve_reference_declines_what_names_no_cell(ref: str) -> None:
    assert _resolve_reference(ref, "Main", _grid()) is None


def test_a_reference_argument_is_a_list_and_a_literal_is_not() -> None:
    wb = _grid()

    assert _resolve_arg("B2", "Main", wb) == [5]
    assert _resolve_arg("A1:A2", "Main", wb) == [1, 4]
    assert _resolve_arg('"B2"', "Main", wb) == "B2"
    assert _resolve_arg("7", "Main", wb) == 7
    assert _resolve_arg("Nowhere!A1", "Main", wb) is _UNRESOLVABLE
    assert _resolve_arg("SUM(A1:A2)", "Main", wb) is _UNRESOLVABLE


# -- how Excel compares --


@pytest.mark.parametrize(
    ("a", "b", "expected"),
    [
        (1, 1, True),
        (1, 1.0, True),
        (1, 2, False),
        (2.5, 2.5, True),
        ("a", "a", True),
        ("A", "a", True),
        ("Straße", "STRASSE", True),
        ("a", "b", False),
        (1, "1", False),
        ("1", 1, False),
        (True, True, True),
        (False, False, True),
        (True, False, False),
        (True, 1, False),
        (1, True, False),
        (False, 0, False),
        (None, None, False),
        (None, 0, False),
        (None, "", False),
        ([1], [1], False),
    ],
)
def test_excel_equality(a: object, b: object, expected: bool) -> None:
    assert _excel_eq(a, b) is expected


@pytest.mark.parametrize(
    ("value", "pattern", "expected"),
    [
        ("Apple", "A*", True),
        ("apple", "A*", True),
        ("A", "A*", True),
        ("bA", "A*", False),
        ("Apple", "*le", True),
        ("Apple", "*pp*", True),
        ("Apple", "*", True),
        ("", "*", True),
        ("Ab", "A?", True),
        ("A", "A?", False),
        ("Abc", "A?", False),
        ("abc", "a?c", True),
        ("ac", "a?c", False),
        ("axc", "a*c", True),
        ("ac", "a*c", True),
        ("a.c", "a.c", True),
        ("abc", "a.c", False),
        ("a+b", "a+b", True),
        ("aab", "a+b", False),
        ("(x)", "(x)", True),
        ("[x]", "[x]", True),
        ("x", "[x]", False),
        (5, "*", False),
        (None, "*", False),
        (True, "*", False),
        ("a\\b", "a\\b", True),
    ],
)
def test_wildcard_match(value: object, pattern: str, expected: bool) -> None:
    assert _wildcard_match(value, pattern) is expected


@pytest.mark.parametrize(
    ("lookup", "array", "mode", "order", "expected"),
    [
        (5, [1, 3, 7, 9], -1, range(4), 1),
        (5, [1, 3, 7, 9], 1, range(4), 2),
        (5, [9, 7, 3, 1], -1, range(4), 2),
        (5, [9, 7, 3, 1], 1, range(4), 1),
        (5, [1, 3, 2], -1, range(3), 1),
        (5, [3, 1, 2], -1, range(3), 0),
        (5, [9, 7, 8], 1, range(3), 1),
        (5, [7, 9, 8], 1, range(3), 0),
        (5, [3, 5, 4], -1, range(3), 1),
        (5, [5, 3], -1, range(2), 0),
        (5, [3, 5], 1, range(2), 1),
        (5, [3, 3], -1, range(2), 0),
        (5, [3, 3], -1, range(1, -1, -1), 1),
        (5, [7, 7], 1, range(2), 0),
        (5, [7, 7], 1, range(1, -1, -1), 1),
        (5, [6], -1, range(1), _NO_MATCH),
        (5, [4], 1, range(1), _NO_MATCH),
        (5, [], 1, range(0), _NO_MATCH),
        (5, ["x", None, True, 4], -1, range(4), 3),
        (5, ["x", None, True, 6], 1, range(4), 3),
        (5, [True, None], -1, range(2), _NO_MATCH),
        (-5, [-9, -7], -1, range(2), 1),
        (2.5, [2, 3], -1, range(2), 0),
        (2.5, [2, 3], 1, range(2), 1),
        ("b", ["a", "c"], -1, range(2), 0),
        ("b", ["a", "c"], 1, range(2), 1),
        ("B", ["a", "C"], -1, range(2), 0),
        ("b", ["c", "a"], -1, range(2), 1),
        ("b", ["a", "B"], -1, range(2), 1),
        ("b", [1, "z"], -1, range(2), _NO_MATCH),
        ("b", [None, "a"], -1, range(2), 1),
        ("b", ["c", "d", "a"], -1, range(3), 2),
        ("b", ["a", "c", "d"], 1, range(3), 1),
        ("b", ["d", "c", "a"], 1, range(3), 1),
        ("b", ["a", "a"], -1, range(2), 0),
        ("b", ["a", "a"], -1, range(1, -1, -1), 1),
    ],
)
def test_nearest_index(
    lookup: object, array: list[object], mode: int, order: range, expected: object
) -> None:
    assert _nearest_index(lookup, array, mode, order) == expected


# -- XMATCH and XLOOKUP --


@pytest.mark.parametrize("search_mode", [1, 2])
def test_forward_searches_return_the_first_of_equal_cells(search_mode: int) -> None:
    assert xmatch_index("x", ["a", "x", "X"], match_mode=0, search_mode=search_mode) == 1
    assert xmatch_index("*", ["a", "b"], match_mode=2, search_mode=search_mode) == 0
    assert xmatch_index(5, [4, 6, 3], match_mode=-1, search_mode=search_mode) == 0


@pytest.mark.parametrize("search_mode", [-1, -2])
def test_backward_searches_return_the_last_of_equal_cells(search_mode: int) -> None:
    assert xmatch_index("x", ["a", "x", "X", "b"], match_mode=0, search_mode=search_mode) == 2
    assert xmatch_index("*", ["a", "b"], match_mode=2, search_mode=search_mode) == 1
    assert xmatch_index(5, [4, 3, 6], match_mode=-1, search_mode=search_mode) == 0
    assert xmatch_index(5, [4, 4], match_mode=-1, search_mode=search_mode) == 1


def test_xmatch_index_wildcards_need_text_and_a_text_pattern() -> None:
    assert xmatch_index(5, [5], match_mode=2, search_mode=1) is _NO_MATCH
    assert xmatch_index(True, [True], match_mode=2, search_mode=1) is _NO_MATCH
    assert xmatch_index("5", [5], match_mode=2, search_mode=1) is _NO_MATCH
    assert xmatch_index("a?", ["a", "ab"], match_mode=2, search_mode=1) == 1


def test_xmatch_index_exact_for_every_kind_of_lookup_value() -> None:
    assert xmatch_index(7, [7.0], match_mode=0, search_mode=1) == 0
    assert xmatch_index(True, [1, True], match_mode=0, search_mode=1) == 1
    assert xmatch_index(None, [None], match_mode=0, search_mode=1) is _NO_MATCH
    assert xmatch_index("a", [], match_mode=0, search_mode=1) is _NO_MATCH
    assert xmatch_index("a", [], match_mode=0, search_mode=-1) is _NO_MATCH
    assert xmatch_index(True, [1, True], match_mode=1, search_mode=1) == 1
    assert xmatch_index(True, [1, 2], match_mode=-1, search_mode=1) is _NO_MATCH
    assert xmatch_index(None, [None, 1], match_mode=1, search_mode=1) is _NO_MATCH


@pytest.mark.parametrize(
    ("value", "default", "expected"),
    [
        (_NOT_GIVEN, 0, 0),
        (_NOT_GIVEN, 1, 1),
        (3, 0, 3),
        (-1, 0, -1),
        (2.9, 0, 2),
        (-1.0, 0, -1),
        (True, 0, None),
        (False, 1, None),
        ("1", 0, None),
        (None, 0, None),
        ([1], 0, None),
    ],
)
def test_as_int_arg(value: object, default: int, expected: object) -> None:
    assert _as_int_arg(value, default) == expected


def _is_error(result: object, code: str) -> bool:
    return isinstance(result, FormulaError) and result.code == code


def test_xmatch_is_one_based_and_defaults_to_an_exact_forward_search() -> None:
    assert evaluate_xmatch(["b", ["a", "b", "b"]]) == 2
    assert evaluate_xmatch(["b", ["a", "b", "b"], 0]) == 2
    assert evaluate_xmatch(["b", ["a", "b", "b"], 0, 1]) == 2
    assert evaluate_xmatch(["b", ["a", "b", "b"], 0, -1]) == 3
    assert evaluate_xmatch(["b", ["a", "b", "b"], 0, 2]) == 2
    assert evaluate_xmatch(["b", ["a", "b", "b"], 0, -2]) == 3
    assert evaluate_xmatch([5, [1, 3, 7], -1]) == 2
    assert evaluate_xmatch([5, [1, 3, 7], 1]) == 3
    assert evaluate_xmatch(["B*", ["a", "bc"], 2]) == 2
    assert evaluate_xmatch([5, [1, 3, 7], 1.0]) == 3
    assert _is_error(evaluate_xmatch([5, [1, 3, 7]]), "#N/A")


@pytest.mark.parametrize("match_mode", [-2, 3, 4, True, "0", None])
def test_xmatch_rejects_a_match_mode_it_does_not_know(match_mode: object) -> None:
    assert _is_error(evaluate_xmatch(["a", ["a"], match_mode]), "#VALUE!")


@pytest.mark.parametrize("search_mode", [-3, 0, 3, True, "1", None])
def test_xmatch_rejects_a_search_mode_it_does_not_know(search_mode: object) -> None:
    assert _is_error(evaluate_xmatch(["a", ["a"], 0, search_mode]), "#VALUE!")


def test_xmatch_needs_a_range_and_says_when_nothing_matches() -> None:
    assert _is_error(evaluate_xmatch(["a", "a"]), "#VALUE!")
    assert _is_error(evaluate_xmatch(["z", ["a"]]), "#N/A")
    assert _is_error(evaluate_xmatch([5, [9], -1]), "#N/A")


def test_xlookup_returns_the_cell_beside_the_match() -> None:
    keys, values = ["a", "b", "b"], [10, 20, 30]

    assert evaluate_xlookup(["b", keys, values]) == 20
    assert evaluate_xlookup(["b", keys, values, "nf"]) == 20
    assert evaluate_xlookup(["b", keys, values, "nf", 0]) == 20
    assert evaluate_xlookup(["b", keys, values, "nf", 0, 1]) == 20
    assert evaluate_xlookup(["b", keys, values, "nf", 0, -1]) == 30
    assert evaluate_xlookup(["b", keys, values, "nf", 0, 2]) == 20
    assert evaluate_xlookup(["b", keys, values, "nf", 0, -2]) == 30
    assert evaluate_xlookup([5, [1, 3, 7], values, "nf", -1]) == 20
    assert evaluate_xlookup([5, [1, 3, 7], values, "nf", 1]) == 30
    assert evaluate_xlookup(["B*", keys, values, "nf", 2]) == 20
    assert evaluate_xlookup([5, [1, 3, 7], values, "nf"]) == "nf"
    assert evaluate_xlookup([5, [1, 3, 7], values, "nf", _NOT_GIVEN]) == "nf"


def test_xlookup_not_found_gives_the_fallback_whatever_it_is() -> None:
    keys, values = ["a"], [1]

    assert _is_error(evaluate_xlookup(["z", keys, values]), "#N/A")
    assert evaluate_xlookup(["z", keys, values, "nf"]) == "nf"
    assert evaluate_xlookup(["z", keys, values, ""]) == ""
    assert evaluate_xlookup(["z", keys, values, 0]) == 0
    assert evaluate_xlookup(["z", keys, values, None]) is None
    assert evaluate_xlookup(["z", keys, values, False]) is False


def test_xlookup_keeps_falsy_matches_but_turns_an_empty_cell_into_zero() -> None:
    keys = ["a", "b", "c", "d"]

    assert evaluate_xlookup(["a", keys, [0, "", False, None]]) == 0
    assert evaluate_xlookup(["b", keys, [0, "", False, None]]) == ""
    assert evaluate_xlookup(["c", keys, [0, "", False, None]]) is False
    assert evaluate_xlookup(["d", keys, [0, "", False, None]]) == 0
    assert evaluate_xlookup(["d", keys, [0, "", False, None]]) is not None


def test_xlookup_value_errors() -> None:
    assert _is_error(evaluate_xlookup(["a", "a", [1]]), "#VALUE!")
    assert _is_error(evaluate_xlookup(["a", ["a"], 1]), "#VALUE!")
    assert _is_error(evaluate_xlookup(["a", ["a", "b"], [1]]), "#VALUE!")
    assert _is_error(evaluate_xlookup(["a", ["a"], [1, 2]]), "#VALUE!")
    for bad in (-2, 3, True, "0", None):
        assert _is_error(evaluate_xlookup(["a", ["a"], [1], "nf", bad]), "#VALUE!")
    for bad in (-3, 0, 3, True, "1", None):
        assert _is_error(evaluate_xlookup(["a", ["a"], [1], "nf", 0, bad]), "#VALUE!")


# -- one formula, from text to value --


def _sheet(rows: list[list[object]]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet"
    for row in rows:
        ws.append(row)
    return wb


@pytest.mark.parametrize(
    "formula",
    [
        "=XMATCH(A1)",
        '=XMATCH("a",A1:A2,0,1,9)',
        "=XLOOKUP(A1,A1:A2)",
        '=XLOOKUP("a",A1:A2,B1:B2,"nf",0,1,9)',
        "=XLOOKUP()",
        "=XLOOKUP(Nowhere!A1,A1:A2,B1:B2)",
        "=SUM(A1:A2)",
        "A1",
    ],
)
def test_evaluate_formula_declines_the_wrong_number_of_arguments_and_unknown_parts(
    formula: str,
) -> None:
    assert evaluate_formula(formula, "Sheet", _sheet([["a", 1], ["b", 2]])) is _UNRESOLVABLE


def test_evaluate_formula_accepts_the_full_argument_count_of_each_function() -> None:
    wb = _sheet([["a", 1], ["b", 2]])

    assert evaluate_formula('=XMATCH("b",A1:A2)', "Sheet", wb) == 2
    assert evaluate_formula('=XMATCH("b",A1:A2,0,1)', "Sheet", wb) == 2
    assert evaluate_formula('=XLOOKUP("b",A1:A2,B1:B2)', "Sheet", wb) == 2
    assert evaluate_formula('=XLOOKUP("b",A1:A2,B1:B2,"nf",0,1)', "Sheet", wb) == 2


def test_a_one_cell_reference_is_that_cell_only_where_a_single_value_goes() -> None:
    wb = _sheet(
        [
            ["b", "a", "b", "q", 0, 1, -1],
            ["x", "y", "z", None, None, None, None],
        ]
    )

    assert evaluate_formula("=XMATCH(A1,A1:C1)", "Sheet", wb) == 1
    assert evaluate_formula("=XMATCH(A1,A1:C1,E1)", "Sheet", wb) == 1
    assert evaluate_formula("=XMATCH(A1,A1:C1,E1,G1)", "Sheet", wb) == 3
    assert evaluate_formula("=XMATCH(A1,A1:C1,E1,F1)", "Sheet", wb) == 1
    assert evaluate_formula('=XMATCH("a",B1)', "Sheet", wb) == 1
    assert evaluate_formula("=XLOOKUP(A1,A1:C1,A2:C2)", "Sheet", wb) == "x"
    assert evaluate_formula('=XLOOKUP("a",B1,B2)', "Sheet", wb) == "y"
    assert evaluate_formula('=XLOOKUP("zz",A1:C1,A2:C2,D1)', "Sheet", wb) == "q"
    assert evaluate_formula('=XLOOKUP("b",A1:C1,A2:C2,D1,E1)', "Sheet", wb) == "x"
    assert evaluate_formula('=XLOOKUP("b",A1:C1,A2:C2,D1,E1,G1)', "Sheet", wb) == "z"
    assert evaluate_formula('=XLOOKUP("b",A1:C1,A2:C2,D1,E1,F1)', "Sheet", wb) == "x"


def test_a_range_stays_a_list_where_a_single_value_goes() -> None:
    wb = _sheet([["b", "b"], ["x", "y"]])

    assert _is_error(evaluate_formula("=XMATCH(A1:B1,A1:B1)", "Sheet", wb), "#N/A")
    assert _is_error(
        evaluate_formula('=XLOOKUP("b",A1:B1,A2:B2,A1:B1,0,A1:B1)', "Sheet", wb), "#VALUE!"
    )
    assert _is_error(
        evaluate_formula('=XLOOKUP("b",A1:B1,A2:B2,"nf",A1:B1)', "Sheet", wb), "#VALUE!"
    )


def test_evaluate_formula_looks_in_the_sheet_it_was_given() -> None:
    wb = _sheet([["a", 1]])
    wb.create_sheet("Other")["A1"] = "b"
    wb["Other"]["B1"] = 2

    assert evaluate_formula('=XLOOKUP("b",A1,B1)', "Other", wb) == 2
    assert evaluate_formula('=XLOOKUP("b",A1,Sheet!B1)', "Other", wb) == 1
    assert evaluate_formula('=XLOOKUP("a",Sheet!A1,Sheet!B1)', "Other", wb) == 1


# -- writing the value into the file --


@pytest.mark.parametrize(
    ("result", "expected"),
    [
        (FormulaError("#N/A"), ("#N/A", "e")),
        (FormulaError("#VALUE!"), ("#VALUE!", "e")),
        (True, ("1", "b")),
        (False, ("0", "b")),
        (5, ("5", "n")),
        (-5, ("-5", "n")),
        (0, ("0", "n")),
        (5.0, ("5", "n")),
        (-0.0, ("0", "n")),
        (1e20, ("100000000000000000000", "n")),
        (2.5, ("2.5", "n")),
        ("text", ("text", "str")),
        ("", ("", "str")),
        ("5", ("5", "str")),
        (None, ("None", "str")),
    ],
)
def test_format_cell_value(result: object, expected: tuple[str, str]) -> None:
    assert _format_cell_value(result) == expected


def _sheet_xml(body: str) -> Any:
    return etree.fromstring(
        f'<worksheet xmlns="{_SML_NS}"><sheetData><row r="1">{body}</row></sheetData></worksheet>'
    )


def _cell(root: Any, address: str) -> Any:
    (cell,) = [c for c in root.iter(f"{{{_SML_NS}}}c") if c.get("r") == address]
    return cell


def _children(cell: Any) -> list[str]:
    return [etree.QName(child).localname for child in cell]


def test_patch_cell_puts_the_value_right_after_the_formula() -> None:
    root = _sheet_xml('<c r="A1"><f>X</f><extra/></c>')

    assert _patch_cell(root, "A1", 7) is True

    cell = _cell(root, "A1")
    assert cell.get("t") == "n"
    assert _children(cell) == ["f", "v", "extra"]
    assert cell.find(f"{{{_SML_NS}}}v").text == "7"


def test_patch_cell_replaces_an_old_value_and_type() -> None:
    root = _sheet_xml(
        '<c r="B2" t="e"><f>X</f><v>#NAME?</v></c><c r="C3" t="n"><f>Y</f><v>1</v></c>'
    )

    assert _patch_cell(root, "B2", "text") is True
    assert _patch_cell(root, "C3", FormulaError("#N/A")) is True

    first, second = _cell(root, "B2"), _cell(root, "C3")
    assert (first.get("t"), first.find(f"{{{_SML_NS}}}v").text) == ("str", "text")
    assert (second.get("t"), second.find(f"{{{_SML_NS}}}v").text) == ("e", "#N/A")
    assert _children(first) == ["f", "v"]
    assert _children(second) == ["f", "v"]


def test_patch_cell_adds_a_type_to_an_untyped_cell() -> None:
    root = _sheet_xml('<c r="A1"><f>X</f></c>')

    assert _patch_cell(root, "A1", True) is True
    assert _cell(root, "A1").get("t") == "b"
    assert _cell(root, "A1").find(f"{{{_SML_NS}}}v").text == "1"


def test_patch_cell_leaves_alone_a_cell_without_a_formula_or_one_that_is_not_there() -> None:
    root = _sheet_xml('<c r="A1" t="s"><v>0</v></c><c r="B1"><f>X</f></c>')
    before = etree.tostring(root)

    assert _patch_cell(root, "A1", 7) is False
    assert _patch_cell(root, "Z9", 7) is False
    assert etree.tostring(root) == before

    assert _patch_cell(root, "B1", 7) is True
    assert _cell(root, "B1").get("t") == "n"


def _archive_entries(sheets: str, rels: str) -> dict[str, bytes]:
    ns_r = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    ns_rel = "http://schemas.openxmlformats.org/package/2006/relationships"
    return {
        "xl/workbook.xml": (
            f'<workbook xmlns="{_SML_NS}" xmlns:r="{ns_r}"><sheets>{sheets}</sheets></workbook>'
        ).encode(),
        "xl/_rels/workbook.xml.rels": (
            f'<Relationships xmlns="{ns_rel}">{rels}</Relationships>'
        ).encode(),
    }


def test_sheet_paths_follow_the_relationships_not_the_file_numbering() -> None:
    entries = _archive_entries(
        '<sheet name="First" sheetId="1" r:id="rId2"/>'
        '<sheet name="Second" sheetId="2" r:id="rId1"/>'
        '<sheet name="Absolute" sheetId="3" r:id="rId3"/>',
        '<Relationship Id="rId1" Target="worksheets/sheet1.xml"/>'
        '<Relationship Id="rId2" Target="worksheets/sheet2.xml"/>'
        '<Relationship Id="rId3" Target="/xl/worksheets/sheet3.xml"/>'
        '<Relationship Id="rId4" Target="theme/theme1.xml"/>',
    )

    assert _sheet_archive_paths(entries) == {
        "First": "xl/worksheets/sheet2.xml",
        "Second": "xl/worksheets/sheet1.xml",
        "Absolute": "xl/worksheets/sheet3.xml",
    }


def test_sheet_paths_skip_a_sheet_that_cannot_be_resolved() -> None:
    entries = _archive_entries(
        '<sheet name="NoRelId" sheetId="1"/>'
        '<sheet name="Dangling" sheetId="2" r:id="rId9"/>'
        '<sheet sheetId="3" r:id="rId1"/>'
        '<sheet name="Fine" sheetId="4" r:id="rId1"/>',
        '<Relationship Id="rId1" Target="worksheets/sheet1.xml"/>',
    )

    assert _sheet_archive_paths(entries) == {"Fine": "xl/worksheets/sheet1.xml"}


# -- a whole file --


def _formula_book(path: Path) -> None:
    wb = Workbook()
    data = wb.active
    data.title = "Data"
    for row in (["Alice", 9], ["Bob", 7], ["Carol", 2.5]):
        data.append(row)
    data["D1"] = '=_xlfn.XLOOKUP("Carol",A1:A3,B1:B3)'
    data["D2"] = '=XMATCH("Bob",A1:A3)'
    data["D3"] = '=XLOOKUP("Nobody",A1:A3,B1:B3)'
    data["D4"] = '=XLOOKUP("Nobody",A1:A3,B1:B3,"none")'
    data["D5"] = "=SUM(B1:B3)"
    data["D6"] = '=XLOOKUP("Bob",A1:A3,A1:A3)'
    data["E1"] = "=SUM(B1:B3)"
    data["F1"] = '=XMATCH("Alice",A1:A3)'
    other = wb.create_sheet("Other")
    other["A1"] = "k"
    other["B1"] = 3.0
    other["C1"] = '=XLOOKUP("k",A1,B1)'
    other["C2"] = "=XMATCH(Data!A2,Data!A1:A3)"
    wb.save(str(path))


def test_every_supported_formula_in_the_file_gets_its_value_and_type(tmp_path: Path) -> None:
    path = tmp_path / "book.xlsx"
    _formula_book(path)

    report = evaluate_and_patch_lookup_formulas(path)

    assert report == {
        "formulas_evaluated": 8,
        "results": {
            "Data!F1": "1",
            "Data!D1": "2.5",
            "Data!D2": "2",
            "Data!D3": "#N/A",
            "Data!D4": "none",
            "Data!D6": "Bob",
            "Other!C1": "3",
            "Other!C2": "2",
        },
    }
    from openpyxl import load_workbook

    values = load_workbook(str(path), data_only=True)
    assert [values["Data"][f"D{n}"].value for n in range(1, 7)] == [
        2.5,
        2,
        "#N/A",
        "none",
        None,
        "Bob",
    ]
    assert [values["Other"][c].value for c in ("C1", "C2")] == [3, 2]
    formulas = load_workbook(str(path))
    assert formulas["Data"]["D1"].value == '=_xlfn.XLOOKUP("Carol",A1:A3,B1:B3)'
    assert formulas["Data"]["D5"].value == "=SUM(B1:B3)"


def test_the_patched_file_keeps_its_other_parts_and_stays_compressed(tmp_path: Path) -> None:
    path = tmp_path / "book.xlsx"
    _formula_book(path)
    with zipfile.ZipFile(path) as archive:
        before = {name: archive.read(name) for name in archive.namelist()}
        names = archive.namelist()

    evaluate_and_patch_lookup_formulas(path)

    with zipfile.ZipFile(path) as archive:
        after = {name: archive.read(name) for name in archive.namelist()}
        assert archive.namelist() == names
        assert {info.compress_type for info in archive.infolist()} == {zipfile.ZIP_DEFLATED}
    assert [n for n in before if before[n] != after[n]] == [
        "xl/worksheets/sheet1.xml",
        "xl/worksheets/sheet2.xml",
    ]
    assert after["xl/worksheets/sheet1.xml"].startswith(
        b"<?xml version='1.0' encoding='UTF-8' standalone='yes'?>"
    )


def test_an_array_formula_is_left_alone(tmp_path: Path) -> None:
    from openpyxl.worksheet.formula import ArrayFormula

    wb = Workbook()
    ws = wb.active
    ws["A1"] = "k"
    ws["B1"] = 1
    ws["D1"] = ArrayFormula("D1", '=XLOOKUP("k",A1,B1)')
    path = tmp_path / "array.xlsx"
    wb.save(str(path))
    before = path.read_bytes()

    assert evaluate_and_patch_lookup_formulas(path) == {"formulas_evaluated": 0, "results": {}}
    assert path.read_bytes() == before


def test_the_sentinels_say_what_they_are() -> None:
    assert repr(_NO_MATCH) == "<no-match>"
    assert repr(_UNRESOLVABLE) == "<unresolvable>"
    assert repr(_Sentinel("fresh")) == "<fresh>"
    assert repr(FormulaError("#N/A")) == "FormulaError('#N/A')"


def test_a_range_of_formula_cells_is_read_through_their_recalculated_values(
    tmp_path: Path,
) -> None:
    wb = Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws["A2"] = "=A1+4"
    ws["C1"] = "=XMATCH(5,A1:A2)"
    path = tmp_path / "cached.xlsx"
    wb.save(str(path))
    with zipfile.ZipFile(path) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    sheet = entries["xl/worksheets/sheet1.xml"].decode()
    assert "<f>A1+4</f><v></v>" in sheet or "<f>A1+4</f><v/>" in sheet
    entries["xl/worksheets/sheet1.xml"] = (
        sheet.replace("<f>A1+4</f><v></v>", "<f>A1+4</f><v>5</v>")
        .replace("<f>A1+4</f><v/>", "<f>A1+4</f><v>5</v>")
        .encode()
    )
    with zipfile.ZipFile(path, "w") as archive:
        for name, payload in entries.items():
            archive.writestr(name, payload)

    report = evaluate_and_patch_lookup_formulas(path)

    assert report == {"formulas_evaluated": 1, "results": {"Sheet!C1": "2"}}
