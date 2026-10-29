"""A Doc written out for Word (decision 0025): the nearest thing the writer's subset can hold."""

from __future__ import annotations

import io
import zipfile

from docx import Document

from coscribe.tools import docx_from_markdown, docx_markdown

DOC = """---
coscribe: doc
---

# Launch plan

We ship *first* in **May**.

> [!WARNING]
> Budget is not final.
>
> Ask finance.

- one
- two

![the chart](plan.doc/assets/chart.png)

---

- [ ] book the room
- [x] send the invite

```chart
bar Sales
| q | n |
| --- | --- |
| Q1 | 3 |
```

<div>raw</div>

```
code line
```
"""


def test_a_doc_becomes_the_markdown_subset_word_takes() -> None:
    out = docx_markdown(DOC)

    assert out.startswith("# Launch plan\n\nWe ship *first* in **May**.")
    assert "§callout:WARNING:Warning§\n\nBudget is not final.\n\nAsk finance.\n\n§/callout§" in out
    assert "- one\n- two" in out
    assert "[Picture: the chart]" in out and "plan.doc/assets" not in out
    assert "§chart:bar:Sales§\n\n| q | n |\n| --- | --- |\n| Q1 | 3 |" in out
    assert "- ☐ book the room\n- ☑ send the invite" in out
    assert "<div>" not in out and "```" not in out and "code line" in out
    assert "coscribe: doc" not in out


def test_the_export_is_a_word_file_with_the_doc_in_it() -> None:
    data = docx_from_markdown(docx_markdown(DOC))

    texts = [p.text for p in Document(io.BytesIO(data)).paragraphs]
    assert "Launch plan" in texts and any("Budget is not final." in t for t in texts)
    assert not Document(io.BytesIO(data)).tables


def _word_file() -> Document:
    return Document(io.BytesIO(docx_from_markdown(docx_markdown(DOC))))


def test_a_callout_is_framed_in_its_colour_with_its_label_and_no_marker_is_left() -> None:
    document = _word_file()

    paragraphs = {p.text: p for p in document.paragraphs}
    label = paragraphs["WARNING"]
    assert "w:pBdr" in label._p.xml and "D98A00" in label._p.xml and "w:shd" in label._p.xml
    assert "w:pBdr" in paragraphs["Ask finance."]._p.xml
    assert not any("§" in p.text for p in document.paragraphs)


def test_a_chart_is_a_native_chart_and_its_table_is_gone() -> None:
    document = _word_file()

    parts = {str(part.partname) for part in document.part.package.iter_parts()}
    assert "/word/charts/chart1.xml" in parts
    assert not document.tables


def test_a_task_item_has_its_box_and_no_bullet() -> None:
    paragraphs = {p.text: p for p in _word_file().paragraphs}

    task = paragraphs["☐ book the room"]
    assert 'w:numId w:val="0"' in task._p.xml
    assert 'w:numId w:val="0"' not in paragraphs["one"]._p.xml


def test_a_chart_without_data_keeps_its_title_over_the_table() -> None:
    source = "---\ncoscribe: doc\n---\n\n```chart\nbar Empty\n| q |\n| --- |\n```\n"
    document = Document(io.BytesIO(docx_from_markdown(docx_markdown(source))))

    assert [p.text for p in document.paragraphs] == ["Empty"]


def test_a_horizontal_rule_is_a_hairline_not_dashes() -> None:
    source = "---\ncoscribe: doc\n---\n\nAbove\n\n---\n\nBelow\n"
    document = Document(io.BytesIO(docx_from_markdown(docx_markdown(source))))

    rule = next(p for p in document.paragraphs if "w:bottom" in p._p.xml)
    assert rule.text == "" and not any("---" in p.text for p in document.paragraphs)


def test_a_callout_is_kept_with_its_label_across_a_page() -> None:
    label, body, last = (p for p in _word_file().paragraphs if "<w:left " in p._p.xml)

    assert "w:keepNext" in label._p.xml and "w:keepNext" in body._p.xml
    assert "w:keepNext" not in last._p.xml


def test_chinese_italic_inline_code_and_links_are_formatting_not_markers() -> None:
    source = (
        "---\ncoscribe: doc\n---\n\n"
        "语气轻一点用*斜体*，字段名用 `order_id`，引用加 [国家统计局](https://www.stats.gov.cn)。\n"
    )
    document = Document(io.BytesIO(docx_from_markdown(docx_markdown(source))))
    (paragraph,) = document.paragraphs
    xml = paragraph._p.xml

    assert "*" not in paragraph.text and "`" not in paragraph.text and "](" not in xml
    assert next(r for r in paragraph.runs if r.text == "斜体").italic
    assert next(r for r in paragraph.runs if r.text == "order_id").font.name == "Consolas"
    assert "<w:hyperlink" in xml and "国家统计局" in xml
    assert "https://www.stats.gov.cn" in [rel.target_ref for rel in document.part.rels.values()]


def test_a_paragraph_ending_in_a_colon_keeps_with_the_chart_it_introduces() -> None:
    source = (
        "---\ncoscribe: doc\n---\n\nAbove.\n\n构成比例用饼图：\n\n"
        "```chart\npie Sources\n| s | n |\n| --- | --- |\n| A | 3 |\n| B | 5 |\n```\n"
    )
    document = Document(io.BytesIO(docx_from_markdown(docx_markdown(source))))
    paragraphs = {p.text: p for p in document.paragraphs}

    assert "w:keepNext" in paragraphs["构成比例用饼图："]._p.xml
    assert "w:keepNext" not in paragraphs["Above."]._p.xml


def test_a_chart_takes_the_doc_panels_colours_and_one_series_has_no_legend() -> None:
    source = (
        "---\ncoscribe: doc\n---\n\n"
        "```chart\nline Orders\n| w | n |\n| --- | --- |\n| W1 | 3 |\n| W2 | 5 |\n```\n"
    )
    data = docx_from_markdown(docx_markdown(source))
    with zipfile.ZipFile(io.BytesIO(data)) as package:
        chart = package.read("word/charts/chart1.xml").decode()

    assert 'val="4F6EF7"' in chart
    assert "<c:legend>" not in chart
