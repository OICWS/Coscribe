"""Slide sections (decision 0015): positioned HTML in, native PowerPoint objects out.

What a section asks for (a wrapped paragraph, a list, a shape, a picture, a drawing, a title, a
background, notes) must come out as the object a user selects and edits in PowerPoint, and what
the section language does not allow must be refused with a message that says what to write.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Emu

from coscribe.tools.office.pptx.engine import check_sections, compile_sections
from coscribe.tools.office.pptx.engine.text_layout import Font, Paragraph, Run, line_count

_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_PX = 6350  # EMU per px on the 1920-wide canvas
_BOX = "left:0;top:0;width:9px;height:9px"

SECTION = """<section style="background:#F7F5F0; color:#1F2937">
  <x-shape id="band" data-shape="rect"
    style="left:0px;top:0px;width:1920px;height:16px;background:#C2410C"/>
  <h1 id="title" style="left:128px;top:96px;width:1664px;height:96px">第三季度经营回顾</h1>
  <p id="lead" style="left:128px;top:220px;width:900px;height:200px;font-size:32px"
    >收入同比增长 <b>23.5%</b>，达到 <span style="color:#C2410C">5,410 万元</span>。</p>
  <ol id="steps" style="left:128px;top:460px;width:900px;height:400px">
    <li>华东区新签客户 42 家</li>
    <li>Renewal rate rose to <i>91%</i></li>
  </ol>
  <img id="photo" src="assets/photo.png" style="left:1132px;top:252px;width:628px;height:300px"/>
  <svg id="bars" style="left:1132px;top:580px;width:628px;height:260px" viewBox="0 0 314 130">
    <rect x="0" y="80" width="60" height="50" fill="#FDBA74"/>
    <text x="200" y="20" font-size="12" fill="#1F2937">5,410</text>
  </svg>
  <aside>先讲收入。
  Then the margin.</aside>
</section>"""


def _deck(tmp_path: Path, *sections: str) -> tuple[object, object]:
    (tmp_path / "assets").mkdir(exist_ok=True)
    Image.new("RGB", (800, 500), (30, 120, 200)).save(tmp_path / "assets" / "photo.png")
    pages = [(f"{i:02d}.html", text) for i, text in enumerate(sections, start=1)]
    report = compile_sections(
        pages, tmp_path / "deck.pptx", resource_root=tmp_path, sections_dir=tmp_path
    )
    check_sections(report)
    return report, Presentation(str(tmp_path / "deck.pptx"))


def _shape(slide: object, name: str) -> object:
    return next(s for s in slide.shapes if s.name == name)  # type: ignore[attr-defined]


def test_a_section_becomes_native_objects_on_a_16_9_slide(tmp_path: Path) -> None:
    report, prs = _deck(tmp_path, SECTION)
    assert report.ok, report.errors()
    assert (prs.slide_width, prs.slide_height) == (Emu(12192000), Emu(6858000))
    slide = prs.slides[0]
    assert slide.shapes.title.text_frame.text == "第三季度经营回顾"
    assert _shape(slide, "photo").shape_type == 13  # a picture
    assert _shape(slide, "band").left == 0 and _shape(slide, "band").width == 1920 * _PX
    assert slide.notes_slide.notes_text_frame.text == "先讲收入。\nThen the margin."
    fill = slide._element.find(f".//{_A}srgbClr")
    assert fill is not None and fill.get("val") == "F7F5F0"


def test_text_wraps_in_powerpoint_with_exact_line_spacing(tmp_path: Path) -> None:
    _report, prs = _deck(tmp_path, SECTION)
    lead = _shape(prs.slides[0], "lead")
    body = lead.text_frame._txBody
    assert body.find(f"{_A}bodyPr").get("wrap") == "square"
    assert body.find(f".//{_A}lnSpc/{_A}spcPts").get("val") == str(round(32 * 1.35 * 50))
    runs = [(r.text, r.font.bold) for r in lead.text_frame.paragraphs[0].runs]
    assert ("23.5%", True) in runs
    assert [p.text for p in _shape(prs.slides[0], "steps").text_frame.paragraphs] == [
        "华东区新签客户 42 家",
        "Renewal rate rose to 91%",
    ]
    assert lead._element.find(f".//{_A}buAutoNum") is None
    assert _shape(prs.slides[0], "steps")._element.find(f".//{_A}buAutoNum") is not None


def test_a_box_taller_than_its_text_is_drawn_at_the_texts_height(tmp_path: Path) -> None:
    section = """<section><h2 id="title" style="left:100px;top:100px;width:800px;height:400px;
      vertical-align:bottom">Short</h2></section>"""
    _report, prs = _deck(tmp_path, section)
    title = prs.slides[0].shapes.title
    height = round(48 * 1.15 * _PX)
    assert abs(title.height - height) <= 1
    assert abs(title.top + title.height - 500 * _PX) <= 1


def test_text_that_needs_more_height_than_its_box_is_an_error(tmp_path: Path) -> None:
    long = "这是一段很长的说明文字，" * 20
    box = "left:100px;top:100px;width:600px;height:80px"
    section = f'<section><p id="title" style="{box}">{long}</p></section>'
    report, _prs = _deck(tmp_path, section)
    assert any("needs" in e and "the box has 80" in e for e in report.errors())


def test_findings_name_the_sections_own_coordinates(tmp_path: Path) -> None:
    section = """<section>
      <p id="title" style="left:100px;top:100px;width:600px;height:60px">Overlapping words here</p>
      <p id="b" style="left:100px;top:110px;width:600px;height:60px">Overlapping words here</p>
    </section>"""
    report, _prs = _deck(tmp_path, section)
    overlap = [e for e in report.errors() if "drawn over each other" in e]
    assert overlap and "x 100" in overlap[0]


@pytest.mark.parametrize(
    ("section", "message"),
    [
        ("<section><script>alert(1)</script></section>", "<script> is not part of a section"),
        (
            '<section><p onclick="x()" style="left:0;top:0;width:9px;height:9px">a</p></section>',
            "'onclick'",
        ),
        (
            '<section><p style="left:0;top:0;width:9em;height:9px">a</p></section>',
            "width must be a number of px",
        ),
        (
            '<section><p style="left:0;top:0;width:9px;height:9px;margin:4px">a</p></section>',
            "'margin'",
        ),
        (
            f'<section><img src="https://x.test/a.png" style="{_BOX}"/></section>',
            "inside the deck folder",
        ),
        (
            f'<section><img src="../secret.png" style="{_BOX}"/></section>',
            "inside the deck folder",
        ),
        (
            f'<section><svg viewBox="0 0 9 9" style="{_BOX}"><script/></svg></section>',
            "<script> is not allowed",
        ),
        (
            '<section><p style="left:0;top:0;width:9px;height:9px">a&nbsp;b</p></section>',
            "not well-formed",
        ),
        (
            '<section><p style="left:0;top:0;width:9px;height:9px"><div>a</div></p></section>',
            "cannot sit in text",
        ),
    ],
)
def test_what_the_language_does_not_allow_is_refused(
    tmp_path: Path, section: str, message: str
) -> None:
    report, _prs = _deck(tmp_path, section)
    assert any(message in e for e in report.errors()), report.errors()


def _paragraph(text: str, size: float = 20.0) -> Paragraph:
    return Paragraph([Run(text, Font("Microsoft YaHei, Arial", size))], size)


def test_lines_break_between_cjk_characters_and_at_spaces() -> None:
    assert line_count(_paragraph("你好" * 10), 420) == 1  # 20 characters of 20 px
    assert line_count(_paragraph("你好" * 10), 200) == 3
    assert line_count(_paragraph("word " * 10), 10_000) == 1
    assert line_count(_paragraph("one\ntwo"), 10_000) == 2


def test_closing_punctuation_never_starts_a_line() -> None:
    # Ten characters fit a line. The comma that would start the second line takes the character
    # before it along, so the second line gets eleven and a third is needed.
    width = 10 * 20 / 0.97 + 1
    plain = _paragraph("你" * 9 + "好你" + "你" * 9, 20)
    tied = _paragraph("你" * 9 + "好，" + "你" * 9, 20)
    assert line_count(plain, width) == 2
    assert line_count(tied, width) == 3


def test_a_word_wider_than_the_line_is_broken_inside() -> None:
    assert line_count(_paragraph("x" * 200, 20), 300) > 1
