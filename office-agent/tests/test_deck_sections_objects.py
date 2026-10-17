"""Slide sections, part 2: tables, charts, icons, transitions and build-ins, and building a
folder of sections the way a folder of SVG pages is built (decision 0015)."""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation

from coscribe.tools.office.pptx.engine import DeckChangedError, build_deck, compile_sections

_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
_PX = 6350
_TITLE = '<h2 id="title" style="left:96px;top:40px;width:1700px;height:80px">Results</h2>'

TABLE = f"""<section>{_TITLE}
  <table id="regions" style="left:100px;top:200px;width:900px;height:300px;font-size:30px">
    <colgroup>
      <col style="width:300px"/><col style="width:300px"/><col style="width:300px"/>
    </colgroup>
    <tr><th style="background:#0F172A;color:#FFFFFF">区域</th><th>营收</th><th>同比</th></tr>
    <tr><td>华东</td><td style="text-align:right">5,410</td>
        <td style="text-align:right;color:#15803D;font-weight:bold">+12.2%</td></tr>
    <tr><td colspan="2">合计</td><td style="text-align:right">7,530</td></tr>
  </table></section>"""

CHART = f"""<section>{_TITLE}
  <x-chart id="revenue" data-build-in="rise" style="left:1000px;top:300px;width:800px;height:600px">
  {{"type": "column", "categories": ["华东", "华北"],
    "series": [{{"name": "营收", "values": [5410, 2120]}}], "font_size": 30}}
  </x-chart>
  <x-icon id="ship" data-icon="ship" data-build-in="pop"
    style="left:100px;top:200px;width:96px;height:96px;color:#C2410C"/>
</section>"""


def _compile(tmp_path: Path, *sections: str) -> tuple[object, object]:
    pages = [(f"{i:02d}.html", text) for i, text in enumerate(sections, start=1)]
    report = compile_sections(pages, tmp_path / "deck.pptx", resource_root=tmp_path)
    return report, Presentation(str(tmp_path / "deck.pptx"))


def _shape(slide: object, name: str) -> object:
    return next(s for s in slide.shapes if s.name == name)  # type: ignore[attr-defined]


def test_a_table_is_a_native_table_with_its_merges_header_and_cell_styles(tmp_path: Path) -> None:
    report, prs = _compile(tmp_path, TABLE)
    assert report.ok, report.errors()
    frame = _shape(prs.slides[0], "regions")
    table = frame.table
    assert (frame.left, frame.width) == (100 * _PX, 900 * _PX)
    assert [c.text for c in table.rows[0].cells] == ["区域", "营收", "同比"]
    assert table.cell(2, 0).is_merge_origin and table.cell(2, 0).span_width == 2
    assert table.cell(2, 1).is_spanned
    growth = table.cell(1, 2)
    assert growth.text_frame.paragraphs[0].alignment == 3  # right
    run = growth.text_frame.paragraphs[0].runs[0]
    assert run.font.bold and str(run.font.color.rgb) == "15803D"
    assert run.font.size.pt == 15  # 30 px on the 1920-wide canvas
    assert [c.width for c in table.columns] == [300 * _PX] * 3


def test_a_chart_is_native_and_the_icon_is_drawn_with_both_built_in_order(tmp_path: Path) -> None:
    report, prs = _compile(tmp_path, CHART)
    assert report.ok, report.errors()
    slide = prs.slides[0]
    chart = _shape(slide, "revenue")
    assert chart.has_chart and list(chart.chart.plots[0].categories) == ["华东", "华北"]
    assert (chart.left, chart.top, chart.width) == (1000 * _PX, 300 * _PX, 800 * _PX)
    assert _shape(slide, "ship").shape_type == 6  # a group of the icon's paths
    effects = slide._element.findall(f".//{_P}cTn[@nodeType='clickEffect']")
    assert [e.get("presetID") for e in effects] == ["37", "23"]  # rise up, then zoom
    builds = slide._element.find(f"{_P}timing/{_P}bldLst")
    assert [(b.tag, b.get("spid")) for b in builds] == [
        (f"{_P}bldGraphic", str(chart.shape_id))
    ]  # the chart builds as one piece; the icon (a group) needs no entry


def test_the_section_names_its_transition(tmp_path: Path) -> None:
    magic = TABLE.replace("<section>", '<section data-transition="magic">', 1)
    push = TABLE.replace("<section>", '<section data-transition="push">', 1)
    report, prs = _compile(tmp_path, magic, push)
    assert report.ok, report.errors()
    first, second = (s._element.find(f"{_P}transition") for s in prs.slides)
    assert first is not None and first[0].tag.endswith("}morph")
    assert second is not None and second[0].tag == f"{_P}push"


@pytest.mark.parametrize(
    ("section", "message"),
    [
        (
            '<section><x-chart style="left:0;top:0;width:9px;height:9px">{"x": 1}</x-chart>'
            "</section>",
            "its box is in its style",
        ),
        (
            '<section><x-chart style="left:0;top:0;width:9px;height:9px">not json</x-chart>'
            "</section>",
            "holds the chart's JSON",
        ),
        (
            '<section><table style="left:0;top:0;width:9px;height:9px"><tr><td>a</td></tr>'
            "<tr><td>a</td><td>b</td></tr></table></section>",
            "same number of columns",
        ),
        (
            '<section><table style="left:0;top:0;width:9px;height:9px"><p>a</p></table></section>',
            "holds rows",
        ),
        (
            '<section><table style="left:0;top:0;width:9px;height:9px"><tr><td><b><i '
            'onclick="x()">a</i></b></td></tr></table></section>',
            "in a cell takes no attributes",
        ),
        (
            '<section><x-icon data-icon="ship" style="left:0;top:0;width:9px;height:9px">'
            "<script/></x-icon></section>",
            "holds no elements",
        ),
        (
            '<section data-transition="spin"><p style="left:0;top:0;width:9px;height:9px">a</p>'
            "</section>",
            "data-transition is one of fade, push, magic",
        ),
    ],
)
def test_malformed_objects_are_refused(tmp_path: Path, section: str, message: str) -> None:
    report, _prs = _compile(tmp_path, section)
    assert any(message in e for e in report.errors()), report.errors()


def test_a_folder_of_sections_builds_and_is_not_overwritten_after_a_change(tmp_path: Path) -> None:
    deck = tmp_path / "q3.slides"
    deck.mkdir()
    (deck / "01.html").write_text(TABLE, encoding="utf-8")
    out = tmp_path / "q3.pptx"
    report = build_deck(deck, out, resource_root=tmp_path)
    assert report.ok and report.changed == ["01.html"]
    assert (deck / "build.json").is_file()
    prs = Presentation(str(out))
    prs.slides[0].shapes.title.text_frame.text = "Edited in PowerPoint"
    prs.save(str(out))
    with pytest.raises(DeckChangedError):
        build_deck(deck, out, resource_root=tmp_path)


def test_a_deck_with_sections_is_not_built_on_a_template(tmp_path: Path) -> None:
    (tmp_path / "01.html").write_text(TABLE, encoding="utf-8")
    (tmp_path / "02.svg").write_text("<svg/>", encoding="utf-8")
    with pytest.raises(ValueError, match="not built on a template"):
        build_deck(tmp_path, tmp_path / "deck.pptx", template=tmp_path / "brand.pptx")
