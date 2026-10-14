"""Shapes inside groups: listed with a path, reached by path or name from every tool that takes
`shape_index`, and placed in slide inches even when the group scales its members."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
from pptx import Presentation
from pptx.util import Emu, Inches

from coscribe.tools.presentations import build_presentation_tools

_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _tools(root: Path) -> dict[str, Any]:
    return {tool.__name__: tool for tool in build_presentation_tools(root)}  # type: ignore[attr-defined]


def _deck(root: Path, *, scale: float = 1.0) -> None:
    """Slide 1: a text box, then a group "benefit-1" holding a circle and two text boxes, the
    group drawn at `scale` times its members' own size, as PowerPoint does after a resize."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(4), Inches(1)).text = "Title"
    group = slide.shapes.add_group_shape()
    group.name = "benefit-1"
    group.shapes.add_shape(9, Inches(1), Inches(2), Inches(0.5), Inches(0.5))
    group.shapes.add_textbox(Inches(2), Inches(2), Inches(3), Inches(0.5)).text = "护眼无频闪"
    detail = group.shapes.add_textbox(Inches(2), Inches(2.6), Inches(3), Inches(0.5))
    detail.text = "光线稳定不频闪"
    detail.name = "detail"
    if scale != 1.0:
        ext = group._element.find(f".//{_A}xfrm/{_A}ext")  # noqa: SLF001
        ext.set("cx", str(round(int(ext.get("cx")) * scale)))
        ext.set("cy", str(round(int(ext.get("cy")) * scale)))
    prs.save(str(root / "deck.pptx"))


def _member(root: Path, *steps: int) -> Any:
    shape: Any = Presentation(str(root / "deck.pptx")).slides[0]
    for step in steps:
        shape = list(shape.shapes)[step]
    return shape


def test_a_group_lists_its_members_with_their_paths(tmp_path: Path) -> None:
    _deck(tmp_path)

    shapes = _tools(tmp_path)["list_pptx_shapes"](path="deck.pptx", slide=1)["shapes"]

    group = shapes[1]
    assert group["name"] == "benefit-1"
    assert [m["path"] for m in group["members"]] == ["1/0", "1/1", "1/2"]
    assert group["members"][1]["text_preview"] == "护眼无频闪"
    assert (group["members"][1]["left_in"], group["members"][1]["top_in"]) == (2.0, 2.0)
    assert "members" not in shapes[0]


def test_members_are_placed_on_the_slide_through_a_scaled_group(tmp_path: Path) -> None:
    _deck(tmp_path, scale=2.0)

    members = _tools(tmp_path)["list_pptx_shapes"](path="deck.pptx", slide=1)["shapes"][1]
    text = members["members"][1]

    # The group's members span 1 to 5 inches across; drawn twice as large from x = 1.
    assert (text["left_in"], text["width_in"]) == (3.0, 6.0)


@pytest.mark.parametrize("ref", ["1/2", "detail"])
def test_a_member_is_reached_by_path_or_by_name(tmp_path: Path, ref: str) -> None:
    _deck(tmp_path)

    _tools(tmp_path)["edit_pptx_shape"](
        path="deck.pptx", slide=1, shape_index=ref, text_color="C2410C"
    )

    run = _member(tmp_path, 1, 2).text_frame.paragraphs[0].runs[0]
    assert str(run.font.color.rgb) == "C2410C"
    assert _member(tmp_path, 1, 1).text_frame.paragraphs[0].runs[0].font.color.type is None


def test_an_index_still_means_a_top_level_shape(tmp_path: Path) -> None:
    _deck(tmp_path)

    _tools(tmp_path)["edit_pptx_shape"](
        path="deck.pptx", slide=1, shape_index=0, text_color="111111"
    )

    assert str(_member(tmp_path, 0).text_frame.paragraphs[0].runs[0].font.color.rgb) == "111111"


def test_moving_a_member_takes_slide_inches_through_a_scaled_group(tmp_path: Path) -> None:
    _deck(tmp_path, scale=2.0)
    tools = _tools(tmp_path)

    tools["edit_pptx_shape"](
        path="deck.pptx", slide=1, shape_index="1/1", left_in=4.0, width_in=2.0
    )

    text = tools["list_pptx_shapes"](path="deck.pptx", slide=1)["shapes"][1]["members"][1]
    assert (text["left_in"], text["width_in"]) == (4.0, 2.0)
    # Stored in the group's own coordinates: half the slide distance from its origin.
    assert _member(tmp_path, 1, 1).left == Emu(Inches(2.5))


def test_a_member_is_deleted_and_its_group_kept(tmp_path: Path) -> None:
    _deck(tmp_path)

    _tools(tmp_path)["delete_pptx_shape"](path="deck.pptx", slide=1, shape_index="1/0")

    group = _member(tmp_path, 1)
    assert [shape.text_frame.text for shape in group.shapes] == ["护眼无频闪", "光线稳定不频闪"]


@pytest.mark.parametrize(
    ("ref", "message"),
    [
        ("1/7", "out of range: slide 1 has 3 shape"),
        ("5", "out of range: slide 1 has 2 shape"),
        ("nothing", "No shape named 'nothing'"),
    ],
)
def test_a_wrong_reference_names_list_pptx_shapes(tmp_path: Path, ref: str, message: str) -> None:
    _deck(tmp_path)

    with pytest.raises(ValueError, match=message) as error:
        _tools(tmp_path)["delete_pptx_shape"](path="deck.pptx", slide=1, shape_index=ref)

    assert "list_pptx_shapes" in str(error.value)


def test_a_name_two_shapes_share_asks_for_the_path(tmp_path: Path) -> None:
    _deck(tmp_path)
    prs = Presentation(str(tmp_path / "deck.pptx"))
    list(prs.slides[0].shapes)[0].name = "detail"
    prs.save(str(tmp_path / "deck.pptx"))

    with pytest.raises(ValueError, match=r"2 shapes are named 'detail' .*paths 0, 1/2"):
        _tools(tmp_path)["delete_pptx_shape"](path="deck.pptx", slide=1, shape_index="detail")


def test_a_member_is_not_animated_on_its_own(tmp_path: Path) -> None:
    _deck(tmp_path)
    animate = _tools(tmp_path)["add_pptx_animation"]

    with pytest.raises(ValueError, match="animate the group"):
        animate(path="deck.pptx", slide=1, shape_index="1/1", animation="fade")

    animate(path="deck.pptx", slide=1, shape_index="benefit-1", animation="fade")
