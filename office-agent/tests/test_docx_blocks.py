"""The Docs editor's reader and writer (decision 0019): a .docx as blocks, and a save that rewrites
only the blocks the person changed, leaving comments, tracked changes and the rest as they were."""

from __future__ import annotations

import base64
import re
import zipfile
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.enum.text import WD_BREAK
from lxml import etree

from coscribe.tools.office.document_tools import build_document_tools
from coscribe.tools.office.docx import (
    BlocksRefused,
    DocChanged,
    read_doc_blocks,
    read_doc_image,
    save_doc_blocks,
)

CONTENT = """\
# Service agreement

## Price

The total price is **100 dollars**, due within 30 days. {{comment: Final price?}}

## Term

The agreement runs for one year, see *schedule A*.

- first point
- second point

1. step one
2. step two

| item | cost |
|---|---|
| a | 1 |
| b | 2 |
"""
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _tools(root: Path) -> dict[str, Any]:
    return {tool.__name__: tool for tool in build_document_tools(root)}


@pytest.fixture
def doc(tmp_path: Path) -> Path:
    _tools(tmp_path)["write_docx"](path="d.docx", content=CONTENT, comment_author="Ana")
    return tmp_path / "d.docx"


def _blocks(path: Path) -> tuple[str, list[dict[str, Any]]]:
    read = read_doc_blocks(path)
    return read["version"], read["blocks"]


def _by_text(blocks: list[dict[str, Any]], text: str) -> dict[str, Any]:
    return next(b for b in blocks if any(text in r["t"] for r in b.get("runs", [])))


def _paragraphs(path: Path) -> list[bytes]:
    tree = etree.fromstring(zipfile.ZipFile(path).read("word/document.xml"))
    return [etree.tostring(p, method="c14n") for p in tree.iter(f"{W}p")]


def _canonical(name: str, data: bytes) -> Any:
    if not name.endswith((".xml", ".rels")):
        return data
    root = etree.fromstring(data)
    if name == "[Content_Types].xml":  # python-docx lists the same entries in its own order
        return sorted(etree.tostring(child, method="c14n") for child in root)
    return etree.tostring(root, method="c14n")


def _parts(path: Path) -> dict[str, Any]:
    """Every part but the body, as XML: python-docx writes the same XML its own way."""
    z = zipfile.ZipFile(path)
    return {n: _canonical(n, z.read(n)) for n in z.namelist() if n != "word/document.xml"}


def test_a_document_reads_as_blocks_with_what_the_editor_cannot_change_locked(doc: Path) -> None:
    _, blocks = _blocks(doc)
    kinds = [(b["kind"], b.get("level")) for b in blocks]
    assert kinds[:3] == [("title", None), ("heading", 2), ("locked", None)]
    assert blocks[2]["what"] == "comment" and "100 dollars" in blocks[2]["text"]
    assert ("bullet", 1) in kinds and ("number", 1) in kinds and kinds[-1] == ("table", None)
    schedule = _by_text(blocks, "schedule A")
    assert [r for r in schedule["runs"] if r.get("i")] == [{"t": "schedule A", "i": True}]
    table = blocks[-1]
    assert table["header"] and table["rows"][1] == [[{"t": "a"}], [{"t": "1"}]]


def test_saving_what_was_read_changes_no_paragraph_and_no_other_part(doc: Path) -> None:
    before_paragraphs, before_parts = _paragraphs(doc), _parts(doc)
    version, blocks = _blocks(doc)
    save_doc_blocks(doc, version, blocks)
    assert _paragraphs(doc) == before_paragraphs
    assert _parts(doc) == before_parts


def test_a_changed_paragraph_is_the_only_one_rewritten(doc: Path) -> None:
    before = _paragraphs(doc)
    version, blocks = _blocks(doc)
    term = _by_text(blocks, "one year")
    term["runs"] = [{"t": "The agreement runs for "}, {"t": "two", "b": True}, {"t": " years."}]
    result = save_doc_blocks(doc, version, blocks)
    after = _paragraphs(doc)
    edited = next(i for i, p in enumerate(before) if b"one year" in p)
    assert [i for i, (x, y) in enumerate(zip(before, after, strict=True)) if x != y] == [edited]
    assert _by_text(result["blocks"], "two")["runs"][1] == {"t": "two", "b": True}


def test_comments_and_tracked_changes_survive_an_edit_elsewhere(tmp_path: Path) -> None:
    tools = _tools(tmp_path)
    tools["write_docx"](path="r.docx", content=CONTENT, comment_author="Ana")
    tools["insert_docx_text"](path="r.docx", text=" (extended)", context_before="30 days")
    tools["delete_docx_text"](path="r.docx", text="one year")
    path = tmp_path / "r.docx"
    version, blocks = _blocks(path)
    assert {b["what"] for b in blocks if b["kind"] == "locked"} >= {"comment", "tracked change"}
    first = blocks[1]
    first["runs"] = [{"t": "Cost"}]
    save_doc_blocks(path, version, blocks)
    document = zipfile.ZipFile(path).read("word/document.xml").decode()
    assert "<w:ins " in document and "<w:del " in document and "commentRangeStart" in document
    assert "word/comments.xml" in zipfile.ZipFile(path).namelist()


def test_a_locked_block_is_kept_whatever_the_editor_sends_for_it(doc: Path) -> None:
    before = _paragraphs(doc)
    version, blocks = _blocks(doc)
    blocks[2] = {**blocks[2], "text": "ignored", "kind": "locked", "runs": [{"t": "nope"}]}
    save_doc_blocks(doc, version, blocks)
    assert _paragraphs(doc) == before


def test_changing_a_kind_new_blocks_and_deletions(doc: Path) -> None:
    version, blocks = _blocks(doc)
    first = _by_text(blocks, "first point")
    second = _by_text(blocks, "second point")
    first["kind"] = "number"
    blocks.remove(second)
    heading = {"kind": "heading", "level": 3, "runs": [{"t": "Appendix"}]}
    quote = {"kind": "quote", "runs": [{"t": "A quoted line"}]}
    note = {
        "kind": "paragraph",
        "runs": [{"t": "tab\there"}, {"t": " site", "link": "https://example.com/a"}],
    }
    blocks[blocks.index(first) :] = [
        first,
        heading,
        quote,
        note,
        *blocks[blocks.index(first) + 1 :],
    ]
    result = save_doc_blocks(doc, version, blocks)
    read = result["blocks"]
    assert _by_text(read, "first point")["kind"] == "number"
    assert not any("second point" in r["t"] for b in read for r in b.get("runs", []))
    assert _by_text(read, "Appendix")["level"] == 3
    assert _by_text(read, "A quoted line")["kind"] == "quote"
    linked = _by_text(read, "site")
    assert linked["runs"][1] == {"t": " site", "link": "https://example.com/a"}
    assert linked["runs"][0]["t"] == "tab\there"
    paragraph = Document(doc).paragraphs
    assert any(p.style.name == "Heading 3" and p.text == "Appendix" for p in paragraph)


def test_table_cells_rows_and_columns(doc: Path) -> None:
    version, blocks = _blocks(doc)
    table = blocks[-1]
    table["rows"][1][1] = [{"t": "10"}]
    table["rows"].append([[{"t": "c"}], [{"t": "3"}]])
    version = save_doc_blocks(doc, version, blocks)["version"]
    _, blocks = _blocks(doc)
    assert [[c[0]["t"] for c in row] for row in blocks[-1]["rows"]] == [
        ["item", "cost"],
        ["a", "10"],
        ["b", "2"],
        ["c", "3"],
    ]
    blocks[-1]["rows"] = [[row[0]] for row in blocks[-1]["rows"]]
    save_doc_blocks(doc, version, blocks)
    _, blocks = _blocks(doc)
    assert [row[0][0]["t"] for row in blocks[-1]["rows"]] == ["item", "a", "b", "c"]
    assert all(len(row) == 1 for row in blocks[-1]["rows"])
    assert len(Document(doc).tables[0].columns) == 1


def test_a_new_table_and_a_list_that_restarts(doc: Path) -> None:
    version, blocks = _blocks(doc)
    blocks.append(
        {
            "kind": "table",
            "header": True,
            "rows": [[[{"t": "x"}], [{"t": "y"}]], [[{"t": "1"}], [{"t": "2"}]]],
        }
    )
    blocks.append({"kind": "number", "level": 1, "runs": [{"t": "again"}]})
    result = save_doc_blocks(doc, version, blocks)
    assert [b["kind"] for b in result["blocks"]][-3:] == ["table", "table", "number"]


def test_a_save_over_a_newer_file_is_refused(doc: Path) -> None:
    version, blocks = _blocks(doc)
    changed = Document(doc)
    changed.add_paragraph("added in Word")
    changed.save(doc)
    with pytest.raises(DocChanged):
        save_doc_blocks(doc, version, blocks)


@pytest.mark.parametrize(
    "bad",
    [
        {"kind": "paragraph", "runs": [{"t": "x", "link": "javascript:alert(1)"}]},
        {"kind": "heading", "level": 9, "runs": []},
        {"kind": "table", "rows": [[[{"t": "a"}]], [[{"t": "b"}], [{"t": "c"}]]]},
        {"kind": "mystery"},
        "text",
    ],
)
def test_a_block_that_cannot_be_written_is_refused(doc: Path, bad: Any) -> None:
    version, blocks = _blocks(doc)
    before = doc.read_bytes()
    with pytest.raises(BlocksRefused):
        save_doc_blocks(doc, version, [*blocks, bad])
    assert doc.read_bytes() == before


def test_a_missing_style_is_said_plainly(tmp_path: Path) -> None:
    plain = Document()
    plain.add_paragraph("hello")
    quote = plain.styles["Quote"].element
    quote.getparent().remove(quote)
    path = tmp_path / "p.docx"
    plain.save(path)
    version, blocks = _blocks(path)
    blocks.append({"kind": "quote", "runs": [{"t": "q"}]})
    with pytest.raises(BlocksRefused, match="no 'Quote' style"):
        save_doc_blocks(path, version, blocks)


def test_contents_pictures_and_charts_are_locked_and_a_save_leaves_them_alone(
    tmp_path: Path,
) -> None:
    from PIL import Image

    Image.new("RGB", (40, 20), "red").save(tmp_path / "pic.png")
    tools = _tools(tmp_path)
    tools["write_docx"](path="m.docx", content="# T\n\n[TOC]\n\n## One\n\ntext\n\n## Two\n\nmore\n")
    tools["add_docx_image"](path="m.docx", image_path="pic.png", after_text="text")
    tools["add_docx_chart"](
        path="m.docx",
        chart_type="bar",
        data="| x | y |\n|---|---|\n| a | 1 |\n| b | 2 |",
        after_text="more",
    )
    path = tmp_path / "m.docx"
    before = _paragraphs(path)
    version, blocks = _blocks(path)
    locked = {b["what"] for b in blocks if b["kind"] == "locked"}
    assert {"field", "image or drawing"} <= locked
    pictured = [b for b in blocks if b["kind"] == "locked" and b.get("images")]
    assert len(pictured) == 1 and re.fullmatch(r"rId\d+", pictured[0]["images"][0])
    found = read_doc_image(path, pictured[0]["images"][0])
    assert found is not None and found[1] == "image/png" and found[0].startswith(b"\x89PNG")
    assert read_doc_image(path, "rId999") is None
    assert read_doc_image(path, "rId1") is None  # a part that is not a picture
    save_doc_blocks(path, version, blocks)
    assert _paragraphs(path) == before
    assert Document(path).paragraphs  # still opens


def _png(width: int = 40, height: int = 20) -> str:
    import io

    from PIL import Image

    out = io.BytesIO()
    Image.new("RGB", (width, height), "blue").save(out, format="PNG")
    return base64.b64encode(out.getvalue()).decode()


def test_a_picture_the_person_adds_is_written_as_a_paragraph_and_read_back_locked(
    doc: Path,
) -> None:
    version, blocks = _blocks(doc)
    before = _paragraphs(doc)
    blocks.insert(2, {"kind": "image", "data": f"data:image/png;base64,{_png()}"})
    saved = save_doc_blocks(doc, version, blocks)["blocks"]
    assert len(saved) == len(blocks)
    added = saved[2]
    assert added["kind"] == "locked" and added["what"] == "image or drawing" and added["images"]
    assert read_doc_image(doc, added["images"][0]) is not None
    assert _paragraphs(doc)[:2] == before[:2]  # what was not touched stays as it was


def test_a_wide_picture_is_scaled_to_the_page_and_a_bad_one_is_refused(doc: Path) -> None:
    version, blocks = _blocks(doc)
    blocks.append({"kind": "image", "data": _png(3000, 1000)})
    save_doc_blocks(doc, version, blocks)
    extent = next(
        iter(
            Document(doc).element.body.iter(
                "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent"
            )
        )
    )
    assert int(extent.get("cx")) == 914400 * 6 and int(extent.get("cy")) == 914400 * 2

    version, blocks = _blocks(doc)
    broken = base64.b64encode(b"\x89PNG\r\n\x1a\n" + b"\xae" * 40).decode()
    for bad in ("not base64 !!", "aGVsbG8=", broken, 5, None):
        with pytest.raises(BlocksRefused):
            save_doc_blocks(doc, version, [*blocks, {"kind": "image", "data": bad}])
    big = "A" * (4 * (8 * 1024 * 1024 // 3 + 10))
    with pytest.raises(BlocksRefused, match="8 MB"):
        save_doc_blocks(doc, version, [*blocks, {"kind": "image", "data": big}])


def test_a_divider_and_a_page_break_are_written_read_back_and_kept(doc: Path) -> None:
    version, blocks = _blocks(doc)
    blocks.insert(3, {"kind": "divider"})
    blocks.insert(5, {"kind": "pagebreak"})
    saved = save_doc_blocks(doc, version, blocks)["blocks"]
    assert [b["kind"] for b in saved[3:6]] == ["divider", saved[4]["kind"], "pagebreak"]
    assert saved[4]["kind"] != "locked"

    version, blocks = _blocks(doc)
    before = _paragraphs(doc)
    save_doc_blocks(doc, version, blocks)  # sending back what was read changes nothing
    assert _paragraphs(doc) == before
    assert [b["kind"] for b in blocks].count("divider") == 1
    assert [b["kind"] for b in blocks].count("pagebreak") == 1

    blocks = [b for b in blocks if b["kind"] not in ("divider", "pagebreak")]
    version = read_doc_blocks(doc)["version"]
    removed = save_doc_blocks(doc, version, blocks)["blocks"]
    assert not {"divider", "pagebreak"} & {b["kind"] for b in removed}


def test_a_page_break_inside_text_and_a_ruled_paragraph_with_text_are_not_breaks(
    tmp_path: Path,
) -> None:
    from docx import Document as Open

    document = Open()
    document.add_paragraph("before the break").add_run().add_break(WD_BREAK.PAGE)
    ruled = document.add_paragraph("ruled text")
    border = ruled._p.get_or_add_pPr()
    from docx.oxml import parse_xml

    border.append(
        parse_xml(
            '<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            '<w:bottom w:val="single" w:sz="6" w:space="1" w:color="auto"/></w:pBdr>'
        )
    )
    document.add_page_break()
    path = tmp_path / "b.docx"
    document.save(path)
    kinds = [(b["kind"], b.get("what")) for b in _blocks(path)[1]]
    assert kinds[0] == ("locked", "page break")
    assert kinds[1][0] == "paragraph"
    assert kinds[2] == ("pagebreak", None)


def test_a_picture_is_resized_and_described_and_nothing_else_changes(doc: Path) -> None:
    version, blocks = _blocks(doc)
    blocks.insert(3, {"kind": "image", "data": _png(400, 200)})
    version = save_doc_blocks(doc, version, blocks)["version"]
    version, blocks = _blocks(doc)
    pictured = next(b for b in blocks if b.get("picture"))
    assert pictured["picture"]["alt"] != "x" and pictured["picture"]["width"] > 0
    before = _paragraphs(doc)

    pictured["picture"] = {"width": 50, "alt": "A blue box"}
    saved = save_doc_blocks(doc, version, blocks)["blocks"]
    again = next(b for b in saved if b.get("picture"))
    assert again["picture"] == {"width": 50, "alt": "A blue box"}
    tree = etree.fromstring(zipfile.ZipFile(doc).read("word/document.xml"))
    wp = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
    extent = next(tree.iter(f"{wp}extent"))
    assert (int(extent.get("cx")), int(extent.get("cy"))) == (914400 * 3, 914400 * 3 // 2)
    assert next(tree.iter(f"{wp}docPr")).get("descr") == "A blue box"
    assert len(_paragraphs(doc)) == len(before)
    assert [p for p in _paragraphs(doc) if b"drawing" not in p] == [
        p for p in before if b"drawing" not in p
    ]


@pytest.mark.parametrize(
    "bad",
    [
        {"width": 2, "alt": ""},
        {"width": 101, "alt": ""},
        {"width": "5", "alt": ""},
        {"width": 50, "alt": "x" * 501},
        {"width": 50},
    ],
)
def test_a_picture_size_or_description_out_of_range_is_refused(doc: Path, bad: Any) -> None:
    version, blocks = _blocks(doc)
    blocks.insert(3, {"kind": "image", "data": _png()})
    save_doc_blocks(doc, version, blocks)
    version, blocks = _blocks(doc)
    next(b for b in blocks if b.get("picture"))["picture"] = bad
    with pytest.raises(BlocksRefused):
        save_doc_blocks(doc, version, blocks)


def test_a_table_the_editor_cannot_change_carries_its_cells_to_be_shown(tmp_path: Path) -> None:
    path = tmp_path / "merged.docx"
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Revenue"
    table.cell(1, 0).text = "12"
    table.cell(1, 1).text = "+3%"
    document.save(path)

    block = read_doc_blocks(path)["blocks"][0]

    assert block["what"] == "table"
    first, second = block["grid"]["rows"]
    assert first["cells"][0]["span"] == 2
    assert first["cells"][0]["paras"][0]["runs"][0]["t"] == "Revenue"
    assert [c["paras"][0]["runs"][0]["t"] for c in second["cells"]] == ["12", "+3%"]


def test_a_chart_is_a_locked_block_that_carries_its_xml_and_printed_size(tmp_path: Path) -> None:
    from coscribe.tools._docx_chart import add_chart_paragraph

    path = tmp_path / "chart.docx"
    document = Document()
    document.add_paragraph("Before")
    add_chart_paragraph(
        document, "bar", ["Q1", "Q2"], {"Revenue": [1.0, 2.0]}, "Revenue", 5000000, 2500000
    )
    document.save(path)

    block = next(b for b in read_doc_blocks(path)["blocks"] if b["kind"] == "locked")

    assert block["what"] == "image or drawing"
    assert "<c:barChart" in block["chart"]["xml"] and "Revenue" in block["chart"]["xml"]
    assert round(block["chart"]["w"]) == round(5000000 / 12700)


def test_a_paragraph_and_its_runs_carry_the_look_the_styles_and_direct_formatting_give(
    tmp_path: Path,
) -> None:
    from docx.shared import Pt, RGBColor

    path = tmp_path / "look.docx"
    document = Document()
    heading = document.add_heading("Plan", level=1)
    heading.paragraph_format.space_before = Pt(18)
    paragraph = document.add_paragraph()
    paragraph.add_run("plain ")
    big = paragraph.add_run("big")
    big.font.size = Pt(20)
    big.font.color.rgb = RGBColor(0x12, 0x34, 0x56)
    document.save(path)

    result = read_doc_blocks(path)
    head, body = result["blocks"][0], result["blocks"][1]

    assert head["look"]["before"] == 18.0 and head["look"]["size"] > 12
    assert body["runs"] == [{"t": "plain "}, {"t": "big", "c": "123456", "sz": 20.0}]
    assert result["page"]["w"] > 400 and result["page"]["base"]["size"] > 0


def test_a_table_the_editor_can_change_carries_widths_and_cell_fills(tmp_path: Path) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    path = tmp_path / "plain-table.docx"
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "a"
    table.cell(0, 1).text = "b"
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), "E4F6EF")
    table.cell(0, 0)._tc.get_or_add_tcPr().append(shade)
    document.save(path)

    block = read_doc_blocks(path)["blocks"][0]

    assert block["kind"] == "table"
    assert block["look"]["rows"][0]["cells"][0]["look"]["fill"] == "E4F6EF"
    assert len(block["look"]["widths"]) == 2


def test_a_runs_colour_size_and_font_survive_an_edit_of_its_paragraph(tmp_path: Path) -> None:
    path = tmp_path / "keep-look.docx"
    document = Document()
    document.add_paragraph("old")
    document.save(path)
    read = read_doc_blocks(path)
    block = {
        **read["blocks"][0],
        "runs": [{"t": "new", "b": True, "c": "147f68", "sz": 20, "f": "Noto Sans"}],
    }

    saved = save_doc_blocks(path, read["version"], [block])

    assert saved["blocks"][0]["runs"] == [
        {"t": "new", "b": True, "c": "147F68", "sz": 20.0, "f": "Noto Sans"}
    ]


def test_a_run_look_that_is_not_a_colour_size_or_font_is_dropped(tmp_path: Path) -> None:
    path = tmp_path / "bad-look.docx"
    Document().save(path)
    read = read_doc_blocks(path)
    block = {"kind": "paragraph", "runs": [{"t": "x", "c": "red;", "sz": 9000, "f": "A\u0007B"}]}

    saved = save_doc_blocks(path, read["version"], [block])

    assert saved["blocks"][-1]["runs"] == [{"t": "x"}]


def test_a_table_carries_what_its_style_gives_it_borders_and_a_header_row(tmp_path: Path) -> None:
    """The house table style puts the borders and the header's fill and white bold text in the
    style, not in each cell, so a reader that took only the cells' own formatting drew a bare grid."""
    target = tmp_path / "t.docx"
    _tools(tmp_path)["write_docx"](
        path="t.docx", content="# T\n\n| a | b |\n| --- | --- |\n| 1 | 2 |\n| 3 | 4 |\n"
    )

    table = next(b for b in read_doc_blocks(target)["blocks"] if b["kind"] == "table")
    look = table["look"]
    first, body = look["rows"][0]["cells"][0], look["rows"][1]["cells"][0]

    assert look["edges"]["insideH"]["c"] == "BFBFBF"
    assert first["look"]["fill"] == "1F4E79"
    assert first["p"]["color"] == "FFFFFF"
    assert "fill" not in body["look"]
    assert "color" not in body["p"]
