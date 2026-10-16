"""Previews drawn without an office program (decision 0012): a deck from the SVG pages it was
built from, a PDF with PDFium."""

from __future__ import annotations

from pathlib import Path

import pytest
from PIL import Image

from coscribe.tools import _thumbnail
from coscribe.tools._svg_preview import page_svg
from coscribe.tools.office.pptx.engine import build_deck

_PAGE = (
    '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1280 720">'
    '<rect x="0" y="0" width="1280" height="720" fill="{fill}"/>'
    '<text id="title" x="64" y="100" font-size="40">第 {n} 页</text>{body}</svg>'
)
_CHART = """<g id="revenue" data-pptx-replace-with="chart"><metadata type="application/json">
{"x": 64, "y": 160, "width": 560, "height": 480, "type": "column", "categories": ["A", "B"],
 "series": [{"name": "Q3 营收", "values": [5, 3]}]}</metadata></g>"""


@pytest.fixture(autouse=True)
def _no_libreoffice(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(_thumbnail, "find_soffice", lambda: None)


def _deck(tmp_path: Path, pages: int = 3) -> Path:
    pages_dir = tmp_path / "pages"
    pages_dir.mkdir()
    for n in range(1, pages + 1):
        fill = "#1F4E79" if n == 2 else "#FFFFFF"
        body = _CHART if n == 1 else ""
        (pages_dir / f"{n}.svg").write_text(
            _PAGE.format(fill=fill, n=n, body=body), encoding="utf-8"
        )
    out = tmp_path / "deck.pptx"
    build_deck(pages_dir, out)
    return out


def test_every_page_of_a_built_deck_is_previewed_from_its_svg(tmp_path: Path) -> None:
    deck, state = _deck(tmp_path), tmp_path / "state"

    names, reason = _thumbnail.render_all_page_previews(deck, state, max_pages=8)

    assert reason is None and len(names) == 3
    second = Image.open(state / "previews" / names[1]).convert("RGB")
    assert second.width == 1280 and second.getpixel((640, 700)) == (31, 78, 121)


def test_one_page_deep_in_the_deck_can_be_previewed(tmp_path: Path) -> None:
    deck, state = _deck(tmp_path), tmp_path / "state"

    name, reason = _thumbnail.render_single_page_preview(deck, state, 2)
    missing, why = _thumbnail.render_single_page_preview(deck, state, 9)

    assert reason is None and name is not None
    assert Image.open(state / "previews" / name).convert("RGB").getpixel((640, 700)) == (
        31,
        78,
        121,
    )
    assert missing is None and "out of range" in str(why)


def test_a_native_chart_is_a_labelled_box_of_its_size(tmp_path: Path) -> None:
    deck = _deck(tmp_path)

    svg = page_svg(deck.parent / "pages" / "1.svg")

    assert 'width="560.0"' in svg and "Column chart: Q3 营收" in svg and "metadata" not in svg


def test_a_deck_edited_since_its_build_is_not_previewed_from_stale_pages(
    tmp_path: Path,
) -> None:
    from pptx import Presentation

    deck = _deck(tmp_path)
    prs = Presentation(str(deck))
    prs.slides[0].shapes.add_textbox(0, 0, 100, 100).text = "added"
    prs.save(str(deck))

    names, reason = _thumbnail.render_all_page_previews(deck, tmp_path / "state")

    assert names == [] and reason == "LibreOffice (soffice) not found"


def test_a_pdf_is_previewed_without_an_office_program(tmp_path: Path) -> None:
    pdf = tmp_path / "doc.pdf"
    Image.new("RGB", (200, 300), "red").save(pdf)
    state = tmp_path / "state"

    thumb, reason = _thumbnail.render_thumbnail(pdf, state)
    pages, _ = _thumbnail.render_all_page_previews(pdf, state)

    assert reason is None and thumb is not None and len(pages) == 1


def test_a_page_edited_since_the_build_is_not_shown_as_the_deck(tmp_path: Path) -> None:
    deck = _deck(tmp_path, pages=1)
    (tmp_path / "pages" / "1.svg").write_text("<svg", encoding="utf-8")

    name, reason = _thumbnail.render_thumbnail(deck, tmp_path / "state")

    assert name is None and reason == "LibreOffice (soffice) not found"


def test_a_page_that_cannot_be_drawn_skips_the_preview(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from coscribe.tools import _svg_preview

    deck = _deck(tmp_path, pages=1)
    monkeypatch.setattr(_svg_preview, "page_svg", lambda _path: "<svg")

    name, reason = _thumbnail.render_thumbnail(deck, tmp_path / "state")

    assert name is None and reason == "page rendering failed"
