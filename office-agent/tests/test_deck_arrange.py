"""The Slides panel's changes to a deck's slides as a whole (design §7): a new slide, a copy, a
move and a delete each renumber the pages, keep their comments on them and rebuild the deck."""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation

from coscribe.tools.office.pptx import (
    DeckChangedError,
    add_comment,
    delete_slide,
    duplicate_slide,
    insert_slide,
    move_slide,
    read_comments,
    read_deck,
)
from coscribe.tools.office.pptx.engine import build_deck

from .test_deck_engine import CARDS, COVER

NEW = (
    '<section style="background:#FFFFFF"><h1 id="title" '
    'style="left:96px;top:96px;width:1200px;height:100px">New</h1></section>'
)


def _deck(tmp_path: Path) -> Path:
    folder = tmp_path / "pages"
    folder.mkdir()
    (folder / "01-cover.svg").write_text(COVER, encoding="utf-8")
    (folder / "02-cards.svg").write_text(CARDS, encoding="utf-8")
    assert build_deck(folder, tmp_path / "deck.pptx", resource_root=tmp_path).ok
    return tmp_path / "deck.pptx"


def _names(deck: Path) -> list[str]:
    return [s["name"] for s in read_deck(deck)["slides"]]


def test_a_new_slide_goes_after_the_current_one(tmp_path: Path) -> None:
    deck = _deck(tmp_path)
    assert insert_slide(deck, 0, NEW, resource_root=tmp_path) == 1
    assert _names(deck) == ["01-cover.svg", "02-slide.html", "03-cards.svg"]
    assert len(Presentation(str(deck)).slides) == 3 and not read_deck(deck)["changed"]
    assert insert_slide(deck, -1, NEW, resource_root=tmp_path) == 0
    assert _names(deck)[0] == "01-slide.html"


def test_a_copy_a_move_and_a_delete_keep_comments_on_their_slides(tmp_path: Path) -> None:
    deck = _deck(tmp_path)
    fields = {"x": 0.5, "y": 0.5, "element": "", "author": "ana", "sent": False}
    add_comment(deck, slide="02-cards", text="on the cards", **fields)
    add_comment(deck, slide="01-cover", text="on the cover", **fields)

    assert duplicate_slide(deck, 0, resource_root=tmp_path) == 1
    assert _names(deck) == ["01-cover.svg", "02-cover-copy.svg", "03-cards.svg"]
    assert move_slide(deck, 2, 0, resource_root=tmp_path) == 0
    assert _names(deck) == ["01-cards.svg", "02-cover.svg", "03-cover-copy.svg"]
    assert {c["text"]: c["slide"] for c in read_comments(deck)} == {
        "on the cards": "01-cards",
        "on the cover": "02-cover",
    }
    assert delete_slide(deck, 1, resource_root=tmp_path) == 1
    assert _names(deck) == ["01-cards.svg", "02-cover-copy.svg"]
    assert [c["text"] for c in read_comments(deck)] == ["on the cards"]
    assert not list((tmp_path / "pages").glob(".arranging-*"))


def test_the_last_slide_stays_and_a_deck_changed_outside_is_left_alone(tmp_path: Path) -> None:
    deck = _deck(tmp_path)
    delete_slide(deck, 1, resource_root=tmp_path)
    with pytest.raises(ValueError):
        delete_slide(deck, 0, resource_root=tmp_path)
    with pytest.raises(IndexError):
        move_slide(deck, 5, 0, resource_root=tmp_path)
    prs = Presentation(str(deck))
    prs.slides[0].shapes.add_textbox(0, 0, 10, 10).text = "typed in PowerPoint"
    prs.save(str(deck))
    with pytest.raises(DeckChangedError):
        insert_slide(deck, 0, NEW, resource_root=tmp_path)
    assert _names(deck) == ["01-cover.svg"]
