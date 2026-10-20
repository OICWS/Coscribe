"""An SVG page converted to a section compiles to the same slide (decision 0016).

The user's first edit turns a slide from its SVG page into a section, so the conversion may change
nothing: every text at the same place with the same runs, every other shape at the same box. The
pages here cover what the model writes: titles and paragraphs, cards in translated groups,
gradients and shadows, native charts and tables, pictures, icons, notes and a background.
"""

from __future__ import annotations

import collections
import re
import zipfile
from pathlib import Path

import pytest
from lxml import etree
from PIL import Image

from coscribe.tools.office.pptx.engine import (
    build_deck,
    compile_pages,
    compile_sections,
    convert_page,
    convert_slide,
)
from coscribe.tools.office.pptx.engine.build import built_pages
from coscribe.tools.office.pptx.engine.section_format import SectionError, parse_section

from .test_deck_engine import CARDS, COVER

_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
_SVG = (
    'xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1280 720" font-family="Microsoft YaHei, Arial"'
)
_CHART = (
    '{"x": 640, "y": 160, "width": 560, "height": 400, "type": "column", "categories": ["华东",'
    ' "华北"], "series": [{"name": "营收", "values": [5410, 2120]}], "font_size": 18}'
)
_TABLE = (
    '{"x": 64, "y": 420, "width": 520, "height": 200, "schema": "ppt-master.semantic-table.v2",'
    ' "columns": ["区域", "营收"],'
    ' "rows": [["华东", {"text": "5,410", "align": "r", "bold": true}],'
    ' ["华北", {"text": "2,120", "align": "r", "color": "#B91C1C"}], ["合计", "7,530"]],'
    ' "column_widths": [2, 1], "style": {"font_size": 16, "header_fill": "#0F172A",'
    ' "band_fill": "#F1F5F9", "border_color": "#E5E7EB"}}'
)
OBJECTS = f"""<svg {_SVG}>
  <rect x="0" y="0" width="1280" height="720" fill="#F7F5F0"/>
  <desc id="notes">先讲收入。
  再讲区域。</desc>
  <text id="title" x="64" y="90" font-size="36" font-weight="bold" fill="#1E2A36">区域表现</text>
  <text x="1216" y="690" font-size="14" fill="#6B7280" text-anchor="end"
    xml:space="preserve">  02 / 06</text>
  <use data-icon="ship" x="64" y="120" width="40" height="40" stroke="#C2410C"/>
  <image href="photo.png" x="64" y="180" width="480" height="220"
    preserveAspectRatio="xMidYMid slice"/>
  <g id="revenue" data-pptx-replace-with="chart">
    <metadata type="application/json">{_CHART}</metadata></g>
  <g id="regions" data-pptx-replace-with="table">
    <metadata type="application/json">{_TABLE}</metadata></g>
</svg>"""


def _texts(deck: Path) -> list[list[tuple[str, int, int, str]]]:
    """Per slide: each text's words, its alignment point and the run properties."""
    slides = []
    for root in _slides(deck):
        items = []
        for sp in root.iter(f"{_P}sp"):
            text = "".join(t.text or "" for t in sp.iter(f"{_A}t")).replace(" ", " ")
            if not text.strip():
                continue
            off, ext = sp.find(f".//{_A}off"), sp.find(f".//{_A}ext")
            ppr = sp.find(f".//{_A}pPr")
            algn = ppr.get("algn", "l") if ppr is not None else "l"
            x, w = int(off.get("x")), int(ext.get("cx"))
            anchor = {"l": x, "ctr": x + w // 2, "r": x + w}.get(algn, x)
            runs = ";".join(
                f"{r.get('sz')},{r.get('b')},{r.get('i')},{r.get('spc')},"
                + "".join(c.get("val", "") for c in r.iter(f"{_A}srgbClr"))
                for r in sp.iter(f"{_A}rPr")
            )
            items.append(
                (text.strip(), round(anchor / 3175), round(int(off.get("y")) / 3175), runs)
            )
        slides.append(sorted(items))
    return slides


def _shapes(deck: Path) -> list[collections.Counter[tuple[str, int, int, int, int]]]:
    """Per slide: every shape that is not text, by kind and box, to a third of a page pixel."""
    slides = []
    for root in _slides(deck):
        found: collections.Counter[tuple[str, int, int, int, int]] = collections.Counter()
        for el in root.iter(f"{_P}sp", f"{_P}pic", f"{_P}graphicFrame", f"{_P}cxnSp"):
            if "".join(t.text or "" for t in el.iter(f"{_A}t")).strip():
                continue
            off, ext = el.find(f".//{_A}off"), el.find(f".//{_A}ext")
            box = (int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy")))
            found[(etree.QName(el).localname, *(round(v / 3175) for v in box))] += 1
        slides.append(found)
    return slides


def _slides(deck: Path) -> list[etree._Element]:
    with zipfile.ZipFile(deck) as package:
        names = [n for n in package.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)]
        names.sort(key=lambda n: int(re.findall(r"\d+", n)[0]))
        return [etree.fromstring(package.read(n)) for n in names]


def _both(tmp_path: Path, *pages: str) -> tuple[Path, Path, list[str]]:
    Image.new("RGB", (960, 440), (30, 120, 200)).save(tmp_path / "photo.png")
    named = [(f"{i:02d}.svg", page) for i, page in enumerate(pages, start=1)]
    conversions = [convert_page(page) for _name, page in named]
    sections = [
        (name.replace(".svg", ".html"), section)
        for (name, _), (section, _kept) in zip(named, conversions, strict=True)
    ]
    original = compile_pages(
        named, tmp_path / "page.pptx", resource_root=tmp_path, svg_dir=tmp_path
    )
    converted = compile_sections(
        sections, tmp_path / "section.pptx", resource_root=tmp_path, sections_dir=tmp_path
    )
    assert original.ok and converted.ok, (original.errors(), converted.errors())
    kept = [k for _section, notes in conversions for k in notes]
    return tmp_path / "page.pptx", tmp_path / "section.pptx", kept


def test_a_converted_page_compiles_to_the_same_slide(tmp_path: Path) -> None:
    page, section, kept = _both(tmp_path, COVER, CARDS, OBJECTS)
    assert kept == []
    assert _texts(section) == _texts(page)
    assert _shapes(section) == _shapes(page)


def test_texts_become_editable_and_the_rest_keeps_its_kind(tmp_path: Path) -> None:
    section = parse_section(convert_page(OBJECTS)[0])
    tags = [el.tag for el, _style in section.elements]
    assert tags.count("p") == 2 and "x-chart" in tags and "table" in tags
    assert "img" in tags and "x-icon" in tags
    assert section.background == "F7F5F0"
    assert section.notes == ["先讲收入。", "再讲区域。"]
    page, converted, _kept = _both(tmp_path, OBJECTS)
    from pptx import Presentation

    slide = Presentation(str(converted)).slides[0]
    assert slide.shapes.title.text_frame.text == "区域表现"
    assert slide.notes_slide.notes_text_frame.text == "先讲收入。\n再讲区域。"


def test_spaces_at_a_line_start_are_kept(tmp_path: Path) -> None:
    section = convert_page(OBJECTS)[0]
    assert "  02 / 06" in section


def test_a_rotated_text_stays_a_drawing_and_is_reported(tmp_path: Path) -> None:
    page = f"""<svg {_SVG}><text id="title" x="64" y="90" font-size="36">标题</text>
      <text x="300" y="300" font-size="24" transform="rotate(30 300 300)">斜写</text></svg>"""
    _section, kept = convert_page(page)
    assert any("rotated" in k for k in kept)
    original, converted, _kept = _both(tmp_path, page)
    assert _shapes(converted) == _shapes(original)


def test_a_slide_on_a_template_is_not_converted() -> None:
    with pytest.raises(SectionError, match="template"):
        convert_page(f'<svg {_SVG} data-template-slide="2"><text x="1" y="1">a</text></svg>')


def test_a_deck_converted_one_slide_at_a_time_stays_the_same(tmp_path: Path) -> None:
    pages = tmp_path / "pages"
    pages.mkdir()
    Image.new("RGB", (960, 440), (30, 120, 200)).save(pages / "photo.png")
    for i, page in enumerate((COVER, CARDS, OBJECTS), start=1):
        (pages / f"{i:02d}.svg").write_text(page, encoding="utf-8")
    deck = tmp_path / "deck.pptx"
    assert build_deck(pages, deck, resource_root=tmp_path).ok
    texts, shapes = _texts(deck), _shapes(deck)
    assert built_pages(deck) is not None

    section, kept = convert_slide(pages / "02.svg")
    assert section.name == "02.html" and kept == []
    assert not (pages / "02.svg").exists()
    report = build_deck(pages, deck, resource_root=tmp_path)
    assert report.ok, [e for p in report.pages for e in p.errors]
    assert [p.name for p in report.pages] == ["01.svg", "02.html", "03.svg"]
    assert report.changed == ["02.html"]
    assert _texts(deck) == texts and _shapes(deck) == shapes
    # The preview draws SVG pages; a deck with a section is drawn from the .pptx instead.
    assert built_pages(deck) is None


def test_a_slide_is_one_file_in_a_mixed_folder(tmp_path: Path) -> None:
    (tmp_path / "01.svg").write_text(COVER, encoding="utf-8")
    (tmp_path / "01.html").write_text(convert_page(COVER)[0], encoding="utf-8")
    with pytest.raises(ValueError, match="a slide is one file"):
        build_deck(tmp_path, tmp_path / "deck.pptx")
    with pytest.raises(ValueError, match="exists already"):
        convert_slide(tmp_path / "01.svg")
    assert (tmp_path / "01.svg").exists()


SHAPES = f"""<svg {_SVG}>
  <rect x="0" y="0" width="1280" height="720" fill="#0F3D3E"/>
  <circle cx="640" cy="360" r="200" fill="none" stroke="#B45F3A" stroke-width="2"/>
  <ellipse cx="300" cy="500" rx="120" ry="60" fill="#1D4ED8" opacity="0.5"/>
  <rect x="64" y="120" width="300" height="160" rx="16" fill="#FFFFFF" stroke="#E5E7EB"
    stroke-width="2"/>
  <line x1="64" y1="300" x2="500" y2="300" stroke="#F59E0B" stroke-width="4"/>
  <line x1="700" y1="100" x2="900" y2="250" stroke="#F59E0B" stroke-width="3"/>
  <line x1="700" y1="600" x2="900" y2="450" stroke="#F59E0B" stroke-width="3"/>
  <g id="card" transform="translate(800, 400)">
    <rect x="0" y="0" width="360" height="200" rx="12" fill="#FFFFFF"/>
    <rect x="24" y="24" width="8" height="48" fill="#C2410C"/>
    <text x="48" y="60" font-size="24" fill="#1E2A36">被保护的一天</text>
  </g>
  <g id="shadowed"><rect x="64" y="400" width="200" height="100" fill="#FFFFFF" filter="url(#s)"/>
    <defs><filter id="s"><feDropShadow dx="0" dy="4" stdDeviation="4" flood-color="#000000"
      flood-opacity="0.3"/></filter></defs></g>
  <rect x="10" y="10" width="100" height="40" fill="url(#g)"/>
  <defs><linearGradient id="g"><stop offset="0" stop-color="#000"/>
    <stop offset="1" stop-color="#FFF"/></linearGradient></defs>
</svg>"""


def test_plain_shapes_become_the_sections_own_shapes_in_the_same_place(tmp_path: Path) -> None:
    section = convert_page(SHAPES)[0]
    tags = [el.tag for el, _style in parse_section(section).elements]
    kinds = re.findall(r'data-shape="([\w-]+)"', section)
    # The ring, the ellipse, the card, three lines (the one going up stays a drawing), the card
    # group's two shapes: everything but the up-going line, the shadowed and the gradient shape.
    assert kinds == ["ellipse", "ellipse", "round-rect", "line", "line", "round-rect", "rect"]
    drawings = [t for t in tags if t.endswith("}svg")]
    assert len(drawings) == 4  # the defs, the up-going line, the shadowed group, the gradient shape
    page, converted, _kept = _both(tmp_path, SHAPES)
    assert _shapes(converted) == _shapes(page)
    assert _texts(converted) == _texts(page)


def _paints(deck: Path) -> list[list[tuple[str, str, str, str]]]:
    """Per slide: each shape's geometry, fill and line (colour, width) and alpha, in order."""
    slides = []
    for root in _slides(deck):
        found = []
        for el in root.iter(f"{_P}sp", f"{_P}cxnSp"):
            if "".join(t.text or "" for t in el.iter(f"{_A}t")).strip():
                continue
            sp_pr = el.find(f"{_P}spPr")
            geom = sp_pr.find(f"{_A}prstGeom")
            ln = sp_pr.find(f"{_A}ln")
            fill = "".join(
                c.get("val", "") for c in sp_pr.findall(f"{_A}solidFill/{_A}srgbClr")
            ) or ("none" if sp_pr.find(f"{_A}noFill") is not None else "")
            alpha = ",".join(a.get("val", "") for a in sp_pr.iter(f"{_A}alpha"))
            line = "" if ln is None else f"{ln.get('w')}:" + "".join(
                c.get("val", "") for c in ln.iter(f"{_A}srgbClr")
            )
            found.append((geom.get("prst") if geom is not None else "", fill, line, alpha))
        slides.append(found)
    return slides


def test_converted_shapes_keep_their_paint(tmp_path: Path) -> None:
    page, converted, _kept = _both(tmp_path, SHAPES)
    assert sorted(_paints(converted)[0]) == sorted(_paints(page)[0])
