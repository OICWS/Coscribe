"""A deck built from pages on the user's own template: its layouts and theme carry over, a page
can start from an example slide's decoration, and the build tool describes the template first."""

from __future__ import annotations

import io
import zipfile
from pathlib import Path
from typing import Any

import pytest
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Emu, Inches

from coscribe.tools.office.pptx.engine import compile_pages
from coscribe.tools.office.pptx.engine.template import describe
from coscribe.tools.office.pptx.engine.tool import build_engine_tools

_NS = (
    'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
)
_SVG = 'xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1280 720" font-family="Microsoft YaHei"'


def _template(path: Path) -> None:
    """Two example slides: a cover with a navy panel and a logo, a content slide with a band."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    cover = prs.slides.add_slide(prs.slide_layouts[0])
    panel = cover.shapes.add_shape(1, 0, 0, Emu(4572000), prs.slide_height)
    panel.name = "panel"
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor(0x1E, 0x27, 0x61)
    logo = io.BytesIO()
    Image.new("RGB", (40, 20), "#C2410C").save(logo, "PNG")
    cover.shapes.add_picture(logo, Inches(11), Inches(0.3), Inches(1), Inches(0.5)).name = "logo"
    content = prs.slides.add_slide(prs.slide_layouts[1])
    band = content.shapes.add_shape(1, 0, Emu(6720840), prs.slide_width, Emu(137160))
    band.name = "band"
    band.fill.solid()
    band.fill.fore_color.rgb = RGBColor(0x1E, 0x27, 0x61)
    prs.save(str(path))


def _page(root: str = "", body: str = "") -> str:
    title = '<text id="title" x="64" y="100" font-size="40">第三季度汇报</text>'
    return f"<svg {_SVG} {root}>{title}{body}</svg>"


def test_the_description_gives_canvas_theme_and_each_example_slides_decoration(
    tmp_path: Path,
) -> None:
    _template(tmp_path / "t.pptx")

    found = describe(tmp_path / "t.pptx")

    assert found["viewBox"] == "0 0 1280 720"
    assert found["colours"]["accent1"]  # type: ignore[index]
    cover, content = found["example_slides"]  # type: ignore[misc]
    assert cover["layout"] == "Title Slide"
    assert cover["decoration"][0] == {
        "kind": "shape",
        "x": 0,
        "y": 0,
        "width": 480,
        "height": 720,
        "fill": "#1E2761",
    }
    assert cover["decoration"][1]["kind"] == "picture"
    assert content["decoration"][0]["y"] == 706


def test_a_page_starts_from_an_example_slides_decoration(tmp_path: Path) -> None:
    _template(tmp_path / "t.pptx")
    out = tmp_path / "deck.pptx"

    report = compile_pages(
        [("cover", _page('data-template-slide="1"')), ("body", _page('data-template-slide="2"'))],
        out,
        template=tmp_path / "t.pptx",
    )

    assert report.ok, report.errors()
    cover, body = Presentation(str(out)).slides
    assert [s.name for s in cover.shapes] == ["panel", "logo", "title"]
    assert cover.slide_layout.name == "Title Slide"
    assert [s.name for s in body.shapes] == ["band", "title"]
    with zipfile.ZipFile(out) as package:
        assert len([n for n in package.namelist() if n.startswith("ppt/slides/slide")]) == 2
    logo = next(s for s in cover.shapes if s.name == "logo")
    assert Image.open(io.BytesIO(logo.image.blob)).getpixel((0, 0)) == (0xC2, 0x41, 0x0C)


def test_a_page_takes_a_layout_by_name_and_a_wrong_name_lists_them(tmp_path: Path) -> None:
    _template(tmp_path / "t.pptx")
    out = tmp_path / "deck.pptx"

    report = compile_pages(
        [("a", _page('data-layout="blank"')), ("b", _page('data-layout="Nope"'))],
        out,
        template=tmp_path / "t.pptx",
    )

    assert Presentation(str(out)).slides[0].slide_layout.name == "Blank"
    (error,) = report.errors()
    assert error.startswith('b: no layout named "Nope" in the template; it has "Title Slide"')


@pytest.mark.parametrize(
    ("root", "template", "message"),
    [
        ('data-template-slide="3"', True, "the template has 2 slides"),
        ('data-template-slide="1"', False, "needs a template"),
        ("4:3", True, "differs from the deck's 1280x720"),
    ],
)
def test_wrong_template_references_are_errors(
    tmp_path: Path, root: str, template: bool, message: str
) -> None:
    _template(tmp_path / "t.pptx")
    page = _page(root) if root != "4:3" else _page().replace("0 0 1280 720", "0 0 960 720")

    report = compile_pages(
        [("p", page)], tmp_path / "deck.pptx", template=tmp_path / "t.pptx" if template else None
    )

    assert any(message in error for error in report.errors()), report.errors()


def test_a_potx_template_is_read(tmp_path: Path) -> None:
    _template(tmp_path / "t.pptx")
    with (
        zipfile.ZipFile(tmp_path / "t.pptx") as source,
        zipfile.ZipFile(tmp_path / "t.potx", "w") as target,
    ):
        for item in source.infolist():
            data = source.read(item)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"presentation.main+xml", b"template.main+xml")
            target.writestr(item, data)

    assert len(describe(tmp_path / "t.potx")["example_slides"]) == 2  # type: ignore[arg-type]
    report = compile_pages(
        [("p", _page('data-template-slide="2"'))],
        tmp_path / "deck.pptx",
        template=tmp_path / "t.potx",
    )
    assert report.ok, report.errors()
    with zipfile.ZipFile(tmp_path / "deck.pptx") as package:
        assert b"presentationml.presentation.main+xml" in package.read("[Content_Types].xml")
    assert [s.name for s in Presentation(str(tmp_path / "deck.pptx")).slides[0].shapes] == [
        "band",
        "title",
    ]


def test_the_tool_describes_the_template_first_and_then_builds_on_it(tmp_path: Path) -> None:
    _template(tmp_path / "t.pptx")
    (tool,) = build_engine_tools(tmp_path)

    first: dict[str, Any] = tool("pages", "deck.pptx", template="t.pptx")
    assert not (tmp_path / "deck.pptx").exists()
    (tmp_path / "pages").mkdir()
    (tmp_path / "pages" / "01.svg").write_text(_page('data-template-slide="2"'), encoding="utf-8")
    built: dict[str, Any] = tool("pages", "deck.pptx", template="t.pptx")

    assert first["template"]["viewBox"] == "0 0 1280 720"
    assert built["ok"], built
    assert "band" in {s.name for s in Presentation(str(tmp_path / "deck.pptx")).slides[0].shapes}


def _rich_template(path: Path) -> None:
    """One example slide with a picture background, a text box linking to a web page, a chart
    and a rectangle, with the low shape ids a slide has once its placeholders are gone."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    picture = io.BytesIO()
    Image.new("RGB", (16, 9), "#1E2761").save(picture, "PNG")
    _image, rel_id = slide.part.get_or_add_image_part(picture)
    slide.element.cSld.insert(
        0,
        etree.fromstring(
            f'<p:bg {_NS}><p:bgPr><a:blipFill><a:blip r:embed="{rel_id}"/><a:stretch>'
            "<a:fillRect/></a:stretch></a:blipFill><a:effectLst/></p:bgPr></p:bg>"
        ),
    )
    link = slide.shapes.add_textbox(Inches(1), Inches(6.5), Inches(4), Inches(0.5))
    link.name = "link"
    run = link.text_frame.paragraphs[0].add_run()
    run.text = "example.com"
    run.hyperlink.address = "https://example.com/"
    data = CategoryChartData()
    data.categories = ["Q1", "Q2"]
    data.add_series("sales", (1, 2))
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6), Inches(1), Inches(4), Inches(3), data
    ).name = "chart"
    slide.shapes.add_shape(1, 0, 0, Inches(1), Inches(1)).name = "corner"
    prs.save(str(path))


def test_a_copied_slide_resolves_every_relationship_and_leaves_out_what_it_cannot_copy(
    tmp_path: Path,
) -> None:
    _rich_template(tmp_path / "t.pptx")
    out = tmp_path / "deck.pptx"

    report = compile_pages(
        [("p", _page('data-template-slide="1"'))], out, template=tmp_path / "t.pptx"
    )

    assert report.ok, report.errors()
    (warning,) = report.pages[0].warnings
    assert '"chart"' in warning
    slide = Presentation(str(out)).slides[0]
    assert [s.name for s in slide.shapes] == ["link", "corner", "title"]
    rels = slide.part.rels
    referenced = {
        value
        for node in slide.element.iter()
        for name, value in node.attrib.items()
        if name.startswith("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}")
    }
    assert referenced and referenced <= set(rels)
    background = slide.element.cSld.bg.xpath(".//a:blip/@r:embed")[0]
    assert Image.open(io.BytesIO(slide.part.related_part(background).blob)).size == (16, 9)
    run = slide.shapes[0].text_frame.paragraphs[0].runs[0]
    assert run.hyperlink.address == "https://example.com/"


def test_copied_shapes_get_ids_of_their_own(tmp_path: Path) -> None:
    _rich_template(tmp_path / "t.pptx")
    out = tmp_path / "deck.pptx"
    body = '<rect id="box" x="100" y="200" width="300" height="100" fill="#1E2761"/>'

    compile_pages(
        [("p", _page('data-template-slide="1"', body))], out, template=tmp_path / "t.pptx"
    )

    ids = [shape.shape_id for shape in Presentation(str(out)).slides[0].shapes]
    assert len(ids) == 4
    assert len(set(ids)) == len(ids)
