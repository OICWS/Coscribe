"""convert_office_file: legacy formats up to OOXML, and PDF export."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from coscribe.runtime import get_tool_metadata
from coscribe.tools import _office_com
from coscribe.tools.convert import build_convert_tools


def _convert(root: Path):  # type: ignore[no-untyped-def]
    (tool,) = build_convert_tools(root)
    return tool


def _fake_office(monkeypatch: pytest.MonkeyPatch, program: str | None = "Word.Application") -> None:
    def convert_(source: Path, target: Path, to_format: str, timeout: float) -> str | None:
        if program is not None:
            target.write_bytes(b"%PDF fake" if to_format == "pdf" else b"converted")
        return program

    monkeypatch.setattr(_office_com, "convert", convert_)


def test_the_tool_writes_files_so_it_is_gated() -> None:
    (tool,) = build_convert_tools(Path("."))
    metadata = get_tool_metadata(tool)

    assert metadata.risk_category == "WRITE_LOCAL"
    assert metadata.requires_approval is True


def test_unsupported_source_and_target_are_refused_with_the_options(tmp_path: Path) -> None:
    (tmp_path / "notes.txt").write_text("x", encoding="utf-8")
    Document().save(str(tmp_path / "a.docx"))
    convert = _convert(tmp_path)

    with pytest.raises(ValueError, match="isn't a type"):
        convert(path="notes.txt", to_format="pdf")
    with pytest.raises(ValueError, match="can become: pdf"):
        convert(path="a.docx", to_format="xlsx")
    with pytest.raises(ValueError, match="does not exist"):
        convert(path="missing.doc", to_format="docx")


def test_an_existing_output_is_never_replaced(tmp_path: Path) -> None:
    Document().save(str(tmp_path / "a.docx"))
    (tmp_path / "a.pdf").write_bytes(b"keep me")
    convert = _convert(tmp_path)

    with pytest.raises(FileExistsError, match="already exists"):
        convert(path="a.docx", to_format="pdf")
    assert (tmp_path / "a.pdf").read_bytes() == b"keep me"


def test_without_an_office_program_it_says_one_is_needed(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    Document().save(str(tmp_path / "a.docx"))
    _fake_office(monkeypatch, program=None)

    with pytest.raises(ValueError, match="needs Microsoft Office"):
        _convert(tmp_path)(path="a.docx", to_format="pdf")


def test_a_legacy_file_becomes_ooxml_next_to_it_and_the_original_stays(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _fake_office(monkeypatch, "Excel.Application")
    (tmp_path / "old.xls").write_bytes(b"legacy")

    result = _convert(tmp_path)(path="old.xls", to_format="xlsx")

    assert result["path"] == "old.xlsx" and result["converted_by"] == "Excel.Application"
    assert (tmp_path / "old.xlsx").read_bytes() == b"converted"
    assert (tmp_path / "old.xls").exists()


def test_a_docx_exports_to_a_pdf_at_the_chosen_path(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _fake_office(monkeypatch)
    Document().save(str(tmp_path / "a.docx"))

    result = _convert(tmp_path)(path="a.docx", to_format="pdf", output_path="out/合同.pdf")

    assert result["path"] == "out/合同.pdf"
    assert (tmp_path / "out" / "合同.pdf").read_bytes().startswith(b"%PDF")
