"""Containers in a section (decision 0018): an `<x-stack>` is laid out into the plain positioned
elements the rest of the engine reads, so a deck built from it holds no containers."""

from __future__ import annotations

from pathlib import Path

import pytest

from coscribe.tools.office.pptx.engine import build_deck
from coscribe.tools.office.pptx.engine.section_format import SectionError, parse_section


def _boxes(source: str) -> list[tuple[str, dict[str, float]]]:
    found = []
    for elem, style in parse_section(source).elements:
        box = dict(zip(("x", "y", "w", "h"), style.box(), strict=True))
        found.append((elem.tag.rsplit("}", 1)[-1], box))
    return found


def _row(*cards: str, extra: str = "") -> str:
    return (
        '<section><x-stack data-layout="row" style="left:100px;top:200px;width:1700px;'
        f'height:400px;gap:40px;padding:30px;{extra}">{"".join(cards)}</x-stack></section>'
    )


CARD = (
    '<x-stack style="flex-grow:1;gap:12px;padding:20px;background:#FFFFFF">'
    "<h3>{}</h3><p>{}</p></x-stack>"
)


def test_cards_in_a_row_share_the_room_and_lay_their_text_out() -> None:
    boxes = _boxes(_row(CARD.format("Books", "Short."), CARD.format("Topics", "Short.")))
    shapes = [b for tag, b in boxes if tag == "x-shape"]
    assert [round(b["w"]) for b in shapes] == [800, 800]
    assert shapes[0]["x"] == 130 and shapes[1]["x"] == 970  # padding 30, then card, gap 40, card
    assert all(b["y"] == 230 and b["h"] == 340 for b in shapes)
    texts = [b for tag, b in boxes if tag in {"h3", "p"}]
    assert texts[0]["x"] == 150 and texts[0]["w"] == 760  # the card's padding is inside it
    assert texts[1]["y"] == pytest.approx(texts[0]["y"] + texts[0]["h"] + 12)


def test_deleting_a_card_gives_the_others_its_room() -> None:
    three = _boxes(_row(*[CARD.format(f"Card {i}", "Short.") for i in range(3)]))
    two = _boxes(_row(*[CARD.format(f"Card {i}", "Short.") for i in range(2)]))
    width = lambda boxes: [round(b["w"]) for tag, b in boxes if tag == "x-shape"]  # noqa: E731
    assert width(three) == [520, 520, 520]
    assert width(two) == [800, 800]


def test_a_fixed_child_keeps_its_width_and_the_others_take_what_is_left() -> None:
    source = _row('<x-shape style="width:200px;background:#111111"/>', CARD.format("A", "Short."))
    shapes = [b for tag, b in _boxes(source) if tag == "x-shape"]
    assert shapes[0]["w"] == 200 and shapes[0]["h"] == 340  # stretched across the row
    assert shapes[1]["x"] == 130 + 200 + 40 and round(shapes[1]["w"]) == 1640 - 240


def test_text_in_a_column_takes_the_height_its_words_need() -> None:
    long = "Every book gets a folder of its own with the PDF and the code that goes with it. " * 6
    short = CARD.format("A", "Short.")
    tall = CARD.format("A", long)
    p_short = [b for tag, b in _boxes(_row(short)) if tag == "p"][0]
    p_tall = [b for tag, b in _boxes(_row(tall)) if tag == "p"][0]
    assert p_tall["h"] > 2 * p_short["h"]


def test_a_container_without_a_height_is_as_tall_as_its_content() -> None:
    source = (
        '<section><x-stack style="left:0px;top:0px;width:800px;height:900px;gap:10px">'
        '<x-stack style="padding:20px;background:#EEEEEE"><p>One line.</p></x-stack>'
        "<p>After it.</p></x-stack></section>"
    )
    boxes = _boxes(source)
    card = boxes[0][1]
    after = [b for tag, b in boxes if tag == "p"][-1]
    assert card["h"] == pytest.approx(40 + 28 * 1.35)  # padding, then one line
    assert after["y"] == pytest.approx(card["y"] + card["h"] + 10)


def test_justify_and_align_place_what_is_left_over() -> None:
    source = (
        '<section><x-stack data-layout="row" style="left:0px;top:0px;width:1000px;height:200px;'
        'justify:between;align:center">'
        '<x-shape style="width:100px;height:50px;background:#111111"/>'
        '<x-shape style="width:100px;height:50px;background:#111111"/></x-stack></section>'
    )
    first, second = (b for _, b in _boxes(source))
    assert first["x"] == 0 and second["x"] == 900
    assert first["y"] == 75  # centred across the row


def test_children_shrink_when_they_do_not_fit() -> None:
    source = _row(
        '<x-shape style="width:1000px;background:#111111"/>',
        '<x-shape style="width:1000px;background:#222222"/>',
    )
    widths = [b["w"] for tag, b in _boxes(source) if tag == "x-shape"]
    assert sum(widths) + 40 == pytest.approx(1640)  # the room, less the gap
    assert widths[0] == pytest.approx(widths[1])


def test_a_card_is_a_shape_behind_its_children() -> None:
    boxes = _boxes(_row(CARD.format("A", "Short.")))
    assert [tag for tag, _ in boxes] == ["x-shape", "h3", "p"]


def test_a_container_that_cannot_be_laid_out_says_why() -> None:
    with pytest.raises(SectionError, match="needs left, top, width and height"):
        parse_section("<section><x-stack><p>x</p></x-stack></section>")
    with pytest.raises(SectionError, match="needs a width or flex-grow"):
        parse_section(_row("<p>No width in a row.</p>"))
    with pytest.raises(SectionError, match="data-layout is row or column"):
        parse_section(_row("").replace("row", "diagonal"))
    with pytest.raises(SectionError, match="not allowed on a container"):
        parse_section(_row("", extra="color:#FF0000;"))
    with pytest.raises(SectionError, match="not text"):
        parse_section(_row("loose words"))


def test_a_deck_built_from_containers_holds_plain_shapes(tmp_path: Path) -> None:
    folder = tmp_path / "pages"
    folder.mkdir()
    (folder / "01-cards.html").write_text(
        _row(*[CARD.format(f"Card {i}", "Short.") for i in range(3)]), encoding="utf-8"
    )
    result = build_deck(folder, tmp_path / "deck.pptx", resource_root=tmp_path)
    assert result.ok, result
    from pptx import Presentation

    shapes = Presentation(str(tmp_path / "deck.pptx")).slides[0].shapes
    assert len(shapes) == 9  # three cards of a rectangle, a heading and a paragraph


def test_the_editor_is_given_the_laid_out_elements(tmp_path: Path) -> None:
    from coscribe.tools.office.pptx import slide_parts

    folder = tmp_path / "pages"
    folder.mkdir()
    (folder / "01-cards.html").write_text(
        _row(*[CARD.format(f"Card {i}", "Short.") for i in range(2)]), encoding="utf-8"
    )
    assert build_deck(folder, tmp_path / "deck.pptx", resource_root=tmp_path).ok
    parts = slide_parts(tmp_path / "deck.pptx", 0, tmp_path)["parts"]
    assert [p["kind"] for p in parts].count("text") == 4
    assert all(p["box"] is not None for p in parts)


def test_the_editor_is_given_one_part_per_element_in_document_order(tmp_path: Path) -> None:
    from coscribe.tools.office.pptx import slide_parts

    folder = tmp_path / "pages"
    folder.mkdir()
    (folder / "01-cards.html").write_text(
        '<section><p style="left:0px;top:0px;width:300px;height:40px">Free</p>'
        + _row(CARD.format("A", "Short."), CARD.format("B", "Short."), extra="background:#EEEEEE;")[
            len("<section>") :
        ],
        encoding="utf-8",
    )
    assert build_deck(folder, tmp_path / "deck.pptx", resource_root=tmp_path).ok
    parts = slide_parts(tmp_path / "deck.pptx", 0, tmp_path)["parts"]
    # The free text, the row (with its card), then each card (with its card) and its two texts.
    assert [p["kind"] for p in parts] == [
        "text",
        "drawing",
        "drawing",
        "text",
        "text",
        "drawing",
        "text",
        "text",
    ]
    assert parts[1]["picture"] is not None and all(p["box"] for p in parts)
    row, card = parts[1]["box"], parts[2]["box"]
    assert card[0] > row[0] and card[2] < row[2]  # the card sits inside the row
