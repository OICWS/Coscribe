"""edit_docx_structure: structural edits to an existing Word file, everything else untouched."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
from docx import Document

from coscribe.tools.office.document_tools import build_document_tools

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

_HANDBOOK = (
    "Staff handbook\n\nThis handbook sets out how we work.\n\n"
    "# Onboarding\n\n- A laptop on day one.\n\n"
    "# Working hours\n\n- Core hours are 10 to 16.\n\n## Overtime\n\nAgreed in advance.\n\n"
    "# Safety\n\n- Report accidents within one hour.\n\n"
    "# Contacts\n\n| Who | Ext |\n| --- | --- |\n| HR | 305 |"
)


def _tools(root: Path) -> dict[str, Any]:
    tools = {tool.__name__: tool for tool in build_document_tools(root)}  # type: ignore[attr-defined]
    tools["write_docx"](path="h.docx", content=_HANDBOOK)
    return tools


def _outline(path: Path) -> list[tuple[str, str]]:
    return [(p.style.name, p.text) for p in Document(str(path)).paragraphs]


def test_insert_after_a_section_uses_the_documents_styles(tmp_path: Path) -> None:
    tools = _tools(tmp_path)

    result = tools["edit_docx_structure"](
        path="h.docx",
        operations=[
            {
                "op": "insert",
                "after_section": "Working hours",
                "content": "# Remote work\n\n- Two days a week.",
            }
        ],
    )

    outline = _outline(tmp_path / "h.docx")
    assert result["done"] == ["operation 1 (insert): inserted 2 block(s)"]
    assert result["document"]["headings"][2:4] == ["## Overtime", "# Remote work"]
    assert result["document"]["tables"] == 1 and not result["document"]["contents_list"]
    position = outline.index(("Heading 1", "Remote work"))
    assert outline[position - 1] == ("Normal", "Agreed in advance.")
    assert outline[position + 1] == ("List Bullet", "Two days a week.")
    assert outline[position + 2] == ("Heading 1", "Safety")


def test_a_section_moves_with_its_subsections_and_the_rest_stays(tmp_path: Path) -> None:
    tools = _tools(tmp_path)
    before = _outline(tmp_path / "h.docx")

    tools["edit_docx_structure"](
        path="h.docx",
        operations=[{"op": "move_section", "section": "Working hours", "after_section": "Safety"}],
    )

    after = _outline(tmp_path / "h.docx")
    assert sorted(after) == sorted(before)
    headings = [text for style, text in after if style.startswith("Heading")]
    assert headings == ["Onboarding", "Safety", "Working hours", "Overtime", "Contacts"]
    assert Document(str(tmp_path / "h.docx")).tables[0].rows[1].cells[0].text == "HR"


def test_delete_section_removes_the_heading_and_everything_under_it(tmp_path: Path) -> None:
    tools = _tools(tmp_path)

    tools["edit_docx_structure"](
        path="h.docx", operations=[{"op": "delete_section", "section": "Working hours"}]
    )

    texts = [text for _, text in _outline(tmp_path / "h.docx")]
    assert "Working hours" not in texts and "Overtime" not in texts
    assert "Agreed in advance." not in texts and "Safety" in texts


def test_a_contents_list_names_every_heading_once_under_its_title(tmp_path: Path) -> None:
    tools = _tools(tmp_path)

    tools["edit_docx_structure"](
        path="h.docx",
        operations=[
            {"op": "insert", "after": "This handbook", "content": "# Contents\n\n[TOC]"},
            {"op": "move_section", "section": "Safety", "before": "Onboarding"},
        ],
    )

    outline = _outline(tmp_path / "h.docx")
    title = outline.index(("TOC Heading", "Contents"))
    entries = [text for _, text in outline[title + 1 : title + 6]]
    assert entries == ["Safety", "Onboarding", "Working hours", "Overtime", "Contacts"]


def test_tracked_insert_and_delete_use_ids_above_the_documents_own(tmp_path: Path) -> None:
    tools = _tools(tmp_path)
    tools["replace_docx_text"](path="h.docx", find="day one", replace="the first day")

    tools["edit_docx_structure"](
        path="h.docx",
        track_changes=True,
        operations=[
            {"op": "insert", "after_section": "Safety", "content": "# Remote work"},
            {"op": "delete_section", "section": "Onboarding"},
        ],
    )

    body = Document(str(tmp_path / "h.docx")).element.body
    ids = [e.get(f"{_W}id") for e in body.iter(f"{_W}ins", f"{_W}del")]
    assert len(ids) == len(set(ids)) > 2
    assert "Onboarding" in "".join(t.text or "" for t in body.iter(f"{_W}delText"))


@pytest.mark.parametrize(
    ("operation", "message"),
    [
        ({"op": "insert", "content": "x", "after": "Nowhere"}, "no paragraph reads"),
        ({"op": "insert", "content": "x"}, "exactly one of"),
        ({"op": "move_section", "section": "Safety", "after": "Report accidents"}, "inside"),
        ({"op": "move_section", "section": "Report accidents", "after": "Onboarding"}, "heading"),
        ({"op": "rename"}, "op must be one of"),
    ],
)
def test_a_bad_operation_saves_nothing(
    tmp_path: Path, operation: dict[str, str], message: str
) -> None:
    tools = _tools(tmp_path)
    original = (tmp_path / "h.docx").read_bytes()

    with pytest.raises(ValueError, match=message):
        tools["edit_docx_structure"](
            path="h.docx",
            operations=[{"op": "delete_section", "section": "Contacts"}, operation],
        )

    assert (tmp_path / "h.docx").read_bytes() == original


def test_tracking_refuses_a_move_and_a_tables_deletion(tmp_path: Path) -> None:
    tools = _tools(tmp_path)

    with pytest.raises(ValueError, match="can't be tracked"):
        tools["edit_docx_structure"](
            path="h.docx",
            track_changes=True,
            operations=[{"op": "move_section", "section": "Safety", "before": "Onboarding"}],
        )
    with pytest.raises(ValueError, match="holds a table"):
        tools["edit_docx_structure"](
            path="h.docx",
            track_changes=True,
            operations=[{"op": "delete_section", "section": "Contacts"}],
        )


def test_output_path_leaves_the_original_and_new_tables_match_its_style(tmp_path: Path) -> None:
    tools = _tools(tmp_path)
    document = Document(str(tmp_path / "h.docx"))
    document.tables[0].style = document.styles["Light Grid Accent 1"]
    document.save(str(tmp_path / "h.docx"))
    original = (tmp_path / "h.docx").read_bytes()

    tools["edit_docx_structure"](
        path="h.docx",
        output_path="h2.docx",
        operations=[
            {"op": "insert", "after_section": "Safety", "content": "| A | B |\n| --- | --- |"}
        ],
    )

    assert (tmp_path / "h.docx").read_bytes() == original
    tables = Document(str(tmp_path / "h2.docx")).tables
    assert [t.style.name for t in tables] == ["Light Grid Accent 1"] * 2
    (new,) = [t for t in tables if t.rows[0].cells[0].text == "A"]
    assert not any(run.bold for cell in new.rows[0].cells for run in cell.paragraphs[0].runs)
    with pytest.raises(FileExistsError):
        tools["edit_docx_structure"](
            path="h.docx",
            output_path="h2.docx",
            operations=[{"op": "delete_section", "section": "Safety"}],
        )


def test_the_structure_view_shows_what_the_markdown_leaves_out(tmp_path: Path) -> None:
    tools = _tools(tmp_path)
    document = Document(str(tmp_path / "h.docx"))
    document.sections[0].header.paragraphs[0].text = "Acme Ltd"
    document.add_paragraph().add_run().add_picture(str(_png(tmp_path)))
    document.save(str(tmp_path / "h.docx"))
    tools["edit_docx_structure"](
        path="h.docx",
        operations=[{"op": "insert", "after": "This handbook", "content": "# Contents\n\n[TOC]"}],
    )

    lines = tools["read_docx"](path="h.docx", view="structure").splitlines()

    assert lines[:4] == [
        "[Normal] Staff handbook",
        "[Normal] This handbook sets out how we work.",
        "[TOC Heading] Contents",
        "contents list (5 entries)",
    ]
    assert "[List Bullet] A laptop on day one." in lines
    assert any(line.startswith("table 2x2 [") and line.endswith("]: Who | Ext") for line in lines)
    assert any("+ picture" in line for line in lines)
    assert lines[-1].startswith("section 1: page ") and "header: Acme Ltd" in lines[-1]
    with pytest.raises(ValueError, match="structure"):
        tools["read_docx"](path="h.docx", view="xml")


def _png(tmp_path: Path) -> Path:
    from PIL import Image

    path = tmp_path / "p.png"
    Image.new("RGB", (40, 20), "white").save(path)
    return path
