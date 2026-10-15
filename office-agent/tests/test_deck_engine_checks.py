"""The deck engine's checks and builds (`tools/office/pptx/engine/checker.py`, `build.py`).

Each check is shown flagging the defect it exists for and staying silent on the look-alike that
model-written pages are full of: a large number with its unit beside it, a translucent tint
behind coloured text, a watermark digit behind a card.
"""

from __future__ import annotations

import json
from decimal import Decimal
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation

from coscribe.tools.office.pptx.engine import (
    MANIFEST,
    DeckChangedError,
    build_deck,
    check_deck,
    compile_pages,
    page_files,
    source_text,
)
from coscribe.tools.office.pptx.engine.checker import (
    is_claim,
    is_known,
    known_numbers,
    number_values,
)

_SVG = (
    'xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1280 720" font-family="Microsoft YaHei, Arial"'
)
_TITLE = '<text id="title" x="64" y="80" font-size="32" fill="#1E2A36">{}</text>'


def _page(body: str) -> str:
    return f"<svg {_SVG}>{_TITLE.format('标题')}{body}</svg>"


def _findings(
    tmp_path: Path, body: str, sources: list[str] | None = None
) -> tuple[list[str], list[str]]:
    report = compile_pages([("P01", _page(body))], tmp_path / "deck.pptx")
    assert report.ok, report.errors()
    check_deck(report, sources)
    page = report.pages[0]
    return page.errors, page.warnings


def test_text_past_the_page_edge_is_an_error_naming_the_text(tmp_path: Path) -> None:
    errors, _ = _findings(
        tmp_path,
        '<text x="1100" y="400" font-size="40" fill="#1E2A36">营收同比增长百分之十二</text>',
    )

    assert len(errors) == 1 and "营收同比增长" in errors[0] and "right edge" in errors[0]


def test_text_inside_the_page_is_not_flagged(tmp_path: Path) -> None:
    errors, _ = _findings(
        tmp_path, '<text x="64" y="400" font-size="40" fill="#1E2A36">营收同比增长百分之十二</text>'
    )

    assert errors == []


def test_two_texts_drawn_over_each_other_are_an_error(tmp_path: Path) -> None:
    errors, _ = _findings(
        tmp_path,
        '<text x="100" y="300" font-size="28" fill="#1E2A36">华东区域增长最快</text>'
        '<text x="160" y="306" font-size="28" fill="#1E2A36">华北唯一负增长</text>',
    )

    assert len(errors) == 1 and "drawn over each other" in errors[0]


def test_a_large_number_with_its_unit_and_label_is_not_an_overlap(tmp_path: Path) -> None:
    # Measured on a research deck: these boxes overlap, their glyphs do not.
    errors, _ = _findings(
        tmp_path,
        '<text x="410" y="190" font-size="13" fill="#5B6472">Q3 营收</text>'
        '<text x="410" y="244" font-size="64" font-weight="bold" fill="#C25E1A">12,710</text>'
        '<text x="632" y="244" font-size="20" fill="#5B6472">万元</text>',
    )

    assert errors == []


def test_text_nobody_can_read_is_an_error(tmp_path: Path) -> None:
    errors, _ = _findings(
        tmp_path,
        '<rect x="0" y="0" width="1280" height="160" fill="#0B2A4A"/>'
        '<text x="88" y="328" font-size="40" fill="#FFFFFF">第三季度经营回顾</text>',
    )

    assert len(errors) == 1 and "cannot be read" in errors[0] and "#FFFFFF on #FFFFFF" in errors[0]


def test_faint_text_is_a_warning_listing_the_worst(tmp_path: Path) -> None:
    errors, warnings = _findings(
        tmp_path,
        '<text x="64" y="600" font-size="14" fill="#94A3B8">数据口径：Q2 与 Q3</text>'
        '<text x="64" y="640" font-size="14" fill="#C3CAD6">第三季度经营回顾</text>',
    )

    assert errors == []
    assert len(warnings) == 1
    assert warnings[0].index("第三季度") < warnings[0].index("数据口径")


def test_a_translucent_tint_is_blended_before_contrast_is_measured(tmp_path: Path) -> None:
    # Blue text on a 12% blue tint: unreadable if the tint were taken as solid blue.
    errors, warnings = _findings(
        tmp_path,
        '<rect x="64" y="240" width="400" height="80" fill="#1A73E8" fill-opacity="0.12"/>'
        '<text x="84" y="292" font-size="28" font-weight="bold" fill="#1A73E8">离座锁屏</text>',
    )

    assert errors == [] and warnings == []


def test_text_over_a_picture_is_left_to_a_person(tmp_path: Path) -> None:
    Image.new("RGB", (64, 20), "#FFFFFF").save(tmp_path / "white.png")
    report = compile_pages(
        [
            (
                "P01",
                _page(
                    '<image href="white.png" x="0" y="200" width="1280" height="400"/>'
                    '<text x="64" y="400" font-size="40" fill="#FFFFFF">照片上的白字</text>'
                ),
            )
        ],
        tmp_path / "deck.pptx",
        resource_root=tmp_path,
    )
    check_deck(report)

    assert report.pages[0].errors == [] and report.pages[0].warnings == []


def test_a_watermark_digit_behind_a_card_is_neither_unreadable_nor_an_overlap(
    tmp_path: Path,
) -> None:
    errors, _ = _findings(
        tmp_path,
        '<rect x="0" y="0" width="1280" height="720" fill="#0F2A44"/>'
        '<text x="320" y="448" font-size="84" font-weight="bold" fill="#FFFFFF" '
        'fill-opacity="0.06" text-anchor="end">01</text>'
        '<text x="110" y="400" font-size="18" fill="#F2B94C">追回北辰制造推迟至 Q4 的订单</text>',
    )

    assert errors == []


def _table(columns: list[int], label: str, y: int, rows: int) -> str:
    table = {
        "x": 64,
        "y": y,
        "width": sum(columns),
        "height": 36 * rows,
        "schema": "ppt-master.semantic-table.v2",
        "columns": ["行业", "增加值"],
        "column_widths": columns,
        "rows": [[label, "5,410"]] * (rows - 1),
        "style": {"font_size": 17},
    }
    return (
        '<g id="industries" data-pptx-replace-with="table">'
        f'<metadata type="application/json">{json.dumps(table, ensure_ascii=False)}</metadata></g>'
        '<text x="64" y="690" font-size="14" fill="#1E2A36">数据来源：国家统计局</text>'
    )


def test_a_table_whose_wrapped_cells_reach_the_footer_is_an_error(tmp_path: Path) -> None:
    errors, _ = _findings(tmp_path, _table([150, 120], "信息传输、软件和信息技术服务业", 380, 8))

    (error,) = errors
    assert "table 'industries'" in error and "数据来源" in error


def test_a_table_whose_cells_fit_is_not_flagged(tmp_path: Path) -> None:
    errors, _ = _findings(tmp_path, _table([420, 120], "信息传输、软件和信息技术服务业", 380, 8))

    assert errors == []


def test_after_two_repair_rounds_the_build_says_to_stop(tmp_path: Path) -> None:
    pages = tmp_path / "pages"
    pages.mkdir()
    broken = '<text x="1100" y="400" font-size="40" fill="#1E2A36">营收同比增长百分之十二</text>'
    (pages / "1.svg").write_text(_page(broken), encoding="utf-8")

    reports = [build_deck(pages, tmp_path / "deck.pptx") for _ in range(3)]

    assert [r.ok for r in reports] == [False, False, False]
    assert not reports[1].warnings
    assert "stop repairing" in reports[2].warnings[0]
    (pages / "1.svg").write_text(_page(""), encoding="utf-8")
    assert build_deck(pages, tmp_path / "deck.pptx").ok
    manifest = json.loads((pages / MANIFEST).read_text(encoding="utf-8"))
    assert manifest["failing_builds"] == 0


def test_a_number_no_source_has_is_a_warning(tmp_path: Path) -> None:
    body = (
        '<text x="64" y="300" font-size="24" fill="#1E2A36">合计 12,710 万元，2024 年</text>'
        '<text x="64" y="360" font-size="24" fill="#1E2A36">第 3 页</text>'
    )

    _, warnings = _findings(tmp_path, body, sources=["Q3 营收合计 12710 万元"])
    _, unchecked = _findings(tmp_path, body)

    assert len(warnings) == 1 and "2024" in warnings[0] and "12710" not in warnings[0]
    assert unchecked == []


def test_source_text_reads_the_office_formats_and_plain_text(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("段落 5410")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "表格 3380"
    document.save(str(tmp_path / "a.docx"))
    book = Workbook()
    book.active.append(["华北", 2870])
    book.save(str(tmp_path / "b.xlsx"))
    (tmp_path / "c.md").write_text("# 640 万元", encoding="utf-8")
    compile_pages([("P01", _page(""))], tmp_path / "d.pptx")

    texts = {name: source_text(tmp_path / name) for name in ("a.docx", "b.xlsx", "c.md", "d.pptx")}

    assert "5410" in texts["a.docx"] and "3380" in texts["a.docx"]
    assert "2870" in texts["b.xlsx"]
    assert "640" in texts["c.md"]
    assert "标题" in texts["d.pptx"]


def _pages(folder: Path, names: list[str]) -> None:
    folder.mkdir(exist_ok=True)
    for name in names:
        (folder / name).write_text(
            f"<svg {_SVG}>{_TITLE.format(name)}</svg>",
            encoding="utf-8",
        )


def test_pages_are_built_in_file_name_order_with_numbers_compared_as_numbers(
    tmp_path: Path,
) -> None:
    _pages(tmp_path / "pages", ["10-end.svg", "2-body.svg", "1-cover.svg"])
    (tmp_path / "pages" / "notes.txt").write_text("not a page", encoding="utf-8")

    report = build_deck(tmp_path / "pages", tmp_path / "deck.pptx")

    titles = [s.shapes.title.text for s in Presentation(str(tmp_path / "deck.pptx")).slides]
    assert titles == ["1-cover.svg", "2-body.svg", "10-end.svg"]
    assert [p.name for p in page_files(tmp_path / "pages")] == titles
    assert report.ok and report.changed == titles


def test_a_rebuild_reports_only_the_pages_that_changed(tmp_path: Path) -> None:
    pages = tmp_path / "pages"
    _pages(pages, ["1.svg", "2.svg"])
    build_deck(pages, tmp_path / "deck.pptx")
    (pages / "2.svg").write_text(
        _page('<text x="64" y="300" fill="#1E2A36">新</text>'), encoding="utf-8"
    )

    report = build_deck(pages, tmp_path / "deck.pptx")

    assert report.changed == ["2.svg"]
    manifest = json.loads((pages / MANIFEST).read_text(encoding="utf-8"))
    assert [entry["page"] for entry in manifest["pages"]] == ["1.svg", "2.svg"]


def test_a_deck_changed_after_its_build_is_not_overwritten_unless_asked(tmp_path: Path) -> None:
    pages, deck = tmp_path / "pages", tmp_path / "deck.pptx"
    _pages(pages, ["1.svg", "2.svg"])
    build_deck(pages, deck)
    prs = Presentation(str(deck))
    prs.slides[1].shapes.title.left += 100
    prs.save(str(deck))

    with pytest.raises(DeckChangedError, match="slides 2 differ"):
        build_deck(pages, deck)
    assert Presentation(str(deck)).slides[1].shapes.title.left > 0

    assert build_deck(pages, deck, replace_changes=True).ok


def test_a_file_these_pages_did_not_build_is_not_overwritten(tmp_path: Path) -> None:
    pages, deck = tmp_path / "pages", tmp_path / "theirs.pptx"
    _pages(pages, ["1.svg"])
    Presentation().save(str(deck))

    with pytest.raises(DeckChangedError, match="not built from these pages"):
        build_deck(pages, deck)


def test_an_image_is_found_next_to_the_pages_but_not_outside_the_workspace(tmp_path: Path) -> None:
    workspace = tmp_path / "work"
    pages = workspace / "deck-pages"
    pages.mkdir(parents=True)
    Image.new("RGB", (40, 20), "#C00000").save(workspace / "photo.png")
    Image.new("RGB", (40, 20), "#00C000").save(tmp_path / "secret.png")
    picture = '<image href="{}" x="10" y="200" width="400" height="200"/>'
    (pages / "1.svg").write_text(_page(picture.format("../photo.png")), encoding="utf-8")
    (pages / "2.svg").write_text(_page(picture.format("../../secret.png")), encoding="utf-8")

    report = build_deck(pages, workspace / "deck.pptx", resource_root=workspace)

    assert report.pages[0].errors == []
    assert report.pages[1].errors


def test_an_unreadable_source_is_reported_and_the_build_goes_on(tmp_path: Path) -> None:
    _pages(tmp_path / "pages", ["1.svg"])
    (tmp_path / "broken.docx").write_bytes(b"not a zip")

    report = build_deck(
        tmp_path / "pages", tmp_path / "deck.pptx", sources=[tmp_path / "broken.docx"]
    )

    assert report.ok
    assert len(report.warnings) == 1 and "broken.docx" in report.warnings[0]


def test_a_picture_replaced_after_the_build_counts_as_a_change(tmp_path: Path) -> None:
    pages, deck = tmp_path / "pages", tmp_path / "deck.pptx"
    pages.mkdir()
    Image.new("RGB", (40, 20), "#C00000").save(pages / "photo.png")
    (pages / "1.svg").write_text(
        _page('<image href="photo.png" x="10" y="200" width="400" height="200"/>'),
        encoding="utf-8",
    )
    build_deck(pages, deck, resource_root=pages)
    prs = Presentation(str(deck))
    picture = next(s for s in prs.slides[0].shapes if s.shape_type == 13)
    image_part = prs.slides[0].part.related_part(picture._element.blipFill.blip.rEmbed)
    Image.new("RGB", (40, 20), "#0000C0").save(tmp_path / "other.png")
    # The same relationship to new bytes: what replacing a picture's content in place leaves.
    image_part._blob = (tmp_path / "other.png").read_bytes()
    prs.save(str(deck))

    with pytest.raises(DeckChangedError):
        build_deck(pages, deck, resource_root=pages)


def test_a_broken_build_record_counts_as_no_build(tmp_path: Path) -> None:
    pages, deck = tmp_path / "pages", tmp_path / "deck.pptx"
    _pages(pages, ["1.svg"])
    build_deck(pages, deck)
    (pages / MANIFEST).write_text('{"deck": "x", "pages": [{"page": "1.svg"}]}', encoding="utf-8")

    with pytest.raises(DeckChangedError, match="not built from these pages"):
        build_deck(pages, deck)
    assert build_deck(pages, deck, replace_changes=True).ok


def test_xml_powerpoint_would_reject_is_an_error(tmp_path: Path) -> None:
    out = tmp_path / "deck.pptx"
    # A run that ends in a space is written with `xml:space`, which is not reported.
    spaced = '<text x="64" y="300" font-size="24">增长 <tspan font-weight="bold">12%</tspan></text>'
    report = compile_pages([("P01", _page("")), ("P02", _page(spaced))], out)
    prs = Presentation(str(out))
    # A transition before the slide's content: the schema wants it after.
    sld = prs.slides[0]._element
    sld.insert(
        0,
        sld.makeelement(
            "{http://schemas.openxmlformats.org/presentationml/2006/main}transition", {}
        ),
    )
    prs.save(str(out))

    check_deck(report)

    assert any("file format" in e for e in report.pages[0].errors)
    assert not any("file format" in e for e in report.pages[1].errors)


def test_a_count_in_another_unit_is_known_and_a_rounded_one_is_not() -> None:
    known = known_numbers("出口 261.5 万辆；产能 15 万辆；GDP 1401879 亿元")

    def known_as(text: str) -> bool:
        ((number, value),) = number_values(text)
        return is_known(number, value, known)

    for written in ["2.615 million", "2.615m", "2,615,000", "261.5", "150k", "150,000"]:
        assert known_as(written), written
    assert not known_as("140 万亿元")
    assert number_values("100km") == [("100", Decimal(100))]


def test_a_unit_spelled_out_as_a_number_is_not_a_claim() -> None:
    assert not is_claim("10000") and not is_claim("1000")
    assert is_claim("10001") and is_claim("2000") and is_claim("100.0")
    assert not is_claim("31")
