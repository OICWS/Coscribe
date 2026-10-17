# ruff: noqa: E402
"""Web tests: the before and after preview on an approval card."""

from pathlib import Path

import pytest

pytest.importorskip("langgraph", reason="needs the langgraph_spike extra installed")

from docx import Document
from langchain_core.messages import (
    AIMessage,
)
from openpyxl import Workbook, load_workbook
from pptx import Presentation

from .helpers import (
    FakeToolCallingChatModel,
    _client_lg,
    _libreoffice_actually_works,
    _receive_until,
    _tool_call,
    _write_test_deck,
)


@pytest.mark.real_libreoffice
@pytest.mark.skipif(
    not _libreoffice_actually_works(),
    reason="LibreOffice not installed/functional in this environment",
)
def test_approval_required_carries_a_before_after_preview_for_a_pptx_edit(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """web/session.py's _build_document_edit_preview dry-runs the exact same
    edit_pptx_shape call against a throwaway copy of the deck and renders
    both states, so the approval card can show what the edit will
    actually produce instead of a raw JSON arguments dump (see
    ROADMAP.md's approval-preview entry). End-to-end through a real WS
    connection and a real LibreOffice/poppler conversion -- not mocked."""
    _write_test_deck(tmp_path / "workspace")
    call = _tool_call(
        "call_1",
        "edit_pptx_shape",
        {"path": "deck.pptx", "slide": 1, "shape_index": 0, "fill_color": "38BDF8"},
    )
    fake_model = FakeToolCallingChatModel(
        responses=[AIMessage(content="", tool_calls=[call]), AIMessage(content="done")]
    )
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect("/ws/t_pptx_preview") as ws:
            ws.receive_json()  # state
            ws.receive_json()  # history
            ws.send_json({"type": "user_message", "text": "recolor that shape"})

            approval = ws.receive_json()
            assert approval["type"] == "approval_required"
            assert approval["tool_name"] == "edit_pptx_shape"
            before_name = approval["before_preview"]
            after_name = approval["after_preview"]
            assert before_name is not None
            assert after_name is not None
            assert before_name != after_name

            # The dry run must never touch the real file -- still
            # un-filled at this point, approval hasn't happened yet.
            real_deck = Presentation(str(tmp_path / "workspace" / "deck.pptx"))
            assert real_deck.slides[0].shapes[0].fill.type is None or str(
                real_deck.slides[0].shapes[0].fill.type
            ).startswith("BACKGROUND")

            ws.send_json({"type": "approval_response", "id": approval["id"], "approved": True})
            _receive_until(ws, "tasks_changed")

    previews_dir = tmp_path / "state" / "previews"
    assert (previews_dir / before_name).is_file()
    assert (previews_dir / after_name).is_file()


def test_approval_preview_is_absent_for_a_non_pptx_tool(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """write_file has nothing to do with PresentationToolkit at all --
    _build_document_edit_preview must fall through to (None, None) instantly,
    not attempt (and fail at) treating its arguments as a pptx edit. No
    LibreOffice needed for this one, so it isn't gated behind the
    real_libreoffice marker."""
    call = _tool_call("call_1", "write_file", {"path": "note.txt", "content": "hi"})
    fake_model = FakeToolCallingChatModel(
        responses=[AIMessage(content="", tool_calls=[call]), AIMessage(content="done")]
    )
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect("/ws/t_no_preview") as ws:
            ws.receive_json()  # state
            ws.receive_json()  # history
            ws.send_json({"type": "user_message", "text": "write hi to note.txt"})

            approval = ws.receive_json()
            assert approval["type"] == "approval_required"
            assert approval["before_preview"] is None
            assert approval["after_preview"] is None

            ws.send_json({"type": "approval_response", "id": approval["id"], "approved": False})
            _receive_until(ws, "agent_message")


def test_approval_preview_is_absent_when_the_target_file_does_not_exist_yet(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """edit_pptx_shape on a path that doesn't exist yet has no "before"
    to show -- _build_document_edit_preview must recognize that up front
    (real_path.is_file() is False) and return (None, None), not attempt a
    dry run that would just fail. No LibreOffice needed: this returns
    before ever shelling out to soffice."""
    call = _tool_call(
        "call_1",
        "edit_pptx_shape",
        {"path": "missing.pptx", "slide": 1, "shape_index": 0, "fill_color": "38BDF8"},
    )
    fake_model = FakeToolCallingChatModel(
        responses=[AIMessage(content="", tool_calls=[call]), AIMessage(content="done")]
    )
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect("/ws/t_missing_preview") as ws:
            ws.receive_json()  # state
            ws.receive_json()  # history
            ws.send_json({"type": "user_message", "text": "recolor that shape"})

            approval = ws.receive_json()
            assert approval["type"] == "approval_required"
            assert approval["before_preview"] is None
            assert approval["after_preview"] is None

            ws.send_json({"type": "approval_response", "id": approval["id"], "approved": False})
            _receive_until(ws, "agent_message")


@pytest.mark.real_libreoffice
@pytest.mark.skipif(
    not _libreoffice_actually_works(),
    reason="LibreOffice not installed/functional in this environment",
)
def test_approval_preview_covers_write_docx_overwriting_an_existing_file(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """_build_document_edit_preview's dispatch table covers docx too,
    dispatched via DocumentToolkit -- write_docx regenerates the *whole*
    document from markdown-like `content`, so calling it against an
    already-existing file (this scenario: revising a one-paragraph status
    report) is exactly the "before/after" case worth previewing. A fresh
    write_docx with no existing target still correctly gets no preview
    (nothing to diff against) -- covered separately below, not repeated
    here."""
    workspace = tmp_path / "workspace"
    workspace.mkdir(parents=True, exist_ok=True)
    original = Document()
    original.add_paragraph("Status: on track.")
    original.save(str(workspace / "status.docx"))

    call = _tool_call(
        "call_1",
        "write_docx",
        {
            "path": "status.docx",
            "content": "Status: **delayed** -- see risks below.",
            "overwrite": True,
        },
    )
    fake_model = FakeToolCallingChatModel(
        responses=[AIMessage(content="", tool_calls=[call]), AIMessage(content="done")]
    )
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect("/ws/t_docx_preview") as ws:
            ws.receive_json()  # state
            ws.receive_json()  # history
            ws.send_json({"type": "user_message", "text": "update the status report"})

            approval = ws.receive_json()
            assert approval["type"] == "approval_required"
            assert approval["tool_name"] == "write_docx"
            before_name = approval["before_preview"]
            after_name = approval["after_preview"]
            assert before_name is not None
            assert after_name is not None
            assert before_name != after_name

            # Untouched until approved.
            real_doc = Document(str(workspace / "status.docx"))
            assert real_doc.paragraphs[0].text == "Status: on track."

            ws.send_json({"type": "approval_response", "id": approval["id"], "approved": True})
            _receive_until(ws, "tasks_changed")

    previews_dir = tmp_path / "state" / "previews"
    assert (previews_dir / before_name).is_file()
    assert (previews_dir / after_name).is_file()
    updated_doc = Document(str(workspace / "status.docx"))
    assert "delayed" in updated_doc.paragraphs[0].text


@pytest.mark.real_libreoffice
@pytest.mark.skipif(
    not _libreoffice_actually_works(),
    reason="LibreOffice not installed/functional in this environment",
)
def test_approval_preview_covers_write_xlsx_cells(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The approval preview's dispatch covers xlsx too, through SpreadsheetToolkit's
    write_xlsx_cells, dry-run on a copy of the existing workbook. Scenario: bolding and
    red-colouring a budget overage cell in an existing sheet."""
    workspace = tmp_path / "workspace"
    workspace.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Budget"
    sheet["A1"] = "Category"
    sheet["B1"] = "Spent"
    sheet["A2"] = "Marketing"
    sheet["B2"] = 15200
    workbook.save(str(workspace / "budget.xlsx"))

    call = _tool_call(
        "call_1",
        "write_xlsx_cells",
        {
            "path": "budget.xlsx",
            "blocks": [
                {
                    "sheet": "Budget",
                    "at": "B2",
                    "rows": [[15200]],
                    "style": {"bold": True, "color": "FF0000"},
                }
            ],
            "overwrite": True,
        },
    )
    fake_model = FakeToolCallingChatModel(
        responses=[AIMessage(content="", tool_calls=[call]), AIMessage(content="done")]
    )
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect("/ws/t_xlsx_preview") as ws:
            ws.receive_json()  # state
            ws.receive_json()  # history
            ws.send_json({"type": "user_message", "text": "flag the marketing overage in red bold"})

            approval = ws.receive_json()
            assert approval["type"] == "approval_required"
            assert approval["tool_name"] == "write_xlsx_cells"
            before_name = approval["before_preview"]
            after_name = approval["after_preview"]
            assert before_name is not None
            assert after_name is not None
            assert before_name != after_name

            real_workbook = load_workbook(str(workspace / "budget.xlsx"))
            assert real_workbook["Budget"]["B2"].font.bold is not True

            ws.send_json({"type": "approval_response", "id": approval["id"], "approved": True})
            _receive_until(ws, "tasks_changed")

    previews_dir = tmp_path / "state" / "previews"
    assert (previews_dir / before_name).is_file()
    assert (previews_dir / after_name).is_file()
    updated_workbook = load_workbook(str(workspace / "budget.xlsx"))
    assert updated_workbook["Budget"]["B2"].font.bold is True


def test_approval_preview_is_absent_for_a_brand_new_write_docx(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """write_docx creating a file that doesn't exist yet has no "before"
    -- same "nothing to diff against" reasoning fill_pptx_template gets,
    now confirmed for docx's own create-vs-overwrite ambiguity (write_docx
    is *both* a create and an overwrite tool depending on whether the
    target already exists). No LibreOffice needed -- this returns before
    ever shelling out to soffice."""
    call = _tool_call(
        "call_1",
        "write_docx",
        {"path": "new_report.docx", "content": "Hello.", "overwrite": True},
    )
    fake_model = FakeToolCallingChatModel(
        responses=[AIMessage(content="", tool_calls=[call]), AIMessage(content="done")]
    )
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect("/ws/t_new_docx") as ws:
            ws.receive_json()  # state
            ws.receive_json()  # history
            ws.send_json({"type": "user_message", "text": "write a new report"})

            approval = ws.receive_json()
            assert approval["type"] == "approval_required"
            assert approval["before_preview"] is None
            assert approval["after_preview"] is None

            ws.send_json({"type": "approval_response", "id": approval["id"], "approved": True})
            _receive_until(ws, "tasks_changed")


def test_approval_preview_leaves_the_real_file_untouched_on_denial(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Denying an edit_pptx_shape call must leave the real file exactly as
    it was -- true by construction for the real *execution* (a denied
    HumanInTheLoopMiddleware decision never runs the tool at all), but
    this is specifically about the *preview* mechanism: confirms the dry
    run itself (which runs regardless of what the human eventually
    decides, since the preview is built before the approval prompt is
    even sent) never wrote through to the real path. No LibreOffice
    needed for the assertion that matters here (the file bytes are
    unchanged); the preview images themselves are covered elsewhere."""
    _write_test_deck(tmp_path / "workspace")
    original_bytes = (tmp_path / "workspace" / "deck.pptx").read_bytes()
    call = _tool_call(
        "call_1",
        "edit_pptx_shape",
        {"path": "deck.pptx", "slide": 1, "shape_index": 0, "fill_color": "38BDF8"},
    )
    fake_model = FakeToolCallingChatModel(
        responses=[AIMessage(content="", tool_calls=[call]), AIMessage(content="denied, ok")]
    )
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect("/ws/t_deny_preview") as ws:
            ws.receive_json()  # state
            ws.receive_json()  # history
            ws.send_json({"type": "user_message", "text": "recolor that shape"})

            approval = ws.receive_json()
            assert approval["type"] == "approval_required"

            ws.send_json({"type": "approval_response", "id": approval["id"], "approved": False})
            _receive_until(ws, "agent_message")

    assert (tmp_path / "workspace" / "deck.pptx").read_bytes() == original_bytes


@pytest.mark.real_libreoffice
@pytest.mark.skipif(
    not _libreoffice_actually_works(),
    reason="LibreOffice not installed/functional in this environment",
)
def test_approval_preview_resolves_a_file_in_an_extra_writable_directory(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """_build_document_edit_preview builds its own WorkspaceScope from
    self.settings.extra_readable_dirs/extra_writable_dirs, the same
    settings the real session's own tools use -- not just the plain
    workspace root. Confirms it actually resolves and renders a file
    that only exists via extra_writable_dirs (a directory outside the
    main workspace), the "file system is the core of this feature" case
    worth nailing explicitly rather than trusting by inspection alone --
    a resolve() bug here would silently degrade to (None, None) instead
    of erroring, so only a real render (non-None previews) proves it
    actually reached the file."""
    extra_dir = tmp_path / "shared_drive"
    _write_test_deck(extra_dir, name="external.pptx")
    call = _tool_call(
        "call_1",
        "edit_pptx_shape",
        {
            "path": str(extra_dir / "external.pptx"),
            "slide": 1,
            "shape_index": 0,
            "fill_color": "38BDF8",
        },
    )
    fake_model = FakeToolCallingChatModel(
        responses=[AIMessage(content="", tool_calls=[call]), AIMessage(content="done")]
    )
    with _client_lg(tmp_path, monkeypatch, fake_model, extra_writable_dirs=[extra_dir]) as client:
        with client.websocket_connect("/ws/t_extra_writable") as ws:
            ws.receive_json()  # state
            ws.receive_json()  # history
            ws.send_json({"type": "user_message", "text": "recolor that shape"})

            approval = ws.receive_json()
            assert approval["type"] == "approval_required"
            assert approval["before_preview"] is not None
            assert approval["after_preview"] is not None

            ws.send_json({"type": "approval_response", "id": approval["id"], "approved": False})
            _receive_until(ws, "agent_message")
