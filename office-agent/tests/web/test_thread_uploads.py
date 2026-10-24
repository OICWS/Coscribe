# ruff: noqa: E402
"""Web tests: uploads, previews and the slide shapes endpoint."""

from pathlib import Path

import pytest

pytest.importorskip("langgraph", reason="needs the langgraph_spike extra installed")

from pptx import Presentation

from coscribe.web.routes import files as files_routes

from .helpers import (
    FakeToolCallingChatModel,
    _client_lg,
    _write_test_deck,
)


def test_upload_writes_file_to_workspace_and_returns_relative_path(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        response = client.post(
            "/api/upload",
            files={"file": ("report.pdf", b"%PDF-1.4 fake content", "application/pdf")},
        )

    assert response.status_code == 200
    body = response.json()
    assert body["path"] == "report.pdf"
    assert body["bytes_written"] == len(b"%PDF-1.4 fake content")
    assert (tmp_path / "workspace" / "report.pdf").read_bytes() == b"%PDF-1.4 fake content"


def test_an_upload_and_a_slide_overlay_use_the_conversations_own_folder_lg(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Real bug: with a folder chosen for a conversation, an attached file
    landed in the app's default workspace -- the path the model was given
    pointed nowhere in its folder -- and the click-a-shape overlay couldn't
    find a deck that lived there."""
    from pptx import Presentation

    chosen = tmp_path / "project-folder"
    chosen.mkdir()
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[5])
    prs.save(str(chosen / "deck.pptx"))
    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        env = client.post("/api/environments", json={"name": "Own", "folders": [str(chosen)]})
        with client.websocket_connect("/ws/t_own_folder") as ws:
            ws.receive_json()
            ws.receive_json()
            ws.send_json({"type": "set_environment", "environment": env.json()["id"]})
            ws.receive_json()
            uploaded = client.post(
                "/api/upload",
                data={"thread_id": "t_own_folder"},
                files={"file": ("brief.txt", b"hello", "text/plain")},
            ).json()
            shapes = client.get(
                "/api/pptx-shapes",
                params={"path": "deck.pptx", "slide": 1, "thread_id": "t_own_folder"},
            )
            without_thread = client.get(
                "/api/pptx-shapes", params={"path": "deck.pptx", "slide": 1}
            )

    assert uploaded["path"] == "brief.txt"
    assert (chosen / "brief.txt").read_bytes() == b"hello"
    assert not (tmp_path / "workspace" / "brief.txt").exists()
    assert shapes.status_code == 200 and shapes.json()["slide_width_in"] > 0
    assert without_thread.status_code == 400


def test_upload_auto_renames_on_name_collision(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        first = client.post("/api/upload", files={"file": ("notes.txt", b"first", "text/plain")})
        second = client.post("/api/upload", files={"file": ("notes.txt", b"second", "text/plain")})

    assert first.json()["path"] == "notes.txt"
    assert second.json()["path"] == "notes (1).txt"
    workspace = tmp_path / "workspace"
    assert (workspace / "notes.txt").read_bytes() == b"first"
    assert (workspace / "notes (1).txt").read_bytes() == b"second"


def test_upload_sanitizes_path_traversal_attempt_in_filename(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        response = client.post(
            "/api/upload", files={"file": ("../../evil.txt", b"data", "text/plain")}
        )

    assert response.status_code == 200
    assert response.json()["path"] == "evil.txt"
    assert (tmp_path / "workspace" / "evil.txt").read_bytes() == b"data"
    assert not (tmp_path / "evil.txt").exists()


def test_upload_rejects_oversized_file(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(files_routes, "MAX_UPLOAD_BYTES", 10)
    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        response = client.post(
            "/api/upload", files={"file": ("big.txt", b"0123456789 -- too big", "text/plain")}
        )

    assert response.status_code == 413
    assert not (tmp_path / "workspace" / "big.txt").exists()


def test_get_preview_serves_an_existing_preview_png(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    previews_dir = tmp_path / "state" / "previews"
    previews_dir.mkdir(parents=True)
    name = "0123456789abcdef0123456789abcdef.png"
    (previews_dir / name).write_bytes(b"\x89PNG\r\n\x1a\n fake png bytes")

    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        response = client.get(f"/api/previews/{name}")

    assert response.status_code == 200
    assert response.headers["content-type"] == "image/png"
    assert response.content == b"\x89PNG\r\n\x1a\n fake png bytes"


def test_get_preview_404s_for_unknown_name(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        response = client.get("/api/previews/0123456789abcdef0123456789abcdef.png")

    assert response.status_code == 404


def test_get_preview_rejects_path_traversal_attempts(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    # Proves the name-shape check, not just a missing-file 404 -- a name
    # that doesn't match render_thumbnail's uuid4().hex pattern is rejected
    # before any filesystem lookup, so it can never escape state_dir/previews/.
    secret = tmp_path / "state" / "secret.txt"
    secret.parent.mkdir(parents=True)
    secret.write_text("do not serve me")

    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        response = client.get("/api/previews/..%2Fsecret.txt")

    assert response.status_code == 404


def test_get_pptx_shapes_returns_shapes_and_slide_dimensions(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Backs the click-a-shape-in-the-preview feature (ChatLog.tsx's
    PptxShapeOverlay) -- a plain UI-facing REST read, not a tool call."""
    _write_test_deck(tmp_path / "workspace")

    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        response = client.get("/api/pptx-shapes", params={"path": "deck.pptx", "slide": 1})

    assert response.status_code == 200
    body = response.json()
    assert body["slide"] == 1
    assert body["shape_count"] == 1
    assert body["slide_width_in"] > 0
    assert body["slide_height_in"] > 0
    [shape] = body["shapes"]
    assert shape["index"] == 0
    assert shape["left_in"] == pytest.approx(1.0)
    assert shape["top_in"] == pytest.approx(1.0)
    assert shape["width_in"] == pytest.approx(2.0)
    assert shape["height_in"] == pytest.approx(1.0)
    assert shape["text_preview"] == "hello"


def test_get_pptx_shapes_400s_for_a_missing_file(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        response = client.get("/api/pptx-shapes", params={"path": "nope.pptx", "slide": 1})

    assert response.status_code == 400


def test_get_pptx_shapes_400s_for_an_out_of_range_slide(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _write_test_deck(tmp_path / "workspace")

    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        response = client.get("/api/pptx-shapes", params={"path": "deck.pptx", "slide": 5})

    assert response.status_code == 400


def test_get_pptx_shapes_400s_for_a_path_outside_the_workspace(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    outside = tmp_path / "outside.pptx"
    Presentation().save(str(outside))

    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        response = client.get("/api/pptx-shapes", params={"path": str(outside), "slide": 1})

    assert response.status_code == 400


def test_an_attached_pdf_can_be_previewed_page_by_page_and_nothing_else_can(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    import io

    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer)
    for number in (1, 2):
        pdf.drawString(100, 700, f"page {number}")
        pdf.showPage()
    pdf.save()

    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        uploaded = client.post(
            "/api/upload", files={"file": ("cv.pdf", buffer.getvalue(), "application/pdf")}
        ).json()
        text = client.post("/api/upload", files={"file": ("a.txt", b"hi", "text/plain")}).json()
        info = client.get("/api/attachment/pdf", params={"path": uploaded["path"]})
        page = client.get(
            "/api/attachment/pdf/page", params={"path": uploaded["path"], "page": 2, "width": 300}
        )
        past_end = client.get(
            "/api/attachment/pdf/page", params={"path": uploaded["path"], "page": 3}
        )
        not_pdf = client.get("/api/attachment/pdf", params={"path": text["path"]})
        outside = client.get("/api/attachment/pdf", params={"path": "../../etc/passwd"})

    assert info.json() == {"pages": 2}
    assert page.headers["content-type"] == "image/png" and page.content.startswith(b"\x89PNG")
    assert past_end.status_code == 404
    assert not_pdf.status_code == 404
    assert outside.status_code in (400, 404)
