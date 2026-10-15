"""The deck engine's compiler (`tools/office/pptx/engine`): SVG pages in, native shapes out.

The copied emitter (`tools/_native_svg/`) is pinned through these pages rather than line by
line: what a page asks for (a run, a line break, a named group, a title, a background, an
image) must come out as the PowerPoint object a user can select and edit.
"""

from __future__ import annotations

import difflib
import io
import os
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation

from coscribe.tools.office.pptx.engine import compile_pages

from .pptx_canonical import canonical_parts

GOLDEN = Path(__file__).parent / "golden" / "engine" / "basic.txt"
_SVG = (
    'xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1280 720" font-family="Microsoft YaHei, Arial"'
)

COVER = f"""<svg {_SVG} data-pptx-page-role="cover">
  <rect x="0" y="0" width="1280" height="720" fill="#0F2A44"/>
  <g id="header" data-pptx-bounds="80 260 1000 200">
    <rect x="80" y="300" width="6" height="120" fill="#F2A541"/>
    <text id="title" x="110" y="350" font-size="48" font-weight="bold"
          fill="#FFFFFF">第三季度经营回顾</text>
    <text x="110" y="405" font-size="24" fill="#C9D6E3">营收增长 <tspan fill="#F2A541"
          font-weight="bold">12%</tspan>，新市场落地<tspan x="110" dy="36">第二行说明</tspan></text>
  </g>
</svg>"""

CARDS = f"""<svg {_SVG} data-pptx-page-role="content">
  <defs>
    <linearGradient id="band" x1="0" y1="0" x2="1" y2="0">
      <stop offset="0%" stop-color="#1A73E8"/><stop offset="100%" stop-color="#0D47A1"/>
    </linearGradient>
    <filter id="soft" x="-15%" y="-20%" width="130%" height="150%">
      <feDropShadow dx="0" dy="6" stdDeviation="8" flood-color="#000000" flood-opacity="0.10"/>
    </filter>
  </defs>
  <text id="title" x="64" y="90" font-size="36" font-weight="bold" fill="#1E2A36">区域表现</text>
  <rect x="64" y="120" width="1152" height="8" fill="url(#band)"/>
  <g id="card-east" transform="translate(64 160)" data-pptx-bounds="64 160 360 200">
    <rect x="0" y="0" width="360" height="200" rx="12" fill="#FFFFFF" filter="url(#soft)"/>
    <text x="24" y="60" font-size="20" fill="#5B6B7B">华东</text>
    <text x="24" y="130" font-size="48" font-weight="bold" fill="#1A73E8">5,410</text>
  </g>
</svg>"""


def _compile(tmp_path: Path, pages: list[tuple[str, str]], **kwargs: object) -> tuple[Path, object]:
    out = tmp_path / "deck.pptx"
    report = compile_pages(pages, out, **kwargs)  # type: ignore[arg-type]
    return out, report


def test_text_keeps_its_runs_and_line_breaks_in_one_frame(tmp_path: Path) -> None:
    out, report = _compile(tmp_path, [("P01", COVER)])

    assert report.ok, report.errors()
    shapes = Presentation(str(out)).slides[0].shapes
    header = next(s for s in shapes if s.name == "header")
    subtitle = next(s for s in header.shapes if s.has_text_frame and "营收" in s.text_frame.text)
    runs = [r for p in subtitle.text_frame.paragraphs for r in p.runs]
    assert [r.text for r in runs if r.font.bold] == ["12%"]
    assert subtitle.text_frame.text.replace("\x0b", "\n").splitlines()[-1] == "第二行说明"


def test_the_title_text_becomes_the_title_placeholder_outside_its_group(tmp_path: Path) -> None:
    out, report = _compile(tmp_path, [("P01", COVER), ("P02", CARDS)])

    assert report.ok, report.errors()
    slides = Presentation(str(out)).slides
    assert [s.shapes.title.text for s in slides] == ["第三季度经营回顾", "区域表现"]
    header = next(s for s in slides[0].shapes if s.name == "header")
    assert all(s.name != "title" for s in header.shapes)


def test_groups_keep_their_ids_and_translate_moves_their_children(tmp_path: Path) -> None:
    out, report = _compile(tmp_path, [("P02", CARDS)])

    assert report.ok, report.errors()
    card = next(s for s in Presentation(str(out)).slides[0].shapes if s.name == "card-east")
    assert len(card.shapes) == 3
    assert card.left == 64 * 9525 and card.top == 160 * 9525


def test_gradient_and_shadow_become_native_fills_and_effects(tmp_path: Path) -> None:
    out, report = _compile(tmp_path, [("P02", CARDS)])

    assert report.ok, report.errors()
    xml = Presentation(str(out)).slides[0]._element.xml
    assert "<a:gradFill" in xml and "<a:outerShdw" in xml


def test_a_full_canvas_solid_rectangle_drawn_first_is_the_background(tmp_path: Path) -> None:
    out, report = _compile(tmp_path, [("P01", COVER)])

    slide = Presentation(str(out)).slides[0]
    assert report.ok, report.errors()
    assert 'srgbClr val="0F2A44"' in slide._element.cSld.bg.xml
    assert not any(s.width == 1280 * 9525 and s.height == 720 * 9525 for s in slide.shapes)


def test_an_image_is_embedded_and_one_outside_the_folder_is_refused(tmp_path: Path) -> None:
    (tmp_path / "pages").mkdir()
    Image.new("RGB", (40, 20), "#C00000").save(tmp_path / "pages" / "red.png")
    Image.new("RGB", (40, 20), "#00C000").save(tmp_path / "green.png")
    picture = '<image href="{}" x="10" y="10" width="400" height="200"/>'
    inside = f"<svg {_SVG}>{picture.format('red.png')}</svg>"
    outside = f"<svg {_SVG}>{picture.format('../green.png')}</svg>"

    out, report = _compile(
        tmp_path, [("in", inside), ("out", outside)], resource_root=tmp_path / "pages"
    )

    slides = Presentation(str(out)).slides
    pictures = [s for s in slides[0].shapes if s.shape_type == 13]
    assert len(pictures) == 1 and pictures[0].image.size == (40, 20)
    assert not report.pages[0].errors
    assert report.pages[1].errors


def test_problems_are_reported_per_page_and_the_other_pages_still_compile(tmp_path: Path) -> None:
    broken = f"""<svg {_SVG}>
      <foreignObject x="0" y="0" width="10" height="10"/>
      <g id="spun" transform="rotate(30)">
        <rect x="0" y="0" width="10" height="10" fill="#000000"/>
      </g>
      <rect x="0" y="0" width="10" height="10" fill="#000000"/>
      <defs><linearGradient id="g" gradientUnits="userSpaceOnUse" x1="64" y1="0" x2="624" y2="0">
        <stop offset="0" stop-color="#000000"/><stop offset="1" stop-color="#FFFFFF"/>
      </linearGradient></defs>
    </svg>"""

    out, report = _compile(tmp_path, [("P01", COVER), ("bad", broken), ("xml", "<svg")])

    assert [bool(p.errors) for p in report.pages] == [False, True, True]
    errors = " ".join(report.pages[1].errors)
    assert "foreignObject" in errors and "spun" in errors and "userSpaceOnUse" in errors
    assert report.pages[1].shapes == 1
    assert len(Presentation(str(out)).slides) == 3


def test_a_pages_notes_become_its_speaker_notes_and_are_not_drawn(tmp_path: Path) -> None:
    notes = '<desc id="notes">\n    先讲华东的增长。\n\n    再说明新市场的进度。\n  </desc>'
    page = COVER.replace("</svg>", f"  {notes}\n</svg>")

    out, report = _compile(tmp_path, [("P01", page), ("P02", CARDS)])

    assert report.ok, report.errors()
    first, second = Presentation(str(out)).slides
    paragraphs = first.notes_slide.notes_text_frame.paragraphs
    assert [p.text for p in paragraphs] == ["先讲华东的增长。", "再说明新市场的进度。"]
    assert not any("华东的增长" in s.text_frame.text for s in first.shapes if s.has_text_frame)
    assert not second.has_notes_slide


def test_an_icon_is_drawn_at_its_box_in_its_colour(tmp_path: Path) -> None:
    icon = (
        '<use id="export" data-icon="ship" x="100" y="200" width="72" height="72" '
        'stroke="#0E3A66"/>'
    )
    page = CARDS.replace("</svg>", f"  {icon}\n</svg>")

    out, report = _compile(tmp_path, [("P01", page)])

    assert report.ok, report.errors()
    shape = next(s for s in Presentation(str(out)).slides[0].shapes if s.name == "export")
    paths = [p for g in shape.shapes for p in getattr(g, "shapes", [g])]
    assert paths and all(str(p.line.color.rgb) == "0E3A66" for p in paths)
    assert all(abs(p.line.width.pt - 4.5) < 0.1 for p in paths)  # 2 px on 24 px, drawn at 72 px
    assert 100 <= shape.left / 9525 and shape.left / 9525 + shape.width / 9525 <= 172 + 2


def test_an_unknown_icon_is_an_error_naming_the_nearest(tmp_path: Path) -> None:
    page = CARDS.replace("</svg>", '  <use data-icon="shipp" x="0" y="0" width="24"/>\n</svg>')

    _out, report = _compile(tmp_path, [("P01", page)])

    (error,) = report.errors()
    assert 'no icon named "shipp"' in error and "ship" in error.split("near:")[1]


def test_a_compiled_deck_matches_its_golden(tmp_path: Path) -> None:
    out, report = _compile(tmp_path, [("P01", COVER), ("P02", CARDS)])
    assert report.ok, report.errors()
    parts = canonical_parts(out)
    actual = "".join(
        f"=== {name}\n{text}\n"
        for name, text in sorted(parts.items())
        if name.startswith("ppt/slides/")
    )

    if os.environ.get("UPDATE_GOLDENS") == "1":
        GOLDEN.parent.mkdir(parents=True, exist_ok=True)
        GOLDEN.write_text(actual, encoding="utf-8")
    expected = GOLDEN.read_text(encoding="utf-8")
    diff = "\n".join(difflib.unified_diff(expected.splitlines(), actual.splitlines(), lineterm=""))
    assert actual == expected, (
        f"UPDATE_GOLDENS=1 pytest {__file__} and commit the diff:\n{diff[:4000]}"
    )


@pytest.mark.parametrize("box", ["0 0 1280", "10 0 1280 720", "0 0 1280.5 720"])
def test_the_canvas_must_be_a_whole_pixel_viewbox_at_the_origin(tmp_path: Path, box: str) -> None:
    page = (
        f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="{box}">'
        '<rect x="1" y="1" width="2" height="2"/></svg>'
    )

    _out, report = _compile(tmp_path, [("P01", page)])

    assert report.pages[0].errors


def test_pictures_in_sibling_groups_keep_their_own_images(tmp_path: Path) -> None:
    Image.new("RGB", (120, 80), "#FF0000").save(tmp_path / "store.png")
    Image.new("RGB", (120, 80), "#0000FF").save(tmp_path / "team.png")
    page = (
        f'<svg {_SVG}><text id="title" x="60" y="80" font-size="30" fill="#000000">门店与团队</text>'
        '<g id="store"><image href="store.png" x="60" y="120" width="500" height="300"/></g>'
        '<g id="team"><image href="team.png" x="660" y="120" width="500" height="300"/></g></svg>'
    )

    out, _report = _compile(tmp_path, [("p1", page)], svg_dir=tmp_path)

    slide = Presentation(str(out)).slides[0]
    colours = {
        group.name: Image.open(io.BytesIO(group.shapes[0].image.blob)).getpixel((1, 1))
        for group in slide.shapes
        if group.name in ("store", "team")
    }
    assert colours == {"store": (255, 0, 0), "team": (0, 0, 255)}
