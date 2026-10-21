"""The Slides panel's view of a deck (decision 0016): the slides of a deck built from pages, and
each slide as one SVG an image can draw, with the native charts and tables drawn as the build
wrote them."""

from __future__ import annotations

import re
from pathlib import Path

import pytest
from PIL import Image

from coscribe.tools.office.pptx import (
    NotBuiltFromPages,
    deck_basic_fonts,
    deck_html,
    deck_pdf,
    read_deck,
    slide_svg,
)
from coscribe.tools.office.pptx.engine import build_deck

from .test_deck_engine import CARDS, COVER
from .test_svg_to_section import OBJECTS


def _deck(tmp_path: Path, *pages: str) -> Path:
    folder = tmp_path / "pages"
    folder.mkdir()
    Image.new("RGB", (96, 44), (30, 120, 200)).save(folder / "photo.png")
    for i, page in enumerate(pages, start=1):
        (folder / f"{i:02d}.svg").write_text(page, encoding="utf-8")
    assert build_deck(folder, tmp_path / "deck.pptx", resource_root=tmp_path).ok
    return tmp_path / "deck.pptx"


def test_the_deck_lists_its_slides_as_built(tmp_path: Path) -> None:
    deck = _deck(tmp_path, COVER, CARDS, OBJECTS)
    view = read_deck(deck)
    assert view["title"] == "deck" and view["folder"] == "pages" and not view["changed"]
    assert [s["name"] for s in view["slides"]] == ["01.svg", "02.svg", "03.svg"]
    assert {s["kind"] for s in view["slides"]} == {"page"}
    third = view["slides"][2]
    assert third["notes"] == "先讲收入。\n再讲区域。" and third["background"] == "#F7F5F0"

    page = tmp_path / "pages" / "02.svg"
    page.write_text(
        page.read_text(encoding="utf-8").replace("</svg>", "<rect/></svg>"), encoding="utf-8"
    )
    assert read_deck(deck)["changed"]


def test_a_slide_is_one_svg_with_nothing_outside_it(tmp_path: Path) -> None:
    deck = _deck(tmp_path, OBJECTS)
    svg = slide_svg(deck, 0, tmp_path)
    assert "data-icon" not in svg and 'href="data:image/png;base64,' in svg
    assert "photo.png" not in svg and "metadata" not in svg
    # The table's cells and the chart's bars, as the build wrote them.
    assert "5,410" in svg and "合计" in svg
    assert 'data-native="table"' in svg and 'data-native="chart"' in svg
    assert svg.count('fill="#4472C4"') == 2
    with pytest.raises(IndexError):
        slide_svg(deck, 1, tmp_path)


def test_a_slide_runs_and_fetches_nothing(tmp_path: Path) -> None:
    # The build refuses these; a page changed after it is drawn all the same.
    deck = _deck(tmp_path, OBJECTS)
    page = tmp_path / "pages" / "01.svg"
    page.write_text(
        OBJECTS.replace(
            "</svg>",
            '<script>alert(1)</script><rect onclick="alert(1)" width="1" height="1"/>'
            '<image href="https://example.com/a.png" width="1" height="1"/>'
            '<image href="../../outside.png" width="1" height="1"/></svg>',
        ),
        encoding="utf-8",
    )
    Image.new("RGB", (4, 4)).save(tmp_path.parent / "outside.png")
    svg = slide_svg(deck, 0, tmp_path)
    assert "script" not in svg and "onclick" not in svg
    assert "example.com" not in svg and "outside.png" not in svg
    assert len(re.findall(r"<image", svg)) == 3 and svg.count('href="data:') == 1


def test_a_deck_not_built_from_pages_is_named_so(tmp_path: Path) -> None:
    from pptx import Presentation

    Presentation().save(str(tmp_path / "own.pptx"))
    with pytest.raises(NotBuiltFromPages):
        read_deck(tmp_path / "own.pptx")


def test_the_web_page_holds_every_slide_and_fetches_nothing(tmp_path: Path) -> None:
    page = deck_html(_deck(tmp_path, COVER, CARDS, OBJECTS), tmp_path)
    assert page.count("data:image/svg+xml;base64,") == 3
    assert "<title>deck</title>" in page
    assert not re.search(r"""(src|href)=["']?https?:""", page)


def test_the_pdf_has_a_page_per_slide(tmp_path: Path) -> None:
    import pypdfium2

    data = deck_pdf(_deck(tmp_path, COVER, OBJECTS), tmp_path)
    pdf = pypdfium2.PdfDocument(data)
    assert len(pdf) == 2
    width, height = pdf[0].get_size()
    assert (round(width), round(height)) == (960, 540)


def test_basic_fonts_swap_only_the_faces_a_computer_may_lack(tmp_path: Path) -> None:
    import zipfile

    page = COVER.replace('font-family="', 'font-family="Noto Serif SC, ', 1)
    deck = _deck(tmp_path, page, CARDS)
    out = tmp_path / "basic.pptx"
    out.write_bytes(deck_basic_fonts(deck))
    faces = set()
    with zipfile.ZipFile(out) as package:
        for name in package.namelist():
            if name.endswith(".xml"):
                xml = package.read(name).decode()
                faces |= set(re.findall(r'<a:(?:latin|ea|cs)\b[^>]*?typeface="([^"]*)"', xml))
    assert faces and not any("Noto" in f for f in faces)
    # The theme's per-script faces (<a:font script="Jpan">) ship with Office and stay.
    assert {f for f in faces if f and not f.startswith("+")} <= {
        "Arial", "Calibri", "Calibri Light", "Times New Roman", "Microsoft YaHei", "SimSun",
    }  # fmt: skip
    from pptx import Presentation

    assert len(Presentation(str(out)).slides) == 2


def test_each_slide_names_the_elements_a_comment_can_be_put_on(tmp_path: Path) -> None:
    slides = read_deck(_deck(tmp_path, OBJECTS))["slides"]
    labels = [e["label"] for e in slides[0]["elements"]]
    assert 'Title "区域表现"' in labels
    assert {"Chart", "Table", "Picture"} <= set(labels)
    chart = next(e for e in slides[0]["elements"] if e["label"] == "Chart")
    # The page's chart frame: x 640, y 160, 560 × 400 on a 1280 × 720 page.
    assert [round(v, 3) for v in chart["box"]] == [0.5, 0.222, 0.438, 0.556]


def test_comments_are_kept_beside_the_pages(tmp_path: Path) -> None:
    from coscribe.tools.office.pptx import (
        CommentNotFound,
        add_comment,
        delete_comment,
        read_comments,
        reply_to_comment,
        resolve_comment,
    )

    deck = _deck(tmp_path, COVER, CARDS)
    first = add_comment(
        deck, slide="01", x=0.2, y=1.4, element="Title", text=" 换个标题 ", author="ana", sent=True
    )
    second = add_comment(
        deck, slide="02", x=0.5, y=0.5, element="Shape", text="ok", author="ana", sent=False
    )
    reply_to_comment(deck, first["id"], text="Changed it.", author="coscribe")
    resolve_comment(deck, first["id"], True)
    delete_comment(deck, second["id"])

    (comment,) = read_comments(deck)
    assert comment["text"] == "换个标题" and comment["y"] == 1.0 and comment["resolved"]
    assert comment["replies"][0]["author"] == "coscribe"
    assert (tmp_path / "pages" / "comments.json").is_file()
    with pytest.raises(CommentNotFound):
        resolve_comment(deck, second["id"], True)
    with pytest.raises(ValueError):
        add_comment(deck, slide="01", x=0, y=0, element="", text="  ", author="a", sent=False)
    (tmp_path / "pages" / "comments.json").write_text("{not json", encoding="utf-8")
    assert read_comments(deck) == []


def test_a_deck_coscribe_did_not_build_is_drawn_read_only_through_the_office_program(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The office program writes the PDF once per version of the file; each slide is its page."""
    import pypdfium2 as pdfium
    from pptx import Presentation

    from coscribe.tools.office.pptx import NotBuiltFromPages, foreign_deck, foreign_slide_svg
    from coscribe.tools.office.pptx.engine import view_foreign

    deck = tmp_path / "theirs.pptx"
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(str(deck))
    calls: list[Path] = []

    def office(source: Path, target: Path, to_format: str, timeout: float) -> str:
        calls.append(source)
        pdf = pdfium.PdfDocument.new()
        for _ in range(2):
            pdf.new_page(960, 540)
        pdf.save(str(target))
        return "PowerPoint.Application"

    monkeypatch.setattr(view_foreign, "convert", office)
    state = tmp_path / "state"
    read = foreign_deck(deck, state)
    assert read["readOnly"] is True and [s["name"] for s in read["slides"]] == [
        "slide 1",
        "slide 2",
    ]
    svg = foreign_slide_svg(deck, 1, state)
    assert 'viewBox="0 0 1280 720"' in svg and "data:image/png;base64," in svg
    assert len(calls) == 1  # the second read found the PDF made for this version of the file
    with pytest.raises(IndexError):
        foreign_slide_svg(deck, 2, state)

    monkeypatch.setattr(view_foreign, "convert", lambda *a: None)
    (tmp_path / "other.pptx").write_bytes(deck.read_bytes())
    with pytest.raises(NotBuiltFromPages):
        foreign_deck(tmp_path / "other.pptx", state)
