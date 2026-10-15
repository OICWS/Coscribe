"""The deterministic parts of the Excel evaluation (`evals/xlsx_checks.py`, the xlsx tasks, the
kind-aware runner, judge and comparison). Nothing here calls a model; the tests that recalculate
through LibreOffice skip without it."""

from __future__ import annotations

import datetime as dt
import json
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from evals import compare, judge, run, tasks, text_checks, xlsx_checks

from .xlsx_reference import SOLUTIONS

# --- small workbooks --------------------------------------------------------------------


def _book(rows: list[list[Any]], title: str = "Data") -> Workbook:
    book = Workbook()
    sheet = book.active
    sheet.title = title
    for row in rows:
        sheet.append(row)
    return book


# --- addresses and values ---------------------------------------------------------------


def test_an_address_names_a_sheet_and_cells_even_when_the_sheet_name_is_quoted() -> None:
    assert xlsx_checks.split_address("Summary!B2:B9") == ("Summary", "B2:B9")
    assert xlsx_checks.split_address("'My sheet'!$B$2") == ("My sheet", "B2")
    assert xlsx_checks.split_address("汇总!A1") == ("汇总", "A1")


def test_a_range_gives_every_cell_and_a_missing_sheet_raises() -> None:
    book = _book([[1, 2], [3, 4]])

    assert [c.value for c in xlsx_checks.cells_in(book, "Data!A1:B2")] == [1, 2, 3, 4]
    assert [c.value for c in xlsx_checks.cells_in(book, "Data!B2")] == [4]
    assert [c.value for c in xlsx_checks.cells_in(book, "Data!A1:A2")] == [1, 3]
    with pytest.raises(KeyError):
        xlsx_checks.cells_in(book, "Nope!A1")


@pytest.mark.parametrize(
    ("got", "spec", "expected"),
    [
        (10.5, {"value": 10.5}, True),
        (10.504, {"value": 10.5}, True),
        (10.51, {"value": 10.5}, False),
        (10.51, {"value": 10.5, "tol": 0.02}, True),
        ("10.5", {"value": 10.5}, False),
        (True, {"value": 1}, False),
        ("Hardware", {"value": "Hardware"}, True),
        ("hardware", {"value": "Hardware"}, False),
        ("Hardware ", {"value": "Hardware"}, False),
        (None, {"value": None}, True),
        ("", {"value": None}, True),
        (0, {"value": None}, False),
        (dt.datetime(2026, 1, 5), {"value": "2026-01-05", "type": "date"}, True),
        ("2026-01-05", {"value": "2026-01-05", "type": "date"}, False),
        (dt.datetime(2026, 1, 6), {"value": "2026-01-05", "type": "date"}, False),
    ],
)
def test_a_cell_value_matches_by_number_text_or_date(
    got: object, spec: dict[str, Any], expected: bool
) -> None:
    assert xlsx_checks.value_matches(got, spec) is expected


def test_the_wrong_cells_are_named_with_what_was_found() -> None:
    values = _book([["a", 1.0], ["b", 2.0]])
    specs = [
        {"at": "Data!B1", "value": 1},
        {"at": "Data!B2", "value": 3},
        {"at": "Gone!A1", "value": 1},
    ]

    report = xlsx_checks.check_cells(values, specs)

    assert report["ok"] is False and report["matched"] == 1
    assert report["wrong"] == [
        {"at": "Data!B2", "expected": 3, "got": 2.0},
        {"at": "Gone!A1", "expected": 1, "got": "no such sheet"},
    ]


def test_sheet_names_are_matched_without_regard_to_case() -> None:
    assert xlsx_checks.check_sheets(["Summary", "Data"], ["summary"])["ok"] is True
    assert xlsx_checks.check_sheets(["Data"], ["Summary"])["missing"] == ["Summary"]


# --- formula or constant ----------------------------------------------------------------


def test_a_total_typed_in_as_a_number_is_not_a_formula() -> None:
    sheet = _book([[1, 2], ["=SUM(A1:B1)", 3]]).active
    book = sheet.parent

    ok = xlsx_checks.check_formulas(book, ["Data!A2"])
    bad = xlsx_checks.check_formulas(book, ["Data!A1:B2"])

    assert ok["ok"] is True
    assert bad["ok"] is False and bad["wrong"] == ["Data!A1", "Data!B1", "Data!B2"]
    assert bad["checked"] == 4


def test_an_input_that_became_a_formula_is_listed() -> None:
    book = _book([[1000, "=B1*2"]])

    assert xlsx_checks.check_constants(book, ["Data!A1"])["ok"] is True
    report = xlsx_checks.check_constants(book, ["Data!A1:B1"])
    assert report["ok"] is False and report["wrong"] == ["Data!B1"]


def test_a_missing_sheet_fails_the_formula_and_constant_checks_instead_of_raising() -> None:
    book = _book([[1]])

    assert xlsx_checks.check_formulas(book, ["Gone!A1"])["ok"] is False
    assert xlsx_checks.check_constants(book, ["Gone!A1"])["ok"] is False


# --- formats ----------------------------------------------------------------------------


def test_a_number_format_must_contain_one_of_the_strings() -> None:
    sheet = _book([[1, 2], [3, 4]]).active
    sheet["A1"].number_format = '#,##0"元"'
    sheet["A2"].number_format = "General"

    report = xlsx_checks.check_number_formats(
        sheet.parent, [{"at": "Data!A1:A2", "any_of": ["元", "万"]}]
    )

    assert report["ok"] is False
    assert report["wrong"] == [{"at": "Data!A2", "format": "General"}]


def test_an_escaped_literal_in_a_number_format_still_matches() -> None:
    book = _book([[dt.datetime(2026, 1, 5)]])
    book["Data"]["A1"].number_format = "yyyy\\-mm\\-dd"

    report = xlsx_checks.check_number_formats(book, [{"at": "Data!A1", "any_of": ["yyyy-mm-dd"]}])

    assert report["ok"] is True


def test_bold_cells_are_found_by_their_font() -> None:
    sheet = _book([["h1", "h2"]]).active
    sheet["A1"].font = Font(bold=True)

    report = xlsx_checks.check_bold(sheet.parent, ["Data!A1:B1"])

    assert report["ok"] is False and report["wrong"] == ["Data!B1"]


def test_a_value_outside_the_allowed_range_is_listed_and_blank_cells_are_not() -> None:
    sheet = _book([[1, 2, 3], [4, 5, 6]]).active
    sheet["A2"] = ""
    sheet["F9"] = "stray"

    report = xlsx_checks.check_used_range(sheet.parent, {"Data": "A1:C2"})

    assert report["ok"] is False and report["outside"] == ["Data!F9"] and report["count"] == 1
    assert xlsx_checks.check_used_range(sheet.parent, {"Data": "A1:F9"})["ok"] is True
    assert xlsx_checks.check_used_range(sheet.parent, {"Gone": "A1:A2"})["ok"] is False


# --- an edit keeps what was there -------------------------------------------------------


def _styled() -> Workbook:
    sheet = _book([["Item", "Cost"], ["a", 5]], "Budget").active
    sheet["A1"].font = Font(bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="1F3864")
    sheet["B2"].number_format = '"$"#,##0'
    sheet.conditional_formatting.add(
        "B2:B2", CellIsRule(operator="greaterThan", formula=["3"], fill=PatternFill("solid"))
    )
    return sheet.parent


def test_untouched_cells_with_the_same_styles_pass() -> None:
    assert xlsx_checks.check_untouched_styles(_styled(), _styled(), ["Budget!A1:B2"])["ok"] is True


@pytest.mark.parametrize(
    ("change", "what"),
    [
        (lambda s: setattr(s["B2"], "number_format", "General"), "number_format"),
        (lambda s: setattr(s["A1"], "font", Font(bold=False, color="FFFFFF")), "bold"),
        (lambda s: setattr(s["A1"], "fill", PatternFill("solid", fgColor="FF0000")), "fill_color"),
        (lambda s: setattr(s["A1"], "fill", PatternFill()), "fill"),
    ],
)
def test_a_changed_style_of_an_untouched_cell_is_reported_with_the_property(
    change: Any, what: str
) -> None:
    produced = _styled()
    change(produced["Budget"])

    report = xlsx_checks.check_untouched_styles(_styled(), produced, ["Budget!A1:B2"])

    assert report["ok"] is False
    assert what in report["changed"][0]["difference"]


def test_a_sheet_that_is_gone_fails_the_style_check() -> None:
    report = xlsx_checks.check_untouched_styles(_styled(), _book([[1]]), ["Budget!A1:B2"])

    assert report["ok"] is False and report["changed"][0]["difference"] == "the sheet is gone"


def test_conditional_formats_and_validations_of_the_original_must_survive() -> None:
    original = _styled()
    plain = _book([["Item", "Cost"], ["a", 5]], "Budget").active.parent

    assert xlsx_checks.check_features_kept(original, _styled(), 1)["ok"] is True
    lost = xlsx_checks.check_features_kept(original, plain, 0)
    assert lost["ok"] is False and lost["before"]["rules"] == 1 and lost["after"]["rules"] == 0
    assert xlsx_checks.check_features_kept(plain, plain, 1)["ok"] is False  # asked for one

    with_list = _book([["x"]]).active
    validation = DataValidation(type="list", formula1='"a,b"')
    with_list.add_data_validation(validation)
    validation.add("A1")
    assert xlsx_checks.check_features_kept(with_list.parent, _book([["x"]]), 0)["ok"] is False


# --- charts -----------------------------------------------------------------------------


def _chart_file(path: Path, chart: Any, rows: int = 12) -> Path:
    book = _book([["Month", "Visitors"], *([f"m{i}", i * 10] for i in range(1, rows + 1))])
    sheet = book.active
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=rows + 1), titles_from_data=True)
    sheet.add_chart(chart, "D2")
    book.save(path)
    return path


def test_a_chart_is_read_for_its_kind_title_axes_and_points(tmp_path: Path) -> None:
    chart = LineChart()
    chart.title = "Monthly visitors"
    chart.x_axis.title = "Month"
    chart.y_axis.title = "Visitors"
    path = _chart_file(tmp_path / "c.xlsx", chart)

    (found,) = xlsx_checks.read_charts(path)

    assert found == {
        "kind": "lineChart",
        "title": "Monthly visitors",
        "axis_titles": ["Month", "Visitors"],
        "points": 12,
    }
    good = {
        "kind": "line",
        "title": "monthly VISITORS",
        "points_min": 12,
        "axis_titles": ["Month", "Visitors"],
    }
    assert xlsx_checks.check_charts([found], [good])["ok"] is True


@pytest.mark.parametrize(
    "spec",
    [
        {"kind": "bar"},
        {"title": "Weekly visitors"},
        {"points_min": 13},
        {"axis_titles": ["Year"]},
    ],
)
def test_a_chart_of_the_wrong_kind_title_size_or_axes_is_missing(
    tmp_path: Path, spec: dict[str, Any]
) -> None:
    chart = LineChart()
    chart.title = "Monthly visitors"
    chart.x_axis.title = "Month"
    found = xlsx_checks.read_charts(_chart_file(tmp_path / "c.xlsx", chart))

    report = xlsx_checks.check_charts(found, [spec])

    assert report["ok"] is False and report["missing"] == [spec]


def test_a_workbook_without_a_chart_has_none_to_find(tmp_path: Path) -> None:
    _book([[1]]).save(tmp_path / "plain.xlsx")

    assert xlsx_checks.read_charts(tmp_path / "plain.xlsx") == []
    assert xlsx_checks.check_charts([], [{"kind": "line"}])["ok"] is False
    assert xlsx_checks.check_charts([], [])["ok"] is True


def test_a_bar_chart_is_a_bar_chart(tmp_path: Path) -> None:
    (found,) = xlsx_checks.read_charts(_chart_file(tmp_path / "b.xlsx", BarChart(), rows=3))

    assert found["kind"] == "barChart" and found["points"] == 3


def test_the_points_of_a_scatter_chart_and_of_a_multi_area_reference_are_counted(
    tmp_path: Path,
) -> None:
    from openpyxl.chart import ScatterChart, Series

    assert xlsx_checks._points_in("(Data!$A$1:$A$3,Data!$B$1:$B$3)") == 6
    assert xlsx_checks._points_in("Data!$B$2:$B$13") == 12
    book = _book([[x, x * 2] for x in range(1, 9)])
    sheet = book.active
    chart = ScatterChart()
    chart.series.append(
        Series(
            Reference(sheet, min_col=2, min_row=1, max_row=8),
            Reference(sheet, min_col=1, min_row=1, max_row=8),
        )
    )
    sheet.add_chart(chart, "D2")
    book.save(tmp_path / "s.xlsx")

    (found,) = xlsx_checks.read_charts(tmp_path / "s.xlsx")

    assert found["kind"] == "scatterChart" and found["points"] == 8


# --- dropdown lists ---------------------------------------------------------------------


def _with_validation(formula: str, cells: str, extra: Any = None) -> Workbook:
    book = _book([["Status"], *([None] for _ in range(4))], "Tickets")
    sheet = book.active
    validation = DataValidation(type="list", formula1=formula)
    sheet.add_data_validation(validation)
    validation.add(cells)
    lists = book.create_sheet("Lists")
    for row in ("Open", "In progress", "Done"):
        lists.append([row])
    if extra:
        extra(sheet)
    return book


SPEC = {"at": "Tickets!A2:A5", "values": ["Open", "In progress", "Done"]}


def test_a_typed_in_list_and_a_list_read_from_cells_both_count() -> None:
    typed = _with_validation('"Open,In progress,Done"', "A2:A5")
    ranged = _with_validation("Lists!$A$1:$A$3", "A2:A5")

    assert xlsx_checks.check_validations(typed, typed, [SPEC])["ok"] is True
    assert xlsx_checks.check_validations(ranged, ranged, [SPEC])["ok"] is True


def test_a_list_with_other_values_or_over_fewer_cells_is_not_the_one_asked_for() -> None:
    other = _with_validation('"Open,Closed"', "A2:A5")
    short = _with_validation('"Open,In progress,Done"', "A2:A4")

    assert xlsx_checks.check_validations(other, other, [SPEC])["ok"] is False
    report = xlsx_checks.check_validations(short, short, [SPEC])
    assert report["ok"] is False and "A5" in report["missing"][0]["reason"]
    assert xlsx_checks.check_validations(short, short, [{**SPEC, "at": "Gone!A1"}])["ok"] is False


def test_a_numeric_list_matches_the_numbers_it_offers() -> None:
    book = _book([["n"], [None]], "Tickets")
    validation = DataValidation(type="list", formula1='"1,2,3"')
    book.active.add_data_validation(validation)
    validation.add("A2")

    assert (
        xlsx_checks.check_validations(book, book, [{"at": "Tickets!A2", "values": [1, 2, 3]}])["ok"]
        is True
    )


def test_a_whole_number_rule_is_not_a_dropdown_list() -> None:
    book = _book([["n"], [1]], "Tickets")
    validation = DataValidation(type="whole", formula1="1", formula2="5")
    book.active.add_data_validation(validation)
    validation.add("A2")

    assert (
        xlsx_checks.check_validations(book, book, [{"at": "Tickets!A2", "values": ["1"]}])["ok"]
        is False
    )


# --- can every number and label be read -------------------------------------------------


def test_a_label_wider_than_its_column_is_cut_off_only_when_the_next_cell_is_taken() -> None:
    sheet = _book([["Order Date", "Customer"], ["Fixed costs", None], ["Net margin", ""]]).active
    sheet.column_dimensions["A"].width = 8.43

    clipped = xlsx_checks.clipped_texts(sheet.parent, sheet.parent)

    assert clipped == ["Data!A1 'Order Date'"]


def test_text_in_a_merged_cell_is_not_cut_off() -> None:
    sheet = _book([["A very long merged title", "x"]]).active
    sheet.merge_cells("A1:B1")

    assert xlsx_checks.clipped_texts(sheet.parent, sheet.parent) == []


def test_wide_columns_wrapped_text_and_chinese_widths_are_measured() -> None:
    sheet = _book(
        [["Order Date", "x"], ["销售额（元）", "x"], ["a long wrapped label", "x"]]
    ).active
    sheet.column_dimensions["A"].width = 10
    sheet["A3"].alignment = Alignment(wrap_text=True)

    clipped = xlsx_checks.clipped_texts(sheet.parent, sheet.parent)

    assert clipped == ["Data!A2 '销售额（元）'"]  # 6 CJK characters take 12


def test_a_column_range_gives_its_width_to_every_column_in_it() -> None:
    labels = ["a longer label", "x", "another long label", "x", "x", "x", "wide label", "x"]
    sheet = _book([labels]).active
    narrow = ["Data!A1 'a longer label'", "Data!C1 'another long label'", "Data!G1 'wide label'"]
    assert xlsx_checks.clipped_texts(sheet.parent, sheet.parent) == narrow

    sheet.column_dimensions["A"].width = 30
    sheet.column_dimensions["C"].width = 30
    sheet.column_dimensions["E"].width = 20
    sheet.column_dimensions["E"].min, sheet.column_dimensions["E"].max = 5, 8

    assert xlsx_checks.clipped_texts(sheet.parent, sheet.parent) == []


def _pdf(path: Path, lines: list[str]) -> Path:
    from reportlab.pdfgen import canvas

    page = canvas.Canvas(str(path))
    for number, line in enumerate(lines):
        page.drawString(72, 750 - 20 * number, line)
    page.save()
    return path


def test_cells_printed_as_hashes_are_counted_and_other_hash_signs_are_not(tmp_path: Path) -> None:
    pdf = _pdf(tmp_path / "p.pdf", ["### ### 5", "see C# and ## not", "2026-01-05 #####"])

    assert xlsx_checks.hash_cells(pdf) == 3


def test_a_sheet_is_legible_when_nothing_is_cut_off_and_nothing_prints_as_hashes(
    tmp_path: Path,
) -> None:
    book = _book([["Order Date", "x"]])
    pdf = _pdf(tmp_path / "p.pdf", ["Order Date x"])

    assert xlsx_checks.check_legibility(book, book, pdf)["ok"] is False
    book.active.column_dimensions["A"].width = 14
    assert xlsx_checks.check_legibility(book, book, pdf)["ok"] is True
    assert (
        xlsx_checks.check_legibility(book, book, _pdf(tmp_path / "h.pdf", ["###"]))["ok"] is False
    )
    assert xlsx_checks.check_legibility(book, book, None)["hash_cells"] is None


# --- the whole check on a file ----------------------------------------------------------


def _fake_recalculation(monkeypatch: pytest.MonkeyPatch) -> None:
    """Skip LibreOffice: the 'recalculated' copy is the file itself, so a workbook of typed-in
    numbers reads as it is and one of formulas has no values."""
    import shutil

    def fake(path: Path, scratch: Path, name: str = "recalc.xlsx") -> tuple[Path, dict[str, Any]]:
        copy = scratch / name
        shutil.copyfile(path, copy)
        return copy, {"ok": True, "status": "success", "formulas": 0, "errors": 0}

    monkeypatch.setattr(xlsx_checks, "recalculate", fake)
    monkeypatch.setattr(xlsx_checks, "render_pages", lambda path, out: {"ok": True, "pages": []})


def _task(name: str) -> tasks.Task:
    return tasks.discover(name, kind="xlsx")[0]


def test_numbers_typed_in_where_formulas_were_asked_pass_the_values_and_fail_the_rest(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _fake_recalculation(monkeypatch)
    sheet = _book([], "Model").active
    for row, (label, value) in enumerate(
        [
            ("Units sold", 1000),
            ("Price per unit", 25),
            ("Cost per unit", 14),
            ("Fixed costs", 6000),
        ],
        start=2,
    ):
        sheet.cell(row=row, column=1, value=label)
        sheet.cell(row=row, column=2, value=value)
    for row, (label, value) in enumerate(
        [
            ("Revenue", 25000),
            ("Variable costs", 14000),
            ("Gross profit", 11000),
            ("Net profit", 5000),
            ("Net margin", 0.2),
        ],
        start=7,
    ):
        sheet.cell(row=row, column=1, value=label)
        sheet.cell(row=row, column=2, value=value)
    sheet["B11"].number_format = "0%"
    path = tmp_path / "ws" / "model.xlsx"
    path.parent.mkdir()
    sheet.parent.save(path)

    report = xlsx_checks.check_workbook(_task("06"), path, path.parent, tmp_path / "render")

    assert report["cells"]["ok"] is True and report["facts"]["ok"] is True
    assert report["constants"]["ok"] is True
    assert report["formulas"]["ok"] is False
    assert report["formulas"]["wrong"] == [f"Model!B{r}" for r in range(7, 12)]
    assert report["probe"]["ok"] is False


def test_a_check_that_raises_is_recorded_and_the_others_still_run(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _fake_recalculation(monkeypatch)

    def boom(*args: Any) -> None:
        raise RuntimeError("boom")

    monkeypatch.setattr(xlsx_checks, "check_cells", boom)
    path = tmp_path / "ws" / "m.xlsx"
    path.parent.mkdir()
    _book([["Units sold", 1]], "Model").save(path)

    report = xlsx_checks.check_workbook(_task("06"), path, path.parent, tmp_path / "r")

    assert report["cells"] == {"ok": False, "error": "RuntimeError: boom"}
    assert "facts" in report and "formulas" in report and "render" in report


def test_a_file_that_is_not_a_workbook_stops_at_opens(tmp_path: Path) -> None:
    ws = tmp_path / "ws"
    ws.mkdir()
    (ws / "broken.xlsx").write_bytes(b"not a zip")

    report = xlsx_checks.check_workbook(_task("06"), ws / "broken.xlsx", ws, tmp_path / "r")

    assert report["opens"]["ok"] is False and "recalc" not in report


# --- the reference solutions (recalculated by LibreOffice) -------------------------------


def _libreoffice_recalculates(tmp_path: Path) -> bool:
    path = tmp_path / "probe.xlsx"
    _book([[1, "=A1+1"]]).save(path)
    _, report = xlsx_checks.recalculate(path, tmp_path)
    return bool(report["ok"])


@pytest.mark.real_libreoffice
@pytest.mark.parametrize("task", [t.name for t in tasks.discover("all", kind="xlsx")])
def test_the_reference_solution_of_every_task_passes_every_check(task: str, tmp_path: Path) -> None:
    if not _libreoffice_recalculates(tmp_path):
        pytest.skip("LibreOffice is not usable here")
    workspace = tmp_path / "ws"
    workspace.mkdir()
    chosen = _task(task)

    output = SOLUTIONS[task](chosen.directory / "inputs", workspace)
    report = xlsx_checks.check_workbook(chosen, output, workspace, tmp_path / "render")

    failed = {k: v for k, v in report.items() if isinstance(v, dict) and v.get("ok") is False}
    assert not failed


@pytest.mark.real_libreoffice
@pytest.mark.parametrize(
    ("task", "original", "untouched"),
    [
        ("02-monthly-summary", "transactions.xlsx", "Data!A1:C15"),
        ("03-lookup-two-sheets", "shop.xlsx", "Orders!A1:C9"),
        ("05-edit-budget", "budget_q1.xlsx", "Budget!A1:C7"),
        ("08-validation-list", "tickets.xlsx", "Tickets!A1:D11"),
    ],
)
def test_a_workbook_that_libreoffice_resaved_keeps_the_styles_the_check_compares(
    task: str, original: str, untouched: str, tmp_path: Path
) -> None:
    if not _libreoffice_recalculates(tmp_path):
        pytest.skip("LibreOffice is not usable here")
    given = _task(task).directory / "inputs" / original
    resaved, _ = xlsx_checks.recalculate(given, tmp_path, original)

    report = xlsx_checks.check_untouched_styles(
        load_workbook(given), load_workbook(resaved), [untouched]
    )
    features = xlsx_checks.check_features_kept(load_workbook(given), load_workbook(resaved), 0)

    assert report["ok"] is True, report["changed"]
    assert features["ok"] is True


@pytest.mark.real_libreoffice
def test_a_model_with_pasted_numbers_fails_where_the_formulas_pass(tmp_path: Path) -> None:
    if not _libreoffice_recalculates(tmp_path):
        pytest.skip("LibreOffice is not usable here")
    workspace = tmp_path / "ws"
    workspace.mkdir()
    chosen = _task("06")
    good = SOLUTIONS["06-budget-model"](chosen.directory / "inputs", workspace)
    pasted = workspace / "pasted.xlsx"
    values = load_workbook(xlsx_checks.recalculate(good, tmp_path)[0], data_only=True)
    values.save(pasted)

    report = xlsx_checks.check_workbook(chosen, pasted, workspace, tmp_path / "render")

    assert report["cells"]["ok"] is True
    assert report["formulas"]["ok"] is False
    assert report["probe"]["ok"] is False


def test_a_figure_is_found_in_any_cell_in_ones_thousands_or_wan() -> None:
    book = _book([["营收", 1800], ["净利润", "270"], ["k", 3492.0]])
    specs = [{"value": 18_000_000}, {"value": 2_700_000}, {"value": 3_492_000}]

    report = xlsx_checks.check_figures(book, specs)

    assert report["missing"] == [2_700_000]


def test_formula_cells_are_counted_over_every_sheet() -> None:
    book = _book([[1, "=A1+1"], ["=SUM(A1:B1)", "text"]])
    book.create_sheet("Other")["A1"] = "=1"

    assert xlsx_checks.check_formula_count(book, 3)["ok"] is True
    assert xlsx_checks.check_formula_count(book, 4) == {"ok": False, "count": 3, "minimum": 4}


@pytest.mark.real_libreoffice
def test_a_forecast_with_pasted_numbers_fails_the_input_probe(tmp_path: Path) -> None:
    if not _libreoffice_recalculates(tmp_path):
        pytest.skip("LibreOffice is not usable here")
    workspace = tmp_path / "ws"
    workspace.mkdir()
    chosen = _task("10")
    good = SOLUTIONS["10-store-forecast"](chosen.directory / "inputs", workspace)
    pasted = workspace / "pasted.xlsx"
    load_workbook(xlsx_checks.recalculate(good, tmp_path)[0], data_only=True).save(pasted)

    report = xlsx_checks.check_workbook(chosen, pasted, workspace, tmp_path / "render")

    assert report["figures"]["ok"] is True
    assert report["input_probe"]["ok"] is False
    assert report["formulas_min"]["ok"] is False


@pytest.mark.real_libreoffice
def test_a_formula_error_after_recalculation_is_reported(tmp_path: Path) -> None:
    if not _libreoffice_recalculates(tmp_path):
        pytest.skip("LibreOffice is not usable here")
    path = tmp_path / "bad.xlsx"
    _book([[1, "=A1/0"]]).save(path)

    _, report = xlsx_checks.recalculate(path, tmp_path)

    assert report["ok"] is False and report["errors"] == 1


# --- task files -------------------------------------------------------------------------


def test_every_excel_task_loads_and_its_files_exist() -> None:
    found = tasks.discover("all", kind="xlsx")

    assert len(found) == 10 and {t.kind for t in found} == {"xlsx"}
    assert set(SOLUTIONS) == {t.name for t in found}
    for task in found:
        assert task.brief and isinstance(task.expect, tasks.XlsxExpect)
        assert task.expect.cells or task.expect.facts or task.expect.figures


def test_the_briefs_name_their_input_files() -> None:
    for task in tasks.discover("all", kind="xlsx"):
        for path in task.inputs:
            assert path.name in task.brief, f"{task.name}: the brief never mentions {path.name}"


def test_the_inputs_are_what_the_generator_makes(tmp_path: Path, monkeypatch: Any) -> None:
    from evals.xlsx import make_inputs

    monkeypatch.setattr(make_inputs, "TASKS", tmp_path)
    make_inputs.main()

    for task in tasks.discover("all", kind="xlsx"):
        for path in task.inputs:
            made = tmp_path / task.name / "inputs" / path.name
            if path.suffix == ".csv":
                assert made.read_text(encoding="utf-8") == path.read_text(encoding="utf-8")
            else:
                assert [
                    [c.value for c in row]
                    for sheet in load_workbook(made).worksheets
                    for row in sheet.iter_rows()
                ] == [
                    [c.value for c in row]
                    for sheet in load_workbook(path).worksheets
                    for row in sheet.iter_rows()
                ]


def test_every_cell_a_task_expects_lies_in_the_range_it_allows() -> None:
    from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries

    for task in tasks.discover("all", kind="xlsx"):
        expect = task.expect
        assert isinstance(expect, tasks.XlsxExpect)
        for spec in expect.cells:
            sheet, cell = xlsx_checks.split_address(str(spec["at"]))
            if sheet in expect.used_range and spec["value"] is not None:
                _, _, last_col, last_row = range_boundaries(expect.used_range[sheet])
                row, col = coordinate_to_tuple(cell)
                assert row <= last_row and col <= last_col, f"{task.name}: {spec['at']}"


def _write_task(directory: Path, expect: dict[str, Any], inputs: tuple[str, ...] = ()) -> None:
    directory.mkdir(parents=True, exist_ok=True)
    (directory / "brief.md").write_text("do it", encoding="utf-8")
    (directory / "expect.json").write_text(json.dumps(expect), encoding="utf-8")
    for name in inputs:
        (directory / "inputs").mkdir(exist_ok=True)
        (directory / "inputs" / name).write_text("x", encoding="utf-8")


@pytest.mark.parametrize(
    ("expect", "message"),
    [
        ({}, "expects nothing"),
        ({"cells": [{"at": "B2", "value": 1}]}, "Sheet!B2"),
        ({"cells": [{"at": "S!B2"}]}, 'needs a "value"'),
        ({"cells": [{"at": "S!", "value": 1}]}, "must look like"),
        ({"cells": [{"at": "S!A1:B2", "value": 1}]}, "one cell, not a range"),
        ({"cells": [{"at": "S!B2", "value": 1}], "formulas": ["B2"]}, "must look like"),
        ({"cells": [{"at": "S!B2", "value": 1}], "untouched": ["S!A1"]}, 'need "original"'),
        ({"cells": [{"at": "S!B2", "value": 1}], "probe": {"set": {}}}, "probe"),
        ({"cells": [{"at": "S!B2", "value": 1}], "used_range": ["S!A1"]}, "used_range"),
        (
            {"cells": [{"at": "S!B2", "value": 1}], "number_formats": [{"at": "S!B2"}]},
            "any_of",
        ),
        ({"facts": "text"}, "facts"),
    ],
)
def test_an_excel_task_with_bad_expectations_is_refused(
    tmp_path: Path, expect: dict[str, Any], message: str
) -> None:
    _write_task(tmp_path, expect)

    with pytest.raises(ValueError, match=message):
        tasks.load_task(tmp_path, "xlsx")


def test_the_original_must_be_one_of_the_inputs(tmp_path: Path) -> None:
    expect = {"cells": [{"at": "S!B2", "value": 1}], "original": "a.xlsx", "untouched": ["S!A1"]}
    _write_task(tmp_path, expect, inputs=("b.xlsx",))

    with pytest.raises(ValueError, match="not among the inputs"):
        tasks.load_task(tmp_path, "xlsx")

    _write_task(tmp_path, expect, inputs=("a.xlsx",))
    assert tasks.load_task(tmp_path, "xlsx").expect.original == "a.xlsx"  # type: ignore[union-attr]


def test_a_kind_without_checks_is_not_offered_by_the_runner() -> None:
    with pytest.raises(SystemExit):
        run.main(["--kind", "odt", "--out", "nowhere"])


def test_a_blank_fact_never_matches() -> None:
    assert text_checks.contains("some text", "   ") is False
    assert text_checks.check_facts("some text", [["  "]], [])["ok"] is False


def test_judging_looks_at_no_more_than_twenty_pages_of_a_file(tmp_path: Path) -> None:
    folder = tmp_path / "a" / "t" / "render"
    folder.mkdir(parents=True)
    for number in range(1, 31):
        (folder / f"page-{number:02d}.png").write_bytes(b"png")

    assert len(judge._pages(tmp_path / "a", "t")) == 20


def test_an_unknown_kind_or_task_is_refused() -> None:
    with pytest.raises(ValueError, match="no such kind"):
        tasks.discover("all", kind="odt")
    with pytest.raises(ValueError, match="no such task: 99"):
        tasks.discover("99", kind="xlsx")
    assert [t.name[:2] for t in tasks.discover("01,05", kind="xlsx")] == ["01", "05"]


# --- text shared by the kinds -----------------------------------------------------------


def test_numbers_that_appear_in_no_input_are_listed_per_unit() -> None:
    result = text_checks.find_inventions(
        ["Q1 1,234 and 777", "only 1,234 and 12.5"], ["total 1234"], ["12.5"], "sheets"
    )

    assert result == {"count": 1, "numbers": [{"number": "777", "sheets": [1]}]}


def test_the_text_of_a_workbook_has_its_sheets_names_dates_and_plain_numbers() -> None:
    book = _book([["Day", "Amount"], [dt.datetime(2026, 1, 5), 1200.5], [None, 3.0]], "Orders")

    (text,) = xlsx_checks.workbook_texts(book)

    assert text.splitlines() == ["Orders", "Day Amount", "2026-01-05 1200.5", "3"]


# --- the run's records ------------------------------------------------------------------


def test_the_output_is_a_file_of_the_tasks_kind_that_changed(tmp_path: Path) -> None:
    (tmp_path / "input.xlsx").write_bytes(b"given")
    (tmp_path / "notes.pptx").write_bytes(b"deck")
    originals = {"input.xlsx": b"given"}
    assert run._find_output(tmp_path, originals, [], ".xlsx") is None

    (tmp_path / "input.xlsx").write_bytes(b"edited")
    assert run._find_output(tmp_path, originals, [], ".xlsx") == tmp_path / "input.xlsx"


def _recorded(tmp_path: Path, kind: str, report: dict[str, Any]) -> Path:
    folder = tmp_path / "run" / "06-budget-model"
    folder.mkdir(parents=True)
    result = {
        "task": "06-budget-model",
        "kind": kind,
        "status": "ok",
        "wall_seconds": 73.1,
        "tool_call_count": 18,
        "input_tokens": 100,
        "cached_tokens": 80,
        "output_tokens": 10,
        "cost": 0.0127,
    }
    (folder / "result.json").write_text(json.dumps(result), encoding="utf-8")
    (folder / "checks.json").write_text(json.dumps(report), encoding="utf-8")
    return tmp_path / "run"


XLSX_REPORT = {
    "opens": {"ok": True},
    "recalc": {"ok": True, "formulas": 9, "errors": 0},
    "facts": {"found": 2, "total": 2},
    "inventions": {"count": 5},
    "cells": {"ok": False, "checked": 15, "matched": 13},
    "formulas": {"ok": False},
    "constants": {"ok": True},
    "probe": {"ok": False},
}


def test_the_summary_of_an_excel_run_shows_what_failed(tmp_path: Path) -> None:
    out = _recorded(tmp_path, "xlsx", XLSX_REPORT)

    run.write_summary(out, tasks.discover("06", kind="xlsx"), "xlsx")

    lines = (out / "summary.md").read_text(encoding="utf-8").splitlines()
    head = [c.strip() for c in lines[0].strip("|").split("|")]
    row = dict(zip(head, [c.strip() for c in lines[2].strip("|").split("|")], strict=True))
    assert row["cells"] == "13/15" and row["formulas"] == "False" and row["probe"] == "False"
    assert row["constants"] == "True" and row["charts"] == "-" and row["cost"] == "0.0127"
    assert "pages" not in head and "overflow" not in head


async def test_a_relative_out_folder_reaches_the_session_as_an_absolute_one(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    seen: list[Path] = []

    class Session:
        async def handle_user_message(self, brief: str, recorder: Any) -> None:
            return None

    def build(
        root: Path,
        workspace: Path,
        model: str,
        thread_id: str,
        effort: str = "low",
    ) -> Session:
        seen.extend([root, workspace])
        return Session()

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(run, "_build_session", build)
    monkeypatch.setitem(run.CHECKERS, "xlsx", lambda task, path, ws, out_dir: {})

    await run.run_task(_task("06"), Path("runs/x"), "deepseek:m", "low", 60)

    assert seen and all(path.is_absolute() for path in seen)


def test_no_output_is_recorded_as_a_file_that_does_not_open(tmp_path: Path) -> None:
    folder = tmp_path / "t"
    folder.mkdir()
    (folder / "result.json").write_text(json.dumps({"output": None}), encoding="utf-8")

    report = run.check_task(_task("06"), folder)

    assert report["opens"] == {"ok": False, "error": "no xlsx file produced"}


# --- the judge and the comparison -------------------------------------------------------


def test_each_kind_is_judged_with_its_own_rubric(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    folder = tmp_path / "a" / "06-budget-model"
    (folder / "render").mkdir(parents=True)
    (folder / "render" / "page-01.png").write_bytes(b"png")
    seen: list[str] = []

    def fake(
        judge_name: str, brief: str, decks: dict, labels: list[str], rubric: str = ""
    ) -> tuple[dict, dict]:
        seen.append(rubric)
        scores = {"design": 4, "content": 4, "coherence": 4, "defects": []}
        return {"decks": {label: scores for label in labels}, "ranking": labels}, {}

    monkeypatch.setattr(judge, "judge_round", fake)

    judge.judge_runs([tmp_path / "a"], ["06-budget-model"], judge.DEFAULT_JUDGE, 1, 7, "xlsx")

    assert seen == [judge.RUBRIC_XLSX] and "spreadsheet" in judge.RUBRIC_XLSX
    assert judge.RUBRIC_XLSX != judge.RUBRIC


def test_the_comparison_of_two_excel_runs_shows_the_workbook_checks(tmp_path: Path) -> None:
    a = _recorded(tmp_path / "x", "xlsx", XLSX_REPORT)
    better = {**XLSX_REPORT, "cells": {"ok": True, "checked": 15, "matched": 15}}
    better["formulas"] = {"ok": True}
    b = _recorded(tmp_path / "y", "xlsx", better)

    table = compare.compare(a, b, kind="xlsx")

    assert "| 06-budget-model | 2/2 | 2/2 | yes (0 errors) | yes (0 errors) |" in table
    assert "| 13/15 | 15/15 |" in table
    assert "NO | yes" in table
    assert "overflow" not in table and "pages" not in table
