"""Canonical text of an `.xlsx`, for golden tests that must be byte-identical across runs.

The generic parts (zip order, attribute order, whitespace, the core-properties timestamps, an
embedded package) come from `pptx_canonical`. A workbook stores its styles as indexes into
`xl/styles.xml`, which makes a diff of the raw XML hard to read, so the canonical text starts
with a `cells` section that resolves every cell to its value, formula, number format, font,
fill, alignment and border, followed by the XML parts themselves. Both are kept: the cells say
what a person sees, the parts catch what the cells do not cover (charts, rules, print setup,
defined names).
"""

from __future__ import annotations

import hashlib
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell

from .pptx_canonical import canonical_part

CELLS_PART = "<cells>"


def _look(cell: Cell) -> str:
    font, fill, align, border = cell.font, cell.fill, cell.alignment, cell.border
    bits = [f"fmt={cell.number_format!r}"]
    if font.b:
        bits.append("bold")
    if font.i:
        bits.append("italic")
    if font.u:
        bits.append(f"underline={font.u}")
    bits.append(f"font={font.name},{font.sz}")
    # A theme or indexed colour has no rgb string; only an explicit one is a choice of the tools.
    if font.color is not None and isinstance(font.color.rgb, str):
        bits.append(f"color={font.color.rgb}")
    if fill.fill_type:
        bits.append(f"fill={fill.fgColor.rgb}")
    if align.horizontal or align.vertical or align.wrap_text or align.indent:
        bits.append(
            f"align={align.horizontal},{align.vertical},wrap={align.wrap_text},indent={align.indent}"
        )
    sides = {n: getattr(border, n).style for n in ("left", "right", "top", "bottom")}
    if any(sides.values()):
        bits.append("border=" + ",".join(f"{n}:{s}" for n, s in sides.items() if s))
    return " ".join(bits)


def _cells_text(path: Path) -> str:
    # Opened twice: once for the formulas as written, once for the values the file caches.
    formulas = load_workbook(path)
    values = load_workbook(path, data_only=True)
    lines: list[str] = []
    for sheet in formulas.worksheets:
        lines.append(f"[{sheet.title}] state={sheet.sheet_state} dimensions={sheet.dimensions}")
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None and not cell.has_style:
                    continue
                shown = repr(cell.value)
                if cell.data_type == "f":
                    shown += f" -> {values[sheet.title][cell.coordinate].value!r}"
                note = f" note={cell.comment.text!r}" if cell.comment else ""
                lines.append(f"{cell.coordinate} {shown}{note} | {_look(cell)}")
    return "\n".join(lines) + "\n"


def canonical_parts(path: Path) -> dict[str, str]:
    """The resolved cells, then every member of the zip at `path`, as canonical text."""
    parts = {CELLS_PART: _cells_text(path)}
    with zipfile.ZipFile(path) as archive:
        for name in sorted(archive.namelist()):
            data = archive.read(name)
            # openpyxl's stock theme, the same in every workbook; its hash still shows a change.
            parts[name] = (
                f"<sha256={hashlib.sha256(data).hexdigest()}>\n"
                if name.startswith("xl/theme/")
                else canonical_part(name, data)
            )
    return parts
