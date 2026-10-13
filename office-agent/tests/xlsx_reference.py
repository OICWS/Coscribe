"""Reference solutions of the Excel evaluation tasks: what a good answer looks like, built with
openpyxl. They show that every task can be solved and that its `expect.json` accepts it."""

from __future__ import annotations

import csv
import datetime as dt
import shutil
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation


def _widen(book: Workbook, width: float = 18) -> None:
    """Room for every number and label, so nothing prints as ### or is cut off."""
    for sheet in book.worksheets:
        for column in range(1, sheet.max_column + 1):
            dimension = sheet.column_dimensions[sheet.cell(row=1, column=column).column_letter]
            dimension.width = max(dimension.width or 0, width)


def clean_csv(inputs: Path, out: Path) -> Path:
    wanted = []
    with (inputs / "orders_export.csv").open(encoding="utf-8", newline="") as handle:
        for row in csv.reader(handle):
            if row and any(c.strip() for c in row):
                wanted.append([c.strip() for c in row])
    header, rows = wanted[0], wanted[1:]
    parsed, seen = [], set()
    for date, customer, category, amount in rows:
        for pattern in ("%Y/%m/%d", "%d %b %Y", "%Y-%m-%d"):
            try:
                day = dt.datetime.strptime(date, pattern)
                break
            except ValueError:
                continue
        key = (day, customer, category.title(), amount)
        if key in seen:
            continue
        seen.add(key)
        number = float(amount.replace("$", "").replace(",", ""))
        parsed.append((day, customer, category.title(), number))
    parsed.sort(key=lambda r: r[0])
    book = Workbook()
    sheet = book.active
    sheet.title = "Orders"
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in parsed:
        sheet.append(list(row))
    for r in range(2, len(parsed) + 2):
        sheet.cell(row=r, column=1).number_format = "yyyy-mm-dd"
        sheet.cell(row=r, column=4).number_format = "#,##0.00"
    path = out / "orders_clean.xlsx"
    _widen(book)
    book.save(path)
    return path


def monthly_summary(inputs: Path, out: Path) -> Path:
    path = out / "transactions.xlsx"
    shutil.copyfile(inputs / "transactions.xlsx", path)
    book = load_workbook(path)
    sheet = book.create_sheet("Summary")
    sheet.append(["Category", "Jan", "Feb", "Mar", "Total"])
    for row, category in enumerate(["Food", "Rent", "Transport"], start=2):
        sheet.cell(row=row, column=1, value=category)
        for column, month in zip("BCD", (1, 2, 3), strict=True):
            sheet[f"{column}{row}"] = (
                f"=SUMIFS(Data!$C$2:$C$15,Data!$B$2:$B$15,$A{row},"
                f'Data!$A$2:$A$15,">="&DATE(2026,{month},1),'
                f'Data!$A$2:$A$15,"<"&DATE(2026,{month + 1},1))'
            )
        sheet[f"E{row}"] = f"=SUM(B{row}:D{row})"
    sheet["A5"] = "Total"
    for column in "BCDE":
        sheet[f"{column}5"] = f"=SUM({column}2:{column}4)"
    _widen(book)
    book.save(path)
    return path


def lookup_two_sheets(inputs: Path, out: Path) -> Path:
    path = out / "shop.xlsx"
    shutil.copyfile(inputs / "shop.xlsx", path)
    book = load_workbook(path)
    sheet = book["Orders"]
    sheet["D1"], sheet["E1"], sheet["F1"] = "Product", "Unit price", "Line total"
    for row in range(2, 10):
        sheet[f"D{row}"] = f"=VLOOKUP(B{row},Prices!$A$2:$C$7,2,FALSE)"
        sheet[f"E{row}"] = f"=VLOOKUP(B{row},Prices!$A$2:$C$7,3,FALSE)"
        sheet[f"F{row}"] = f"=E{row}*C{row}"
    sheet["E11"], sheet["F11"] = "Total", "=SUM(F2:F9)"
    _widen(book)
    book.save(path)
    return path


def line_chart(inputs: Path, out: Path) -> Path:
    book = Workbook()
    sheet = book.active
    sheet.title = "Data"
    with (inputs / "monthly_visits.csv").open(encoding="utf-8", newline="") as handle:
        for index, row in enumerate(csv.reader(handle)):
            sheet.append(row if index == 0 else [row[0], int(row[1])])
    chart = LineChart()
    chart.title = "Monthly visitors"
    chart.x_axis.title = "Month"
    chart.y_axis.title = "Visitors"
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=13), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=13))
    sheet.add_chart(chart, "D2")
    path = out / "visits.xlsx"
    _widen(book)
    book.save(path)
    return path


def edit_budget(inputs: Path, out: Path) -> Path:
    path = out / "budget_q1.xlsx"
    shutil.copyfile(inputs / "budget_q1.xlsx", path)
    book = load_workbook(path)
    sheet = book["Budget"]
    sheet["D1"] = "Variance"
    sheet["D1"]._style = sheet["C1"]._style
    for row in range(2, 8):
        sheet[f"D{row}"] = f"=C{row}-B{row}"
        sheet[f"D{row}"]._style = sheet[f"C{row}"]._style
    sheet["A8"] = "Total"
    for column in "BCD":
        sheet[f"{column}8"] = f"=SUM({column}2:{column}7)"
        sheet[f"{column}8"]._style = sheet["C2"]._style
    _widen(book)
    book.save(path)
    return path


def budget_model(inputs: Path, out: Path) -> Path:
    book = Workbook()
    sheet = book.active
    sheet.title = "Model"
    for row, (label, value) in enumerate(
        [
            ("Units sold", 1000),
            ("Price per unit", 25),
            ("Cost per unit", 14),
            ("Fixed costs", 6000),
        ],
        start=2,
    ):
        sheet.cell(row=row, column=1, value=label)
        sheet.cell(row=row, column=2, value=value)
    results = [
        ("Revenue", "=B2*B3"),
        ("Variable costs", "=B2*B4"),
        ("Gross profit", "=B7-B8"),
        ("Net profit", "=B9-B5"),
        ("Net margin", "=B10/B7"),
    ]
    for row, (label, formula) in enumerate(results, start=7):
        sheet.cell(row=row, column=1, value=label)
        sheet.cell(row=row, column=2, value=formula)
    sheet["B11"].number_format = "0.0%"
    path = out / "budget_model.xlsx"
    _widen(book)
    book.save(path)
    return path


def chinese_formats(inputs: Path, out: Path) -> Path:
    book = Workbook()
    sheet = book.active
    sheet.title = "汇总"
    sheet.append(["区域", "销售额（元）", "销售额（万元）", "占比"])
    with (inputs / "各区域销售.csv").open(encoding="utf-8", newline="") as handle:
        rows = list(csv.reader(handle))[1:]
    for index, (region, amount) in enumerate(rows, start=2):
        sheet.append([region, int(amount), f"=B{index}/10000", f"=B{index}/$B$7"])
    sheet.append(["合计", "=SUM(B2:B6)", "=B7/10000", "=B7/$B$7"])
    for row in range(2, 8):
        sheet[f"B{row}"].number_format = '#,##0"元"'
        sheet[f"C{row}"].number_format = '#,##0.00"万"'
        sheet[f"D{row}"].number_format = "0.0%"
    path = out / "区域销售汇总.xlsx"
    _widen(book)
    book.save(path)
    return path


def validation_list(inputs: Path, out: Path) -> Path:
    path = out / "tickets.xlsx"
    shutil.copyfile(inputs / "tickets.xlsx", path)
    book = load_workbook(path)
    sheet = book["Tickets"]
    validation = DataValidation(type="list", formula1='"Open,In progress,Done"', allow_blank=False)
    sheet.add_data_validation(validation)
    validation.add("C2:C11")
    for row in range(2, 12):
        sheet[f"C{row}"] = "Open"
    sheet["G1"] = "Open tickets"
    sheet["H1"] = '=COUNTIF(C2:C11,"Open")'
    _widen(book)
    book.save(path)
    return path


SOLUTIONS = {
    "01-clean-csv": clean_csv,
    "02-monthly-summary": monthly_summary,
    "03-lookup-two-sheets": lookup_two_sheets,
    "04-line-chart": line_chart,
    "05-edit-budget": edit_budget,
    "06-budget-model": budget_model,
    "07-chinese-formats": chinese_formats,
    "08-validation-list": validation_list,
}
