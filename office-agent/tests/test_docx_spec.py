"""The spec a Word document is written from (decision 0024): its head and components, the facts it
defines once, and what the compiler writes for each."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
from docx import Document

from coscribe.tools.office.document_tools import build_document_tools
from coscribe.tools.office.docx import read_doc_blocks
from coscribe.tools.office.docx.spec import SpecError, parse_spec

SPEC = """---
title: Quarterly report
subtitle: Product and operations
author: Ana Lee
date: October 2026
cover: band
toc: true
header: Quarterly report
footer: Internal
facts:
  revenue: 12.8 M
  growth: 18%
---

# Overview

Revenue reached {fact:revenue}, up {fact:growth}.

:::stats
Revenue | {fact:revenue} | +{fact:growth}
Users | 4.2 M | +31%
On-time rate | 96.5% | -1.2 pt
:::

:::callout warn
The queue is still long.
- A third of the requests repeat
:::

:::pagebreak
:::

# Closing

| Item | Q2 | Q3 |
| --- | --- | --- |
| Revenue | 10.9 | 12.8 |
"""


def _tools(tmp_path: Path) -> dict[str, Any]:
    return {t.__name__: t for t in build_document_tools(tmp_path)}  # type: ignore[attr-defined]


def _write(tmp_path: Path, content: str, **kwargs: Any) -> Document:
    _tools(tmp_path)["write_docx"](path="out.docx", content=content, **kwargs)
    return Document(str(tmp_path / "out.docx"))


def _xml(document: Any) -> str:
    return document.element.xml


def test_a_content_without_head_or_component_is_plain_markdown() -> None:
    assert parse_spec("# Title\n\nSome text\n\n- a\n- b").is_plain
    assert not parse_spec("---\ntitle: x\n---\n\ntext").is_plain
    assert not parse_spec(":::pagebreak\n:::").is_plain


def test_a_fact_is_defined_once_and_expanded_everywhere_it_is_used() -> None:
    head = "---\nfacts:\n  total: 6.7 M\n---\n\n"
    spec = parse_spec(head + "Total {fact:total}.\n\n:::callout info\n{fact:total}\n:::")
    assert spec.items[0] == "\nTotal 6.7 M.\n"
    assert spec.items[1].body == "6.7 M"  # type: ignore[union-attr]


@pytest.mark.parametrize(
    ("content", "message"),
    [
        ("---\ncolor: red\n---\ntext", "Unknown document head key 'color'"),
        ("---\ncover: fancy\n---\ntext", "cover is one of"),
        ("---\ntitle: x", "no closing ---"),
        ("text {fact:missing}", "not defined in the head's facts"),
        (":::radar\n:::", "Unknown component 'radar'"),
        (":::stats\na | 1", "no closing :::"),
        (":::callout loud\ntext\n:::", "A callout is one of"),
    ],
)
def test_a_spec_the_compiler_cannot_write_is_refused_with_what_to_change(
    content: str, message: str
) -> None:
    with pytest.raises(SpecError, match=message):
        parse_spec(content)


def test_the_compiler_writes_the_cover_contents_components_and_running_text(
    tmp_path: Path,
) -> None:
    document = _write(tmp_path, SPEC)
    xml = _xml(document)

    assert "Quarterly report" in [p.text for p in document.paragraphs if p.style.name == "Title"]
    for name in ("cs_cover", "cs_stats", "cs_callout_warn"):
        assert f'w:name="{name}"' in xml
    assert xml.count('w:type="page"') == 3  # after the cover, after the contents, the break
    assert "TOC \\o" in xml
    assert "Revenue reached 12.8 M, up 18%." in [p.text for p in document.paragraphs]
    section = document.sections[0]
    assert section.different_first_page_header_footer
    assert "Quarterly report" in section.header.paragraphs[0].text
    assert "Internal" in section.footer.paragraphs[0].text
    assert section.first_page_footer.paragraphs[0].text == ""


def test_the_cards_and_the_call_out_take_their_colours_from_the_theme_not_the_spec(
    tmp_path: Path,
) -> None:
    xml = _xml(_write(tmp_path, SPEC))

    assert 'w:fill="1F4E79"' in xml  # the cover band, the theme's accent
    assert 'w:fill="FFF4DC"' in xml  # the warn call-out's tint
    assert 'w:color="D98A00"' in xml  # its bar
    assert xml.count("C0392B") >= 1  # a falling figure is red


def test_a_figure_is_a_centred_picture_with_its_caption(tmp_path: Path) -> None:
    from PIL import Image

    Image.new("RGB", (120, 60), "tomato").save(tmp_path / "chart.png")
    document = _write(tmp_path, ":::figure chart.png\nFigure 1: **Revenue**\n:::")

    assert len(document.inline_shapes) == 1
    caption = next(p for p in document.paragraphs if p.style.name == "Caption")
    assert caption.text == "Figure 1: Revenue"


def test_the_document_reads_back_as_blocks_the_editor_can_draw(tmp_path: Path) -> None:
    _write(tmp_path, SPEC)
    blocks = read_doc_blocks(tmp_path / "out.docx")["blocks"]

    assert [b["kind"] for b in blocks].count("pagebreak") == 3
    assert any(b["kind"] == "heading" and b["runs"][0]["t"] == "Overview" for b in blocks)
    assert any(b["kind"] == "table" for b in blocks)  # the closing table, editable


def test_a_cover_needs_a_title_and_components_cannot_be_tracked_changes(tmp_path: Path) -> None:
    tools = _tools(tmp_path)
    with pytest.raises(SpecError, match="A cover needs a title"):
        tools["write_docx"](path="a.docx", content="---\ncover: band\n---\n\ntext")
    with pytest.raises(SpecError, match="cannot be tracked changes"):
        tools["write_docx"](path="b.docx", content=":::pagebreak\n:::", track_changes=True)
    assert not (tmp_path / "a.docx").exists()


def test_a_stats_row_has_two_to_four_cards(tmp_path: Path) -> None:
    with pytest.raises(SpecError, match="two to four lines"):
        _write(tmp_path, ":::stats\nOnly | one\n:::")


CHART = """\
:::chart bar Revenue by quarter
| Quarter | Revenue |
| --- | --- |
| Q1 | 8.6 |
| Q2 | 10.9 |
| Q3 | 12.8 |
caption: Figure 1: **Revenue** by quarter
:::
"""


def test_a_chart_is_a_native_chart_with_its_caption(tmp_path: Path) -> None:
    document = _write(tmp_path, CHART)

    parts = [str(part.partname) for part in document.part.package.iter_parts()]
    assert "/word/charts/chart1.xml" in parts
    caption = next(p for p in document.paragraphs if p.style.name == "Caption")
    assert caption.text == "Figure 1: Revenue by quarter"


@pytest.mark.parametrize(
    ("content", "message"),
    [
        (":::chart radar T\n| a | b |\n| --- | --- |\n| x | 1 |\n:::", "the type one of"),
        (":::chart bar T\n| a | b |\n| --- | --- |\n| x | many |\n:::", "not numeric"),
        (
            ":::chart pie T\n| a | b | c |\n| --- | --- | --- |\n| x | 1 | 2 |\n:::",
            "one data column",
        ),
    ],
)
def test_a_chart_that_cannot_be_drawn_is_refused_with_what_to_change(
    tmp_path: Path, content: str, message: str
) -> None:
    with pytest.raises(SpecError, match=message):
        _write(tmp_path, content)


def test_every_component_is_found_again_when_the_file_is_read_as_blocks(tmp_path: Path) -> None:
    from PIL import Image

    Image.new("RGB", (120, 60), "tomato").save(tmp_path / "pic.png")
    content = SPEC + "\n" + CHART + "\n:::figure pic.png\nA picture\n:::\n"
    _write(tmp_path, content)

    blocks = read_doc_blocks(tmp_path / "out.docx")["blocks"]
    components = {b["component"] for b in blocks if "component" in b}

    assert components == {"cover", "stats", "callout_warn", "chart", "figure"}


def _check(tmp_path: Path, content: str) -> list[dict[str, str]]:
    result = _tools(tmp_path)["write_docx"](path="out.docx", content=content)
    return list(result.get("check", []))


def test_a_clean_spec_has_no_findings(tmp_path: Path) -> None:
    assert _check(tmp_path, SPEC) == []


def test_plain_markdown_is_not_checked(tmp_path: Path) -> None:
    assert _check(tmp_path, "Text without a title\n\n### deep") == []


@pytest.mark.parametrize(
    ("content", "code"),
    [
        ("---\ntitle: T\n---\n\n# A\n\n### C", "heading-skip"),
        ("---\ntitle: T\nfacts:\n  a: 1\n---\n\n# A", "unused-fact"),
        ("---\ntitle: T\n---\n\n:::callout info\n:::", "empty-callout"),
        (
            "---\ntitle: T\n---\n\n:::chart bar\n| a | b |\n| --- | --- |\n| x | 1 |\n:::",
            "chart-untitled",
        ),
        ("---\ntoc: true\n---\n\ntext", "no-title"),
        ("---\ntitle: T\ntoc: true\n---\n\ntext", "empty-contents"),
    ],
)
def test_each_finding_has_a_code_a_place_and_a_fix(tmp_path: Path, content: str, code: str) -> None:
    found = {f["code"]: f for f in _check(tmp_path, content)}
    assert code in found
    assert found[code]["where"] and found[code]["fix"]


def test_a_chart_is_read_back_against_its_table(tmp_path: Path) -> None:
    from coscribe.tools.office.docx.spec_check import check_document

    content = (
        "---\ntitle: T\n---\n\n:::chart bar Sales\n"
        "| q | s |\n| --- | --- |\n| Q1 | 3 |\n| Q2 | 5 |\n:::"
    )
    document = _write(tmp_path, content)
    assert check_document(document, parse_spec(content)) == []
    changed = parse_spec(content.replace("| Q2 | 5 |", "| Q2 | 6 |"))
    assert [f.code for f in check_document(document, changed)] == ["chart-data"]


def test_a_document_missing_a_component_is_reported(tmp_path: Path) -> None:
    from coscribe.tools.office.docx.spec_check import check_document

    document = _write(tmp_path, "---\ntitle: T\n---\n\n# A")
    asked = parse_spec("---\ntitle: T\n---\n\n:::callout info\nx\n:::")
    assert [f.code for f in check_document(document, asked)] == ["components-lost"]
