"""The deterministic parts of the PowerPoint evaluation (`evals/`): reading a deck, the fact
and invention checks, the editability probe, task files, the run's records and its summary.
Nothing here calls a model; the render test needs LibreOffice and skips without it."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from coscribe.tools.office.pptx import shape_tools
from evals import pptx_checks, prices, run, tasks

# --- building small decks -------------------------------------------------------------


def _deck(path: Path, slides: list[tuple[str, list[str]]]) -> Path:
    prs = Presentation()
    for title, lines in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text_frame.text = "\n".join(lines)
    prs.save(path)
    return path


def _deck_with_extras(path: Path) -> Path:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Revenue"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Region"
    table.cell(1, 1).text = "5,410"
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(
        Inches(1), Inches(4), Inches(3), Inches(1)
    ).text_frame.text = "In a group"
    data = CategoryChartData()
    data.categories = ["East", "South"]
    data.add_series("Q3", (5410, 3380))
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5), Inches(2), Inches(4), Inches(3), data
    )
    prs.save(path)
    return path


# --- reading ----------------------------------------------------------------------------


def test_read_deck_reaches_tables_groups_and_chart_data(tmp_path: Path) -> None:
    deck = pptx_checks.read_deck(_deck_with_extras(tmp_path / "d.pptx"))

    text = deck.all_text()
    assert deck.slides[0].title == "Revenue"
    assert "5,410" in text and "In a group" in text
    assert "East" in text and "3380" in text


# --- facts ------------------------------------------------------------------------------


@pytest.mark.parametrize(
    "deck_text",
    ["营收 12,710 万", "营收 12710 万", "营收 １２，７１０ 万"],
)
def test_a_number_matches_whatever_its_separators(deck_text: str) -> None:
    assert pptx_checks.check_facts(deck_text, [["12,710"]], [])["ok"] is True


def test_a_fact_may_list_alternative_spellings_and_a_missing_one_is_named() -> None:
    result = pptx_checks.check_facts(
        "Win + L locks the screen", [["Win+L", "Windows+L"], ["8000"]], []
    )

    assert result["ok"] is False
    assert (result["found"], result["total"], result["missing"]) == (1, 2, ["8000"])


def test_text_that_must_be_gone_fails_the_facts_when_it_is_still_there() -> None:
    result = pptx_checks.check_facts("营收 1,280 万，现为 1,365 万", [["1,365"]], ["1,280"])

    assert result["ok"] is False
    assert result["forbidden_present"] == ["1,280"]


# --- invented numbers -------------------------------------------------------------------


def test_numbers_in_no_input_are_listed_with_their_slides(tmp_path: Path) -> None:
    deck = pptx_checks.read_deck(
        _deck(
            tmp_path / "d.pptx",
            [
                ("Q3", ["营收 12,710 万，增长 16.3%"]),
                ("Outlook", ["目标 2024 年翻倍", "增长 25.5%"]),
            ],
        )
    )

    result = pptx_checks.check_inventions(deck, ["营收 12710 万元，环比 +16.3%"], [])

    assert [(n["number"], n["slides"]) for n in result["numbers"]] == [("2024", [2]), ("25.5", [2])]


def test_a_count_converted_to_another_unit_is_not_an_invention(tmp_path: Path) -> None:
    deck = pptx_checks.read_deck(
        _deck(
            tmp_path / "d.pptx", [("Exports", ["2.615 million units", "SGMW 267k", "约 178 万辆"])]
        )
    )

    result = pptx_checks.check_inventions(deck, ["出口 261.5 万辆，五菱 26.7 万辆"], [])

    assert [n["number"] for n in result["numbers"]] == ["178"]


def test_small_integers_and_allowed_numbers_are_not_claims(tmp_path: Path) -> None:
    deck = pptx_checks.read_deck(
        _deck(tmp_path / "d.pptx", [("Page 3 of 12", ["30% 对冲，第 7 条", "99.9%"])])
    )

    result = pptx_checks.check_inventions(deck, [], ["30%", "99.9%"])

    assert result["count"] == 0


# --- pages, pictures, template ----------------------------------------------------------


def test_page_count_is_inside_or_outside_the_range() -> None:
    assert pptx_checks.check_page_count(6, 5, 7)["ok"] is True
    assert pptx_checks.check_page_count(8, 5, 7)["ok"] is False
    assert pptx_checks.check_page_count(15, 15, 15)["ok"] is True


def test_pictures_are_counted_inside_groups_too(tmp_path: Path) -> None:
    from PIL import Image

    Image.new("RGB", (20, 20), "red").save(tmp_path / "p.png")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(tmp_path / "p.png"), 0, 0)
    slide.shapes.add_group_shape().shapes.add_picture(str(tmp_path / "p.png"), 0, 0)
    prs.save(tmp_path / "d.pptx")

    count = pptx_checks.count_pictures(tmp_path / "d.pptx")

    assert count == 2
    assert pptx_checks.check_pictures(count, 3)["ok"] is False


def test_a_deck_without_the_templates_layouts_did_not_keep_it() -> None:
    assert pptx_checks.check_template_kept({"Title", "Body"}, {"Title"})["ok"] is True
    kept = pptx_checks.check_template_kept({"Blank"}, {"Title", "Body"})
    assert kept == {"ok": False, "missing_layouts": ["Body", "Title"]}


# --- findings of the review ---------------------------------------------------------------


def test_a_numeric_fact_is_not_found_inside_a_longer_number() -> None:
    def found(deck_text: str, fact: str) -> bool:
        return pptx_checks.check_facts(deck_text, [[fact]], [])["ok"]

    assert not found("需要 13 周", "3 周")
    assert not found("增长 142%", "42%")
    assert not found("共 1,200 篇", "120")
    assert not found("占 0.42%", "42%")
    assert found("需要 3 周完成", "3 周") and found("增长 42%，", "42%") and found("120 篇", "120")
    assert found("Q3 12,710 万", "12,710") and not found("Q3 12,710 万", "312,710")
    assert found("先 8000 后", "8000") and found("按 Win+L 锁屏", "Win+L")


def test_a_large_chart_value_is_read_as_written(tmp_path: Path) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    data = CategoryChartData()
    data.categories = ["East", "South"]
    data.add_series("Q3", (12710000, 0.5))
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(4), Inches(3), data
    )
    prs.save(tmp_path / "d.pptx")

    text = pptx_checks.read_deck(tmp_path / "d.pptx").all_text()

    assert "12710000" in text and "0.5" in text and "e+" not in text


def test_speaker_notes_are_read(tmp_path: Path) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "T"
    slide.notes_slide.notes_text_frame.text = "say 99 clearly"
    prs.save(tmp_path / "d.pptx")

    assert "say 99 clearly" in pptx_checks.read_deck(tmp_path / "d.pptx").all_text()


def test_a_shape_that_reuses_the_titles_id_is_still_read(tmp_path: Path) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Title"
    box = slide.shapes.add_textbox(0, 0, Inches(2), Inches(1))
    box.text_frame.text = "same id body"
    box._element.nvSpPr.cNvPr.set("id", str(slide.shapes.title.shape_id))
    prs.save(tmp_path / "d.pptx")

    deck = pptx_checks.read_deck(tmp_path / "d.pptx")

    assert deck.slides[0].title == "Title" and "same id body" in deck.slides[0].texts


def test_the_pdf_is_kept_when_only_the_page_images_fail(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from coscribe.tools import _thumbnail

    def make_pdf(path: Path, out_dir: Path) -> Path:
        pdf = out_dir / "d.pdf"
        pdf.write_bytes(b"%PDF")
        return pdf

    def no_images(*args: Any) -> list[Path]:
        raise RuntimeError("no renderer")

    monkeypatch.setattr(_thumbnail, "_convert_to_pdf", make_pdf)
    monkeypatch.setattr(_thumbnail, "_render_pdf_pages", no_images)

    result = pptx_checks.render_deck(tmp_path / "d.pptx", tmp_path / "out")

    assert result["ok"] is False and result["pdf"] == "deck.pdf"
    assert (tmp_path / "out" / "deck.pdf").is_file()


def test_numbers_in_word_and_excel_inputs_are_not_inventions(tmp_path: Path) -> None:
    from docx import Document
    from openpyxl import Workbook

    document = Document()
    document.add_paragraph("Budget 86 万")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Headcount 245"
    document.save(tmp_path / "in.docx")
    workbook = Workbook()
    workbook.active.append(["Sales", 18650])
    workbook.save(tmp_path / "in.xlsx")

    assert "86 万" in tasks.input_text(tmp_path / "in.docx")
    assert "245" in tasks.input_text(tmp_path / "in.docx")
    assert "18650" in tasks.input_text(tmp_path / "in.xlsx")
    assert tasks.input_text(tmp_path / "x.png") == ""


# --- the editability probe --------------------------------------------------------------


def test_a_titled_deck_whose_shapes_list_is_editable(tmp_path: Path) -> None:
    path = _deck(tmp_path / "d.pptx", [("One", ["a"]), ("Two", ["b"])])

    assert pptx_checks.check_editability(path, tmp_path) == {
        "untitled_slides": [],
        "shape_list_errors": [],
        "ok": True,
    }


def test_a_slide_without_a_title_is_listed(tmp_path: Path) -> None:
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[1]).shapes.title.text = "One"
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(tmp_path / "d.pptx")

    result = pptx_checks.check_editability(tmp_path / "d.pptx", tmp_path)

    assert result["untitled_slides"] == [2]
    assert result["ok"] is False


def test_a_crash_of_the_shape_list_is_recorded_per_slide(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = _deck(tmp_path / "d.pptx", [("One", ["a"]), ("Two", ["b"])])

    def broken(index: int, shape: Any) -> dict[str, object]:
        raise TypeError("a group shape cannot have a click action")

    monkeypatch.setattr(shape_tools, "_describe_shape", broken)

    result = pptx_checks.check_editability(path, tmp_path)

    assert [e["slide"] for e in result["shape_list_errors"]] == [1, 2]
    assert "TypeError" in result["shape_list_errors"][0]["error"]
    assert result["ok"] is False


# --- task files -------------------------------------------------------------------------


def test_every_task_loads_and_its_files_exist() -> None:
    found = tasks.discover("all")

    assert len(found) >= 10
    for task in found:
        assert task.brief and task.expect.facts
        assert task.expect.pages_min <= task.expect.pages_max
        if task.expect.template:
            assert task.expect.template in {p.name for p in task.inputs}


def test_the_briefs_name_the_input_files_that_exist() -> None:
    for task in tasks.discover("all"):
        for path in task.inputs:
            assert path.name in task.brief, f"{task.name}: the brief never mentions {path.name}"


def test_a_selector_takes_numbers_names_and_rejects_unknown_ones() -> None:
    assert [t.name[:2] for t in tasks.discover("01,03")] == ["01", "03"]
    assert [t.name for t in tasks.discover("02-security-training")] == ["02-security-training"]
    with pytest.raises(ValueError, match="no such task: 99"):
        tasks.discover("99")


def test_a_task_with_bad_expectations_is_refused(tmp_path: Path) -> None:
    (tmp_path / "brief.md").write_text("make a deck", encoding="utf-8")
    (tmp_path / "expect.json").write_text(
        json.dumps({"facts": ["x"], "pages": [5, 3]}), encoding="utf-8"
    )

    with pytest.raises(ValueError, match="pages"):
        tasks.load_task(tmp_path)


def test_source_texts_include_text_inputs_and_input_decks(tmp_path: Path) -> None:
    (tmp_path / "inputs").mkdir()
    (tmp_path / "brief.md").write_text("brief 100", encoding="utf-8")
    (tmp_path / "inputs" / "data.csv").write_text("a,200", encoding="utf-8")
    _deck(tmp_path / "inputs" / "old.pptx", [("Old", ["300"])])
    (tmp_path / "expect.json").write_text(
        json.dumps({"facts": ["x"], "pages": [1, 2]}), encoding="utf-8"
    )

    texts = "\n".join(tasks.load_task(tmp_path).source_texts())

    assert "brief 100" in texts and "a,200" in texts and "300" in texts


# --- the run's records ------------------------------------------------------------------


def test_usage_adds_tokens_cache_and_reasoning_per_model_call() -> None:
    usage = run.Usage()
    usage.add(
        {
            "input_tokens": 1000,
            "output_tokens": 50,
            "input_token_details": {"cache_read": 900},
            "output_token_details": {"reasoning": 20},
        }
    )
    usage.add({"input_tokens": 1100, "output_tokens": 10})

    assert (usage.model_calls, usage.input_tokens, usage.cached_tokens) == (2, 2100, 900)
    assert (usage.output_tokens, usage.reasoning_tokens) == (60, 20)


def test_cost_prices_cached_uncached_and_output_tokens() -> None:
    value = prices.cost("deepseek:deepseek-flash", 1_000_000, 800_000, 100_000)

    assert value == pytest.approx(0.8 * 0.003 + 0.2 * 0.15 + 0.1 * 0.6)
    assert prices.cost("other:unknown", 1, 0, 1) is None


class _FakeSession:
    def __init__(self) -> None:
        self.answers: dict[str, Any] = {}
        self.approvals: dict[str, bool] = {}

    def resolve_question(self, request_id: str, answer: Any) -> None:
        self.answers[request_id] = answer

    def resolve_approval(self, request_id: str, approved: bool) -> None:
        self.approvals[request_id] = approved


async def test_the_recorder_keeps_what_each_deck_build_found() -> None:
    recorder = run.Recorder(session=_FakeSession())
    report = {
        "ok": False,
        "slides": 2,
        "pages": [
            {"page": "01.svg", "shapes": 3},
            {"page": "02.svg", "shapes": 5, "errors": ["text overlaps"], "warnings": ["faint"]},
        ],
    }

    await recorder.send_json(
        {
            "type": "tool_result",
            "tool_name": "build_pptx_from_pages",
            "is_error": False,
            "result": json.dumps(report),
        }
    )
    await recorder.send_json(
        {
            "type": "tool_result",
            "tool_name": "build_pptx_from_pages",
            "is_error": False,
            "result": {"ok": True, "pages": [], "warnings": ["stop repairing"]},
        }
    )

    assert recorder.builds == [
        {
            "ok": False,
            "pages": [{"page": "02.svg", "errors": ["text overlaps"], "warnings": ["faint"]}],
            "warnings": [],
        },
        {"ok": True, "pages": [], "warnings": ["stop repairing"]},
    ]


async def test_the_recorder_sees_a_deferred_tool_by_its_own_name() -> None:
    recorder = run.Recorder(session=_FakeSession())
    called = {"name": "build_pptx_from_pages", "arguments": {}}

    await recorder.send_json({"type": "tool_started", "tool_name": "use_tool", "arguments": called})
    await recorder.send_json(
        {
            "type": "tool_result",
            "tool_name": "use_tool",
            "arguments": called,
            "is_error": False,
            "result": {"ok": True, "pages": []},
        }
    )

    assert recorder.builds == [{"ok": True, "pages": [], "warnings": []}]
    assert set(recorder.tool_seconds) == {"build_pptx_from_pages"}


async def test_the_recorder_counts_calls_and_answers_what_would_wait_for_a_person() -> None:
    session = _FakeSession()
    recorder = run.Recorder(session=session)

    await recorder.send_json({"type": "tool_started", "tool_name": "write_pptx", "arguments": {}})
    await recorder.send_json(
        {"type": "tool_started", "tool_name": "use_tool", "arguments": {"name": "add_pptx_chart"}}
    )
    await recorder.send_json(
        {"type": "tool_result", "tool_name": "write_pptx", "is_error": True, "result": "boom"}
    )
    await recorder.send_json({"type": "question_required", "id": "q1", "questions": [{}, {}]})
    await recorder.send_json({"type": "plan_ready", "id": "q2", "plan": "p"})
    await recorder.send_json({"type": "approval_required", "id": "a1"})
    await recorder.send_json({"type": "agent_message", "text": "done"})
    await recorder.send_json({"type": "error", "message": "bad"})

    assert dict(recorder.tool_calls) == {"write_pptx": 1, "use_tool:add_pptx_chart": 1}
    assert recorder.tool_errors == ["write_pptx: boom"]
    assert len(session.answers["q1"]) == 2 and isinstance(session.answers["q2"], str)
    assert recorder.questions == 2 and session.approvals == {"a1": True}
    assert (recorder.messages, recorder.errors) == (["done"], ["bad"])


def test_the_produced_deck_is_the_newest_one_that_is_not_an_untouched_input(
    tmp_path: Path,
) -> None:
    _deck(tmp_path / "input.pptx", [("In", ["x"])])
    originals = {"input.pptx": (tmp_path / "input.pptx").read_bytes()}
    assert run._find_output(tmp_path, originals, [], ".pptx") is None

    _deck(tmp_path / "made.pptx", [("Out", ["y"])])
    assert run._find_output(tmp_path, originals, [], ".pptx") == tmp_path / "made.pptx"


def test_a_deck_the_agent_names_wins_over_a_newer_scratch_one(tmp_path: Path) -> None:
    import os

    real = _deck(tmp_path / "report.pptx", [("Real", ["a"])])
    scratch = _deck(tmp_path / "thumb-test.pptx", [("Scratch", ["b"])])
    os.utime(real, (1, 1))
    os.utime(scratch, (2, 2))

    assert run._find_output(tmp_path, {}, ["done"], ".pptx") == scratch
    assert run._find_output(tmp_path, {}, ["saved as report.pptx"], ".pptx") == real


async def test_each_overlapping_task_gets_its_own_usage_and_the_resolver_is_restored(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    import asyncio

    from coscribe.conversation import session as session_module

    class Model:
        callbacks: Any = None
        reasoning_effort: str | None = None

    monkeypatch.setattr(session_module, "resolve_chat_model", lambda name, providers=None: Model())
    original = session_module.resolve_chat_model
    seen: dict[str, Any] = {}

    async def task(name: str, delay: float) -> None:
        run._CURRENT_USAGE.set(run.Usage())
        await asyncio.sleep(delay)
        seen[name] = session_module.resolve_chat_model("deepseek:m")

    with run.measured_models():
        await asyncio.gather(task("fast", 0.0), task("slow", 0.05))

    assert session_module.resolve_chat_model is original
    assert seen["fast"].callbacks[0] is not seen["slow"].callbacks[0]


def test_the_effort_levels_are_the_products() -> None:
    from coscribe.runtime_lg.effort import LEVELS

    assert run.EFFORT_LEVELS == LEVELS


def test_a_run_reaches_the_session_at_its_effort_level_as_the_products_default(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("DEEPSEEK_API_KEY", "k")

    session = run._build_session(
        tmp_path, tmp_path / "ws", "deepseek:deepseek-flash", "t", effort="extra"
    )

    assert session._effort == "extra"
    assert session._effort_runs_as == "xhigh"
    assert session.model.reasoning_effort == "xhigh"


def test_the_summary_has_one_row_per_task_that_ran(tmp_path: Path) -> None:
    task = tasks.discover("01")[0]
    folder = tmp_path / task.name
    folder.mkdir()
    result = {
        "task": task.name,
        "status": "ok",
        "wall_seconds": 12.5,
        "tool_call_count": 4,
        "input_tokens": 100,
        "cached_tokens": 80,
        "output_tokens": 10,
        "cost": 0.0123,
    }
    report = {
        "opens": {"ok": True},
        "facts": {"found": 18, "total": 19},
        "pages": {"ok": True, "count": 6},
        "editability": {"ok": False},
        "inventions": {"count": 2},
        "layout": {"overflow_words": 0, "text_overlaps": 1},
    }
    (folder / "result.json").write_text(json.dumps(result), encoding="utf-8")
    (folder / "checks.json").write_text(json.dumps(report), encoding="utf-8")

    run.write_summary(tmp_path, tasks.discover("all"))

    rows = json.loads((tmp_path / "summary.json").read_text(encoding="utf-8"))
    assert [(r["task"], r["facts"], r["pages"], r["editable"], r["overlaps"]) for r in rows] == [
        (task.name, "18/19", 6, False, 1)
    ]
    assert "| 01-q3-review | ok |" in (tmp_path / "summary.md").read_text(encoding="utf-8")


# --- the render (needs LibreOffice) -----------------------------------------------------


@pytest.mark.real_libreoffice
def test_a_deck_renders_to_one_png_per_page(tmp_path: Path) -> None:
    path = _deck(tmp_path / "d.pptx", [("One", ["a"]), ("Two", ["b"])])

    result = pptx_checks.render_deck(path, tmp_path / "render")

    if not result["ok"]:
        pytest.skip(f"LibreOffice is not usable here: {result['reason']}")
    assert result["pages"] == ["page-01.png", "page-02.png"]
    assert (
        pptx_checks.check_layout(path, tmp_path / "render" / result["pdf"])["overflow_words"] == 0
    )


def test_a_picture_is_kept_when_re_saved_and_missing_when_swapped() -> None:
    from PIL import Image

    red, blue = Image.new("RGB", (32, 32), "#C03030"), Image.new("RGB", (32, 32), "#3050C0")
    red_again = Image.new("RGB", (32, 32), "#C23131")

    assert pptx_checks.check_pictures_kept([red, blue], [blue, red_again])["ok"]
    assert pptx_checks.check_pictures_kept([red, blue], [blue, blue])["missing"] == [1]
