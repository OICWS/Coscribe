"""The Excel tools as the agent calls them: read_xlsx, recalc_xlsx, write_xlsx_cells and
modify_xlsx, bound to a workspace. What the blocks and operations do is in test_xlsx_blocks.py."""

import concurrent.futures
import threading
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from coscribe.runtime import get_tool_metadata
from coscribe.tools.office.xlsx import build_spreadsheet_tools

_SCORES = [["name", "score"], ["Alice", 9], ["Bob", 7]]


def _tools(root: Path, **kwargs: object) -> dict[str, object]:
    return {tool.__name__: tool for tool in build_spreadsheet_tools(root, **kwargs)}  # type: ignore[arg-type, attr-defined]


def _scores(tools: dict[str, object], path: str, sheet: str = "Scores") -> None:
    tools["write_xlsx_cells"](path=path, blocks=[{"sheet": sheet, "rows": _SCORES}])  # type: ignore[operator]


def test_the_agent_has_exactly_the_four_excel_tools(tmp_path: Path) -> None:
    tools = _tools(tmp_path)

    assert sorted(tools) == ["modify_xlsx", "read_xlsx", "recalc_xlsx", "write_xlsx_cells"]
    assert get_tool_metadata(tools["read_xlsx"]).risk_category == "READ"
    for name in ("recalc_xlsx", "write_xlsx_cells", "modify_xlsx"):
        metadata = get_tool_metadata(tools[name])
        assert metadata.risk_category == "WRITE_LOCAL" and metadata.requires_approval is True


def test_a_written_table_reads_back(tmp_path: Path) -> None:
    tools = _tools(tmp_path)
    _scores(tools, "report.xlsx")

    text = tools["read_xlsx"](path="report.xlsx", sheet="Scores")  # type: ignore[operator]

    assert "| name | score |" in text and "| Alice | 9 |" in text and "| Bob | 7 |" in text


def test_reading_without_a_sheet_shows_every_sheet_under_its_name(tmp_path: Path) -> None:
    tools = _tools(tmp_path)
    _scores(tools, "report.xlsx")
    tools["write_xlsx_cells"](path="report.xlsx", blocks=[{"sheet": "Raw", "rows": [["x"], [1]]}])  # type: ignore[operator]

    text = tools["read_xlsx"](path="report.xlsx")  # type: ignore[operator]

    assert "## Scores" in text and "## Raw" in text and "| 1 |" in text


def test_an_unknown_sheet_is_refused_with_the_real_names(tmp_path: Path) -> None:
    tools = _tools(tmp_path)
    _scores(tools, "report.xlsx")

    with pytest.raises(ValueError, match="Scores"):
        tools["read_xlsx"](path="report.xlsx", sheet="Nope")  # type: ignore[operator]


def test_recalc_xlsx_computes_a_workbook_written_some_other_way(tmp_path: Path) -> None:
    book = Workbook()
    book.active.append(["a", "b", "sum"])
    book.active.append([3, 4, "=SUM(A2:B2)"])
    book.save(tmp_path / "report.xlsx")

    result = _tools(tmp_path)["recalc_xlsx"](path="report.xlsx")  # type: ignore[operator]

    assert result["status"] == "success" and result["total_formulas"] == 1
    assert load_workbook(tmp_path / "report.xlsx", data_only=True)["Sheet"]["C2"].value == 7


def test_recalc_xlsx_on_a_missing_file_says_so(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="does not exist"):
        _tools(tmp_path)["recalc_xlsx"](path="missing.xlsx")  # type: ignore[operator]


def test_concurrent_writes_to_one_workbook_all_land(tmp_path: Path) -> None:
    """One model message's tool calls run at once; unserialized, two saves of the same file
    raced and one corrupted it or undid the other."""
    tools = _tools(tmp_path)
    names = [f"Sheet{i}" for i in range(8)]

    def write(name: str) -> object:
        return tools["write_xlsx_cells"](  # type: ignore[operator]
            path="concurrent.xlsx", blocks=[{"sheet": name, "rows": [[name, 1]]}]
        )

    with concurrent.futures.ThreadPoolExecutor(max_workers=8) as pool:
        list(pool.map(write, names))

    assert set(load_workbook(tmp_path / "concurrent.xlsx").sheetnames) >= set(names)


def test_a_read_during_a_write_never_sees_a_half_written_file(tmp_path: Path) -> None:
    tools = _tools(tmp_path)
    tools["write_xlsx_cells"](path="r.xlsx", blocks=[{"sheet": "S1", "rows": [[1, "=A1*2"]]}])  # type: ignore[operator]
    failures: list[str] = []
    done = threading.Event()

    def reader() -> None:
        while not done.is_set():
            try:
                tools["read_xlsx"](path="r.xlsx", sheet="S1")  # type: ignore[operator]
            except Exception as exc:  # noqa: BLE001 -- recording what a reader would see
                failures.append(f"{type(exc).__name__}: {exc}")

    thread = threading.Thread(target=reader)
    thread.start()
    for index in range(4):
        tools["write_xlsx_cells"](  # type: ignore[operator]
            path="r.xlsx", blocks=[{"sheet": f"T{index}", "rows": [[1, "=A1*2"]]}]
        )
    done.set()
    thread.join()

    assert failures == []


def test_an_extra_writable_folder_takes_a_workbook_and_a_readable_one_refuses_it(
    tmp_path: Path,
) -> None:
    shared, downloads = tmp_path / "shared", tmp_path / "downloads"
    shared.mkdir()
    downloads.mkdir()
    tools = _tools(tmp_path / "ws", extra_writable=[shared], extra_readable=[downloads])

    _scores(tools, str(shared / "report.xlsx"))

    assert "| Alice | 9 |" in tools["read_xlsx"](path=str(shared / "report.xlsx"))  # type: ignore[operator]
    with pytest.raises(PermissionError):
        _scores(tools, str(downloads / "report.xlsx"))


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("00123", "00123"),
        ("6222021234567890123", "6222021234567890123"),
        ("1e5", "1e5"),
        ("1_000", "1_000"),
        ("NaN", "NaN"),
        ("inf", "inf"),
        ("１２３", "１２３"),
        ("42", 42),
        ("  42  ", 42),
        ("-7", -7),
        ("3.10", 3.1),
        ("0.5", 0.5),
        ("123456789012345", 123456789012345),
    ],
)
def test_cells_only_become_numbers_when_they_are_plain_numbers(text: str, expected: object) -> None:
    """Real data loss: an ID like 00123 became 123 and a 19-digit card
    number a rounded float."""
    from coscribe.tools.office.xlsx.helpers import _coerce_cell

    assert _coerce_cell(text) == expected
    assert type(_coerce_cell(text)) is type(expected)
