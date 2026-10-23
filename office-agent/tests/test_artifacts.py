"""What the Artifacts page shows of a file: its kind, its time and a document's opening lines."""

import json
from pathlib import Path
from typing import Any

from docx import Document
from langchain_core.messages import AIMessage, BaseMessage, ToolMessage

from coscribe.conversation import describe_artifact
from coscribe.conversation.artifacts import artifact_paths
from coscribe.runtime import ToolMetadata
from coscribe.tools import WorkspaceScope

_CATALOG = {
    "read_pptx": ToolMetadata("READ"),
    "edit_pptx": ToolMetadata("WRITE_LOCAL"),
    "write_docx": ToolMetadata("WRITE_LOCAL"),
}


def _calls(*calls: tuple[str, dict[str, Any], Any]) -> list[BaseMessage]:
    messages: list[BaseMessage] = []
    for i, (name, args, result) in enumerate(calls):
        call = {"id": f"c{i}", "name": name, "args": args}
        messages += [
            AIMessage(content="", tool_calls=[call]),
            ToolMessage(
                content=result if isinstance(result, str) else json.dumps(result),
                tool_call_id=f"c{i}",
            ),
        ]
    return messages


def test_a_document_is_described_with_its_opening_lines(tmp_path: Path) -> None:
    document = Document()
    document.add_heading("Quarterly report", level=1)
    document.add_paragraph("")
    document.add_paragraph("Revenue grew in every region.")
    document.add_paragraph("Costs held steady.")
    document.save(str(tmp_path / "report.docx"))
    scope = WorkspaceScope(tmp_path)

    described = describe_artifact(scope, "report.docx")

    assert described is not None
    assert (described["name"], described["kind"], described["path"]) == (
        "report",
        "docs",
        "report.docx",
    )
    assert described["excerpt"] == {
        "heading": "Quarterly report",
        "text": "Revenue grew in every region. Costs held steady.",
    }


def test_what_is_not_an_artifact_or_not_there_is_not_described(tmp_path: Path) -> None:
    for name in ("notes.md", "~$deck.pptx", "broken.docx"):
        (tmp_path / name).write_text("x", encoding="utf-8")
    scope = WorkspaceScope(tmp_path)

    assert describe_artifact(scope, "notes.md") is None
    assert describe_artifact(scope, "~$deck.pptx") is None
    assert describe_artifact(scope, "gone.pptx") is None
    assert describe_artifact(scope, "../outside.docx") is None
    # A document that does not open is still listed, without an excerpt.
    broken = describe_artifact(scope, "broken.docx")
    assert broken is not None and broken["excerpt"] is None


def test_a_file_the_person_brought_is_an_artifact_once_a_script_or_write_tool_replaced_it(
    tmp_path: Path,
) -> None:
    for name in ("brought.pptx", "template.pptx", "edited.pptx", "report.docx", "made.pptx"):
        (tmp_path / name).write_bytes(b"x")
    scope = WorkspaceScope(tmp_path)
    script = {"script": "...", "description": "rewrite"}
    messages = _calls(
        ("read_pptx", {"path": "brought.pptx"}, "ok"),
        (
            "run_python_script",
            script,
            {"exit_code": 0, "files_written": ["brought.pptx", "made.pptx"]},
        ),
        ("read_pptx", {"path": "template.pptx"}, "ok"),
        ("read_pptx", {"path": "edited.pptx"}, "ok"),
        ("edit_pptx", {"path": "edited.pptx"}, {"lines_removed": 2}),
        ("read_pptx", {"path": "report.docx"}, "ok"),
        ("write_docx", {"path": "report.docx"}, {"lines_removed": 5}),
    )

    found = artifact_paths(messages, _CATALOG, scope)

    # A deck only read, or only changed in place by an edit tool, is not coscribe's.
    assert sorted(found) == ["brought.pptx", "made.pptx", "report.docx"]
