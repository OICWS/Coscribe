"""The deterministic parts of the Word evaluation (`evals/docx_checks.py`, the docx tasks and the
kind-aware summary, judge and comparison). Nothing here calls a model; the tests that render
through LibreOffice skip without it."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.shared import Inches

from evals import compare, docx_checks, judge, run, tasks

from .docx_reference import SOLUTIONS, redline

# --- reading a document -----------------------------------------------------------------


def _save(document: Any, tmp_path: Path, name: str = "doc.docx") -> Path:
    path = tmp_path / name
    document.save(path)
    return path


def _redlined(tmp_path: Path) -> Path:
    document = Document()
    document.add_paragraph("The fee is 2,400 GBP a month.")
    redline(document.paragraphs[0], "2,400", "2,650", 1)
    return _save(document, tmp_path)


def test_the_accepted_view_has_the_insertions_and_the_rejected_view_the_deletions(
    tmp_path: Path,
) -> None:
    read = docx_checks.read_document(_redlined(tmp_path))

    assert read.accepted_text() == "The fee is 2,650 GBP a month."
    assert read.rejected_text() == "The fee is 2,400 GBP a month."
    assert (read.inserted, read.deleted) == (1, 1)


def test_python_docx_alone_would_miss_the_inserted_text(tmp_path: Path) -> None:
    assert "2,650" not in Document(str(_redlined(tmp_path))).paragraphs[0].text


def test_a_paragraph_style_is_read_by_its_name(tmp_path: Path) -> None:
    document = Document()
    document.add_heading("Plan", level=1)
    document.add_heading("Report", level=0)
    document.add_paragraph("text")

    read = docx_checks.read_document(_save(document, tmp_path))

    assert [(p.style, p.accepted) for p in read.paragraphs] == [
        ("heading 1", "Plan"),
        ("title", "Report"),
        ("normal", "text"),
    ]


def test_a_file_that_is_not_a_document_does_not_open(tmp_path: Path) -> None:
    path = tmp_path / "broken.docx"
    path.write_bytes(b"not a zip")

    assert docx_checks.check_opens(path)["ok"] is False


# --- headings ---------------------------------------------------------------------------


def _headed(tmp_path: Path, titles: list[str], bold_only: tuple[str, ...] = ()) -> Any:
    document = Document()
    for title in titles:
        if title in bold_only:
            document.add_paragraph().add_run(title).bold = True
        else:
            document.add_heading(title, level=1)
    return docx_checks.read_document(_save(document, tmp_path))


def test_headings_are_found_in_order_with_other_headings_allowed_between(tmp_path: Path) -> None:
    read = _headed(tmp_path, ["Summary", "Background", "Costs"])

    assert docx_checks.check_headings(read, [{"text": "Summary"}, {"text": "Costs"}])["ok"]


def test_headings_in_the_wrong_order_are_reported(tmp_path: Path) -> None:
    read = _headed(tmp_path, ["Costs", "Summary"])

    report = docx_checks.check_headings(read, [{"text": "Summary"}, {"text": "Costs"}])

    assert report["ok"] is False and report["missing_or_out_of_order"] == ["Costs"]


def test_bold_text_is_not_a_heading(tmp_path: Path) -> None:
    read = _headed(tmp_path, ["Summary", "Costs"], bold_only=("Costs",))

    assert docx_checks.check_headings(read, [{"text": "Costs"}])["ok"] is False


def test_a_heading_level_is_checked_when_asked(tmp_path: Path) -> None:
    read = _headed(tmp_path, ["Summary"])

    assert docx_checks.check_headings(read, [{"text": "Summary", "level": 1}])["ok"]
    assert not docx_checks.check_headings(read, [{"text": "Summary", "level": 2}])["ok"]


# --- table of contents, tables, pictures, captions --------------------------------------


def _field(document: Any, instruction: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    paragraph = document.add_paragraph()
    for kind in ("begin", "instr", "end"):
        run = OxmlElement("w:r")
        node = OxmlElement("w:instrText" if kind == "instr" else "w:fldChar")
        if kind == "instr":
            node.text = instruction
        else:
            node.set(qn("w:fldCharType"), kind)
        run.append(node)
        paragraph._p.append(run)


def test_a_table_of_contents_is_a_field_not_the_words_contents(tmp_path: Path) -> None:
    typed = Document()
    typed.add_paragraph("Table of contents")
    fielded = Document()
    _field(fielded, ' TOC \\o "1-3" ')

    assert docx_checks.read_document(_save(typed, tmp_path, "a.docx")).has_toc is False
    assert docx_checks.read_document(_save(fielded, tmp_path, "b.docx")).has_toc is True
    assert docx_checks.check_toc(docx_checks.read_document(tmp_path / "b.docx"))["ok"]


def _tabled(tmp_path: Path, shapes: list[tuple[int, int]], header: str = "") -> Any:
    document = Document()
    for rows, cols in shapes:
        table = document.add_table(rows=rows, cols=cols)
        if header:
            for column, word in enumerate(header.split()):
                table.cell(0, column).text = word
    return docx_checks.read_document(_save(document, tmp_path))


def test_a_table_is_big_enough_by_rows_and_columns(tmp_path: Path) -> None:
    read = _tabled(tmp_path, [(4, 2)])

    assert docx_checks.check_tables(read, [{"rows_min": 4, "cols_min": 2}])["ok"]
    assert not docx_checks.check_tables(read, [{"rows_min": 5, "cols_min": 2}])["ok"]
    assert not docx_checks.check_tables(read, [{"rows_min": 4, "cols_min": 3}])["ok"]


def test_each_wanted_table_needs_a_table_of_its_own(tmp_path: Path) -> None:
    wanted = [{"rows_min": 3, "cols_min": 2}, {"rows_min": 3, "cols_min": 2}]

    assert not docx_checks.check_tables(_tabled(tmp_path, [(3, 2)]), wanted)["ok"]
    assert docx_checks.check_tables(_tabled(tmp_path, [(3, 2), (3, 2)]), wanted)["ok"]


def test_a_table_header_must_hold_the_asked_words(tmp_path: Path) -> None:
    read = _tabled(tmp_path, [(2, 3)], header="Region Q1 Q4")

    assert docx_checks.check_tables(read, [{"rows_min": 2, "cols_min": 3, "header": ["Q4"]}])["ok"]
    assert not docx_checks.check_tables(
        read, [{"rows_min": 2, "cols_min": 3, "header": ["Total"]}]
    )["ok"]


def _pictured(tmp_path: Path, width_inches: float, caption: str | None = None) -> Any:
    from PIL import Image

    Image.new("RGB", (400, 300), "#336699").save(tmp_path / "p.png")
    document = Document()
    document.add_paragraph("before")
    document.add_picture(str(tmp_path / "p.png"), width=Inches(width_inches))
    if caption:
        document.add_paragraph(caption, style="Caption")
    return docx_checks.read_document(_save(document, tmp_path))


def test_a_picture_wider_than_the_text_area_is_reported_with_its_width(tmp_path: Path) -> None:
    fitting = docx_checks.check_images(_pictured(tmp_path, 6.0), 1, True)
    wide = docx_checks.check_images(_pictured(tmp_path, 7.5), 1, True)

    assert fitting["ok"] is True
    assert wide["ok"] is False and wide["wider_than_text_cm"] == [19.1]


def test_fewer_pictures_than_asked_fail(tmp_path: Path) -> None:
    assert not docx_checks.check_images(_pictured(tmp_path, 3), 2, False)["ok"]


def test_a_caption_is_in_its_style_and_directly_below_the_picture(tmp_path: Path) -> None:
    read = _pictured(tmp_path, 3, "Figure 1: Intake")
    spec = {"style": "Caption", "starts_with": "Figure 1", "below_image": True}

    assert docx_checks.check_captions(read, [spec])["ok"]
    assert not docx_checks.check_captions(read, [{**spec, "starts_with": "Figure 2"}])["ok"]
    assert not docx_checks.check_captions(read, [{**spec, "style": "Heading 1"}])["ok"]


def test_a_caption_above_the_picture_is_not_below_it(tmp_path: Path) -> None:
    from PIL import Image

    Image.new("RGB", (40, 30)).save(tmp_path / "p.png")
    document = Document()
    document.add_paragraph("Figure 1: Intake", style="Caption")
    document.add_picture(str(tmp_path / "p.png"), width=Inches(2))
    read = docx_checks.read_document(_save(document, tmp_path))

    spec = {"style": "Caption", "starts_with": "Figure 1", "below_image": True}
    assert not docx_checks.check_captions(read, [spec])["ok"]


# --- tracked changes --------------------------------------------------------------------


def _original(tmp_path: Path) -> Any:
    document = Document()
    document.add_paragraph("The fee is 2,400 GBP a month.")
    return docx_checks.read_document(_save(document, tmp_path, "original.docx"))


def test_a_redline_that_rejects_back_to_the_input_passes(tmp_path: Path) -> None:
    read = docx_checks.read_document(_redlined(tmp_path))
    spec = {"min_ins": 1, "min_del": 1, "rejected_equals_original": True}

    assert docx_checks.check_tracked(read, spec, _original(tmp_path))["ok"]


def test_an_edit_made_without_tracking_has_no_marks_to_show(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("The fee is 2,650 GBP a month.")
    read = docx_checks.read_document(_save(document, tmp_path, "silent.docx"))
    spec = {"min_ins": 1, "min_del": 1, "rejected_equals_original": True}

    report = docx_checks.check_tracked(read, spec, _original(tmp_path))

    assert report["ok"] is False and report["rejected_equals_original"] is False


def test_a_tracked_edit_that_also_changed_other_text_silently_fails(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("The fee is 2,400 GBP a month, payable monthly.")
    redline(document.paragraphs[0], "2,400", "2,650", 1)
    read = docx_checks.read_document(_save(document, tmp_path, "sneaky.docx"))
    spec = {"min_ins": 1, "min_del": 1, "rejected_equals_original": True}

    assert not docx_checks.check_tracked(read, spec, _original(tmp_path))["ok"]


# --- the user's template ----------------------------------------------------------------


def _template_expect(**given: Any) -> tasks.DocxExpect:
    return tasks.DocxExpect(facts=[], original="letterhead.docx", **given)


def _letter(tmp_path: Path, keep_header: bool = True) -> tuple[Any, Any]:
    source = tasks.discover("04", kind="docx")[0].directory / "inputs" / "letterhead.docx"
    document = Document(str(source))
    document.add_paragraph("Dear Ms Okafor,", style="Letter Body")
    if not keep_header:
        document.sections[0].header.paragraphs[0].text = ""
    return (
        docx_checks.read_document(source),
        docx_checks.read_document(_save(document, tmp_path, "letter.docx")),
    )


def test_a_letter_in_the_letterheads_styles_header_and_footer_passes(tmp_path: Path) -> None:
    original, read = _letter(tmp_path)
    expect = _template_expect(
        styles=["Letter Body"],
        header_has=["ACME Corporation"],
        footer_has=["04417281"],
        style_uses={"Letter Body": 1},
    )

    assert docx_checks.check_template(read, original, expect)["ok"]


def test_a_missing_header_is_named(tmp_path: Path) -> None:
    original, read = _letter(tmp_path, keep_header=False)

    report = docx_checks.check_template(
        read, original, _template_expect(header_has=["ACME Corporation"])
    )

    assert report["ok"] is False and "header lacks 'ACME Corporation'" in report["problems"]


def test_text_set_in_the_default_style_is_not_in_the_letterheads_style(tmp_path: Path) -> None:
    original, _ = _letter(tmp_path)
    document = Document(
        str(tasks.discover("04", kind="docx")[0].directory / "inputs" / "letterhead.docx")
    )
    document.add_paragraph("Dear Ms Okafor,")
    read = docx_checks.read_document(_save(document, tmp_path, "plain.docx"))

    report = docx_checks.check_template(
        read, original, _template_expect(style_uses={"Letter Body": 1})
    )

    assert (
        report["ok"] is False
        and "'Letter Body' used by 0 paragraphs, expected 1" in report["problems"]
    )


def test_a_style_the_original_used_that_no_paragraph_uses_any_more_is_named(tmp_path: Path) -> None:
    document = Document()
    document.add_heading("Title", level=1)
    document.add_paragraph("a", style="List Bullet")
    original = docx_checks.read_document(_save(document, tmp_path, "o.docx"))
    flat = Document()
    flat.add_heading("Title", level=1)
    flat.add_paragraph("a")
    read = docx_checks.read_document(_save(flat, tmp_path, "f.docx"))

    report = docx_checks.check_template(read, original, _template_expect(keep_styles=True))

    assert report["ok"] is False and "list bullet" in report["problems"][0]


# --- markup typed into the text ---------------------------------------------------------


@pytest.mark.parametrize(
    "text",
    ["# Heading", "- a bullet", "**bold** words", "| a | b |", "|---|---|", "use `code` here"],
)
def test_markdown_left_in_the_text_is_found(text: str, tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph(text)

    report = docx_checks.find_markdown(docx_checks.read_document(_save(document, tmp_path)))

    assert report["ok"] is False and report["count"] == 1


def test_plain_sentences_with_dashes_and_stars_are_not_markdown(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("A 3-4 week move - with a * footnote, costs 5 * 2 GBP.")

    assert docx_checks.find_markdown(docx_checks.read_document(_save(document, tmp_path)))["ok"]


# --- the tasks --------------------------------------------------------------------------


def _task(name: str) -> tasks.Task:
    return tasks.discover(name, kind="docx")[0]


def test_there_are_eight_word_tasks_each_with_its_inputs() -> None:
    found = tasks.discover("all", kind="docx")

    assert [t.name[:2] for t in found] == ["01", "02", "03", "04", "05", "06", "07", "08"]
    assert all(t.inputs and isinstance(t.expect, tasks.DocxExpect) for t in found)


def test_every_input_a_brief_needs_is_named_in_it() -> None:
    for task in tasks.discover("all", kind="docx"):
        for path in task.inputs:
            assert path.name in task.brief, f"{task.name}: the brief never mentions {path.name}"


def _write_task(directory: Path, expect: dict[str, Any], inputs: tuple[str, ...] = ()) -> None:
    directory.mkdir(parents=True, exist_ok=True)
    (directory / "brief.md").write_text("do it", encoding="utf-8")
    (directory / "expect.json").write_text(json.dumps(expect), encoding="utf-8")
    for name in inputs:
        (directory / "inputs").mkdir(exist_ok=True)
        (directory / "inputs" / name).write_text("x", encoding="utf-8")


@pytest.mark.parametrize(
    ("expect", "message"),
    [
        ({}, "expects nothing"),
        ({"headings": [{"level": 1}]}, 'needs "text"'),
        ({"tables": [{"rows_min": 2}]}, "rows_min"),
        ({"captions": [{"style": "Caption"}]}, "starts_with"),
        ({"headings": [{"text": "A"}], "pages": [3, 1]}, "pages"),
        ({"headings": [{"text": "A"}], "tracked": 3}, "tracked"),
        ({"headings": [{"text": "A"}], "style_uses": {"Body": "two"}}, "style_uses"),
        ({"headings": [{"text": "A"}], "header_has": ["x"]}, 'need "original"'),
        ({"facts": "text"}, "facts"),
    ],
)
def test_a_word_task_with_bad_expectations_is_refused(
    tmp_path: Path, expect: dict[str, Any], message: str
) -> None:
    _write_task(tmp_path, expect)

    with pytest.raises(ValueError, match=message):
        tasks.load_task(tmp_path, "docx")


def test_the_original_must_be_one_of_the_inputs(tmp_path: Path) -> None:
    expect = {"headings": [{"text": "A"}], "original": "a.docx", "keep_styles": True}
    _write_task(tmp_path, expect, inputs=("b.docx",))

    with pytest.raises(ValueError, match="not among the inputs"):
        tasks.load_task(tmp_path, "docx")


def test_the_inputs_are_what_the_generator_makes(tmp_path: Path, monkeypatch: Any) -> None:
    from evals.docx import make_inputs

    monkeypatch.setattr(make_inputs, "TASKS", tmp_path)
    make_inputs.main()

    for task in tasks.discover("all", kind="docx"):
        for path in task.inputs:
            made = tmp_path / task.name / "inputs" / path.name
            if path.suffix == ".docx":
                assert docx_checks.read_document(made).accepted_text() == (
                    docx_checks.read_document(path).accepted_text()
                )
            elif path.suffix in {".md", ".csv"}:
                assert made.read_text(encoding="utf-8") == path.read_text(encoding="utf-8")
            else:
                assert made.exists()


# --- the reference solutions ------------------------------------------------------------


def _check(task: str, tmp_path: Path, fake_pages: bool) -> dict[str, Any]:
    workspace = tmp_path / "ws"
    workspace.mkdir()
    chosen = _task(task)
    output = SOLUTIONS[task](chosen.directory / "inputs", workspace)
    return docx_checks.check_document(chosen, output, workspace, tmp_path / "render")


def _failed(report: dict[str, Any]) -> dict[str, Any]:
    return {k: v for k, v in report.items() if isinstance(v, dict) and v.get("ok") is False}


@pytest.mark.parametrize("task", [t.name for t in tasks.discover("all", kind="docx")])
def test_the_reference_solution_of_every_task_passes_every_check_but_the_page_count(
    task: str, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(
        docx_checks, "render_pages", lambda path, out: {"ok": True, "pages": ["page-01.png"]}
    )

    assert _failed(_check(task, tmp_path, True)) == {}


def _libreoffice_renders(tmp_path: Path) -> bool:
    from evals.render import render_pages

    document = Document()
    document.add_paragraph("probe")
    return bool(render_pages(_save(document, tmp_path, "probe.docx"), tmp_path / "probe")["ok"])


@pytest.mark.real_libreoffice
@pytest.mark.parametrize("task", [t.name for t in tasks.discover("all", kind="docx")])
def test_the_reference_solution_of_every_task_passes_with_the_pages_really_counted(
    task: str, tmp_path: Path
) -> None:
    if not _libreoffice_renders(tmp_path):
        pytest.skip("LibreOffice is not usable here")

    report = _check(task, tmp_path, False)

    assert _failed(report) == {}
    assert report["render"]["ok"] is True


@pytest.mark.real_libreoffice
def test_a_redline_opens_in_libreoffice(tmp_path: Path) -> None:
    if not _libreoffice_renders(tmp_path):
        pytest.skip("LibreOffice is not usable here")
    from evals.render import render_pages

    assert render_pages(_redlined(tmp_path), tmp_path / "out")["ok"] is True


# --- each check fails on a wrong answer -------------------------------------------------


def _broken(task: str, tmp_path: Path, monkeypatch: pytest.MonkeyPatch, change: Any) -> dict:
    monkeypatch.setattr(
        docx_checks, "render_pages", lambda path, out: {"ok": True, "pages": ["page-01.png"]}
    )
    workspace = tmp_path / "ws"
    workspace.mkdir()
    chosen = _task(task)
    output = SOLUTIONS[task](chosen.directory / "inputs", workspace)
    document = Document(str(output))
    change(document)
    document.save(output)
    return docx_checks.check_document(chosen, output, workspace, tmp_path / "render")


def test_a_table_flattened_into_text_fails_the_table_check(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def flatten(document: Any) -> None:
        document.tables[0]._tbl.getparent().remove(document.tables[0]._tbl)

    report = _broken("05-table-heavy", tmp_path, monkeypatch, flatten)

    assert report["tables"]["ok"] is False


def test_a_changed_figure_fails_the_facts(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    def change(document: Any) -> None:
        document.tables[0].cell(5, 5).text = "247,000"

    report = _broken("05-table-heavy", tmp_path, monkeypatch, change)

    assert report["facts"]["missing"] == ["247,500"]


def test_a_heading_turned_into_plain_text_fails_the_heading_check(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def demote(document: Any) -> None:
        for paragraph in document.paragraphs:
            if paragraph.text == "Risks":
                paragraph.style = document.styles["Normal"]

    report = _broken("01-report-from-notes", tmp_path, monkeypatch, demote)

    assert report["headings"]["missing_or_out_of_order"] == ["Risks"]


def test_a_letter_with_the_letterheads_header_removed_fails_the_template(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def strip(document: Any) -> None:
        document.sections[0].header.paragraphs[0].text = ""

    report = _broken("04-letter-on-template", tmp_path, monkeypatch, strip)

    assert report["template"]["ok"] is False


def test_a_picture_too_wide_for_the_page_fails_the_fit(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def widen(document: Any) -> None:
        document.inline_shapes[0].width = Inches(9)

    report = _broken("08-image-caption", tmp_path, monkeypatch, widen)

    assert report["images"]["ok"] is False


def test_a_page_count_outside_the_range_fails(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(
        docx_checks, "render_pages", lambda path, out: {"ok": True, "pages": ["a", "b"]}
    )
    workspace = tmp_path / "ws"
    workspace.mkdir()
    chosen = _task("02-memo-from-pdf")
    output = SOLUTIONS["02-memo-from-pdf"](chosen.directory / "inputs", workspace)

    report = docx_checks.check_document(chosen, output, workspace, tmp_path / "render")

    assert report["pages"] == {"ok": False, "count": 2, "expected": [1, 1]}


def test_a_check_that_cannot_run_says_why_and_the_others_still_run(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def boom(*args: Any) -> None:
        raise RuntimeError("no table")

    monkeypatch.setattr(docx_checks, "check_tables", boom)

    report = _broken("05-table-heavy", tmp_path, monkeypatch, lambda document: None)

    assert report["tables"] == {"ok": False, "error": "RuntimeError: no table"}
    assert report["facts"]["ok"] is True


# --- the run's records, summary, judge, comparison --------------------------------------


def test_the_output_is_a_docx_that_changed(tmp_path: Path) -> None:
    (tmp_path / "input.docx").write_bytes(b"given")
    (tmp_path / "notes.xlsx").write_bytes(b"x")
    originals = {"input.docx": b"given"}
    assert run._find_output(tmp_path, originals, [], ".docx") is None

    (tmp_path / "new.docx").write_bytes(b"made")
    assert run._find_output(tmp_path, originals, [], ".docx") == tmp_path / "new.docx"


DOCX_REPORT = {
    "opens": {"ok": True},
    "facts": {"found": 8, "total": 9},
    "inventions": {"count": 2},
    "pages": {"ok": False, "count": 4, "expected": [1, 2]},
    "headings": {"ok": True},
    "tracked": {"ok": False},
    "markdown": {"ok": True},
}


def _recorded(tmp_path: Path, report: dict[str, Any]) -> Path:
    folder = tmp_path / "run" / "03-tracked-contract"
    folder.mkdir(parents=True)
    result = {
        "task": "03-tracked-contract",
        "kind": "docx",
        "status": "ok",
        "wall_seconds": 40.0,
        "tool_call_count": 9,
        "input_tokens": 100,
        "cached_tokens": 80,
        "output_tokens": 10,
        "cost": 0.0111,
    }
    (folder / "result.json").write_text(json.dumps(result), encoding="utf-8")
    (folder / "checks.json").write_text(json.dumps(report), encoding="utf-8")
    return tmp_path / "run"


def test_the_summary_of_a_word_run_shows_what_failed(tmp_path: Path) -> None:
    out = _recorded(tmp_path, DOCX_REPORT)

    run.write_summary(out, tasks.discover("03", kind="docx"), "docx")

    lines = (out / "summary.md").read_text(encoding="utf-8").splitlines()
    head = [c.strip() for c in lines[0].strip("|").split("|")]
    row = dict(zip(head, [c.strip() for c in lines[2].strip("|").split("|")], strict=True))
    assert row["facts"] == "8/9" and row["pages"] == "4" and row["tracked"] == "False"
    assert row["headings"] == "True" and row["tables"] == "-" and row["cost"] == "0.0111"
    assert "overflow" not in head and "recalc" not in head


def test_a_word_file_is_judged_with_the_word_rubric(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    folder = tmp_path / "a" / "03-tracked-contract"
    (folder / "render").mkdir(parents=True)
    (folder / "render" / "page-01.png").write_bytes(b"png")
    seen: list[str] = []

    def fake(
        judge_name: str, brief: str, decks: dict, labels: list[str], rubric: str = ""
    ) -> tuple[dict, dict]:
        seen.append(rubric)
        scores = {"design": 4, "content": 4, "coherence": 4, "defects": []}
        return {"decks": {label: scores for label in labels}, "ranking": labels}, {}

    monkeypatch.setattr(judge, "judge_round", fake)

    judge.judge_runs([tmp_path / "a"], ["03-tracked-contract"], judge.DEFAULT_JUDGE, 1, 7, "docx")

    assert seen == [judge.RUBRIC_DOCX] and "Word documents" in judge.RUBRIC_DOCX


def test_the_comparison_of_two_word_runs_shows_the_document_checks(tmp_path: Path) -> None:
    a = _recorded(tmp_path / "x", DOCX_REPORT)
    b = _recorded(tmp_path / "y", {**DOCX_REPORT, "tracked": {"ok": True}})

    table = compare.compare(a, b, kind="docx")

    assert "| yes | yes | - | - | - | - | - | - | - | - | NO | yes |" in table
    assert "recalculates" not in table and "overflow" not in table


def test_a_title_is_level_zero(tmp_path: Path) -> None:
    document = Document()
    document.add_heading("Report", level=0)
    read = docx_checks.read_document(_save(document, tmp_path))

    assert docx_checks.check_headings(read, [{"text": "Report", "level": 0}])["ok"]
    assert not docx_checks.check_headings(read, [{"text": "Report", "level": 1}])["ok"]


def test_a_redline_with_fewer_marks_than_asked_fails(tmp_path: Path) -> None:
    read = docx_checks.read_document(_redlined(tmp_path))

    assert not docx_checks.check_tracked(read, {"min_ins": 2}, None)["ok"]
    assert not docx_checks.check_tracked(read, {"min_del": 2}, None)["ok"]


def test_the_marks_of_an_inserted_paragraph_count_once_not_twice(tmp_path: Path) -> None:
    chosen = _task("03-tracked-contract")
    output = SOLUTIONS[chosen.name](chosen.directory / "inputs", tmp_path)

    read = docx_checks.read_document(output)

    assert (read.inserted, read.deleted) == (4, 2)


def test_a_style_the_letterhead_defined_but_the_letter_lost_is_named(tmp_path: Path) -> None:
    original, read = _letter(tmp_path)

    report = docx_checks.check_template(
        read, original, _template_expect(styles=["Letter Body", "No Such Style"])
    )

    assert report["problems"] == ["styles no longer defined: ['No Such Style']"]


# --- what a first review found ------------------------------------------------------------


def _without(path: Path, member: str) -> Path:
    import zipfile

    stripped = path.with_name("stripped.docx")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(stripped, "w") as target:
        for item in source.infolist():
            if item.filename != member:
                target.writestr(item, source.read(item.filename))
    return stripped


def test_a_document_without_a_styles_part_still_opens_with_every_paragraph_normal(
    tmp_path: Path,
) -> None:
    document = Document()
    document.add_heading("Plan", level=1)
    path = _without(_save(document, tmp_path), "word/styles.xml")

    read = docx_checks.read_document(path)

    assert [(p.style, p.accepted) for p in read.paragraphs] == [("normal", "Plan")]


def test_a_table_inside_a_table_is_not_a_second_table(tmp_path: Path) -> None:
    document = Document()
    outer = document.add_table(rows=2, cols=2)
    outer.cell(0, 0).add_table(rows=3, cols=3)
    read = docx_checks.read_document(_save(document, tmp_path))
    wanted = [{"rows_min": 2, "cols_min": 2}, {"rows_min": 2, "cols_min": 2}]

    assert len(read.tables) == 1
    assert not docx_checks.check_tables(read, wanted)["ok"]


def test_the_fallback_copy_of_a_text_box_is_not_read_twice(tmp_path: Path) -> None:
    from lxml import etree

    document = Document()
    paragraph = document.add_paragraph()
    xml = (
        "<mc:AlternateContent "
        'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        "<mc:Choice><w:p><w:r><w:t>boxed</w:t></w:r></w:p></mc:Choice>"
        "<mc:Fallback><w:p><w:r><w:t>boxed</w:t></w:r></w:p></mc:Fallback></mc:AlternateContent>"
    )
    paragraph._p.append(etree.fromstring(xml))

    read = docx_checks.read_document(_save(document, tmp_path))

    assert [p.accepted for p in read.paragraphs].count("boxed") == 1


def test_a_tab_stop_definition_is_not_a_space_in_the_text(tmp_path: Path) -> None:
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.shared import Cm

    document = Document()
    paragraph = document.add_paragraph("Q3")
    paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(5), WD_TAB_ALIGNMENT.RIGHT)
    paragraph.add_run("\t12,710")

    read = docx_checks.read_document(_save(document, tmp_path))

    assert read.paragraphs[0].accepted == "Q3 12,710"


def test_a_title_does_not_stand_in_for_a_wanted_section(tmp_path: Path) -> None:
    document = Document()
    document.add_heading("Summary of the warehouse move", level=0)
    document.add_heading("Costs", level=1)
    read = docx_checks.read_document(_save(document, tmp_path))

    report = docx_checks.check_headings(read, [{"text": "Summary"}, {"text": "Costs"}])

    assert report["missing_or_out_of_order"] == ["Summary"]


def test_a_redline_check_without_the_original_to_compare_with_fails(tmp_path: Path) -> None:
    read = docx_checks.read_document(_redlined(tmp_path))

    report = docx_checks.check_tracked(read, {"rejected_equals_original": True}, None)

    assert report["ok"] is False


def test_a_task_that_asks_for_the_rejected_view_must_give_the_original(tmp_path: Path) -> None:
    _write_task(tmp_path, {"tracked": {"rejected_equals_original": True}})

    with pytest.raises(ValueError, match="needs .original"):
        tasks.load_task(tmp_path, "docx")


@pytest.mark.parametrize(
    "expect",
    [
        {"captions": [{"style": "Caption", "starts_with": "Figure 1"}]},
        {"toc": True},
        {"pages": [1, 2]},
        {"tracked": {"min_ins": 1}},
        {"original": "a.docx", "keep_styles": True},
    ],
)
def test_a_task_whose_only_expectation_is_any_check_the_runner_has_is_accepted(
    tmp_path: Path, expect: dict[str, Any]
) -> None:
    _write_task(tmp_path, expect, inputs=("a.docx",))

    assert tasks.load_task(tmp_path, "docx").expect.facts == []  # type: ignore[union-attr]


def test_an_original_that_cannot_be_read_is_reported_and_the_other_checks_go_on(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(
        docx_checks, "render_pages", lambda path, out: {"ok": True, "pages": ["page-01.png"]}
    )
    chosen = _task("03-tracked-contract")
    workspace = tmp_path / "ws"
    workspace.mkdir()
    output = SOLUTIONS[chosen.name](chosen.directory / "inputs", workspace)
    real = docx_checks.read_document

    def read(path: Path) -> Any:
        if path.name == "service_agreement.docx":
            raise ValueError("damaged")
        return real(path)

    monkeypatch.setattr(docx_checks, "read_document", read)

    report = docx_checks.check_document(chosen, output, workspace, tmp_path / "render")

    assert report["original"] == {"ok": False, "error": "ValueError: damaged"}
    assert report["tracked"]["ok"] is False and report["template"]["ok"] is False
    assert report["facts"]["ok"] is True


def test_the_comparison_of_word_runs_shows_the_contents_pictures_and_captions(
    tmp_path: Path,
) -> None:
    report = {**DOCX_REPORT, "toc": {"ok": True}, "images": {"ok": False}, "captions": {"ok": True}}
    a = _recorded(tmp_path / "x", report)
    b = _recorded(tmp_path / "y", report)

    table = compare.compare(a, b, kind="docx")

    assert "| toc " in table and "| images " in table and "| captions " in table
