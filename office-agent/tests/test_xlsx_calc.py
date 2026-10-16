"""Formula values computed in-process (decision 0012), written into the file's cells."""

from __future__ import annotations

import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook

from coscribe.tools.office.xlsx.calc import recalculate


def test_values_land_on_the_right_sheet_whatever_its_name_or_order(tmp_path: Path) -> None:
    path = tmp_path / "b.xlsx"
    book = Workbook()
    book.active.title = "汇总"
    book["汇总"]["A1"] = "='数据 1'!A1*2"
    book.create_sheet("数据 1")["A1"] = 21
    book.move_sheet("数据 1", offset=-1)
    book.save(path)

    report = recalculate(path)

    assert report["status"] == "success" and report["total_formulas"] == 1
    assert load_workbook(path, data_only=True)["汇总"]["A1"].value == 42


def test_errors_are_counted_by_code_and_stored_as_errors(tmp_path: Path) -> None:
    path = tmp_path / "b.xlsx"
    book = Workbook()
    book.active.append(["=1/0", "=1/0", "=NA()", "=1+1"])
    book.save(path)

    report = recalculate(path)

    assert report["status"] == "errors_found" and report["total_errors"] == 3
    assert report["error_locations"] == {"#DIV/0!": ["Sheet!A1", "Sheet!B1"], "#N/A": ["Sheet!C1"]}
    with zipfile.ZipFile(path) as archive:
        assert b'<c r="A1" t="e"><f>1/0</f><v>#DIV/0!</v>' in archive.read(
            "xl/worksheets/sheet1.xml"
        )


def test_a_recalculated_value_replaces_a_stale_one(tmp_path: Path) -> None:
    path = tmp_path / "b.xlsx"
    book = Workbook()
    book.active.append([5, "=A1>3"])
    book.save(path)
    recalculate(path)
    book = load_workbook(path)
    book.active["A1"] = 1
    book.save(path)

    recalculate(path)

    assert load_workbook(path, data_only=True).active["B1"].value is False


def test_a_file_the_engine_cannot_read_is_skipped_untouched(tmp_path: Path) -> None:
    path = tmp_path / "b.xlsx"
    path.write_bytes(b"not a workbook")

    report = recalculate(path)

    assert report["status"] == "skipped" and report["skipped_reason"]
    assert path.read_bytes() == b"not a workbook"


def test_a_function_the_engine_lacks_is_not_checked_rather_than_an_error(tmp_path: Path) -> None:
    path = tmp_path / "b.xlsx"
    book = Workbook()
    book.active.append(['=CUBEVALUE("x")', "=A1+1", "=NOSUCHFN(1)", "=Totl*2", "=1+1"])
    book.save(path)

    report = recalculate(path)

    assert report["unchecked"]["cells"] == ["Sheet!A1", "Sheet!B1"]
    assert report["error_locations"] == {"#NAME?": ["Sheet!C1", "Sheet!D1"]}
    assert load_workbook(path, data_only=True).active["E1"].value == 2


def test_a_defined_name_is_a_name_excel_knows(tmp_path: Path) -> None:
    from openpyxl.workbook.defined_name import DefinedName

    path = tmp_path / "b.xlsx"
    book = Workbook()
    book.active.append([3, '=Rate*CUBEVALUE("x")'])
    book.defined_names["Rate"] = DefinedName("Rate", attr_text="Sheet!$A$1")
    book.save(path)

    assert recalculate(path)["unchecked"]["cells"] == ["Sheet!B1"]
