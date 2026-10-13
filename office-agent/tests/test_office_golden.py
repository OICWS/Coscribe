"""Golden (characterization) tests of the XML the pptx tools write.

`tools/presentations.py` is to be split into a package by pure moves (the scale assessment,
step 11). A pure move must leave every file the tools write unchanged, and these tests are
how that is shown: each deck is built through the public tools, reduced to canonical text
(`tests/pptx_canonical.py`), and compared with a checked-in golden under
`tests/golden/pptx/`. A change that alters what a deck contains edits its golden in the same
pull request, so the reviewer sees the XML that changed.
`UPDATE_GOLDENS=1 pytest tests/test_office_golden.py` rewrites them; read the diff before
committing it.
"""

from __future__ import annotations

import contextlib
import difflib
import os
import re
import shutil
from collections.abc import Callable
from pathlib import Path

import pytest

from coscribe.tools.presentations import build_presentation_tools

from .pptx_canonical import canonical_parts
from .pptx_fixtures import write_smartart_fixture

GOLDEN_DIR = Path(__file__).parent / "golden" / "pptx"


def _tools(root: Path) -> dict[str, Callable[..., object]]:
    return {tool.__name__: tool for tool in build_presentation_tools(root)}  # type: ignore[attr-defined]


def _bullets(tools: dict[str, Callable[..., object]]) -> None:
    tools["write_pptx"](
        path="deck.pptx",
        content="# Quarterly review\n- Revenue up 12%\n- **Two** new markets\n- *Hiring* on track",
    )


def _table(tools: dict[str, Callable[..., object]]) -> None:
    tools["write_pptx"](
        path="deck.pptx",
        content="# Results\n| Region | Q1 | Q2 |\n| --- | --- | --- |\n| North | 10 | 12 |\n"
        "| South | 8 | 9 |",
    )


def _two_column(tools: dict[str, Callable[..., object]]) -> None:
    tools["write_pptx"](
        path="deck.pptx",
        content="layout: two-column\n# Before vs after\n- Manual tracking\n- No alerts\n>>>\n"
        "- Live dashboards\n- Proactive alerts",
    )


def _transition_and_animation(tools: dict[str, Callable[..., object]]) -> None:
    tools["write_pptx"](path="deck.pptx", content="# Intro\n- one\n- two\n\n---\n# Next\n- three")
    tools["set_pptx_transition"](path="deck.pptx", slide=1, transition="fade", duration=0.8)
    tools["add_pptx_animation"](path="deck.pptx", slide=2, shape_index=0, animation="fade")


def _theme(tools: dict[str, Callable[..., object]]) -> None:
    tools["write_pptx"](
        path="deck.pptx", content="# Dark deck\n- body text", theme="bg=111827,text=F8FAFC"
    )
    tools["edit_pptx_theme_colors"](path="deck.pptx", colors="accent1=0F6B5C,accent2=C2410C")


# Committed files, not written by Pillow at test time: a PNG's bytes depend on the encoder and
# the zlib it links, which can differ between machines and versions.
ASSETS = GOLDEN_DIR / "assets"


def _png(name: str) -> None:
    shutil.copyfile(ASSETS / name, name)


def _wav(name: str) -> None:
    shutil.copyfile(ASSETS / "narration.wav", name)


def _image(tools: dict[str, Callable[..., object]]) -> None:
    tools["write_pptx"](path="deck.pptx", content="# Our logo\n- placed top right")
    _png("logo.png")
    tools["add_pptx_image"](
        path="deck.pptx", slide=1, image_path="logo.png", left=8.0, top=0.5, width=2.0
    )


def _chart(chart_type: str, data: str) -> Callable[[dict[str, Callable[..., object]]], None]:
    def build(tools: dict[str, Callable[..., object]]) -> None:
        tools["write_pptx"](path="deck.pptx", content=f"# A {chart_type} chart\n- see below")
        tools["add_pptx_chart"](
            path="deck.pptx", slide=1, chart_type=chart_type, data=data, title="Results"
        )

    return build


_ONE_SERIES = "| Quarter | Revenue |\n| --- | --- |\n| Q1 | 100 |\n| Q2 | 120 |\n| Q3 | 90 |"
_THREE_SERIES = (
    "| Quarter | North | South | West |\n| --- | --- | --- | --- |\n"
    "| Q1 | 10 | 8 | 5 |\n| Q2 | 12 | 9 | 7 |"
)
_SHARES = "| Region | Share |\n| --- | --- |\n| North | 60 |\n| South | 40 |"


def _icon_list(tools: dict[str, Callable[..., object]]) -> None:
    tools["write_pptx"](
        path="deck.pptx",
        content="layout: icon-list FF0000\n# Feature highlights\n- [*] Real-time tracking\n"
        "- Mobile app control\n- [!] Alerts",
    )


def _stat_callout(tools: dict[str, Callable[..., object]]) -> None:
    tools["write_pptx"](
        path="deck.pptx",
        content="layout: stat-callout 028090\n# Q4 highlights\n| Stat | Label |\n| --- | --- |\n"
        "| 3x | Revenue growth |\n| 12 | New markets |\n| 98% | Retention |",
    )


def _svg(tools: dict[str, Callable[..., object]]) -> None:
    svg = (
        '<svg width="960" height="540" viewBox="0 0 960 540">'
        '<rect x="0" y="0" width="960" height="540" fill="#1E2761"/>'
        '<circle cx="120" cy="270" r="40" fill="FF0000"/>'
        '<rect x="300" y="380" width="200" height="60" rx="12" fill="#F59E0B"/>'
        '<path d="M 600 400 L 700 300 L 800 400 Z" fill="#10B981"/>'
        '<text x="220" y="260" font-size="36" fill="#FFFFFF">Real-time tracking</text>'
        "</svg>"
    )
    tools["write_pptx"](path="deck.pptx", content=f"layout: svg\n{svg}\n")


def _icon(tools: dict[str, Callable[..., object]]) -> None:
    tools["write_pptx"](path="deck.pptx", content="# With an icon\n- a star, top right")
    tools["add_pptx_icon"](
        path="deck.pptx", slide=1, icon_name="star", left_in=8.0, top_in=0.5, color="F59E0B"
    )


def _shapes(tools: dict[str, Callable[..., object]]) -> None:
    tools["write_pptx"](path="deck.pptx", content="# Process\n- three steps")
    for index, (shape_type, fill) in enumerate(
        (("RECTANGLE", "1D4ED8"), ("ROUNDED_RECTANGLE", "059669"), ("OVAL", "DC2626"))
    ):
        tools["add_pptx_shape"](
            path="deck.pptx",
            slide=1,
            shape_type=shape_type,
            left_in=1.0 + 3.0 * index,
            top_in=4.0,
            width_in=2.5,
            height_in=1.2,
            text=f"Step {index + 1}",
            fill_color=fill,
            line_color="111827",
        )


def _picture_index(tools: dict[str, Callable[..., object]]) -> int:
    shapes = tools["list_pptx_shapes"](path="deck.pptx", slide=1)["shapes"]  # type: ignore[index]
    return next(shape["index"] for shape in shapes if shape["is_picture"])


def _image_cropped_and_replaced(tools: dict[str, Callable[..., object]]) -> None:
    _image(tools)
    index = _picture_index(tools)
    tools["crop_pptx_image"](
        path="deck.pptx", slide=1, shape_index=index, crop_left=0.1, crop_right=0.2
    )
    _png("logo2.png")
    tools["replace_pptx_image"](
        path="deck.pptx", slide=1, shape_index=index, image_path="logo2.png"
    )


def _image_effects(tools: dict[str, Callable[..., object]]) -> None:
    _image(tools)
    index = _picture_index(tools)
    for effect in ("soft_edge", "shadow", "glow"):
        tools["add_pptx_shape_effect"](path="deck.pptx", slide=1, shape_index=index, effect=effect)


def _background_image(tools: dict[str, Callable[..., object]]) -> None:
    tools["write_pptx"](path="deck.pptx", content="# Over a photo\n- with a scrim")
    _png("photo.png")
    tools["set_pptx_background_image"](path="deck.pptx", slide=1, image_path="photo.png")
    tools["add_pptx_scrim"](path="deck.pptx", slide=1, opacity=0.4, color="000000")


def _template_tools(root: Path) -> dict[str, Callable[..., object]]:
    """Tools with a custom template distilled from a fixed reference deck.

    The reference is built here from python-pptx's stock template (title, two content slides,
    closing), not read from a bundled template, so editing a shipped template does not fail
    these goldens."""
    from pptx import Presentation

    prs = Presentation()
    for layout_index in (0, 1, 1, 0):
        slide = prs.slides.add_slide(prs.slide_layouts[layout_index])
        slide.shapes.title.text = "Reference"
    prs.save(str(root / "reference.pptx"))
    tools = {
        tool.__name__: tool  # type: ignore[attr-defined]
        for tool in build_presentation_tools(root, custom_templates_dir=root / "_templates")
    }
    tools["extract_pptx_template"](
        source_path="reference.pptx", template_id="fixed", name="Fixed", description="a fixture"
    )
    return tools


def _template_fill(content: str) -> Callable[[dict[str, Callable[..., object]]], None]:
    def build(_: dict[str, Callable[..., object]]) -> None:
        tools = _template_tools(Path.cwd())
        tools["fill_pptx_template"](path="deck.pptx", template_id="fixed", content=content)

    return build


# Four template slides (title, two content, closing) against three chunks: a content slide is
# dropped; against six: the repeatable content slide is duplicated.
_FILL_SHRUNK = "# Annual review\nA year of growth\n---\n# Highlights\n- Revenue up\n---\n# Thanks"
_FILL_GROWN = (
    "# Annual review\nA year of growth\n---\n# One\n- a\n---\n# Two\n- b\n---\n# Three\n- c"
    "\n---\n# Four\n- d\n---\n# Thanks"
)
_FILL_FIXED_ONLY = "# Annual review\nA year of growth\n---\n# Thanks"


def _write_with_template_path(tools: dict[str, Callable[..., object]]) -> None:
    tools["write_pptx"](
        path="template.pptx",
        content="# Old\n- old bullet",
        theme="bg=111827,text=F8FAFC",
    )
    tools["write_pptx"](
        path="deck.pptx", content="# New title\n- new bullet", template_path="template.pptx"
    )


def _smartart_text_edit(tools: dict[str, Callable[..., object]]) -> None:
    write_smartart_fixture(Path("deck.pptx"))
    tools["edit_pptx_xml"](
        path="deck.pptx",
        slide=1,
        shape_index=0,
        xpath=".//dgm:t[.//a:t='Plan']//a:t",
        new_text="PLANNED START",
    )


def _smartart_shape_attributes(tools: dict[str, Callable[..., object]]) -> None:
    # The same escape hatch on the graphicFrame itself (part="shape"), schema-validated.
    write_smartart_fixture(Path("deck.pptx"))
    tools["edit_pptx_xml"](
        path="deck.pptx",
        slide=1,
        shape_index=0,
        xpath="//p:cNvPr",
        set_attributes="name=Roadmap\ndescr=Three phases",
        part="shape",
    )


_CJK_CONTENT = "# 季度回顾\n- 收入增长 **12%**\n- 新增两个市场\n- 招聘进展顺利"


_CJK_FILL = "# 年度回顾\n增长的一年\n---\n# 亮点\n- 收入增长\n---\n# 谢谢"


def _cjk_bullets(tools: dict[str, Callable[..., object]]) -> None:
    tools["write_pptx"](path="deck.pptx", content=_CJK_CONTENT)


def _cjk_table_and_chart(tools: dict[str, Callable[..., object]]) -> None:
    tools["write_pptx"](
        path="deck.pptx",
        content="# 区域业绩\n| 区域 | 第一季度 | 第二季度 |\n| --- | --- | --- |\n"
        "| 华北 | 10 | 12 |\n| 华南 | 8 | 9 |",
    )
    tools["add_pptx_chart"](
        path="deck.pptx",
        slide=1,
        chart_type="bar",
        data="| 季度 | 收入 |\n| --- | --- |\n| 一 | 100 |\n| 二 | 120 |",
        title="收入",
    )


def _audio_deck(**options: object) -> Callable[[dict[str, Callable[..., object]]], None]:
    def build(tools: dict[str, Callable[..., object]]) -> None:
        tools["write_pptx"](path="deck.pptx", content="# Narrated\n- first\n- second")
        _wav("narration.wav")
        tools["add_pptx_audio"](path="deck.pptx", slide=1, audio_path="narration.wav", **options)

    return build


def _audio_after_animation(tools: dict[str, Callable[..., object]]) -> None:
    # The audio shares the slide's existing timing tree instead of making a second one.
    tools["write_pptx"](path="deck.pptx", content="# Narrated\n- first\n- second")
    tools["add_pptx_animation"](path="deck.pptx", slide=1, shape_index=0, animation="fade")
    _wav("narration.wav")
    tools["add_pptx_audio"](path="deck.pptx", slide=1, audio_path="narration.wav", hidden=True)


DECKS: dict[str, Callable[[dict[str, Callable[..., object]]], None]] = {
    "audio_after_animation": _audio_after_animation,
    "audio_auto_visible": _audio_deck(left_in=8.0, top_in=0.5, start_delay=1.5),
    "audio_on_click": _audio_deck(trigger="on-click"),
    "cjk_bullets": _cjk_bullets,
    "cjk_table_and_chart": _cjk_table_and_chart,
    "cjk_template_fill": _template_fill(_CJK_FILL),
    "background_image": _background_image,
    "bullets": _bullets,
    "chart_bar": _chart("bar", _ONE_SERIES),
    "chart_bar_three_series": _chart("bar", _THREE_SERIES),
    "chart_line": _chart("line", _ONE_SERIES),
    "chart_pie": _chart("pie", _SHARES),
    "icon": _icon,
    "icon_list": _icon_list,
    "image": _image,
    "image_cropped_and_replaced": _image_cropped_and_replaced,
    "image_effects": _image_effects,
    "shapes": _shapes,
    "stat_callout": _stat_callout,
    "svg": _svg,
    "smartart_shape_attributes": _smartart_shape_attributes,
    "smartart_text_edit": _smartart_text_edit,
    "table": _table,
    "template_fill_fixed_slides_only": _template_fill(_FILL_FIXED_ONLY),
    "template_fill_grown": _template_fill(_FILL_GROWN),
    "template_fill_shrunk": _template_fill(_FILL_SHRUNK),
    "two_column": _two_column,
    "write_with_template_path": _write_with_template_path,
    "transition_and_animation": _transition_and_animation,
    "theme": _theme,
}


def build_deck(name: str, root: Path) -> Path:
    """Build deck `name` in `root` and return its path."""
    root.mkdir(parents=True, exist_ok=True)
    # Builders copy their source images next to the deck, by relative path.
    with contextlib.chdir(root):
        DECKS[name](_tools(root))
    return root / "deck.pptx"


@pytest.fixture(autouse=True)
def _no_libreoffice(monkeypatch: pytest.MonkeyPatch) -> None:
    # The LibreOffice QA pass only reads the deck; skipping it keeps these tests fast and the
    # same on machines with and without LibreOffice.
    monkeypatch.setattr("shutil.which", lambda name: None)


# The slide master, layouts and theme are the same in nearly every deck; stored once, they
# would otherwise be most of every golden. There are two masters: the 16:9 one `write_pptx`
# makes, and python-pptx's stock 4:3 one (SmartArt and template decks). A deck whose copy
# differs from both (the theme deck) still shows its own full text, so theme code stays covered.
SHARED_GOLDEN = GOLDEN_DIR / "_shared.txt"
SHARED_GOLDEN_4X3 = GOLDEN_DIR / "_shared_4x3.txt"
_SHARED_PREFIXES = ("ppt/slideMasters/", "ppt/slideLayouts/", "ppt/theme/")
_SHARED_SOURCE_DECK = "bullets"


def _stock_4x3_deck(root: Path) -> Path:
    from pptx import Presentation

    root.mkdir(parents=True, exist_ok=True)
    Presentation().save(str(root / "deck.pptx"))
    return root / "deck.pptx"


def _shared_parts(parts: dict[str, str]) -> dict[str, str]:
    return {name: text for name, text in parts.items() if name.startswith(_SHARED_PREFIXES)}


def _render(parts: dict[str, str]) -> str:
    return "".join(f"===== {name}\n{text}" for name, text in parts.items())


def _parse(text: str) -> dict[str, str]:
    parts: dict[str, str] = {}
    for chunk in re.split(r"^===== ", text, flags=re.MULTILINE)[1:]:
        name, _, body = chunk.partition("\n")
        parts[name] = body
    return parts


def _deck_golden_text(parts: dict[str, str], shared: dict[Path, dict[str, str]]) -> str:
    def reduced(name: str, text: str) -> str:
        for golden, golden_parts in shared.items():
            if golden_parts.get(name) == text:
                return f"<same as {golden.name}>\n"
        return text

    return _render({name: reduced(name, text) for name, text in parts.items()})


def _expected_and_diff(golden: Path, actual: str) -> tuple[str, str]:
    if os.environ.get("UPDATE_GOLDENS") == "1":
        golden.parent.mkdir(parents=True, exist_ok=True)
        golden.write_text(actual, encoding="utf-8", newline="\n")
    expected = golden.read_text(encoding="utf-8")
    diff = "".join(
        difflib.unified_diff(
            expected.splitlines(keepends=True),
            actual.splitlines(keepends=True),
            f"golden/pptx/{golden.name}",
            "built now",
            n=2,
        )
    )
    return expected, (
        f"{golden.name} differs from what the tools build; if the change is intended, run "
        f"UPDATE_GOLDENS=1 pytest tests/test_office_golden.py and commit the diff:\n{diff[:4000]}"
    )


def _built_shared(tmp_path: Path) -> dict[Path, dict[str, str]]:
    plain = build_deck(_SHARED_SOURCE_DECK, tmp_path / "_shared")
    stock = _stock_4x3_deck(tmp_path / "_shared_4x3")
    return {
        SHARED_GOLDEN: _shared_parts(canonical_parts(plain)),
        SHARED_GOLDEN_4X3: _shared_parts(canonical_parts(stock)),
    }


@pytest.mark.parametrize("golden", [SHARED_GOLDEN, SHARED_GOLDEN_4X3], ids=lambda p: p.stem)
def test_shared_golden_matches_the_plain_deck(golden: Path, tmp_path: Path) -> None:
    actual = _render(_built_shared(tmp_path)[golden])
    expected, message = _expected_and_diff(golden, actual)
    assert actual == expected, message


@pytest.mark.parametrize("name", sorted(DECKS))
def test_deck_matches_its_golden(name: str, tmp_path: Path) -> None:
    parts = canonical_parts(build_deck(name, tmp_path / name))
    # When regenerating, the shared goldens may not be written yet (test order is not fixed).
    shared = (
        _built_shared(tmp_path)
        if os.environ.get("UPDATE_GOLDENS") == "1"
        else {
            golden: _parse(golden.read_text(encoding="utf-8"))
            for golden in (SHARED_GOLDEN, SHARED_GOLDEN_4X3)
        }
    )
    actual = _deck_golden_text(parts, shared)
    expected, message = _expected_and_diff(GOLDEN_DIR / f"{name}.txt", actual)
    assert actual == expected, message


@pytest.mark.parametrize("golden", [SHARED_GOLDEN, SHARED_GOLDEN_4X3], ids=lambda p: p.stem)
def test_the_shared_golden_round_trips_through_the_parser(golden: Path) -> None:
    text = golden.read_text(encoding="utf-8")
    assert _render(_parse(text)) == text


def test_every_golden_belongs_to_a_deck() -> None:
    stale = sorted(
        path.stem
        for path in GOLDEN_DIR.glob("*.txt")
        if path.stem not in DECKS and path not in (SHARED_GOLDEN, SHARED_GOLDEN_4X3)
    )
    assert stale == [], f"goldens with no deck that builds them: {stale}"
