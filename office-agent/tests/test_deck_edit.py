"""The Slides panel's edits (decision 0016): a slide read as the editor's parts, and saved back
as a section, after which the deck is rebuilt; a save over a newer slide or a deck changed
outside coscribe is refused."""

from __future__ import annotations

import base64
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation

from coscribe.tools.office.pptx import (
    DeckChangedError,
    SectionError,
    SlideChanged,
    read_deck,
    save_slide,
    slide_parts,
)
from coscribe.tools.office.pptx.engine import build_deck

from .test_deck_engine import COVER
from .test_deck_sections import SECTION
from .test_deck_sections_objects import CHART


def _deck(tmp_path: Path, *slides: tuple[str, str]) -> Path:
    folder = tmp_path / "pages"
    (folder / "assets").mkdir(parents=True)
    Image.new("RGB", (96, 44), (30, 120, 200)).save(folder / "assets" / "photo.png")
    for name, text in slides:
        (folder / name).write_text(text, encoding="utf-8")
    assert build_deck(folder, tmp_path / "deck.pptx", resource_root=tmp_path).ok
    return tmp_path / "deck.pptx"


def _picture(part: dict) -> str:
    return base64.b64decode(part["picture"].split(",", 1)[1]).decode("utf-8")


def test_a_section_reads_as_its_parts(tmp_path: Path) -> None:
    deck = _deck(tmp_path, ("01.html", SECTION), ("02.html", CHART))
    first = slide_parts(deck, 0, tmp_path)
    assert first["kind"] == "section" and first["source"] == SECTION
    kinds = [p["kind"] for p in first["parts"]]
    assert kinds == ["drawing", "text", "text", "text", "drawing", "drawing"]
    band, title = first["parts"][0], first["parts"][1]
    assert band["box"] == [0, 0, 1920, 16] and title["picture"] is None
    assert 'fill="#C2410C"' in _picture(band)
    assert "data:image/png;base64," in _picture(first["parts"][4])
    chart = slide_parts(deck, 1, tmp_path)["parts"]
    assert 'data-native="chart"' in _picture(chart[1]) and "data-icon" not in _picture(chart[2])


def test_an_svg_page_reads_as_a_section_and_is_left_as_it_is(tmp_path: Path) -> None:
    deck = _deck(tmp_path, ("01.svg", COVER))
    parts = slide_parts(deck, 0, tmp_path)
    assert parts["kind"] == "page" and parts["source"].startswith("<section")
    assert (tmp_path / "pages" / "01.svg").is_file()
    assert not (tmp_path / "pages" / "01.html").exists()


def test_a_save_writes_the_section_and_builds_the_deck(tmp_path: Path) -> None:
    deck = _deck(tmp_path, ("01.svg", COVER), ("02.html", SECTION))
    read = slide_parts(deck, 0, tmp_path)
    moved = read["source"].replace("<section", '<section data-transition="fade"', 1)
    saved = save_slide(deck, 0, source=moved, version=read["version"], resource_root=tmp_path)
    assert saved["name"] == "01.html" and saved["version"] != read["version"]
    assert not (tmp_path / "pages" / "01.svg").exists()
    view = read_deck(deck)
    assert [s["kind"] for s in view["slides"]] == ["section", "section"]
    assert view["slides"][0]["transition"] == "fade" and not view["changed"]

    again = save_slide(deck, 0, source=moved, version=saved["version"], resource_root=tmp_path)
    assert again["version"] == saved["version"]
    with pytest.raises(SlideChanged):
        save_slide(deck, 0, source=moved, version=read["version"], resource_root=tmp_path)


def test_a_save_that_would_lose_work_or_is_not_a_section_is_refused(tmp_path: Path) -> None:
    deck = _deck(tmp_path, ("01.html", SECTION))
    version = slide_parts(deck, 0, tmp_path)["version"]
    with pytest.raises(SectionError):
        save_slide(
            deck, 0, source="<section><script/></section>", version=version, resource_root=tmp_path
        )
    prs = Presentation(str(deck))
    prs.slides[0].shapes.add_textbox(0, 0, 100, 100).text = "typed in PowerPoint"
    prs.save(str(deck))
    with pytest.raises(DeckChangedError):
        save_slide(deck, 0, source=SECTION, version=version, resource_root=tmp_path)
    assert (tmp_path / "pages" / "01.html").read_text(encoding="utf-8") == SECTION
